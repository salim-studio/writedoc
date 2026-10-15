"""Original font-metric controls derived from the repository's OFL Bebas Neue.
Only names, selected OS/2 metrics, checksums and DSIG removal differ.
"""
from pathlib import Path
import struct,json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures/script-metric-controls';ROOT.mkdir(exist_ok=True)
SOURCE=Path(__file__).resolve().parents[1]/'public/fonts/BebasNeue-Regular.ttf'
original=SOURCE.read_bytes()
def checksum(b):
 b=b+b'\0'*((-len(b))%4);return sum(struct.unpack('>'+str(len(b)//4)+'I',b))&0xffffffff
def build(family,changes):
 tables={}
 for i in range(struct.unpack_from('>H',original,4)[0]):
  tag=original[12+i*16:16+i*16];off,length=struct.unpack_from('>II',original,20+i*16);tables[tag]=bytearray(original[off:off+length])
 name=tables[b'name'];count=struct.unpack_from('>H',name,2)[0];start=struct.unpack_from('>H',name,4)[0];records=[];strings=bytearray()
 for i in range(count):
  platform,encoding,language,ident,length,offset=struct.unpack_from('>6H',name,6+i*12);value=bytes(name[start+offset:start+offset+length])
  if ident in [1,3,4,6,16]:value=(family.replace(' ','') if ident==6 else family).encode('utf-16-be' if platform in [0,3] else 'mac_roman')
  if ident in [2,17]:value='Regular'.encode('utf-16-be' if platform in [0,3] else 'mac_roman')
  records.append(struct.pack('>6H',platform,encoding,language,ident,len(value),len(strings)));strings.extend(value)
 tables[b'name']=struct.pack('>3H',0,len(records),6+12*len(records))+b''.join(records)+strings
 for offset,value in changes.items():struct.pack_into('>h',tables[b'OS/2'],int(offset),value)
 struct.pack_into('>I',tables[b'head'],8,0)
 tables.pop(b'DSIG',None);n=len(tables);power=2**(n.bit_length()-1);header=struct.pack('>I4H',0x10000,n,power*16,power.bit_length()-1,n*16-power*16)
 directory=bytearray();data=bytearray();offset=12+n*16;headOffset=0
 for tag,b in sorted(tables.items()):
  directory.extend(struct.pack('>4sIII',tag,checksum(b),offset,len(b)))
  if tag==b'head':headOffset=offset
  padded=bytes(b)+b'\0'*((-len(b))%4);data.extend(padded);offset+=len(padded)
 result=bytearray(header+directory+data);struct.pack_into('>I',result,headOffset+8,(0xb1b0afba-checksum(result))&0xffffffff)
 path=ROOT/(family.replace(' ','')+'.ttf');path.write_bytes(result);return path
variants=[('Base',{})]+[(f'Sup {v}',{24:v}) for v in [0,75,150,250,450,700]]+[(f'Sub {v}',{16:v}) for v in [0,50,100,150,250,350]]+[('Sup Half',{20:500}),('Sup Wide',{18:850}),('Sup Tall',{20:850}),('Sup X',{22:200}),('Sub Tall',{12:850})]
doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(12);doc.styles['Normal'].paragraph_format.space_after=Pt(0)
paths=[]
for i,(variant,changes) in enumerate(variants):
 family='Script Metric '+variant;paths.append(build(family,changes).name)
 p=doc.add_paragraph(family);p.paragraph_format.page_break_before=i>0
 for script in ['baseline','superscript','subscript']:
  p=doc.add_paragraph();p.paragraph_format.line_spacing=1
  for text,active in [('START ',False),('TEST',True),(' END',False)]:
   r=p.add_run(text);r.font.name=family;r.font.size=Pt(20)
   if active and script!='baseline':E.SubElement(r._r.get_or_add_rPr(),qn('w:vertAlign'),{qn('w:val'):script})
doc.save(ROOT/'controls.docx');(ROOT/'fonts.json').write_text(json.dumps(paths))
print('Generated',len(paths),'independent metric font variants')
