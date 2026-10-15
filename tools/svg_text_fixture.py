"""Original SVG text controls with a contrasting magenta raster fallback."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import struct, zlib
from lxml import etree as E
from docx import Document
from docx.shared import Inches
from canvas_group_fixture import NS, tag, add

ROOT=Path(__file__).resolve().parents[1]
FOLDER=ROOT/'tests/fixtures/svg-text'
FOLDER.mkdir(parents=True,exist_ok=True)
def chunk(t,b): return struct.pack('>I',len(b))+t+b+struct.pack('>I',zlib.crc32(t+b)&0xffffffff)
png=b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('>IIBBBBB',2,2,8,2,0,0,0))+chunk(b'IDAT',zlib.compress(b'\0\xff\0\xff\xff\0\xff'*2))+chunk(b'IEND',b'')
texts=[
 '<text x="20" y="45" font-size="24">Alpha Beta</text>',
 '<text x="20" y="45" font-family="Arial" font-size="24">Alpha Beta</text>',
 '<text x="20" y="45" font-family="Calibri" font-size="24">Alpha Beta</text>',
 '<text x="20" y="45" font-family="Times New Roman" font-size="24">Alpha Beta</text>',
 '<text x="20" y="45" font-family="MissingControlFont92837" font-size="24">Alpha Beta</text>',
 '<text x="20" y="45" font-family="Arial" font-size="24" font-weight="bold" font-style="italic">Alpha Beta</text>',
 '<text transform="translate(20,40)" font-family="Arial" font-size="24"><tspan x="0" y="0">Alpha</tspan><tspan x="10" y="42" font-size="18">Beta</tspan></text>',
 '<style>.label {font-family:Arial; font-size:24px; text-anchor:middle; letter-spacing:2px;}</style><text class="label" x="150" y="45">Alpha Beta</text>',
 '<text x="20" y="45" font-family="Arial">Alpha Beta</text>',
 '<text x="20" y="45" font-family="Arial" font-size="24" letter-spacing="4">AlphaBeta</text>',
 '<text x="20" y="45" font-family="Arial" font-size="24" letter-spacing="4">Alpha Beta</text>',
 '<text x="200" y="45" font-family="Arial" font-size="24" letter-spacing="-1" text-anchor="end" xml:space="preserve">A  B</text>',
 '<text x="20" y="45" font-family="Arial" font-size="24" letter-spacing="2" word-spacing="3">A B C</text>',
 '<text x="20" y="45" font-family="Arial" font-size="24" letter-spacing="2">A\u00a0B</text>',
]
doc=Document(); doc.settings.element.find(tag('w:zoom')).set(tag('w:percent'),'100')
section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
for key in ['left_margin','right_margin','top_margin','bottom_margin']:setattr(section,key,Inches(1))
for i,content in enumerate(texts):
 p=doc.add_paragraph('SVG text control '+str(i+1));p.paragraph_format.page_break_before=bool(i)
 inline=doc.add_paragraph().add_run().add_picture(BytesIO(png),width=Inches(300/96),height=Inches(120/96))._inline
 blip=inline.find('.//'+tag('a:blip')); ext=add(add(blip,'a:extLst'),'a:ext',uri='{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
 E.SubElement(ext,'{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',{tag('r:embed'):'rIdSvg'+str(i)},nsmap={'asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
 doc.add_paragraph('Following paragraph')
out=BytesIO();doc.save(out)
with ZipFile(out) as z:parts={n:z.read(n) for n in z.namelist()}
rels=E.fromstring(parts['word/_rels/document.xml.rels'])
for i,content in enumerate(texts):
 E.SubElement(rels,rels.tag.replace('Relationships','Relationship'),Id='rIdSvg'+str(i),Type=NS['r']+'/image',Target=f'media/text{i}.svg')
 parts[f'word/media/text{i}.svg']=('<svg xmlns="http://www.w3.org/2000/svg" width="300" height="120" viewBox="0 0 300 120"><rect width="300" height="120" fill="#EEEEEE"/><g fill="#E02030">'+content+'</g></svg>').encode()
parts['word/_rels/document.xml.rels']=E.tostring(rels)
types=E.fromstring(parts['[Content_Types].xml']);E.SubElement(types,types.tag.replace('Types','Default'),Extension='svg',ContentType='image/svg+xml');parts['[Content_Types].xml']=E.tostring(types)
target=FOLDER/'svg-text.docx'; old={}
if target.exists():
 with ZipFile(target) as z:old={n:z.read(n) for n in z.namelist()}
if old!=parts:
 with ZipFile(target,'w',ZIP_DEFLATED) as z:
  for n,b in parts.items():z.writestr(n,b)
