"""Original no-gap cell-spacing controls, independent of engine code."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'tests/fixtures/zero-cell-spacing'; OUT.mkdir(exist_ok=True)
cases=[('omitted',None,None,None),('row-zero','row','dxa','0'),('row-default-unit','row',None,'0'),('row-auto','row','auto','240'),('row-percent','row','pct','500'),('row-nil','row','nil','240'),('table-zero','table','dxa','0'),('table-auto','table','auto','240'),('table-percent','table','pct','500'),('exception-zero','exception','dxa','0')]
d=Document();d.sections[0].page_width=Inches(8.5);d.sections[0].page_height=Inches(11)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
d.sections[0].left_margin=d.sections[0].right_margin=Inches(1)
d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
d.styles['Normal'].paragraph_format.space_after=Pt(0);d.styles['Normal'].paragraph_format.line_spacing=Pt(14)
def table(parent,nested=False):
 t=parent.add_table(rows=2,cols=2);t.autofit=False
 for col in t.columns:col.width=Inches(.75 if nested else 2)
 for row in t.rows:
  for j,cell in enumerate(row.cells):
   cell.width=Inches(.75 if nested else 2);cell.text=('Inner' if nested else 'Outer')+str(j+1)
   shd=OxmlElement('w:shd');shd.set(qn('w:val'),'clear');shd.set(qn('w:color'),'auto');shd.set(qn('w:fill'),'C9E7B5' if nested else 'DDEEFF');cell._tc.get_or_add_tcPr().append(shd)
 if not nested:table(t.cell(1,0),True)
 return t
for i,(name,level,kind,value) in enumerate(cases):
 p=d.add_paragraph(name);p.paragraph_format.page_break_before=bool(i);p.paragraph_format.space_after=Pt(6)
 t=table(d)
 if level:
  for tbl in t._tbl.iter(qn('w:tbl')):
   targets=[tbl.find(qn('w:tblPr'))] if level=='table' else []
   if level!='table':
    for row in tbl.findall(qn('w:tr')):
     tag='w:trPr' if level=='row' else 'w:tblPrEx'; pr=OxmlElement(tag);row.insert(0,pr);targets.append(pr)
   for pr in targets:
    spacing=OxmlElement('w:tblCellSpacing');spacing.set(qn('w:w'),value)
    if kind:spacing.set(qn('w:type'),kind)
    pr.insert_element_before(spacing,'w:tblInd','w:tblBorders','w:shd','w:tblLayout','w:tblCellMar','w:tblLook')
d.save(OUT/'source.docx');(OUT/'cases.json').write_text(json.dumps([dict(name=n,level=l,type=t,value=v) for n,l,t,v in cases],indent=2)+'\n',encoding='utf-8')
print('Created',len(cases),'controls')
