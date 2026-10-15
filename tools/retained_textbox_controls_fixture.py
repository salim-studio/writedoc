"""Original retained drawing controls: block, inline and nested saved content."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import canvas_group_fixture as g
out=Path(__file__).resolve().parents[1]/'tests/fixtures/retained-textbox-controls';out.mkdir(exist_ok=True)
doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12);normal.paragraph_format.space_after=Pt(0)
ident=0
for grouped in [False,True]:
 for kind in ['block','inline','nested']:
  ident+=1;p=doc.add_paragraph(f'Case {ident}: {kind}, group={grouped}');p.paragraph_format.page_break_before=ident>1
  frame=g.add(g.add(doc.add_paragraph().add_run()._r,'w:drawing'),'wp:inline',distT=0,distB=0,distL=0,distR=0)
  width,height=240*12700,90*12700;g.add(frame,'wp:extent',cx=width,cy=height);g.add(frame,'wp:docPr',id=ident,name=f'Control {ident}');g.add(frame,'wp:cNvGraphicFramePr')
  data=g.add(g.add(frame,'a:graphic'),'a:graphicData',uri=g.NS['wpg' if grouped else 'wps']);parent=data
  if grouped:
   parent=g.add(data,'wpg:wgp');g.add(parent,'wpg:cNvGrpSpPr');props=g.add(parent,'wpg:grpSpPr');g.transform(props,width,height,group=True);props[0].find(g.tag('a:chOff')).attrib.update({'x':'0','y':'0'})
  shape=g.shape(parent,ident+100,width,height,0,0,'DDEEFF',f'Saved label {ident}')
  content=shape.find('.//'+g.tag('w:txbxContent'));paragraph=content.find(g.tag('w:p'))
  def wrap(node,extra):
   at=node.getparent();idx=at.index(node);at.remove(node);sdt=g.add(at,'w:sdt');at.remove(sdt);at.insert(idx,sdt);pr=g.add(sdt,'w:sdtPr');g.add(pr,'w:alias',**{'w:val':'Retained control'});g.add(pr,'w:id',**{'w:val':ident*10+extra})
   if extra==1:g.add(pr,'w:lock',**{'w:val':'sdtContentLocked'});g.add(pr,'w:richText')
   else:date=g.add(pr,'w:date');g.add(date,'w:dateFormat',**{'w:val':'MMMM d, yyyy'});g.add(date,'w:lid',**{'w:val':'en-US'})
   g.add(sdt,'w:sdtContent').append(node)
  if kind in ['inline','nested']:wrap(paragraph.find(g.tag('w:r')),2)
  if kind in ['block','nested']:wrap(paragraph,1)
  for run in content.iter(g.tag('w:r')):
   rp=run.find(g.tag('w:rPr'))
   if rp is None:rp=g.add(run,'w:rPr');run.remove(rp);run.insert(0,rp)
   fonts=rp.find(g.tag('w:rFonts'))
   if fonts is None:fonts=g.add(rp,'w:rFonts',**{'w:ascii':'Arial','w:hAnsi':'Arial'})
  doc.add_paragraph('After drawing')
doc.save(out/'control.docx')

# Same independently authored content without wrappers, to separate envelope
# projection from pre-existing text rasterization.
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
with ZipFile(out/'control.docx') as z:parts={n:z.read(n) for n in z.namelist()}
r=E.fromstring(parts['word/document.xml'])
for control in reversed(r.findall('.//'+g.tag('w:sdt'))):
 parent=control.getparent();at=parent.index(control)
 for node in list(control.find(g.tag('w:sdtContent'))):parent.insert(at,node);at+=1
 parent.remove(control)
parts['word/document.xml']=E.tostring(r,xml_declaration=True,encoding='utf-8',standalone=True)
with ZipFile(out/'unwrapped.docx','w',ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
