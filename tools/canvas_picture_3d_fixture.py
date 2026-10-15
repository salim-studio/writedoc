"""Original anchored 3D picture probe; python-docx/lxml and our own PNG."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile,ZIP_DEFLATED
from docx import Document
from docx.shared import Inches,Pt
from lxml import etree as E
root=Path(__file__).resolve().parents[1]
folder=root/'tests/fixtures/picture-3d';folder.mkdir(exist_ok=True)
with ZipFile(root/'public/fixtures/shape-picture-fills.docx') as z:png=z.read('word/media/image1.png')
d=Document();d.settings.element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}zoom').set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}percent','100');d.sections[0].page_width=Inches(8.5);d.sections[0].page_height=Inches(11)
cases=[('front','orthographicFront',0,0,'rect',144,108),('right-flat','perspectiveRight',0,0,'rect',144,108),('right-extrusion','perspectiveRight',20,0,'rect',144,108),('right-ellipse','perspectiveRight',20,4.5,'ellipse',144,108),('right-large','perspectiveRight',20,0,'rect',288,216),('left-extrusion','perspectiveLeft',20,0,'rect',144,108),('ellipse-contour-only','perspectiveRight',0,4.5,'ellipse',144,108),('ellipse-extrusion-only','perspectiveRight',20,0,'ellipse',144,108),('rectangle-contour','perspectiveRight',20,4.5,'rect',144,108)]
cases.extend([('right-flat-ellipse','perspectiveRight',0,0,'ellipse',144,108),('front-flat-ellipse','orthographicFront',0,0,'ellipse',144,108),('front-contour-ellipse','orthographicFront',0,4.5,'ellipse',144,108),('front-contour-rectangle','orthographicFront',0,4.5,'rect',144,108)])
A='{http://schemas.openxmlformats.org/drawingml/2006/main}';P='{http://schemas.openxmlformats.org/drawingml/2006/picture}'
for i,(name,camera,depth,contour,geom,w,h) in enumerate(cases):
 if i:d.add_page_break()
 d.add_paragraph(name)
 inline=d.add_paragraph().add_run().add_picture(BytesIO(png),width=Pt(w),height=Pt(h))._inline
 props=inline.find('.//'+P+'spPr');props.find(A+'prstGeom').set('prst',geom)
 scene=E.SubElement(props,A+'scene3d');E.SubElement(scene,A+'camera',prst=camera);E.SubElement(scene,A+'lightRig',rig='threePt',dir='t')
 sp=E.SubElement(props,A+'sp3d',extrusionH=str(round(depth*12700)),contourW=str(round(contour*12700)))
 if contour:E.SubElement(E.SubElement(sp,A+'contourClr'),A+'srgbClr',val='202060')
d.save(root/'public/fixtures/picture-3d.docx')
WP='{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
with ZipFile(root/'public/fixtures/picture-3d.docx') as z:parts={n:z.read(n) for n in z.namelist()}
r=E.fromstring(parts['word/document.xml'])
for inline in r.iter(WP+'inline'):
 inline.tag=WP+'anchor';inline.attrib.update(dict(distT='0',distB='0',distL='0',distR='0',simplePos='0',relativeHeight='1',behindDoc='0',locked='0',layoutInCell='1',allowOverlap='1'))
 off=E.Element(WP+'simplePos',x='0',y='0');inline.insert(0,off)
 for j,axis,val in [(1,'H',100),(2,'V',140)]:
  pos=E.Element(WP+'position'+axis,relativeFrom='page');E.SubElement(pos,WP+'posOffset').text=str(val*12700);inline.insert(j,pos)
 extent=inline.find(WP+'extent');idx=list(inline).index(extent)+1
 inline.insert(idx,E.Element(WP+'effectExtent',l='508000',t='508000',r='508000',b='508000'));inline.insert(idx+1,E.Element(WP+'wrapNone'))
parts['word/document.xml']=E.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(root/'public/fixtures/picture-3d.docx','w',ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
import json
(folder/'cases.json').write_text(json.dumps(cases,indent=2))
