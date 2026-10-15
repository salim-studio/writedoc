"""Independent DOCX check using python-docx (MIT), installed in a separate venv.

python tools/verify-canvas-docx.py --generate-fixture
python tools/verify-canvas-docx.py path/to/edited.docx
"""
import argparse
import sys
sys.dont_write_bytecode = True
from pathlib import Path
from docx import Document
from docx.table import _Cell
from docx.shared import Pt, Inches, Mm
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

parser = argparse.ArgumentParser()
parser.add_argument("paths", nargs="*")
parser.add_argument("--generate-fixture", action="store_true")
parser.add_argument("--generate-formatting-fixture", action="store_true")
parser.add_argument("--paragraph-formatting", action="store_true")
parser.add_argument("--generate-font-fixture", action="store_true")
parser.add_argument("--font-formatting", action="store_true")
parser.add_argument("--generate-page-fixture", action="store_true")
parser.add_argument("--page-settings", action="store_true")
parser.add_argument("--generate-indent-fixture", action="store_true")
parser.add_argument("--indents-tabs", action="store_true")
parser.add_argument("--generate-pagination-fixture", action="store_true")
parser.add_argument("--pagination", action="store_true")
parser.add_argument("--generate-numbering-fixture", action="store_true")
parser.add_argument("--numbering", action="store_true")
parser.add_argument("--new-lists", action="store_true")
parser.add_argument("--generate-styles-fixture", action="store_true")
parser.add_argument("--styles", action="store_true")
parser.add_argument("--new-styles", action="store_true")
parser.add_argument("--generate-character-styles-fixture", action="store_true")
parser.add_argument("--character-styles", action="store_true")
parser.add_argument("--new-character-styles", action="store_true")
parser.add_argument("--generate-tables-fixture", action="store_true")
parser.add_argument("--tables", action="store_true")
parser.add_argument("--new-tables", action="store_true")
parser.add_argument("--generate-autofit-fixture", action="store_true")
parser.add_argument("--autofit-tables", action="store_true")
parser.add_argument("--new-autofit-tables", action="store_true")
parser.add_argument("--generate-merged-cells-fixture", action="store_true")
parser.add_argument("--merged-cells", action="store_true")
parser.add_argument("--new-merged-cells", action="store_true")
parser.add_argument("--generate-vertical-cells-fixture", action="store_true")
parser.add_argument("--vertical-cells", action="store_true")
parser.add_argument("--new-vertical-cells", action="store_true")
parser.add_argument("--generate-vertical-pages-fixture", action="store_true")
parser.add_argument("--vertical-pages", action="store_true")
parser.add_argument("--new-vertical-pages", action="store_true")
parser.add_argument("--generate-table-headers-fixture", action="store_true")
parser.add_argument("--table-headers", action="store_true")
parser.add_argument("--new-table-headers", action="store_true")
parser.add_argument("--generate-inline-images-fixture", action="store_true")
parser.add_argument("--inline-images", action="store_true")
parser.add_argument("--new-inline-images", action="store_true")
parser.add_argument("--generate-cropped-images-fixture", action="store_true")
parser.add_argument("--cropped-images", action="store_true")
parser.add_argument("--new-cropped-images", action="store_true")
parser.add_argument("--generate-flipped-images-fixture", action="store_true")
parser.add_argument("--flipped-images", action="store_true")
parser.add_argument("--new-flipped-images", action="store_true")
parser.add_argument("--generate-rotated-images-fixture", action="store_true")
parser.add_argument("--rotated-images", action="store_true")
parser.add_argument("--new-rotated-images", action="store_true")
parser.add_argument("--generate-outline-fixture", action="store_true")
parser.add_argument("--outline-levels", action="store_true")
parser.add_argument("--new-outline-levels", action="store_true")
parser.add_argument("--generate-colors-fixture", action="store_true")
parser.add_argument("--text-colors", action="store_true")
parser.add_argument("--new-text-colors", action="store_true")
parser.add_argument("--generate-highlights-fixture", action="store_true")
parser.add_argument("--text-highlights", action="store_true")
parser.add_argument("--new-text-highlights", action="store_true")
parser.add_argument("--generate-toggle-fixtures", action="store_true")
parser.add_argument("--style-toggles-off", action="store_true")
parser.add_argument("--style-toggles-on", action="store_true")
parser.add_argument("--generate-script-fixture", action="store_true")
parser.add_argument("--script-text", action="store_true")
parser.add_argument("--new-script-text", action="store_true")
parser.add_argument("--generate-position-fixture", action="store_true")
parser.add_argument("--baseline-shifts", action="store_true")
parser.add_argument("--new-baseline-shifts", action="store_true")
parser.add_argument("--generate-spacing-fixture", action="store_true")
parser.add_argument("--character-spacing", action="store_true")
parser.add_argument("--new-character-spacing", action="store_true")
parser.add_argument("--generate-width-fixture", action="store_true")
parser.add_argument("--text-width", action="store_true")
parser.add_argument("--new-text-width", action="store_true")
parser.add_argument("--generate-contextual-fixtures", action="store_true")
parser.add_argument("--contextual-spacing", choices=["body", "table", "new"])
parser.add_argument("--generate-kerning-fixture", action="store_true")
parser.add_argument("--kerning", action="store_true")
parser.add_argument("--new-kerning", action="store_true")
parser.add_argument("--generate-spaces-fixtures", action="store_true")
parser.add_argument("--spaces", choices=["trailing", "wrapped", "new"])
parser.add_argument("--generate-strike-fixture", action="store_true")
parser.add_argument("--strikethrough", action="store_true")
parser.add_argument("--new-strikethrough", action="store_true")
parser.add_argument("--generate-caps-fixture", action="store_true")
parser.add_argument("--capitalization", action="store_true")
parser.add_argument("--new-capitalization", action="store_true")
parser.add_argument("--generate-underline-fixture", action="store_true")
parser.add_argument("--underlines", action="store_true")
parser.add_argument("--new-underlines", action="store_true")
parser.add_argument("--generate-hyperlink-fixture", action="store_true")
parser.add_argument("--hyperlinks", action="store_true")
parser.add_argument("--new-hyperlinks", action="store_true")
parser.add_argument("--generate-paragraph-border-fixture", action="store_true")
parser.add_argument("--generate-cell-alignment-fixture", action="store_true")
parser.add_argument("--generate-cell-shading-fixture", action="store_true")
parser.add_argument("--generate-cell-border-fixture", action="store_true")
parser.add_argument("--generate-cell-margin-fixture", action="store_true")
parser.add_argument("--generate-justification-fixture", action="store_true")
parser.add_argument("--generate-tab-leader-fixture", action="store_true")
parser.add_argument("--generate-row-height-fixture", action="store_true")
parser.add_argument("--generate-line-spacing-fixture", action="store_true")
parser.add_argument("--line-spacing", choices=['main', 'new'])
parser.add_argument("--generate-style-spacing-fixture", action="store_true")
parser.add_argument("--style-spacing", choices=['main', 'new'])
parser.add_argument("--generate-search-fixture", action="store_true")
parser.add_argument("--search", choices=['main', 'new'])
parser.add_argument("--generate-clear-format-fixture", action="store_true")
parser.add_argument("--clear-format", choices=['main', 'new'])
parser.add_argument("--generate-paragraph-reset-fixture", action="store_true")
parser.add_argument("--paragraph-reset", choices=['main', 'new'])
parser.add_argument("--row-heights", choices=['main', 'new'])
parser.add_argument("--tab-leaders", choices=['main', 'new'])
parser.add_argument("--justification", choices=['main', 'soft', 'new'])
parser.add_argument("--cell-margins", action="store_true")
parser.add_argument("--new-cell-margins", action="store_true")
parser.add_argument("--cell-borders", action="store_true")
parser.add_argument("--new-cell-borders", action="store_true")
parser.add_argument("--cell-shading", action="store_true")
parser.add_argument("--new-cell-shading", action="store_true")
parser.add_argument("--cell-alignment", action="store_true")
parser.add_argument("--new-cell-alignment", action="store_true")
parser.add_argument("--paragraph-borders", action="store_true")
parser.add_argument("--new-paragraph-borders", action="store_true")
parser.add_argument("--generate-paragraph-shading-fixture", action="store_true")
parser.add_argument("--paragraph-shading", action="store_true")
parser.add_argument("--new-paragraph-shading", action="store_true")
parser.add_argument("--generate-bookmark-fixture", action="store_true")
parser.add_argument("--bookmarks", action="store_true")
parser.add_argument("--new-bookmarks", action="store_true")
args = parser.parse_args()

def fixture_document():
    doc = Document()
    # The python-docx default template omits a required CT_Zoom attribute.
    # Keep the independent fixtures valid against the official transitional XSD.
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    return doc


if args.generate_justification_fixture:
    from canvas_justification_fixture import generate
    generate(fixture_document)

if args.generate_tab_leader_fixture:
    from canvas_tab_leader_fixture import generate
    generate(fixture_document)

if args.generate_row_height_fixture:
    from canvas_row_height_fixture import generate
    generate(fixture_document)

if args.generate_line_spacing_fixture:
    from canvas_line_spacing_fixture import generate
    generate(fixture_document)

if args.generate_style_spacing_fixture:
    from canvas_style_spacing_fixture import generate
    generate(fixture_document)

if args.generate_search_fixture:
    from canvas_search_fixture import generate
    generate(fixture_document)

if args.generate_clear_format_fixture:
    from canvas_clear_format_fixture import generate
    generate(fixture_document)

if args.generate_paragraph_reset_fixture:
    from canvas_paragraph_reset_fixture import generate
    generate(fixture_document)

if args.generate_cell_margin_fixture:
    from canvas_cell_margin_fixture import generate
    generate(fixture_document)

if args.generate_cell_border_fixture:
    from canvas_cell_border_fixture import generate
    generate(fixture_document)

if args.generate_cell_shading_fixture:
    from canvas_cell_shading_fixture import generate
    generate(fixture_document)

if args.generate_cell_alignment_fixture:
    from canvas_cell_alignment_fixture import generate
    generate(fixture_document)

if args.generate_paragraph_border_fixture:
    from canvas_paragraph_border_fixture import generate
    generate(fixture_document)

if args.generate_paragraph_shading_fixture:
    from canvas_paragraph_shading_fixture import generate
    generate(fixture_document)

if args.generate_bookmark_fixture:
    from canvas_bookmark_fixture import generate
    generate(fixture_document)

if args.generate_hyperlink_fixture:
    from canvas_hyperlink_fixture import generate
    generate(fixture_document)

if args.generate_underline_fixture:
    from canvas_underline_fixture import generate
    generate(fixture_document)

if args.generate_caps_fixture:
    from canvas_caps_fixture import generate
    generate(fixture_document)

if args.generate_strike_fixture:
    from canvas_strike_fixture import generate
    generate(fixture_document)

if args.generate_spaces_fixtures:
    from canvas_spaces_fixture import generate
    generate(fixture_document)

if args.generate_kerning_fixture:
    from canvas_kerning_fixture import generate
    generate(fixture_document)

if args.generate_contextual_fixtures:
    from canvas_contextual_fixture import generate
    generate(fixture_document)

if args.generate_width_fixture:
    from canvas_width_fixture import generate
    generate(fixture_document)

if args.generate_spacing_fixture:
    from canvas_spacing_fixture import generate
    generate(fixture_document)

if args.generate_position_fixture:
    from canvas_position_fixture import generate
    generate(fixture_document)

if args.generate_script_fixture:
    from canvas_vertical_fixture import generate
    generate(fixture_document)

if args.generate_toggle_fixtures:
    from canvas_toggle_fixture import generate
    generate(fixture_document)

if args.generate_highlights_fixture:
    from canvas_highlights_fixture import generate
    generate(fixture_document)

if args.generate_colors_fixture:
    from canvas_colors_fixture import generate
    generate(fixture_document)

if args.generate_fixture:
    doc = fixture_document()
    doc.styles["Normal"].font.name = "Noto Sans"
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_paragraph("A small beginning. Click anywhere in this paragraph and start writing.")
    doc.add_paragraph("Select words across these two paragraphs, apply bold, then undo. Download the DOCX and open it here again to check that your changes survived.")
    target = Path(__file__).resolve().parents[1] / "public/fixtures/two-paragraphs.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    print(f"Generated independent fixture: {target}")

if args.generate_formatting_fixture:
    doc = fixture_document()
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(9)
    title = doc.styles.add_style("Canvas Heading", WD_STYLE_TYPE.PARAGRAPH)
    title.base_style = normal
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.5
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(18)
    doc.add_paragraph("Formatting survives editing, undo, and reopening.", style=title)
    doc.add_paragraph("Every page is another chance to make the editor a little better. " * 60)
    target = Path(__file__).resolve().parents[1] / "public/fixtures/paragraph-formatting.docx"
    doc.save(target)
    print(f"Generated independent formatting fixture: {target}")

if args.generate_font_fixture:
    doc = fixture_document()
    normal = doc.styles['Normal']
    normal.font.name = 'Noto Sans'
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    heading = doc.styles.add_style('Canvas Font Heading', WD_STYLE_TYPE.PARAGRAPH)
    heading.base_style = normal
    heading.font.size = Pt(24)
    heading.font.italic = True
    doc.add_paragraph(style=heading).add_run('A 24-point heading').underline = True
    body = doc.add_paragraph()
    body.add_run('Small ')
    run = body.add_run('mixed')
    run.font.size = Pt(18)
    run.italic = True
    run.underline = True
    run = body.add_run(' text at 13.5 pt.')
    run.font.size = Pt(13.5)
    empty = doc.add_paragraph()
    mark = OxmlElement('w:rPr')
    for tag, value in [('i', '1'), ('sz', '36'), ('u', 'single')]:
        prop = OxmlElement('w:' + tag)
        prop.set(qn('w:val'), value)
        mark.append(prop)
    empty._p.get_or_add_pPr().append(mark)
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/font-formatting.docx'
    doc.save(target)
    print(f'Generated independent font fixture: {target}')

if args.generate_page_fixture:
    doc = fixture_document()
    normal = doc.styles['Normal']
    normal.font.name = 'Noto Sans'
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(9)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(297), Mm(210)
    section.left_margin, section.right_margin = Inches(.75), Inches(1)
    section.top_margin, section.bottom_margin = Inches(.5), Inches(.75)
    section.gutter = Inches(.125)
    p = doc.add_paragraph()
    run = p.add_run('Before break')
    run.bold = True
    run.add_break(WD_BREAK.PAGE)
    run = p.add_run('After break')
    run.italic = True
    run.underline = True
    run.font.size = Pt(18)
    run.add_break(WD_BREAK.LINE)
    run.add_text('Same paragraph')
    style = doc.styles.add_style('Start New Page', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = normal
    style.paragraph_format.page_break_before = True
    doc.add_paragraph('New page from style', style=style)
    doc.add_paragraph('Last paragraph')
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/page-settings.docx'
    doc.save(target)
    print(f'Generated independent page fixture: {target}')
    for enabled in [False, True]:
        check = fixture_document()
        normal = check.styles['Normal']
        normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
        normal.paragraph_format.line_spacing = 1.0
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        first = check.add_paragraph()
        first.paragraph_format.space_before = Pt(50)
        run = first.add_run()
        run.add_break(WD_BREAK.PAGE)
        run.add_text('Leading')
        check.add_paragraph('Trailing').add_run().add_break(WD_BREAK.PAGE)
        check.add_paragraph('Following')
        if enabled:
            compat = check.settings.element.find(qn('w:compat'))
            # Insert in CT_Compat order, ahead of useFELayout / compatSetting.
            compat.insert(0, OxmlElement('w:suppressSpBfAfterPgBrk'))
            index = next(i for i, child in enumerate(compat) if child.tag == qn('w:compatSetting'))
            compat.insert(index, OxmlElement('w:splitPgBreakAndParaMark'))
        check.save(target.with_name('page-breaks-compat.docx' if enabled else 'page-breaks-default.docx'))

if args.generate_indent_fixture:
    doc = fixture_document()
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(9)
    doc.sections[0].left_margin = doc.sections[0].right_margin = Inches(1)
    default = doc.settings.element.find(qn('w:defaultTabStop'))
    default.set(qn('w:val'), '720')
    base = doc.styles.add_style('Canvas Indents', WD_STYLE_TYPE.PARAGRAPH)
    base.base_style = normal
    pf = base.paragraph_format
    pf.left_indent, pf.right_indent, pf.first_line_indent = Inches(.75), Inches(.25), Inches(-.25)
    for pos, align in [(1.5, WD_TAB_ALIGNMENT.LEFT), (3, WD_TAB_ALIGNMENT.CENTER), (4.5, WD_TAB_ALIGNMENT.RIGHT), (6, WD_TAB_ALIGNMENT.DECIMAL)]:
        pf.tab_stops.add_tab_stop(Inches(pos), align)
    p = doc.add_paragraph(style=base)
    p.add_run('Lead\tLeft\tCenter\tRight\t')
    p.add_run('123').bold = True
    p.add_run('.45').italic = True
    derived = doc.styles.add_style('Canvas Derived Indents', WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = base
    derived.paragraph_format.right_indent = Inches(.5)
    derived.paragraph_format.first_line_indent = Inches(.375)
    p = doc.add_paragraph('Inherited stops\t42.50', style=derived)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1.5), WD_TAB_ALIGNMENT.CLEAR)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3), WD_TAB_ALIGNMENT.RIGHT)
    p = doc.add_paragraph('A\tHanging indent continues after a manual line break.', style=base)
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run('Wrap this paragraph with its indents intact. ' * 45)
    empty = doc.add_paragraph(style=normal)
    empty.paragraph_format.left_indent = Inches(.5)
    empty.paragraph_format.first_line_indent = Inches(.375)
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/indents-tabs.docx'
    doc.save(target)
    print(f'Generated independent indentation/tab fixture: {target}')

if args.generate_pagination_fixture:
    doc = fixture_document()
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing, pf.space_before, pf.space_after = Pt(18), Pt(0), Pt(0)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    base = doc.styles.add_style('Canvas Keep Base', WD_STYLE_TYPE.PARAGRAPH)
    base.base_style = normal
    base.paragraph_format.keep_with_next = True
    base.paragraph_format.keep_together = True
    base.paragraph_format.widow_control = False
    heading = doc.styles.add_style('Canvas Keep Heading', WD_STYLE_TYPE.PARAGRAPH)
    heading.base_style = base
    doc.add_paragraph('\n'.join(f'Filler line {i}' for i in range(1, 11)))
    doc.add_paragraph('Keep this heading', style=heading).runs[0].bold = True
    doc.add_paragraph('Keep this subheading', style=heading)
    doc.add_paragraph('First body line\nSecond body line\nThird body line')
    kept = doc.add_paragraph('\n'.join(f'Kept line {i}' for i in range(1, 9)))
    kept.paragraph_format.keep_together = True
    kept.paragraph_format.keep_with_next = False
    kept.paragraph_format.widow_control = False
    doc.add_paragraph('\n'.join(f'Widow line {i}' for i in range(1, 6)))
    doc.add_paragraph()  # Empty paragraph remains an editable line.
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/paragraph-pagination.docx'
    doc.save(target)
    print(f'Generated independent pagination fixture: {target}')

if args.generate_numbering_fixture:
    doc = fixture_document()
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.line_spacing = Pt(18)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    root = doc.part.numbering_part.element
    def element(tag, value=None):
        node = OxmlElement('w:' + tag)
        if value is not None:
            node.set(qn('w:val'), str(value))
        return node
    for abstract_id, bullet in [(41, False), (42, True)]:
        abstract = element('abstractNum')
        abstract.set(qn('w:abstractNumId'), str(abstract_id))
        abstract.append(element('multiLevelType', 'multilevel'))
        for i in range(9):
            lvl = element('lvl')
            lvl.set(qn('w:ilvl'), str(i))
            lvl.append(element('start', 3 if i == 0 and not bullet else 1))
            lvl.append(element('numFmt', 'bullet' if bullet else 'lowerLetter' if i == 1 else 'decimal'))
            if i == 0 and not bullet:
                lvl.append(element('pStyle', 'CanvasList'))
            lvl.append(element('suff', 'tab'))
            text = '\uf0b7' if bullet else '.'.join('%' + str(j + 1) for j in range(i + 1)) + (')' if i == 1 else '.')
            lvl.append(element('lvlText', text))
            lvl.append(element('lvlJc', 'left'))
            ppr = element('pPr')
            tabs = element('tabs')
            tab = element('tab', 'num'); tab.set(qn('w:pos'), str((i + 1) * 720)); tabs.append(tab); ppr.append(tabs)
            ind = element('ind'); ind.set(qn('w:left'), str((i + 1) * 720)); ind.set(qn('w:hanging'), '360'); ppr.append(ind)
            lvl.append(ppr)
            if bullet:
                rpr = element('rPr'); font = element('rFonts')
                font.set(qn('w:ascii'), 'Symbol'); font.set(qn('w:hAnsi'), 'Symbol'); rpr.append(font); lvl.append(rpr)
            abstract.append(lvl)
        first_num = root.find(qn('w:num'))
        first_num.addprevious(abstract)
    for num_id, abstract_id, start in [(42, 41, None), (43, 41, 7), (44, 42, None)]:
        num = element('num'); num.set(qn('w:numId'), str(num_id)); num.append(element('abstractNumId', abstract_id))
        if start is not None:
            override = element('lvlOverride'); override.set(qn('w:ilvl'), '0'); override.append(element('startOverride', start)); num.append(override)
        root.append(num)
    style = doc.styles.add_style('Canvas List', WD_STYLE_TYPE.PARAGRAPH); style.base_style = normal
    ref = element('numPr'); ref.append(element('numId', 42)); style.element.get_or_add_pPr().append(ref)
    derived = doc.styles.add_style('Canvas Derived List', WD_STYLE_TYPE.PARAGRAPH); derived.base_style = style
    style.font.italic = True; derived.font.italic = True  # Nearest declaration, not two toggles.
    def numbered(text, num_id, level=0):
        p = doc.add_paragraph(text)
        ref = element('numPr'); ref.append(element('ilvl', level)); ref.append(element('numId', num_id)); p._p.get_or_add_pPr().append(ref)
        return p
    doc.add_paragraph('Numbering starts at three', style=derived).runs[0].bold = True
    doc.add_paragraph('Second top item', style=derived)
    numbered('Nested item', 42, 1)
    numbered('Nested continuation', 42, 1)
    doc.add_paragraph('Top after nested', style=derived)
    numbered('Restart at seven', 43)
    numbered('Continue eight', 43)
    numbered('Bullet item', 44)
    numbered('A bullet that wraps onto more than one line while its continuation stays aligned with the paragraph text. ' * 2, 44)
    plain = doc.add_paragraph('Not a list', style=derived)
    ref = element('numPr'); ref.append(element('numId', 0)); plain._p.get_or_add_pPr().append(ref)
    numbered('', 44)
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/numbering.docx'
    doc.save(target)
    print(f'Generated independent numbering fixture: {target}')

if args.generate_styles_fixture:
    doc = fixture_document()
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = 1.0
    section = doc.sections[0]; section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    body = doc.styles.add_style('Canvas Body', WD_STYLE_TYPE.PARAGRAPH); body.base_style = normal
    body.font.size = Pt(11); body.paragraph_format.space_after = Pt(3)
    following = OxmlElement('w:next'); following.set(qn('w:val'), body.style_id); body.element.find(qn('w:basedOn')).addnext(following)
    heading = doc.styles.add_style('Canvas Heading', WD_STYLE_TYPE.PARAGRAPH); heading.base_style = body
    heading.font.size = Pt(20); heading.font.bold = True
    heading.paragraph_format.keep_with_next = True; heading.paragraph_format.keep_together = True
    heading.paragraph_format.space_before = Pt(12); heading.paragraph_format.space_after = Pt(6)
    heading.next_paragraph_style = body
    sub = doc.styles.add_style('Canvas Subheading', WD_STYLE_TYPE.PARAGRAPH); sub.base_style = heading
    sub.font.size = Pt(16); sub.next_paragraph_style = body
    emphasis = doc.styles.add_style('Canvas Emphasis', WD_STYLE_TYPE.CHARACTER); emphasis.font.italic = True
    listing = doc.styles.add_style('Canvas List', WD_STYLE_TYPE.PARAGRAPH); listing.base_style = body
    ref = OxmlElement('w:numPr'); num = OxmlElement('w:numId'); num.set(qn('w:val'), '1'); ref.append(num); listing.element.get_or_add_pPr().append(ref)
    doc.add_paragraph('Style heading', style=heading)
    p = doc.add_paragraph('Body ', style=body); p.add_run('bold').bold = True; p.add_run(' and '); p.add_run('emphasis').style = emphasis
    p = doc.add_paragraph('Direct size and spacing', style=sub); p.runs[0].font.size = Pt(24); p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph('Style list item', style=listing)
    doc.add_paragraph('', style=body)
    doc.add_paragraph('Text across pages. ' * 28, style=body)
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/styles.docx'; doc.save(target)
    print(f'Generated independent styles fixture: {target}')

if args.generate_outline_fixture:
    doc = fixture_document(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(11)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    def outline(ppr, level):
        item = OxmlElement('w:outlineLvl'); item.set(qn('w:val'), str(level)); ppr.append(item)
    defaults = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:pPrDefault')).find(qn('w:pPr')); outline(defaults, 9)
    base = doc.styles.add_style('Canvas Outline Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    base.font.bold = True; base.font.size = Pt(18); base.next_paragraph_style = normal
    base.paragraph_format.keep_with_next = True; outline(base.element.get_or_add_pPr(), 0)
    child = doc.styles.add_style('Canvas Outline Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    child.font.italic = True; outline(child.element.get_or_add_pPr(), 2)
    inherited = doc.styles.add_style('Canvas Outline Inherited', WD_STYLE_TYPE.PARAGRAPH); inherited.base_style = child
    mask = doc.styles.add_style('Canvas Outline Body', WD_STYLE_TYPE.PARAGRAPH); mask.base_style = base; outline(mask.element.get_or_add_pPr(), 9)
    deep = doc.styles.add_style('Canvas Outline Deep', WD_STYLE_TYPE.PARAGRAPH); deep.base_style = normal; outline(deep.element.get_or_add_pPr(), 8)
    doc.add_paragraph('Outline levels Á😀', style=base)
    doc.add_paragraph('Body introduction.')
    doc.add_paragraph('Nested Á😀', style=child)
    doc.add_paragraph('More nested inherited', style=inherited)
    doc.add_paragraph('Hidden from outline', style=mask)
    p = doc.add_paragraph('Direct body override', style=base); outline(p._p.get_or_add_pPr(), 9)
    p = doc.add_paragraph('Direct level two'); outline(p._p.get_or_add_pPr(), '+01')
    doc.add_paragraph('Deep heading', style=deep)
    doc.add_paragraph('', style=child)
    doc.add_paragraph('Body across pages. ' * 90)
    p = doc.add_paragraph('Later heading', style=base); p.paragraph_format.page_break_before = True
    table = doc.add_table(rows=3, cols=2, style='Table Grid'); table.autofit = False
    for column in table.columns: column.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).text = 'Header heading'; table.cell(0, 0).paragraphs[0].style = child
    table.cell(0, 1).text = 'Header detail'; table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Table body. ' * 65; table.cell(1, 1).text = 'Beside body'
    table.cell(2, 0).text = 'Cell heading'; outline(table.cell(2, 0).paragraphs[0]._p.get_or_add_pPr(), 4)
    table.cell(2, 1).text = 'Cell detail'
    p = doc.add_paragraph('Numbered heading', style=child)
    numpr = OxmlElement('w:numPr'); level = OxmlElement('w:ilvl'); level.set(qn('w:val'), '0'); numpr.append(level)
    num = OxmlElement('w:numId'); num.set(qn('w:val'), '7'); numpr.append(num); p._p.get_or_add_pPr().append(numpr)
    # The reference application fixes built-in heading IDs despite conflicting, valid declarations.
    doc.styles['Heading 1'].element.pPr.find(qn('w:outlineLvl')).set(qn('w:val'), '9')
    p = doc.add_paragraph('Built-in heading', style='Heading 1'); outline(p._p.get_or_add_pPr(), 9)
    doc.add_paragraph('Tail.')
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/outline-levels.docx'; doc.save(target)
    print(f'Generated independent outline fixture: {target}')

if args.generate_inline_images_fixture or args.generate_cropped_images_fixture or args.generate_flipped_images_fixture or args.generate_rotated_images_fixture:
    # Original test pixels, encoded with Python's standard library. No downloaded
    # artwork or image-codec dependency is used to generate these PNG assets.
    import struct, zlib
    from io import BytesIO
    def png(transparent=False):
        def chunk(kind, data): return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xffffffff)
        rows = []
        for y in range(32):
            row = bytearray([0])
            for x in range(64):
                color = (30, 90, 180, 255) if x < 32 else (245, 150, 35, 255)
                if (args.generate_flipped_images_fixture or args.generate_rotated_images_fixture) and y >= 16: color = (40, 160, 85, 255) if x < 32 else (155, 60, 175, 255)
                if transparent and (x - 32)**2 + (y - 16)**2 > 200: color = (0, 0, 0, 0)
                row.extend(color)
            rows.append(bytes(row))
        return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 64, 32, 8, 6, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(b''.join(rows))) + chunk(b'IEND', b'')
    opaque, transparent = png(), png(True)
    doc = fixture_document(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(11)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    style = doc.styles.add_style('Canvas Picture Text', WD_STYLE_TYPE.PARAGRAPH); style.base_style = normal
    style.font.italic = True; style.paragraph_format.line_spacing = 1.5
    doc.add_paragraph('Inline pictures Á😀')
    p = doc.add_paragraph(style=style); p.add_run('Before ')
    picture = p.add_run().add_picture(BytesIO(opaque), width=Inches(1.25)); picture._inline.docPr.set('descr', 'Blue and orange panels')
    picture._inline.set('distL', '914400'); picture._inline.set('distB', '457200')
    p.add_run(' after.').bold = True
    p = doc.add_paragraph('Transparent: '); p.add_run().add_picture(BytesIO(transparent), width=Inches(1.5)); p.add_run(' tail')
    p = doc.add_paragraph('Shared asset: '); p.add_run().add_picture(BytesIO(opaque), width=Inches(.75))
    table = doc.add_table(rows=3, cols=2, style='Table Grid'); table.autofit = False
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(BytesIO(transparent), width=Inches(.5)); table.cell(0, 1).text = 'Picture header'
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Body text. ' * 80; table.cell(1, 1).text = 'Beside picture header'
    table.cell(2, 0).text = 'Last row'; table.cell(2, 1).paragraphs[0].add_run().add_picture(BytesIO(opaque), width=Inches(1))
    doc.add_paragraph('AutoFit picture')
    auto = doc.add_table(rows=1, cols=2, style='Table Grid'); auto.autofit = True
    auto.cell(0, 0).paragraphs[0].add_run().add_picture(BytesIO(opaque), width=Inches(2)); auto.cell(0, 1).text = 'AutoFit label'
    for col in auto.columns: col.width = Inches(0)
    for row in auto.rows:
        for tc in row._tr.tc_lst:
            width = tc.tcPr.find(qn('w:tcW')); width.set(qn('w:type'), 'auto'); width.set(qn('w:w'), '0')
    doc.add_paragraph('After inline pictures.')
    name = 'inline-images'
    if args.generate_cropped_images_fixture or args.generate_flipped_images_fixture or args.generate_rotated_images_fixture:
        from lxml import etree
        name = 'cropped-images'; doc.paragraphs[0].text = 'Cropped pictures Á😀'
        doc.paragraphs[3].runs[0].text = 'Outside source: '
        offsets = [dict(l='25000', t='12.5%', b='12500'), dict(l='-25000', t='-25000', r='-25000', b='-25000'), dict(l='125000', r='-50000'), dict(l='50%'), dict(l='12.345%'), None]
        a14 = 'http://schemas.microsoft.com/office/drawing/2010/main'
        for index, (picture, crop) in enumerate(zip(doc.inline_shapes, offsets)):
            fill = picture._inline.xpath('.//pic:blipFill')[0]
            if crop:
                rect = OxmlElement('a:srcRect')
                for key, value in crop.items(): rect.set(key, value)
                fill.insert(1, rect)
            blip = fill.find(qn('a:blip')); blip.set('cstate', 'print' if index % 2 else 'screen')
            ext_list = OxmlElement('a:extLst'); ext = OxmlElement('a:ext'); ext.set('uri', '{28A0092B-C50C-407e-A947-70E740481C1C}')
            dpi = etree.Element(f'{{{a14}}}useLocalDpi', nsmap={'a14': a14})
            if index % 3 != 2: dpi.set('val', str(index % 3))
            ext.append(dpi); ext_list.append(ext); blip.append(ext_list)
        if args.generate_flipped_images_fixture or args.generate_rotated_images_fixture:
            name = 'flipped-images'; doc.paragraphs[0].text = 'Flipped pictures Á😀'; doc.paragraphs[3].runs[0].text = 'Both axes: '
            third_crop = doc.inline_shapes[2]._inline.xpath('.//a:srcRect')[0]; third_crop.attrib.clear(); third_crop.set('r', '25%')
            flips = [('true', '0'), ('0', '1'), ('1', 'true'), ('1', '1'), None, ('false', 'false')]
            for picture, axes in zip(doc.inline_shapes, flips):
                transform = picture._inline.xpath('.//a:xfrm')[0]
                if axes is None: transform.getparent().remove(transform)
                else:
                    transform.set('flipH', axes[0]); transform.set('flipV', axes[1]); transform.set('rot', '0')
    if args.generate_rotated_images_fixture:
        from math import radians, sin, cos, ceil
        name = 'rotated-images'; doc.paragraphs[0].text = 'Rotated pictures Á😀'
        # Derive extents independently from transformed corner coordinates.
        for index, (picture, degrees) in enumerate(zip(doc.inline_shapes, [30, -90, 450, 45, 0, 315])):
            transforms = picture._inline.xpath('.//a:xfrm')
            if transforms: transforms[0].set('rot', str(degrees * 60000))
            angle = radians(degrees); w, h = picture.width, picture.height
            points = [(x * cos(angle) - y * sin(angle), x * sin(angle) + y * cos(angle)) for x in [-w / 2, w / 2] for y in [-h / 2, h / 2]]
            horizontal = max(0, ceil(max(x for x, y in points) - w / 2 - 1e-7))
            vertical = max(0, ceil(max(y for x, y in points) - h / 2 - 1e-7))
            extents = dict(l=horizontal, t=vertical, r=horizontal, b=vertical)
            if index == 0: extents['l'] += Inches(.05); extents['b'] += Inches(.03)
            effects = OxmlElement('wp:effectExtent')
            for side, value in extents.items(): effects.set(side, str(value))
            picture._inline.insert(1, effects)
    target = Path(__file__).resolve().parents[1] / f'public/fixtures/{name}.docx'; doc.save(target)
    print(f'Generated independent {name} fixture: {target}')

if args.generate_table_headers_fixture:
    doc = fixture_document(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(11)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    style = doc.styles.add_style('Canvas Header Table', WD_STYLE_TYPE.TABLE)
    style.base_style = doc.styles['Table Grid']; style.font.italic = True
    trpr = OxmlElement('w:trPr'); trpr.append(OxmlElement('w:tblHeader')); style.element.append(trpr)
    doc.add_paragraph('Repeated table headers.')
    doc.add_paragraph('\n'.join(f'Lead {i}' for i in range(1, 6)))
    table = doc.add_table(rows=7, cols=3, style=style); table.autofit = False
    for col, width in zip(table.columns, [2, 2, 1]): col.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, [2, 2, 1]): cell.width = Inches(width)
    table.cell(0, 0).merge(table.cell(1, 0)).text = 'Product Á😀'
    table.cell(0, 1).merge(table.cell(0, 2)).text = 'Details'
    table.cell(1, 1).text = 'Description'; table.cell(1, 2).text = 'Code'
    for row in table.rows[:2]:
        for tc in row._tr.tc_lst:
            shade = OxmlElement('w:shd'); shade.set(qn('w:val'), 'clear'); shade.set(qn('w:fill'), 'EAF2FC'); tc.tcPr.append(shade)
            for p in _Cell(tc, table).paragraphs:
                for run in p.runs: run.bold = True
    for r in range(2, 7):
        for cell, value in zip(table.rows[r].cells, [f'Item {r - 1}', 'Body text. ' * (40 if r == 3 else 7), str(r - 1)]): cell.text = value
        header = OxmlElement('w:tblHeader'); header.set(qn('w:val'), '1' if r == 6 else '0'); table.rows[r]._tr.get_or_add_trPr().append(header)
    table.cell(3, 0).merge(table.cell(4, 0)).text = 'Merged body. ' * 45
    table.rows[3]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    doc.add_paragraph('AutoFit header table.')
    auto = doc.add_table(rows=4, cols=2, style='Table Grid'); auto.autofit = True
    auto.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for cell, value in zip(auto.rows[0].cells, ['Auto label', 'Auto value']): cell.text = value
    for r in range(1, 4):
        auto.cell(r, 0).text = f'Auto {r}'; auto.cell(r, 1).text = 'Automatic content. ' * 30
    for col in auto.columns: col.width = Inches(0)
    for row in auto.rows:
        for tc in row._tr.tc_lst:
            width = tc.tcPr.find(qn('w:tcW')); width.set(qn('w:type'), 'auto'); width.set(qn('w:w'), '0')
    doc.add_paragraph('After repeated headers.')
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/table-headers.docx'
    doc.save(target); print(f'Generated independent repeated-header fixture: {target}')

if args.generate_vertical_pages_fixture:
    doc = fixture_document(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(11)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    style = doc.styles.add_style('Canvas Flow Table', WD_STYLE_TYPE.TABLE)
    style.base_style = doc.styles['Table Grid']; style.font.italic = True
    row_style = OxmlElement('w:trPr'); row_style.append(OxmlElement('w:cantSplit')); style.element.append(row_style)
    doc.add_paragraph('Merged cells can continue across pages.')
    doc.add_paragraph('\n'.join(f'Lead {i}' for i in range(1, 7)))
    table = doc.add_table(rows=4, cols=3, style=style); table.autofit = False
    for col, width in zip(table.columns, [2, 2, 1]): col.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, [2, 2, 1]): cell.width = Inches(width)
    left = table.cell(0, 0).merge(table.cell(2, 0)); left.text = 'Left merged text. ' * 70
    left.add_paragraph('Left end Á😀').paragraph_format.keep_with_next = False
    right = table.cell(1, 1).merge(table.cell(3, 1)); right.text = 'Right merged text. ' * 50
    right.paragraphs[0].runs[0].font.size = Pt(17)
    right.add_paragraph('Right end'); tail = right.add_paragraph('Merged list tail', style='List Number')
    contextual = OxmlElement('w:contextualSpacing'); contextual.set(qn('w:val'), '0'); tail._p.get_or_add_pPr().append(contextual)
    for cell, value in [(table.cell(0, 1), 'Upper middle'), (table.cell(3, 0), 'Lower left')]: cell.text = value
    for r, value in enumerate(['Side zero', 'Side text. ' * 25, 'Side two', 'Side three']): table.cell(r, 2).text = value
    for r, row in enumerate(table.rows):
        if r in [0, 2]:
            keep = OxmlElement('w:cantSplit'); keep.set(qn('w:val'), '0'); row._tr.get_or_add_trPr().append(keep)
    height = OxmlElement('w:trHeight'); height.set(qn('w:val'), '600'); height.set(qn('w:hRule'), 'atLeast'); table.rows[2]._tr.get_or_add_trPr().append(height)
    for cell, color in [(left, 'EAF2FC'), (right, 'FFF0CF')]:
        shade = OxmlElement('w:shd'); shade.set(qn('w:val'), 'clear'); shade.set(qn('w:fill'), color); cell._tc.tcPr.append(shade)
    doc.add_paragraph('AutoFit across pages.')
    auto = doc.add_table(rows=2, cols=2, style='Table Grid'); auto.autofit = True
    merged = auto.cell(0, 0).merge(auto.cell(1, 1)); merged.text = 'Auto merged content. ' * 100
    merged.add_paragraph('Auto ending')
    for col in auto.columns: col.width = Inches(0)
    for row in auto.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for tc in row._tr.tc_lst:
            preferred = tc.tcPr.find(qn('w:tcW')); preferred.set(qn('w:type'), 'auto'); preferred.set(qn('w:w'), '0')
    doc.add_paragraph('Minimum height across pages.')
    tall = doc.add_table(rows=1, cols=2, style='Table Grid'); tall.autofit = False
    tall.cell(0, 0).text = 'Height only'; tall.cell(0, 1).text = 'First\nLast'
    height = OxmlElement('w:trHeight'); height.set(qn('w:val'), '5760'); height.set(qn('w:hRule'), 'atLeast'); tall.rows[0]._tr.get_or_add_trPr().append(height)
    doc.add_paragraph('After all flowing tables.')
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/vertical-pages.docx'
    doc.save(target); print(f'Generated independent flowing vertical-cell fixture: {target}')

if args.generate_vertical_cells_fixture:
    doc = fixture_document(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(11)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    style = doc.styles.add_style('Canvas Vertical Table', WD_STYLE_TYPE.TABLE)
    style.base_style = doc.styles['Table Grid']; style.font.italic = True
    doc.add_paragraph('Vertical cells across rows.')
    doc.add_paragraph('\n'.join(f'Lead line {i}' for i in range(1, 8)))
    table = doc.add_table(rows=6, cols=3, style=style); table.autofit = False
    for col in table.columns: col.width = Inches(5 / 3)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(5 / 3)
    for r, values in enumerate([['Top', 'B0', 'C0'], ['', 'B1', 'C1'], ['D0', 'Wide', ''], ['D1', '', ''], ['F0', 'F1', 'F2'], ['G0', 'G1', 'G2']]):
        for cell, value in zip(table.rows[r].cells, values): cell.text = value
    anchor = table.cell(0, 0).merge(table.cell(1, 0)); anchor.add_paragraph('Second Á😀'); anchor.paragraphs[0].runs[0].bold = True
    shade = OxmlElement('w:shd'); shade.set(qn('w:val'), 'clear'); shade.set(qn('w:fill'), 'EAF2FC'); anchor._tc.tcPr.append(shade)
    table.cell(2, 1).merge(table.cell(3, 2))
    for row in table.rows[:4]:
        height = OxmlElement('w:trHeight'); height.set(qn('w:val'), '864'); height.set(qn('w:hRule'), 'atLeast'); row._tr.get_or_add_trPr().append(height)
    continuation = table.rows[1]._tr.tc_lst[0]
    ppr = continuation.p_lst[0].get_or_add_pPr(); spacing = OxmlElement('w:spacing'); spacing.set(qn('w:after'), '123'); ppr.append(spacing)
    table.rows[3]._tr.tc_lst[1].tcPr.find(qn('w:vMerge')).set(qn('w:val'), 'continue')
    doc.add_paragraph('AutoFit vertical cells.')
    auto = doc.add_table(rows=3, cols=3, style='Table Grid'); auto.autofit = True
    auto.cell(0, 0).merge(auto.cell(1, 1)).text = 'Auto Á😀'
    auto.cell(0, 2).text = 'Auto side'; auto.cell(1, 2).text = 'Auto lower'
    for cell, value in zip(auto.rows[2].cells, ['Auto end', 'M', 'Z']): cell.text = value
    auto.cell(2, 2)._tc.tcPr.append(OxmlElement('w:vMerge'))  # Orphan continue is ignored by §17.18.57.
    for col in auto.columns: col.width = Inches(0)
    for row in auto.rows:
        for tc in row._tr.tc_lst:
            preferred = tc.tcPr.find(qn('w:tcW')); preferred.set(qn('w:type'), 'auto'); preferred.set(qn('w:w'), '0')
    doc.add_paragraph('After vertical tables.')
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/vertical-cells.docx'
    doc.save(target); print(f'Generated independent vertical-cell fixture: {target}')

if args.generate_merged_cells_fixture:
    doc = fixture_document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(11)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    style = doc.styles.add_style('Canvas Span Table', WD_STYLE_TYPE.TABLE)
    style.base_style = doc.styles['Table Grid']; style.font.italic = True
    doc.add_paragraph('Cells can span different grid columns.')
    table = doc.add_table(rows=4, cols=3, style=style); table.autofit = False
    for col, value in zip(table.columns, [1, 2, 2]): col.width = Inches(value)
    for row in table.rows:
        for cell, value in zip(row.cells, [1, 2, 2]): cell.width = Inches(value)
    header = table.cell(0, 0).merge(table.cell(0, 1)); header.text = 'Wide heading'
    header.paragraphs[0].runs[0].bold = True
    header._tc.tcPr.remove(header._tc.tcPr.find(qn('w:tcW')))
    flags = OxmlElement('w:cnfStyle'); flags.set(qn('w:val'), '000000000000'); header._tc.tcPr.insert(0, flags)
    shade = OxmlElement('w:shd'); shade.set(qn('w:val'), 'clear'); shade.set(qn('w:fill'), 'EAF2FC'); header._tc.get_or_add_tcPr().append(shade)
    table.cell(0, 2).text = 'Side'
    table.cell(1, 0).text = 'L'
    table.cell(1, 1).merge(table.cell(1, 2)).text = 'Long merged cell. ' * 45
    table.rows[1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for cell, value in zip(table.rows[2].cells, ['Left', 'Middle', 'Right']): cell.text = value
    table.cell(3, 0).merge(table.cell(3, 2)).text = 'End Á😀'
    doc.add_paragraph('AutoFit merged cells.')
    auto = doc.add_table(rows=3, cols=3, style='Table Grid'); auto.autofit = True
    auto.cell(0, 0).merge(auto.cell(0, 1)).text = 'Extraordinarylongword'; auto.cell(0, 2).text = 'X'
    auto.cell(1, 0).text = 'A'; auto.cell(1, 1).merge(auto.cell(1, 2)).text = 'Otherlongword'
    for cell, value in zip(auto.rows[2].cells, ['a', 'b', 'c']): cell.text = value
    for col in auto.columns: col.width = Inches(0)
    for row in auto.rows:
        for cell in row.cells:
            preferred = cell._tc.tcPr.find(qn('w:tcW')); preferred.set(qn('w:type'), 'auto'); preferred.set(qn('w:w'), '0')
    doc.add_paragraph('After merged tables.')
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/merged-cells.docx'
    doc.save(target); print(f'Generated independent merged-cell fixture: {target}')

if args.generate_autofit_fixture:
    doc = fixture_document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(11)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    doc.add_paragraph('AutoFit columns respond to content.')
    table = doc.add_table(rows=3, cols=3, style='Table Grid')
    # Omit tblLayout: AutoFit is the OOXML default. Explicit percent lexical form.
    width = table._tbl.tblPr.find(qn('w:tblW')); width.set(qn('w:type'), 'pct'); width.set(qn('w:w'), '100%')
    for col, value in zip(table.columns, [.75, 2.25, 2]): col.width = Inches(value)
    for row in table.rows:
        for cell, value in zip(row.cells, [.75, 2.25, 2]): cell.width = Inches(value)
    values = [['Extraordinary', 'Many short words for measuring and wrapping. ' * 8, 'Notes'], ['A\u00a0B', 'Value\t12.5', 'Á😀'], ['', '', 'Last']]
    for row, values in zip(table.rows, values):
        for cell, value in zip(row.cells, values): cell.text = value
    table.cell(0, 0).paragraphs[0].runs[0].bold = True
    table.cell(1, 1).paragraphs[0].paragraph_format.tab_stops.add_tab_stop(Inches(.8), WD_TAB_ALIGNMENT.DECIMAL)
    doc.add_paragraph('A compact table.')
    compact = doc.add_table(rows=2, cols=2, style='Table Grid')
    compact.autofit = True
    for col in compact.columns: col.width = Inches(0)
    for row, values in zip(compact.rows, [['One', 'Two'], ['Three', '']]):
        for cell, value in zip(row.cells, values):
            cell.text = value
            preferred = cell._tc.tcPr.find(qn('w:tcW')); preferred.set(qn('w:type'), 'auto'); preferred.set(qn('w:w'), '0')
    doc.add_paragraph('After both tables.')
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/autofit-tables.docx'
    doc.save(target); print(f'Generated independent AutoFit fixture: {target}')

if args.generate_tables_fixture:
    doc = fixture_document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Noto Sans', Pt(11)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(3)
    style = doc.styles.add_style('Canvas Table', WD_STYLE_TYPE.TABLE)
    style.base_style = doc.styles['Table Grid']
    style.font.italic = True
    doc.add_paragraph('A table with editable cells.')
    table = doc.add_table(rows=4, cols=2, style=style)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Inches(3), Inches(2)
    for row in table.rows:
        row.cells[0].width, row.cells[1].width = Inches(3), Inches(2)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for cell, text in zip(table.rows[0].cells, ['Product', 'Details']):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
        shade = OxmlElement('w:shd'); shade.set(qn('w:val'), 'clear'); shade.set(qn('w:fill'), 'EAF2FC'); cell._tc.get_or_add_tcPr().append(shade)
    table.cell(1, 0).text = 'Alpha café Á😀'
    table.cell(1, 0).add_paragraph('Second paragraph')
    table.cell(1, 1).text = 'Before\tAfter'
    table.cell(2, 0).text = 'Text inside a tall row. ' * 45
    table.cell(2, 1).text = 'Short cell'
    table.cell(3, 0).text = 'Last row'
    table.cell(3, 1).text = ''
    height = OxmlElement('w:trHeight'); height.set(qn('w:hRule'), 'atLeast'); height.set(qn('w:val'), '480'); table.rows[3]._tr.get_or_add_trPr().append(height)
    doc.add_paragraph('After the table.')
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/tables.docx'
    doc.save(target)
    print(f'Generated independent table fixture: {target}')

if args.generate_character_styles_fixture:
    doc = fixture_document(); normal = doc.styles['Normal']
    normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = 1.0
    section = doc.sections[0]; section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    accent = doc.styles.add_style('Canvas Paragraph', WD_STYLE_TYPE.PARAGRAPH); accent.base_style = normal
    accent.font.bold = True; accent.font.italic = True
    base = doc.styles.add_style('Canvas Character Base', WD_STYLE_TYPE.CHARACTER); base.font.italic = True; base.font.size = Pt(14)
    child = doc.styles.add_style('Canvas Character Child', WD_STYLE_TYPE.CHARACTER); child.base_style = base
    child.font.bold = True; child.font.italic = False
    grand = doc.styles.add_style('Canvas Character Grandchild', WD_STYLE_TYPE.CHARACTER); grand.base_style = child; grand.font.underline = True
    decorative = doc.styles.add_style('Canvas Decorative', WD_STYLE_TYPE.CHARACTER)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'double'); decorative.element.get_or_add_rPr().append(underline)
    p = doc.add_paragraph('Plain '); p.add_run('base').style = base; p.add_run(' child').style = child
    direct = p.add_run(' direct'); direct.style = child; direct.bold = False; direct.font.size = Pt(22); direct.italic = True
    p = doc.add_paragraph('Paragraph ', style=accent); p.add_run('inherited child').style = child
    p = doc.add_paragraph(); p.add_run('Clear this style').style = grand
    p = doc.add_paragraph(); ref = OxmlElement('w:rStyle'); ref.set(qn('w:val'), grand.style_id); rpr = OxmlElement('w:rPr'); rpr.append(ref); p._p.get_or_add_pPr().append(rpr)
    p = doc.add_paragraph(); p.add_run('Styled text across pages. ' * 25).style = base
    listing = doc.styles.add_style('Canvas Character List', WD_STYLE_TYPE.PARAGRAPH); listing.base_style = normal
    ref = OxmlElement('w:numPr'); num = OxmlElement('w:numId'); num.set(qn('w:val'), '1'); ref.append(num); listing.element.get_or_add_pPr().append(ref)
    p = doc.add_paragraph(style=listing); p.add_run('Styled bullet item').style = grand
    p = doc.add_paragraph(); p.add_run('A\u0301😀 symbol').style = base
    target = Path(__file__).resolve().parents[1] / 'public/fixtures/character-styles.docx'; doc.save(target)
    print(f'Generated independent character style fixture: {target}')

for file in args.paths:
    doc = Document(file)
    if args.justification:
        from canvas_justification_fixture import verify
        verify(file, args.justification)
        continue
    if args.tab_leaders:
        from canvas_tab_leader_fixture import verify
        verify(file, args.tab_leaders)
        continue
    if args.row_heights:
        from canvas_row_height_fixture import verify
        verify(file, args.row_heights)
        continue
    if args.line_spacing:
        from canvas_line_spacing_fixture import verify
        verify(file, args.line_spacing)
        continue
    if args.style_spacing:
        from canvas_style_spacing_fixture import verify
        verify(file, args.style_spacing)
        continue
    if args.search:
        from canvas_search_fixture import verify
        verify(file, args.search)
        continue
    if args.clear_format:
        from canvas_clear_format_fixture import verify
        verify(file, args.clear_format)
        continue
    if args.paragraph_reset:
        from canvas_paragraph_reset_fixture import verify
        verify(file, args.paragraph_reset)
        continue
    if args.cell_margins or args.new_cell_margins:
        from canvas_cell_margin_fixture import verify
        verify(file, args.new_cell_margins)
        continue
    if args.cell_borders or args.new_cell_borders:
        from canvas_cell_border_fixture import verify
        verify(file, args.new_cell_borders)
        continue
    if args.cell_shading or args.new_cell_shading:
        from canvas_cell_shading_fixture import verify
        verify(file, args.new_cell_shading)
        continue
    if args.cell_alignment or args.new_cell_alignment:
        from canvas_cell_alignment_fixture import verify
        verify(file, args.new_cell_alignment)
        continue
    if args.paragraph_borders or args.new_paragraph_borders:
        from canvas_paragraph_border_fixture import verify
        verify(file, args.new_paragraph_borders)
        continue
    if args.paragraph_shading or args.new_paragraph_shading:
        from canvas_paragraph_shading_fixture import verify
        verify(file, args.new_paragraph_shading)
        continue
    if args.bookmarks or args.new_bookmarks:
        from canvas_bookmark_fixture import verify
        verify(file, args.new_bookmarks)
        continue
    if args.hyperlinks or args.new_hyperlinks:
        from canvas_hyperlink_fixture import verify
        verify(file, args.new_hyperlinks)
        continue
    if args.underlines or args.new_underlines:
        from canvas_underline_fixture import verify
        verify(file, args.new_underlines)
        continue
    if args.capitalization or args.new_capitalization:
        from canvas_caps_fixture import verify
        verify(file, args.new_capitalization)
        continue
    if args.strikethrough or args.new_strikethrough:
        from canvas_strike_fixture import verify
        verify(file, args.new_strikethrough)
        continue
    if args.spaces:
        from canvas_spaces_fixture import verify
        verify(file, args.spaces)
        continue
    if args.kerning or args.new_kerning:
        from canvas_kerning_fixture import verify
        verify(file, args.new_kerning)
        continue
    if args.contextual_spacing:
        from canvas_contextual_fixture import verify
        verify(file, args.contextual_spacing)
        continue
    if args.text_width or args.new_text_width:
        from canvas_width_fixture import verify
        verify(file, args.new_text_width)
        continue
    if args.character_spacing or args.new_character_spacing:
        from canvas_spacing_fixture import verify
        verify(file, args.new_character_spacing)
        continue
    if args.baseline_shifts or args.new_baseline_shifts:
        from canvas_position_fixture import verify
        verify(file, args.new_baseline_shifts)
        continue
    if args.script_text or args.new_script_text:
        from canvas_vertical_fixture import verify
        verify(file, args.new_script_text)
        continue
    if args.style_toggles_off or args.style_toggles_on:
        from canvas_toggle_fixture import verify
        verify(file, args.style_toggles_on)
        continue
    if args.text_highlights or args.new_text_highlights:
        from canvas_highlights_fixture import verify
        verify(file, args.new_text_highlights)
        continue
    if args.text_colors or args.new_text_colors:
        from canvas_colors_fixture import verify
        verify(file, args.new_text_colors)
        continue
    if args.outline_levels or args.new_outline_levels:
        from zipfile import ZipFile
        def declared_level(ppr):
            nodes = [] if ppr is None else ppr.findall(qn('w:outlineLvl'))
            assert len(nodes) <= 1, 'Duplicate outline declarations'
            return None if not nodes else int(nodes[0].get(qn('w:val')))
        def inherited_level(style):
            seen = set()
            while style is not None:
                assert style.style_id not in seen, 'Cyclic outline inheritance'
                seen.add(style.style_id)
                value = declared_level(style.element.find(qn('w:pPr')))
                if value is not None: return value
                style = style.base_style
            defaults = doc.styles.element.find(qn('w:docDefaults'))
            nodes = [] if defaults is None else defaults.findall('/'.join(qn(tag) for tag in ['w:pPrDefault', 'w:pPr', 'w:outlineLvl']))
            return int(nodes[0].get(qn('w:val'))) if nodes else 9
        def level(paragraph):
            if paragraph.style.style_id in [f'Heading{i}' for i in range(1, 10)]: return int(paragraph.style.style_id[-1]) - 1
            value = declared_level(paragraph._p.pPr)
            return inherited_level(paragraph.style) if value is None else value
        if args.outline_levels:
            source_path = Path(__file__).resolve().parents[1] / 'public/fixtures/outline-levels.docx'; original_doc = Document(source_path)
            expected_text = [p.text for p in original_doc.paragraphs]; expected_text[10] = 'Revised Later heading'
            assert [p.text for p in doc.paragraphs] == expected_text, 'Outline edits/navigation changed source text'
            assert [declared_level(p._p.pPr) for p in doc.paragraphs] == [None, None, 9, None, None, None, 3, None, None, None, None, 8, 9, None], 'Direct levels or cleared body override changed'
            assert [level(p) for p in doc.paragraphs] == [0, 9, 9, 5, 9, 0, 3, 8, 5, 9, 0, 8, 0, 9], 'Outline inheritance or masking changed'
            assert declared_level(doc.styles['Canvas Outline Child'].element.pPr) == 5, 'Edited style level lost'
            for p, old in zip(doc.paragraphs, original_doc.paragraphs):
                assert p.style.style_id == old.style.style_id, 'Outline editing changed paragraph styles'
                assert [(r.bold, r.italic, r.underline, r.font.size) for r in p.runs] == [(r.bold, r.italic, r.underline, r.font.size) for r in old.runs], 'Outline editing changed run appearance'
            header, cell = doc.tables[0].cell(0, 0).paragraphs[0], doc.tables[0].cell(2, 0).paragraphs[0]
            assert level(header) == 5 and level(cell) == 4, 'Table outline inheritance changed'
            assert len(doc.tables[0].rows) == 3 and doc.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None, 'Repeated header was expanded or removed'
            assert doc.paragraphs[11]._p.pPr.numPr.numId.val == 7 and doc.paragraphs[11]._p.pPr.numPr.ilvl.val == 0, 'Outline editing changed numbering'
            with ZipFile(source_path) as original, ZipFile(file) as exported:
                assert original.namelist() == exported.namelist(), 'Outline editing changed package parts'
                for name in original.namelist():
                    if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), f'Unrelated part changed: {name}'
            def properties(style):
                # Compare all style nodes/attributes, excluding only the edited outline value.
                return [(n.tag, sorted(n.attrib.items()), (n.text or '').strip()) for n in style.element.iter() if n.tag != qn('w:outlineLvl')]
            exported_styles = {style.style_id: style for style in doc.styles}
            for old in original_doc.styles:
                assert properties(exported_styles[old.style_id]) == properties(old), 'Unrelated style data changed'
                if old.style_id != 'CanvasOutlineChild': assert declared_level(exported_styles[old.style_id].element.find(qn('w:pPr'))) == declared_level(old.element.find(qn('w:pPr'))), 'Unrelated style outline changed'
        else:
            assert [p.text for p in doc.paragraphs] == ['First heading', 'Body text', 'Second heading', ''], 'Authored outline source text changed'
            assert [p.style.style_id for p in doc.paragraphs] == ['Heading1', 'Normal', 'Heading2', 'Normal'], 'Following style or body override changed'
            assert [level(p) for p in doc.paragraphs] == [0, 3, 1, 8], 'Authored outline levels lost'
            assert [declared_level(p._p.pPr) for p in doc.paragraphs] == [None, 3, None, 8], 'Body masking undo or direct levels changed'
            styles_by_id = {style.style_id: style for style in doc.styles}
            assert declared_level(styles_by_id['Heading1'].element.pPr) == 0 and declared_level(styles_by_id['Heading2'].element.pPr) == 1, 'Built-in heading semantics missing'
        print(f'PASS: {file}: outline levels, style/default inheritance, direct body overrides, clear-to-style, source text, numbering, table headers and unrelated parts')
        continue
    if args.rotated_images or args.new_rotated_images:
        from zipfile import ZipFile
        from math import radians, sin, cos, ceil
        shapes = list(doc.inline_shapes)
        source_path = Path(__file__).resolve().parents[1] / 'public/fixtures/rotated-images.docx'
        if args.rotated_images:
            original_doc = Document(source_path)
            assert [p.text for p in doc.paragraphs] == ['Rotated pictures Á😀', 'Before  rotated after.', 'Transparent:  tail', 'Both axes: ', 'AutoFit picture', 'After inline pictures.'], 'Rotated picture source positions changed'
            assert len(shapes) == 6, 'Image count changed or repeated copies were serialized'
            assert [(s.width, s.height) for s in shapes] == [(Inches(w), Inches(h)) for w, h in [(1.5, .75), (1.5, .75), (.75, .375), (.5, .25), (1, .5), (2, 1)]], 'Rotation overwrote base dimensions'
            assert [s._inline.xpath('.//a:xfrm')[0].get('rot') for s in shapes] == ['16200000', '-5400000', '27000000', '5400000', '1350000', '18900000'], 'Edited angle or untouched signed/multi-turn spelling lost'
            for shape, original in zip(shapes, original_doc.inline_shapes):
                for tag in ['a:srcRect', 'a:blip']:
                    assert [dict(n.attrib) for n in shape._inline.xpath(f'.//{tag}')] == [dict(n.attrib) for n in original._inline.xpath(f'.//{tag}')], 'Crop or media reference changed during rotation'
                assert [(n.tag, dict(n.attrib)) for n in shape._inline.xpath('.//a:blip/a:extLst')[0].iter()] == [(n.tag, dict(n.attrib)) for n in original._inline.xpath('.//a:blip/a:extLst')[0].iter()], 'Local-DPI metadata changed'
                old_transform = original._inline.xpath('.//a:xfrm')
                new_transform = shape._inline.xpath('.//a:xfrm')[0]
                for key in ['flipH', 'flipV']: assert new_transform.get(key) == (old_transform[0].get(key) if old_transform else None), 'Rotation changed a flip'
            assert shapes[0]._inline.docPr.get('descr') == 'Rotated cropped quadrants', 'Rotated description changed'
            assert doc.paragraphs[1].style.style_id == 'CanvasPictureText' and doc.paragraphs[1].style.font.italic and doc.paragraphs[1].style.paragraph_format.line_spacing == 1.5, 'Inherited formatting changed'
            assert doc.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None and doc.tables[1].autofit, 'Header/AutoFit structure changed'
            with ZipFile(source_path) as original, ZipFile(file) as exported:
                assert original.namelist() == exported.namelist(), 'Rotation changed package parts'
                for name in original.namelist():
                    if name != 'word/document.xml': assert original.read(name) == exported.read(name), f'Unrelated part changed: {name}'
        else:
            assert [p.text for p in doc.paragraphs] == [' after'] and len(shapes) == 1, 'Authored rotation position or deletion undo changed'
            shape = shapes[0]; transform = shape._inline.xpath('.//a:xfrm')[0]
            assert (shape.width, shape.height) == (Inches(1.5), Inches(.75)), 'Authored rotation changed base dimensions'
            assert (transform.get('rot'), transform.get('flipH'), transform.get('flipV')) == ('16200000', '1', '0'), 'Authored transform lost'
            assert dict(shape._inline.xpath('.//a:srcRect')[0].attrib) == dict(l='-50000', t='-100000', r='0', b='0'), 'Source crop changed during rotation'
            assert shape._inline.docPr.get('descr') == 'Rotated outsets', 'Authored description changed'
            part = doc.part.related_parts[shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))]
            with ZipFile(source_path) as original: assert part.blob == original.read('word/media/image1.png'), 'Rotation changed bitmap bytes'
        for index, shape in enumerate(shapes):
            transform = shape._inline.xpath('.//a:xfrm')[0]
            assert [n.tag for n in transform] == [qn('a:off'), qn('a:ext')], 'Transform order/children changed'
            assert dict(transform.find(qn('a:off')).attrib) == dict(x='0', y='0'), 'Origin changed'
            extent = transform.find(qn('a:ext'))
            assert (int(extent.get('cx')), int(extent.get('cy'))) == (shape.width, shape.height), 'Inner and outer dimensions differ'
            angle = radians(int(transform.get('rot')) / 60000); w, h = shape.width, shape.height
            points = [(x * cos(angle) - y * sin(angle), x * sin(angle) + y * cos(angle)) for x in [-w / 2, w / 2] for y in [-h / 2, h / 2]]
            horizontal = max(0, ceil(max(x for x, y in points) - w / 2 - 1e-7))
            vertical = max(0, ceil(max(y for x, y in points) - h / 2 - 1e-7))
            expected = dict(l=horizontal, t=vertical, r=horizontal, b=vertical)
            if args.rotated_images and index == 0: expected['l'] += Inches(.05); expected['b'] += Inches(.03)
            assert dict(shape._inline.find(qn('wp:effectExtent')).attrib) == {k: str(v) for k, v in expected.items()}, 'Rotation padding or original additional space lost'
        print(f'PASS: {file}: signed rotations, transform order, base dimensions, corner padding, crops/flips, local-DPI metadata, source bytes and package preservation')
        continue
    if args.flipped_images or args.new_flipped_images:
        from zipfile import ZipFile
        shapes = list(doc.inline_shapes)
        source_path = Path(__file__).resolve().parents[1] / 'public/fixtures/flipped-images.docx'
        original_doc = Document(source_path)
        if args.flipped_images:
            assert [p.text for p in doc.paragraphs] == ['Flipped pictures Á😀', 'Before  flipped after.', 'Transparent:  tail', 'Both axes: ', 'AutoFit picture', 'After inline pictures.'], 'Flipped picture positions changed'
            assert len(shapes) == 6, 'Image count changed or repeated copies were serialized'
            assert [(s.width, s.height) for s in shapes] == [(Inches(w), Inches(h)) for w, h in [(1.5, .75), (1.5, .75), (.75, .375), (.5, .25), (1, .5), (2, 1)]], 'Flipped frame dimensions changed'
            expected = [('0', '1'), ('0', '1'), ('1', 'true'), ('0', '0'), ('1', None), ('false', 'false')]
            for index, (shape, original) in enumerate(zip(shapes, original_doc.inline_shapes)):
                transform = shape._inline.xpath('.//a:xfrm')[0]
                assert (transform.get('flipH'), transform.get('flipV')) == expected[index], 'Flip edit or untouched boolean spelling lost'
                assert transform.get('rot') == (None if index == 4 else '0'), 'Rotation changed while flipping'
                for tag in ['a:srcRect', 'a:blip']:
                    old_nodes, new_nodes = original._inline.xpath(f'.//{tag}'), shape._inline.xpath(f'.//{tag}')
                    assert len(old_nodes) == len(new_nodes) and [dict(n.attrib) for n in old_nodes] == [dict(n.attrib) for n in new_nodes], 'Cropping or compression attributes changed during flip'
                old_ext, new_ext = original._inline.xpath('.//a:blip/a:extLst')[0], shape._inline.xpath('.//a:blip/a:extLst')[0]
                assert [(n.tag, dict(n.attrib)) for n in old_ext.iter()] == [(n.tag, dict(n.attrib)) for n in new_ext.iter()], 'Local-DPI extension changed'
            assert shapes[0]._inline.docPr.get('descr') == 'Vertically flipped quadrants', 'Flipped description changed'
            assert doc.paragraphs[1].style.style_id == 'CanvasPictureText' and doc.paragraphs[1].style.font.italic and doc.paragraphs[1].style.paragraph_format.line_spacing == 1.5, 'Inherited picture formatting changed'
            assert doc.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None and doc.tables[1].autofit, 'Table header or AutoFit lost'
            with ZipFile(source_path) as original, ZipFile(file) as exported:
                assert original.namelist() == exported.namelist(), 'Flipping changed package parts'
                for name in original.namelist():
                    if name != 'word/document.xml': assert original.read(name) == exported.read(name), f'Unrelated part changed: {name}'
        else:
            assert [p.text for p in doc.paragraphs] == [' after'] and len(shapes) == 1, 'Authored flip position or deletion undo changed'
            shape = shapes[0]; assert (shape.width, shape.height) == (Inches(1.5), Inches(.75)), 'Authored flipped frame changed'
            transform = shape._inline.xpath('.//a:xfrm')[0]
            assert (transform.get('flipH'), transform.get('flipV')) == ('1', '1'), 'Authored flip flags lost'
            assert dict(shape._inline.xpath('.//a:srcRect')[0].attrib) == dict(l='-50000', t='-100000', r='0', b='0'), 'Source crop edges were exchanged instead of mirroring the frame'
            assert shape._inline.docPr.get('descr') == 'Both axes with outsets', 'Authored flip description changed'
            part = doc.part.related_parts[shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))]
            with ZipFile(source_path) as original: assert part.blob == original.read('word/media/image1.png'), 'Flipping changed the source bitmap'
        for shape in shapes:
            transform = shape._inline.xpath('.//a:xfrm')[0]
            assert [n.tag for n in transform] == [qn('a:off'), qn('a:ext')], 'Transform children missing or out of schema order'
            extent = transform.find(qn('a:ext'))
            assert (int(extent.get('cx')), int(extent.get('cy'))) == (shape.width, shape.height), 'Inner and outer frame dimensions differ'
        print(f'PASS: {file}: independent flip axes, preserved crop/extension metadata, missing-transform creation, frame dimensions, source bytes and unrelated parts')
        continue
    if args.cropped_images or args.new_cropped_images:
        from zipfile import ZipFile
        shapes = list(doc.inline_shapes)
        def crop_attributes(shape):
            nodes = shape._inline.xpath('.//a:srcRect')
            return dict(nodes[0].attrib) if nodes else None
        if args.cropped_images:
            assert [p.text for p in doc.paragraphs] == ['Cropped pictures Á😀', 'Before  edited after.', 'Transparent:  tail', 'Outside source: ', 'AutoFit picture', 'After inline pictures.'], 'Cropped picture source positions changed'
            assert len(shapes) == 6, 'Source picture count changed or repeated copies were serialized'
            assert [(s.width, s.height) for s in shapes] == [(Inches(w), Inches(h)) for w, h in [(1.5, .75), (1.5, .75), (.75, .375), (.5, .25), (1, .5), (2, 1)]], 'Cropped frame dimensions changed'
            assert [crop_attributes(s) for s in shapes] == [dict(l='50000', t='0', r='0', b='0'), dict(l='-25000', t='-25000', r='-25000', b='-25000'), dict(l='125000', r='-50000'), None, dict(l='12.345%'), None], 'Crop edits/reset or untouched percentage spellings changed'
            assert shapes[0]._inline.docPr.get('descr') == 'Cropped orange panel', 'Cropped description changed'
            for index, shape in enumerate(shapes):
                blip = shape._inline.xpath('.//a:blip')[0]
                assert blip.get('cstate') == ('print' if index % 2 else 'screen'), 'Compression preference changed'
                ext = blip.find(qn('a:extLst')).find(qn('a:ext'))
                assert dict(ext.attrib) == {'uri': '{28A0092B-C50C-407e-A947-70E740481C1C}'}, 'Extension was overwritten by resize'
                dpi = list(ext)[0]
                assert dpi.tag == '{http://schemas.microsoft.com/office/drawing/2010/main}useLocalDpi' and dpi.get('val') == (None if index % 3 == 2 else str(index % 3)), 'Local-DPI extension changed'
            assert doc.paragraphs[1].style.style_id == 'CanvasPictureText' and doc.paragraphs[1].style.font.italic and doc.paragraphs[1].style.paragraph_format.line_spacing == 1.5, 'Inherited image paragraph formatting changed'
            assert doc.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None, 'Repeated header lost'
            with ZipFile(Path(__file__).resolve().parents[1] / 'public/fixtures/cropped-images.docx') as original, ZipFile(file) as exported:
                assert original.namelist() == exported.namelist(), 'Package parts changed during cropping'
                for name in original.namelist():
                    if name != 'word/document.xml': assert original.read(name) == exported.read(name), f'Unrelated part changed: {name}'
        else:
            assert [p.text for p in doc.paragraphs] == [' after'], 'Authored crop position changed'
            assert len(shapes) == 1 and (shapes[0].width, shapes[0].height) == (Inches(1), Inches(.5)), 'Authored cropped frame changed'
            assert crop_attributes(shapes[0]) == dict(l='0', t='25000', r='50000', b='25000'), 'Authored crop or reset undo failed'
            assert shapes[0]._inline.docPr.get('descr') == 'Cropped blue panel', 'Authored crop description changed'
            part = doc.part.related_parts[shapes[0]._inline.xpath('.//a:blip')[0].get(qn('r:embed'))]
            with ZipFile(Path(__file__).resolve().parents[1] / 'public/fixtures/cropped-images.docx') as original:
                assert part.blob == original.read('word/media/image1.png'), 'Crop destructively changed the image binary'
        for shape in shapes:
            extent = shape._inline.xpath('.//a:xfrm/a:ext')[0]
            assert (int(extent.get('cx')), int(extent.get('cy'))) == (shape.width, shape.height), 'Inner and outer frame dimensions differ'
        print(f'PASS: {file}: source crop percentages, fixed frames, reset/undo, descriptions, DPI extensions, shared original bytes and unrelated parts')
        continue
    if args.inline_images:
        assert [p.text for p in doc.paragraphs] == ['Inline pictures Á😀', 'Before  edited after.', 'Transparent:  tail', 'Shared asset: ', 'AutoFit picture', 'After inline pictures.'], 'Body text or image insertion position changed'
        shapes = list(doc.inline_shapes)
        assert len(shapes) == 6, 'Image count changed or repeated header copies were serialized'
        assert [(shape.width, shape.height) for shape in shapes] == [(Inches(w), Inches(h)) for w, h in [(1.5, .75), (1.5, .75), (.75, .375), (.5, .25), (1, .5), (2, 1)]], 'Image extents changed'
        assert shapes[0]._inline.docPr.get('descr') == 'Updated blue and orange panels', 'Image description edit lost'
        assert shapes[0]._inline.get('distL') == '914400' and shapes[0]._inline.get('distB') == '457200', 'Ignored inline distances were removed'
        assert doc.paragraphs[1].style.style_id == 'CanvasPictureText' and doc.paragraphs[1].style.font.italic, 'Inherited picture paragraph style changed'
        assert doc.paragraphs[1].style.paragraph_format.line_spacing == 1.5 and doc.paragraphs[1].runs[-1].bold, 'Line spacing or neighboring bold formatting changed'
        ids = [shape._inline.docPr.get('id') for shape in shapes]; assert len(set(ids)) == len(ids), 'Drawing IDs are not unique'
        rel_ids = [shape._inline.xpath('.//a:blip')[0].get(qn('r:embed')) for shape in shapes]
        assert len(set(rel_ids)) == 2 and rel_ids[0] == rel_ids[2] == rel_ids[4] == rel_ids[5] and rel_ids[1] == rel_ids[3], 'Shared image relationships changed'
        source = Document(Path(__file__).resolve().parents[1] / 'public/fixtures/inline-images.docx')
        for rel_id in rel_ids: assert doc.part.related_parts[rel_id].blob == source.part.related_parts[rel_id].blob, 'Original image binary changed'
        table, auto = doc.tables
        assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None and table.cell(1, 0).text == 'Body text. ' * 80, 'Header or body text changed'
        assert auto.autofit and [col.width.twips for col in auto.columns] == [0, 0], 'AutoFit preferences changed'
        print(f'PASS: {file}: inline picture count/positions/extents, descriptions, shared media bytes, inherited spacing, table header and AutoFit')
        continue
    if args.new_inline_images:
        assert [p.text for p in doc.paragraphs] == ['Start ', 'Second  after.'], 'Authored picture positions or adjacent text changed'
        shapes = list(doc.inline_shapes)
        assert len(shapes) == 2, 'PNG/JPEG insertion or deletion undo lost an image'
        assert [(shape.width, shape.height) for shape in shapes] == [(Inches(2), Inches(1)), (Inches(1), Inches(.5))], 'Authored image sizes or oversized-edit rollback failed'
        assert [shape._inline.docPr.get('descr') for shape in shapes] == ['Red panel', 'Blue and orange panels'], 'Authored descriptions changed'
        assert len({shape._inline.docPr.get('id') for shape in shapes}) == 2, 'New drawing IDs collided'
        images = [doc.part.related_parts[shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))] for shape in shapes]
        assert [part.content_type for part in images] == ['image/jpeg', 'image/png'], 'Image content types changed'
        assert [(part.image.px_width, part.image.px_height) for part in images] == [(80, 40), (64, 32)], 'Image binary dimensions changed'
        source = Document(Path(__file__).resolve().parents[1] / 'public/fixtures/inline-images.docx')
        original_id = source.inline_shapes[0]._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
        assert images[1].blob == source.part.related_parts[original_id].blob, 'Inserted PNG bytes changed'
        for shape in shapes:
            extent = shape._inline.xpath('.//a:xfrm/a:ext')[0]
            assert int(extent.get('cx')) == shape.width and int(extent.get('cy')) == shape.height, 'Inner and outer extents disagree'
        print(f'PASS: {file}: authored PNG/JPEG relationships, media types/dimensions, inline extents, descriptions, text and resize/deletion undo')
        continue
    if args.table_headers:
        assert [p.text for p in doc.paragraphs] == ['Repeated table headers.', '\n'.join(f'Lead {i}' for i in range(1, 6)), 'AutoFit header table.', 'After repeated headers.'], 'Body text changed'
        table, auto = doc.tables
        assert table.autofit is False and table.style.style_id == 'CanvasHeaderTable' and table.style.font.italic, 'Inherited table style/layout lost'
        assert table.style.element.find(qn('w:trPr')).find(qn('w:tblHeader')) is not None, 'Inherited header declaration lost'
        assert [None if row._tr.trPr is None or row._tr.trPr.find(qn('w:tblHeader')) is None else row._tr.trPr.find(qn('w:tblHeader')).get(qn('w:val')) for row in table.rows] == [None, None, '0', '0', '0', '0', '1'], 'Header inheritance, overrides or ignored flag changed'
        expected = [['New Product Á😀', 'Details'], ['', 'Description', 'Code']]
        for r in range(2, 7): expected.append(['Merged body. ' * 45 if r == 3 else '' if r == 4 else f'Item {r - 1}', 'Body text. ' * (40 if r == 3 else 7), str(r - 1)])
        assert [[_Cell(tc, table).text for tc in row._tr.tc_lst] for row in table.rows] == expected, 'Header/body text changed or display copies were serialized'
        assert [[tc.grid_span for tc in row._tr.tc_lst] for row in table.rows] == [[1, 2], *([[1, 1, 1]] * 6)], 'Header spans changed'
        assert [[tc.vMerge for tc in row._tr.tc_lst] for row in table.rows] == [['restart', None], ['continue', None, None], [None, None, None], ['restart', None, None], ['continue', None, None], [None, None, None], [None, None, None]], 'Header/body vertical merges changed'
        assert all(run.underline for run in table.cell(1, 1).paragraphs[0].runs), 'Header formatting edit lost'
        assert table.cell(0, 0)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')) == 'EAF2FC', 'Header fill lost'
        assert table.rows[3]._tr.trPr.find(qn('w:cantSplit')) is not None, 'Body row keep lost'
        assert auto.autofit is True and [col.width.twips for col in auto.columns] == [0, 0], 'AutoFit preferences changed'
        assert [[cell.text for cell in row.cells] for row in auto.rows] == [['Auto label', 'Auto value'], *[[f'Auto {r}', 'Automatic content. ' * 30] for r in range(1, 4)]], 'AutoFit contents changed'
        assert auto.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None, 'AutoFit header flag lost'
        print(f'PASS: {file}: inherited and direct headers, ignored noncontiguous flag, single source text, formatting, vertical/horizontal merges and AutoFit')
        continue
    if args.new_table_headers:
        assert len(doc.tables) == 1, 'New table count changed'
        table = doc.tables[0]
        assert table.autofit is False and len(table.rows) == 5 and len(table.columns) == 2, 'New table geometry changed'
        assert [[_Cell(tc, table).text for tc in row._tr.tc_lst] for row in table.rows] == [['Header\nValue'], ['Inserted body'], ['One', 'Long body content. ' * 100], ['Two', 'Last body. ' * 20], ['Three', 'End']], 'Header merge, body insertion or contents changed'
        assert [[tc.grid_span for tc in row._tr.tc_lst] for row in table.rows] == [[2], [2], [1, 1], [1, 1], [1, 1]], 'New header/body spans changed'
        assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')).get(qn('w:val')) == '1', 'New repeated header lost'
        assert table.rows[1]._tr.trPr.find(qn('w:tblHeader')).get(qn('w:val')) == '0', 'Boundary insertion became a header'
        assert all(row._tr.trPr is None or row._tr.trPr.find(qn('w:tblHeader')) is None for row in table.rows[2:]), 'Body row flags changed'
        assert all(tc.p_lst for row in table.rows for tc in row._tr.tc_lst), 'Required cell paragraphs missing'
        print(f'PASS: {file}: authored header, clear/undo, header merge, body insertion, single stored header and retained cell text')
        continue
    if args.vertical_pages:
        assert [p.text for p in doc.paragraphs] == ['Merged cells can continue across pages.', '\n'.join(f'Lead {i}' for i in range(1, 7)), 'AutoFit across pages.', 'Minimum height across pages.', 'After all flowing tables.'], 'Body text changed'
        table, auto, tall = doc.tables
        assert table.autofit is False and table.style.style_id == 'CanvasFlowTable' and table.style.font.italic, 'Flowing table style/layout changed'
        assert [[tc.vMerge for tc in row._tr.tc_lst] for row in table.rows] == [['restart', None, None], ['continue', 'restart', None], ['continue', 'continue', None], [None, 'continue', None]], 'Overlapping merge chains changed'
        assert table.cell(0, 0).text == 'Left merged text. ' * 70 + '\nLeft end Á😀 revised', 'Long merged text changed'
        assert table.cell(1, 1).text == 'Right merged text. ' * 50 + '\nRight end\nMerged list tail', 'Second merge text changed'
        assert table.cell(1, 1).paragraphs[0].runs[0].font.size.pt == 17, 'Mixed font size lost'
        assert all(run.bold for run in table.cell(1, 1).paragraphs[1].runs), 'Direct formatting lost'
        assert table.cell(1, 1).paragraphs[2].style.style_id == 'ListNumber', 'List style lost'
        assert table.style.element.find(qn('w:trPr')).find(qn('w:cantSplit')) is not None, 'Inherited row keep lost'
        assert [None if row._tr.trPr is None or row._tr.trPr.find(qn('w:cantSplit')) is None else row._tr.trPr.find(qn('w:cantSplit')).get(qn('w:val')) for row in table.rows] == ['0', None, '0', None], 'Row keep inheritance/override undo failed'
        assert table.rows[2]._tr.trPr.find(qn('w:trHeight')).get(qn('w:val')) == '600', 'Row minimum changed'
        assert [[tc.grid_span for tc in row._tr.tc_lst] for row in auto.rows] == [[2], [2]], 'AutoFit rectangle changed'
        assert auto.cell(0, 0).text == 'Auto merged content. ' * 100 + '\nAuto ending', 'AutoFit content lost'
        assert auto.autofit is True and [col.width.twips for col in auto.columns] == [0, 0], 'AutoFit preferences changed'
        assert [cell.text for cell in tall.rows[0].cells] == ['Height only edited', 'First\nLast'], 'Tall row text edit lost'
        height = tall.rows[0]._tr.trPr.find(qn('w:trHeight'))
        assert height.get(qn('w:val')) == '5760' and height.get(qn('w:hRule')) == 'atLeast', 'Multi-page minimum height lost'
        print(f'PASS: {file}: overlapping multi-page merges, row keeps, inherited/direct formatting, list style, AutoFit and tall row minima')
        continue
    if args.new_vertical_pages:
        assert len(doc.tables) == 1, 'New table count changed'
        table = doc.tables[0]
        assert table.autofit is False and len(table.rows) == 3 and len(table.columns) == 2, 'New table geometry changed'
        assert [[tc.vMerge for tc in row._tr.tc_lst] for row in table.rows] == [['restart', None], ['continue', None], ['continue', None]], 'Inserted continuation chain changed'
        assert [[_Cell(tc, table).text for tc in row._tr.tc_lst] for row in table.rows] == [['New merged line. ' * 160 + '\nLower paragraph', 'Side top'], ['', 'Inserted row'], ['', 'Side bottom']], 'Long merge text or side rows lost'
        assert all(row._tr.trPr.find(qn('w:cantSplit')) is not None for row in table.rows), 'Authored row keeps lost'
        assert all(tc.p_lst for row in table.rows for tc in row._tr.tc_lst), 'Required continuation paragraph missing'
        print(f'PASS: {file}: newly authored multi-page merge, inserted continuation row, preserved text, required paragraphs and row keeps')
        continue
    if args.vertical_cells:
        assert [p.text for p in doc.paragraphs] == ['Vertical cells across rows.', '\n'.join(f'Lead line {i}' for i in range(1, 8)), 'AutoFit vertical cells.', 'After vertical tables.'], 'Body text changed'
        table, auto = doc.tables
        assert table.autofit is False and table.style.style_id == 'CanvasVerticalTable' and table.style.font.italic is True, 'Table layout or inheritance lost'
        assert [[_Cell(tc, table).text for tc in row._tr.tc_lst] for row in table.rows] == [['Top revised\n\nSecond Á😀', 'B0', 'C0'], ['', 'B1', 'C1'], ['D0', 'Wide\n\n\n'], ['D1', ''], ['F0\nG0', 'F1', 'F2'], ['', 'G1', 'G2']], 'Vertical merge text/order changed'
        assert [[tc.vMerge for tc in row._tr.tc_lst] for row in table.rows] == [['restart', None, None], ['continue', None, None], [None, 'restart'], [None, 'continue'], ['restart', None, None], ['continue', None, None]], 'Vertical markers changed'
        assert [[tc.grid_span for tc in row._tr.tc_lst] for row in table.rows] == [[1, 1, 1], [1, 1, 1], [1, 2], [1, 2], [1, 1, 1], [1, 1, 1]], 'Grid spans changed'
        assert all(run.bold for run in table.cell(4, 0).paragraphs[0].runs), 'Merged formatting lost'
        assert table.rows[1]._tr.tc_lst[0].p_lst[0].pPr.find(qn('w:spacing')).get(qn('w:after')) == '123', 'Continuation paragraph properties lost'
        assert table.cell(0, 0)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')) == 'EAF2FC', 'Merged fill lost'
        assert auto.autofit is True and [col.width.twips for col in auto.columns] == [0, 0, 0], 'AutoFit declarations changed'
        assert [[_Cell(tc, auto).text for tc in row._tr.tc_lst] for row in auto.rows] == [['Auto Á😀', 'Auto side'], ['', 'Auto lower'], ['Auto end', 'M', 'Z']], 'AutoFit content or orphan continuation lost'
        print(f'PASS: {file}: vertical/rectangular merge markers, physical cell text, formatting, continuation preservation and orphan semantics')
        continue
    if args.new_vertical_cells:
        assert len(doc.tables) == 1, 'Table count changed'
        table = doc.tables[0]
        assert table.autofit is True, 'AutoFit missing'
        assert [[_Cell(tc, table).text for tc in row._tr.tc_lst] for row in table.rows] == [['One\nThree', 'Two'], ['', 'Inserted side'], ['', 'Four'], ['Five\nSix\nLast row'], ['']], 'New merged content or row insertion lost'
        assert [[tc.vMerge for tc in row._tr.tc_lst] for row in table.rows] == [['restart', None], ['continue', None], ['continue', None], ['restart'], ['continue']], 'Vertical merge chains changed'
        assert [[tc.grid_span for tc in row._tr.tc_lst] for row in table.rows] == [[1, 1], [1, 1], [1, 1], [2], [2]], 'Rectangle grid spans changed'
        assert table._tbl.tblPr.find(qn('w:tblW')).get(qn('w:type')) == 'auto', 'Fit to contents preference missing'
        assert all(tc.p_lst and tc.tcPr.find(qn('w:tcW')).get(qn('w:type')) == 'auto' for row in table.rows for tc in row._tr.tc_lst), 'Required paragraphs or AutoFit cell preferences missing'
        print(f'PASS: {file}: new vertical merges, inserted continuation rows, rectangular merge/split undo, AutoFit and required cell content')
        continue
    if args.merged_cells:
        assert [p.text for p in doc.paragraphs] == ['Cells can span different grid columns.', 'AutoFit merged cells.', 'After merged tables.'], 'Body paragraphs changed'
        assert len(doc.tables) == 2, 'Table count changed'
        table, auto = doc.tables
        assert table.autofit is False and table.style.style_id == 'CanvasSpanTable', 'Fixed layout or table style lost'
        assert table.style.font.italic is True, 'Shared table formatting changed'
        assert [[tc.grid_span for tc in row._tr.tc_lst] for row in table.rows] == [[2, 1], [1, 2], [1, 2], [3]], 'Merged cell boundaries changed'
        assert [[_Cell(tc, table).text for tc in row._tr.tc_lst] for row in table.rows] == [['Wide heading revised', 'Side'], ['L', 'Long merged cell. ' * 45], ['Left\nMiddle\nRight', 'Split middle\nSplit right'], ['End Á😀']], 'Merge/split text changed'
        assert [col.width.twips for col in table.columns] == [1440, 2880, 2160], 'Grid column edit lost'
        assert [[_Cell(tc, table).width.twips for tc in row._tr.tc_lst] for row in table.rows] == [[4320, 2160], [1440, 5040], [1440, 5040], [6480]], 'Merged preferred widths do not sum grid widths'
        assert all(r.bold is True for r in table.cell(2, 0).paragraphs[1].runs), 'Formatting inside merged cell lost'
        assert table.cell(0, 0)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')) == 'EAF2FC', 'Merged shading lost'
        assert [node.tag for node in table.cell(0, 0)._tc.tcPr][:3] == [qn('w:cnfStyle'), qn('w:tcW'), qn('w:gridSpan')], 'Cell property order or retained flags changed'
        assert table.rows[1]._tr.trPr.find(qn('w:cantSplit')) is not None, 'Oversized merged row keep lost'
        assert auto.autofit is True and [col.width.twips for col in auto.columns] == [0, 0, 0], 'Unedited AutoFit grid changed'
        assert [[tc.grid_span for tc in row._tr.tc_lst] for row in auto.rows] == [[2, 1], [1, 2], [1, 1, 1]], 'AutoFit spans changed'
        print(f'PASS: {file}: grid spans, merge/split contents, grid and summed cell preferences, inherited style, direct bold, shading and row keeps')
        continue
    if args.new_merged_cells:
        assert len(doc.tables) == 1, 'Table count changed'
        table = doc.tables[0]
        assert table.autofit is True, 'AutoFit missing'
        assert [[tc.grid_span for tc in row._tr.tc_lst] for row in table.rows] == [[3], [3], [1, 1, 1]], 'Inserted row or grid-column edits lost spans'
        assert [[_Cell(tc, table).text for tc in row._tr.tc_lst] for row in table.rows] == [['One\nTwo\nThree'], ['New merged row'], ['Four', 'Five', 'Six']], 'Merged content or new row changed'
        assert [col.width.twips for col in table.columns] == [2340, 2340, 2340], 'Column insertion/deletion preferences changed'
        assert table._tbl.tblPr.find(qn('w:tblW')).get(qn('w:type')) == 'auto', 'Fit to contents table preference missing'
        assert all(tc.tcPr.find(qn('w:tcW')).get(qn('w:type')) == 'auto' for row in table.rows for tc in row._tr.tc_lst), 'Fit to contents retained cell preferences'
        print(f'PASS: {file}: new merged cells, retained paragraphs, span-aware row/column editing, AutoFit and source width preferences')
        continue
    if args.autofit_tables:
        assert [p.text for p in doc.paragraphs] == ['AutoFit columns respond to content.', 'A compact table.', 'After both tables.'], 'Body paragraphs changed'
        assert len(doc.tables) == 2, 'Table count changed'
        table, compact = doc.tables
        assert table.autofit is True and len(table.rows) == 3 and len(table.columns) == 3, 'AutoFit structure changed'
        assert table.cell(0, 0).text == 'Extraordinary revised' and all(run.bold is True for run in table.cell(0, 0).paragraphs[0].runs), 'Cell edit or direct bold lost'
        assert table.cell(0, 1).text == 'Many short words for measuring and wrapping. ' * 8, 'Long cell text changed'
        assert [cell.text for cell in table.rows[1].cells] == ['A\u00a0B', 'Value\t12.5', 'Á😀'], 'Nonbreaking space, tab or graphemes changed'
        assert [cell.text for cell in table.rows[2].cells] == ['', '', 'Last'], 'Empty cells changed'
        assert [col.width.twips for col in table.columns] == [1800, 3240, 2880], 'Column preferences lost'
        for row in table.rows:
            assert [cell.width.twips for cell in row.cells] == [1800, 3240, 2880], 'Cell preferences lost'
        w = table._tbl.tblPr.find(qn('w:tblW'))
        assert w.get(qn('w:type')) == 'pct' and w.get(qn('w:w')) == '83.33%', 'Fractional percentage lost'
        assert table._tbl.tblPr.find(qn('w:jc')).get(qn('w:val')) == 'center', 'Alignment lost'
        assert table.style.style_id == 'TableGrid' and compact.style.style_id == 'TableGrid', 'Styles changed'
        assert compact.autofit is True and [col.width.twips for col in compact.columns] == [0, 0], 'Preview measurements replaced the zero grid'
        assert [[cell.text for cell in row.cells] for row in compact.rows] == [['One', 'Two'], ['Three', '']], 'Compact cells changed'
        print(f'PASS: {file}: AutoFit, fractional percentage, declared grid/cell widths, alignment, text, styles and zero grid preservation')
        continue
    if args.new_autofit_tables:
        assert len(doc.tables) == 1, 'Table count changed'
        table = doc.tables[0]
        assert table.autofit is True and [[cell.text for cell in row.cells] for row in table.rows] == [['One', 'Two'], ['', '']], 'New AutoFit table text or structure changed'
        assert [col.width.twips for col in table.columns] == [1440, 4320], 'Column preferences changed'
        w = table._tbl.tblPr.find(qn('w:tblW'))
        assert w.get(qn('w:type')) == 'pct' and w.get(qn('w:w')) == '5000', 'Fit to page width missing'
        indent = table._tbl.tblPr.find(qn('w:tblInd'))
        assert indent.get(qn('w:type')) == 'dxa' and indent.get(qn('w:w')) == '360', 'Indentation lost'
        for row in table.rows:
            assert all(cell._tc.tcPr.find(qn('w:tcW')).get(qn('w:type')) == 'auto' for cell in row.cells), 'Fit to page retained fixed cell requests'
        print(f'PASS: {file}: new AutoFit table, page percentage, automatic cells, retained column preferences and indentation')
        continue
    if args.tables:
        assert [p.text for p in doc.paragraphs] == ['A table with editable cells.', 'After the table.'], 'Body paragraphs changed'
        assert len(doc.tables) == 1, 'Table count changed'
        table = doc.tables[0]
        assert len(table.rows) == 4 and len(table.columns) == 2 and table.autofit is False, 'Table structure/layout changed'
        assert [col.width.twips for col in table.columns] == [4320, 2880], 'Grid widths changed'
        assert table.style.style_id == 'CanvasTable' and table.style.font.italic is True, 'Table style reference or inherited declaration changed'
        expected = [['Product', 'Details'], ['Alpha revised café Á😀\nSecond paragraph\nAdded paragraph', 'Before\tAfter'], ['Text inside a tall row. ' * 45, 'Short cell'], ['Last row', '']]
        assert [[cell.text for cell in row.cells] for row in table.rows] == expected, 'Cell text or boundaries changed'
        for row in table.rows:
            assert [cell.width.twips for cell in row.cells] == [4320, 2880], 'Cell preferred widths changed'
        assert table.rows[2]._tr.trPr.find(qn('w:cantSplit')).get(qn('w:val')) == '0', 'Row splitting override missing'
        assert table.rows[3]._tr.trPr.find(qn('w:trHeight')).get(qn('w:val')) == '480', 'Minimum row height changed'
        assert table.cell(0, 0)._tc.tcPr.find(qn('w:shd')).get(qn('w:fill')) == 'EAF2FC', 'Cell shading changed'
        for p in table.cell(1, 0).paragraphs[1:]:
            assert all(run.bold is True and run.italic is None and run.font.size is None for run in p.runs), 'Direct bold or unflattened table inheritance lost'
        print(f'PASS: {file}: table/cell structure, grid/preferred widths, inherited style, row split override, cell edits, formatting and retained properties')
        continue
    if args.new_tables:
        assert [p.text for p in doc.paragraphs] == ['Before', 'After'], 'Table insertion lost surrounding text'
        assert len(doc.tables) == 1, 'Table count changed'
        table = doc.tables[0]
        assert len(table.rows) == 4 and len(table.columns) == 3 and table.autofit is False, 'Row/column editing or fixed layout lost'
        assert [[cell.text for cell in row.cells] for row in table.rows] == [['One', '', 'Two'], ['Three', '', 'Four'], ['Five', 'Added column', ''], ['Last', '', '']], 'Cell editing or empty cells changed'
        assert table.style.style_id == 'CanvasTable1' and table.style.type == WD_STYLE_TYPE.TABLE, 'Generated table style missing'
        assert [col.width.twips for col in table.columns] == [3120, 3120, 3120], 'Column insertion changed the overall width'
        assert table._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w')) == '9360', 'Preferred table width changed'
        for row in table.rows:
            assert row._tr.trPr.find(qn('w:cantSplit')).get(qn('w:val')) in (None, '1'), 'New row keep property missing'
            assert all(cell.width.twips == 3120 and len(cell.paragraphs) >= 1 for cell in row.cells), 'Cell widths/required paragraphs missing'
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            assert table._tbl.tblPr.find(qn('w:tblBorders')).find(qn('w:' + side)).get(qn('w:val')) == 'single', 'Table border missing'
        print(f'PASS: {file}: inserted table, row/column operations, required cell paragraphs, content, dimensions, row keeps, style and borders')
        continue
    if args.character_styles:
        assert [p.text for p in doc.paragraphs] == ['Plain base child direct', 'Paragraph inherited child', 'Clear this style', 'Empty styled', 'Styled text across pages. ' * 25, 'Styled bullet item', 'A\u0301😀 symbol'], 'Character-style text changed'
        by_id = {style.style_id: style for style in doc.styles}
        base = by_id['CanvasCharacterBase']; child = by_id['CanvasCharacterChild']; grand = by_id['CanvasCharacterGrandchild']
        assert base.font.size == Pt(18) and base.font.italic is None, 'Shared base edit or removed declaration missing'
        assert child.base_style.style_id == base.style_id and child.font.bold is True and child.font.italic is False, 'Character inheritance changed'
        assert grand.base_style.style_id == child.style_id and grand.font.underline is True, 'Grandchild declarations changed'
        assert doc.paragraphs[0].runs[0].text == 'Plain' and doc.paragraphs[0].runs[0].style.style_id == grand.style_id, 'Selected style was not applied'
        direct = next(r for r in doc.paragraphs[0].runs if r.text == ' direct')
        assert direct._r.rPr.find(qn('w:rStyle')) is None, 'Removed reference remains'
        assert direct.font.size == Pt(22) and direct.bold is False and direct.italic is True, 'Direct formatting was not preserved'
        assert doc.paragraphs[2].runs[0]._r.rPr.find(qn('w:rStyle')) is None, 'Clear style did not remove rStyle'
        assert doc.paragraphs[3].runs[0].style.style_id == grand.style_id, 'Empty-paragraph typing lost its character style'
        assert doc.paragraphs[3]._p.pPr.find(qn('w:rPr')).find(qn('w:rStyle')).get(qn('w:val')) == grand.style_id, 'Paragraph-mark reference changed'
        assert doc.paragraphs[3].runs[0].font.size is None, 'Inherited size was flattened into direct formatting'
        assert by_id['CanvasCharacterList'].element.pPr.numPr.numId.val == 1, 'Inherited list reference changed'
        assert by_id['CanvasDecorative'].element.rPr.find(qn('w:u')).get(qn('w:val')) == 'double', 'Unused style was not preserved'
        print(f'PASS: {file}: character references, base edit, removed declarations, direct overrides, paragraph marks, text and inherited numbering')
        continue
    if args.new_character_styles:
        assert [p.text for p in doc.paragraphs] == ['Seed Café', 'Plain again', ''], 'Composed or plain text changed'
        custom = doc.styles['Review characters']
        assert custom.type == WD_STYLE_TYPE.CHARACTER and custom.style_id == 'CanvasCharacter1', 'Custom character identity changed'
        assert custom.base_style.style_id == 'Emphasis' and custom.font.size == Pt(16) and custom.font.underline is True, 'Custom character declarations changed'
        assert ''.join(r.text for r in doc.paragraphs[0].runs if r.style is not None and r.style.style_id == custom.style_id) == 'Café', 'Composition lost its character reference'
        assert all(r._r.rPr is None or r._r.rPr.find(qn('w:rStyle')) is None for r in doc.paragraphs[1].runs), 'Plain typing kept the prior style'
        assert doc.paragraphs[2]._p.pPr.find(qn('w:rPr')).find(qn('w:rStyle')).get(qn('w:val')) == custom.style_id, 'Styled empty paragraph was not saved'
        assert not doc.element.xpath('./w:body/w:p/w:r/w:rPr/w:sz | ./w:body/w:p/w:pPr/w:rPr/w:sz'), 'Inherited character sizes were flattened'
        print(f'PASS: {file}: custom character definition, composition, future typing, cleared references, empty mark and unflattened sizes')
        continue
    if args.styles:
        assert [p.style.style_id for p in doc.paragraphs] == ['CanvasSubheading', 'CanvasBody', 'CanvasSubheading', 'CanvasList', 'CanvasBody', 'CanvasBody'], 'Paragraph style references changed'
        assert [p.text for p in doc.paragraphs[:5]] == ['Style heading', 'Body bold and emphasis', 'Direct size and spacing', 'Style list item', ''], 'Styled text changed'
        assert doc.paragraphs[5].text == 'Text across pages. ' * 28, 'Paginated body text changed'
        body = doc.styles['Canvas Body']
        assert body.font.size == Pt(14) and body.font.italic is True, 'Shared base style edit missing'
        assert body.paragraph_format.space_after == Pt(8), 'Style paragraph spacing edit missing'
        assert body.element.find(qn('w:next')).get(qn('w:val')) == 'CanvasBody', 'Explicit self-next reference was lost'
        assert doc.styles['Canvas Subheading'].base_style.style_id == 'CanvasHeading', 'basedOn changed'
        assert doc.styles['Canvas Heading'].next_paragraph_style.style_id == 'CanvasBody', 'Following style changed'
        assert doc.styles['Canvas Subheading'].font.size == Pt(16), 'Derived size changed'
        assert doc.paragraphs[0].runs[0].font.size is None, 'Inherited size was flattened into direct formatting'
        assert doc.paragraphs[1].runs[1].bold is True, 'Direct bold was lost'
        assert doc.paragraphs[1].runs[3].style.style_id == 'CanvasEmphasis', 'Character style reference was lost'
        assert doc.styles['Canvas Emphasis'].font.italic is True, 'Character style definition changed'
        assert doc.paragraphs[2].runs[0].font.size == Pt(24) and doc.paragraphs[2].paragraph_format.space_after == Pt(0), 'Direct size or spacing override changed'
        assert doc.styles['Canvas List'].element.pPr.numPr.numId.val == 1, 'Inherited list reference changed'
        print(f'PASS: {file}: paragraph styles, shared base edit, basedOn/next, character style, direct overrides, text and inherited numbering')
        continue
    if args.new_styles:
        assert [p.text for p in doc.paragraphs] == ['Heading text', 'Body text', 'Next heading'], 'New styled text changed'
        assert [p.style.style_id for p in doc.paragraphs] == ['Heading1', 'CanvasStyle1', 'Heading2'], 'Style transitions changed'
        custom = doc.styles['Review body']
        assert custom.style_id == 'CanvasStyle1' and custom.base_style.style_id == 'Normal', 'Custom style identity or base changed'
        assert custom.next_paragraph_style.style_id == 'Heading2', 'Custom following style changed'
        assert custom.font.size == Pt(15) and custom.font.underline is True, 'Custom text properties changed'
        assert custom.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER, 'Custom alignment changed'
        assert all(p.runs[0].font.size is None for p in doc.paragraphs), 'Style font sizes were flattened into direct formatting'
        assert not doc.element.xpath('./w:body/w:p/w:pPr/w:rPr/w:sz'), 'Inherited paragraph-mark sizes were flattened into direct formatting'
        by_id = {style.style_id: style for style in doc.styles}
        assert by_id['Heading1'].font.size == Pt(20) and by_id['Heading2'].font.size == Pt(16), 'Heading definitions changed'
        print(f'PASS: {file}: custom style identity, inheritance, next-style transitions, shared text/paragraph declarations and unflattened runs')
        continue
    if args.new_lists:
        assert [p.text for p in doc.paragraphs] == ['First item', 'Second item', 'Bullet item'], 'New list text changed'
        refs = [p._p.pPr.numPr for p in doc.paragraphs]
        assert all(ref.ilvl.val == 0 for ref in refs), 'New list levels changed'
        assert refs[0].numId.val == refs[1].numId.val != refs[2].numId.val, 'New sequences merged'
        root = doc.part.numbering_part.element
        for i, expected in [(0, 'decimal'), (2, 'bullet')]:
            instance = root.xpath(f'./w:num[@w:numId="{refs[i].numId.val}"]')[0]
            abstract_id = instance.xpath('./w:abstractNumId')[0].get(qn('w:val'))
            level = root.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="0"]')[0]
            assert level.find(qn('w:numFmt')).get(qn('w:val')) == expected, 'List kind changed'
            assert level.find(qn('w:pPr')).find(qn('w:ind')).get(qn('w:left')) == '720', 'List hanging geometry missing'
        print(f'PASS: {file}: new number and bullet definitions, references, text, levels and hanging indents')
        continue
    if args.numbering:
        assert len(doc.paragraphs) == 11, 'Paragraph boundaries changed'
        assert doc.paragraphs[0].text == 'Edited Numbering starts at three', 'Text edit missing'
        assert all(r.bold for r in doc.paragraphs[0].runs), 'Original bold lost'
        root = doc.part.numbering_part.element
        refs = [p._p.pPr.numPr if p._p.pPr is not None else None for p in doc.paragraphs]
        new_id = refs[1].numId.val
        assert new_id not in [0, 42, 43, 44], 'Restart did not allocate a distinct instance'
        assert [refs[i].numId.val for i in [1, 2, 3, 4]] == [new_id] * 4, 'Restart did not follow the list continuation'
        assert refs[3].ilvl.val == 2, 'Nested level edit missing'
        instance = root.xpath(f'./w:num[@w:numId="{new_id}"]')[0]
        assert instance.xpath('./w:lvlOverride/w:startOverride')[0].get(qn('w:val')) == '1', 'Start override missing'
        assert refs[5].numId.val == refs[6].numId.val == 43, 'Independent sequence changed'
        assert refs[9].numId.val == 0, 'Explicit list removal lost'
        assert doc.paragraphs[0].style.base_style.name == 'Canvas List', 'Based-on style lost'
        assert doc.paragraphs[0].style.font.italic and doc.paragraphs[0].style.base_style.font.italic, 'Nearest italic declarations changed'
        assert doc.paragraphs[7].text == 'Bullet item' and doc.paragraphs[10].text == '', 'Bullet text or empty item changed'
        assert len(root.xpath('./w:abstractNum[@w:abstractNumId="41"]/w:lvl')) == 9, 'Multilevel definition changed'
        print(f'PASS: {file}: list references, nested level, restart instance, independent sequence, numId=0, styles, text and bold')
        continue
    if args.pagination:
        assert len(doc.paragraphs) == 7, 'Paragraph count changed'
        first, heading, subheading, body, kept, widow, empty = doc.paragraphs
        assert first.text == '\n'.join(f'Filler line {i}' for i in range(1, 11)), 'Filler changed'
        assert heading.text == 'Edited Keep this heading' and all(r.bold for r in heading.runs), 'Heading edit or original bold lost'
        assert subheading.text == 'Keep this subheading', 'Subheading changed'
        for p in [heading, subheading]:
            assert p.paragraph_format.keep_with_next is False, 'Direct keepNext off missing'
            assert p.paragraph_format.keep_together is False, 'Direct keepLines off missing'
            assert p.paragraph_format.widow_control is True, 'Direct widowControl on missing'
            assert p.style.name == 'Canvas Keep Heading', 'Style reference lost'
            inherited = p.style.base_style.paragraph_format
            assert inherited.keep_with_next and inherited.keep_together and inherited.widow_control is False, 'Inherited style changed'
        assert body.text == 'First body line\nSecond body line\nThird body line', 'Body changed'
        assert kept.paragraph_format.keep_together is True and kept.paragraph_format.widow_control is False, 'Untouched direct flags changed'
        assert widow.text == '\n'.join(f'Widow line {i}' for i in range(1, 6)) and empty.text == '', 'Ending content changed'
        assert doc.sections[0].page_height.inches == 4, 'Page setup changed'
        print(f'PASS: {file}: pagination overrides, based-on styles, text, manual lines, bold, empty paragraph and pages')
        continue
    if args.indents_tabs:
        assert len(doc.paragraphs) == 4, 'Paragraph count changed'
        p = doc.paragraphs[0]
        assert p.text == 'New Lead\tLeft\tCenter\tRight\t123.45', 'Tab characters or edited text lost'
        pf = p.paragraph_format
        assert [pf.left_indent.twips, pf.right_indent.twips, pf.first_line_indent.twips] == [720, 720, 360], 'Indent overrides missing'
        assert [(t.position.twips, t.alignment) for t in pf.tab_stops] == [(2160, WD_TAB_ALIGNMENT.CLEAR), (2520, WD_TAB_ALIGNMENT.LEFT), (4320, WD_TAB_ALIGNMENT.CENTER), (6480, WD_TAB_ALIGNMENT.RIGHT), (7920, WD_TAB_ALIGNMENT.DECIMAL), (8640, WD_TAB_ALIGNMENT.CLEAR)], 'Custom stops or inherited clear missing'
        assert doc.settings.element.find(qn('w:defaultTabStop')).get(qn('w:val')) == '360', 'Default tab interval missing'
        assert p.style.paragraph_format.first_line_indent.twips == -360, 'Base style changed'
        assert p.runs[-2].text == '123' and p.runs[-2].bold, 'Mixed bold run lost'
        assert p.runs[-1].text == '.45' and p.runs[-1].italic, 'Mixed italic run lost'
        assert doc.paragraphs[1].text == 'Inherited stops\t42.50', 'Untouched paragraph changed'
        assert doc.paragraphs[2]._p.xpath('.//w:br'), 'Manual break lost'
        assert doc.paragraphs[3].text == '', 'Empty paragraph lost'
        print(f'PASS: {file}: indents, custom stops, inherited clear, default interval, text, mixed runs, empty paragraph and manual break')
        continue
    if args.page_settings:
        assert len(doc.sections) == 1 and len(doc.paragraphs) == 3, 'Section or paragraph boundaries changed'
        s = doc.sections[0]
        assert s.orientation == WD_ORIENT.PORTRAIT, 'Page orientation not saved'
        assert (s.page_width.twips, s.page_height.twips) == (12240, 15840), 'Letter dimensions not saved'
        assert [s.top_margin.twips, s.right_margin.twips, s.bottom_margin.twips, s.left_margin.twips, s.gutter.twips] == [720, 1080, 1080, 720, 180], 'Margins/gutter changed'
        assert (s.header_distance.twips, s.footer_distance.twips) == (720, 720), 'Header/footer distances changed'
        # python-docx intentionally omits page breaks from .text; inspect those elements separately.
        assert doc.paragraphs[0].text == 'Before breakAfter break\nSame paragraph', 'Break editing damaged text'
        breaks = doc.paragraphs[0]._p.xpath('.//w:br')
        assert [b.get(qn('w:type'), 'textWrapping') for b in breaks] == ['page', 'page', 'textWrapping'], 'Authored break kinds/order changed'
        prefix = []
        break_positions = []
        for item in doc.paragraphs[0]._p.iter():
            if item.tag == qn('w:t'):
                prefix.append(item.text or '')
            elif item.tag == qn('w:br'):
                break_positions.append(len(''.join(prefix)))
        assert break_positions == [7, 12, 23], 'Break source positions changed'
        assert all(r.bold for r in doc.paragraphs[0].runs if r.text and 'After' not in r.text), 'Original bold lost'
        mixed = [r for r in doc.paragraphs[0].runs if 'After break' in r.text][0]
        assert mixed.italic and mixed.underline and mixed.font.size.pt == 18, 'Original mixed formatting lost'
        assert doc.paragraphs[1].paragraph_format.page_break_before is False, 'Explicit pagination override missing'
        assert doc.paragraphs[1].style.paragraph_format.page_break_before is True, 'Inherited style was changed'
        assert doc.paragraphs[2].text == 'Last paragraph', 'Untouched paragraph changed'
        print(f'PASS: {file}: dimensions, orientation, margins, exact break positions, paragraph structure, styles and mixed runs')
        continue
    if args.font_formatting:
        assert len(doc.paragraphs) == 3, 'Expected three paragraphs'
        assert doc.paragraphs[0].text == 'Edited A 24-point heading', 'Edited heading missing'
        for run in doc.paragraphs[0].runs:
            assert run.font.size.pt == 32 and run.italic is False and run.underline is False, 'Explicit font overrides missing'
        assert doc.paragraphs[0].style.font.size.pt == 24, 'Original style size changed'
        assert doc.paragraphs[0].style.font.italic is True, 'Original italic style changed'
        body = doc.paragraphs[1].runs
        assert [r.text for r in body] == ['Small ', 'mixed', ' text at 13.5 pt.'], 'Body text changed'
        assert body[1].font.size.pt == 18 and body[1].italic and body[1].underline, 'Mixed body formatting changed'
        assert body[2].font.size.pt == 13.5, 'Half-point size changed'
        mark = doc.paragraphs[2]._p.pPr.find(qn('w:rPr'))
        assert mark.find(qn('w:sz')).get(qn('w:val')) == '36', 'Empty paragraph size lost'
        print(f'PASS: {file}: sizes, explicit italic/underline off, inherited styles, mixed runs, empty paragraph')
        continue
    assert len(doc.paragraphs) == 2, "Expected two paragraphs"
    if args.paragraph_formatting:
        assert doc.paragraphs[0].text.startswith("Edited Formatting survives"), "Edited text missing"
        assert doc.paragraphs[1].text.startswith("Every page is another chance"), "Second paragraph missing"
        for p in doc.paragraphs:
            assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER, "Center alignment missing"
            assert p.paragraph_format.line_spacing == 2.0, "Double line spacing missing"
            assert p.paragraph_format.space_before.pt == 12, "Before spacing missing"
            assert p.paragraph_format.space_after.pt == 18, "After spacing missing"
        assert doc.paragraphs[0].style.name == "Canvas Heading", "Original style changed"
        assert doc.sections[0].page_width.inches == 8.5, "Original page settings changed"
        print(f"PASS: {file}: edited text, alignment, spacing, original style and page settings")
        continue
    assert doc.paragraphs[0].text.startswith("New A small beginning"), "Edited text missing"
    assert doc.paragraphs[1].text.startswith("Select words across"), "Second paragraph missing"
    assert all(run.bold is True for p in doc.paragraphs for run in p.runs if run.text), "Bold formatting missing"
    assert doc.sections[0].page_width.inches == 8.5, "Original page settings changed"
    print(f"PASS: {file}: two paragraphs, edited text, bold formatting, page settings")
