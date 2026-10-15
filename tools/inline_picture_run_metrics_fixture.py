from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_LINE_SPACING
import json
out=Path(__file__).resolve().parents[1]/'tests/fixtures/inline-picture-run-metrics';out.mkdir(exist_ok=True)
d=Document();d.sections[0].left_margin=d.sections[0].right_margin=Inches(1);d.sections[0].page_width=Inches(8.5);d.sections[0].page_height=Inches(11);d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
s=d.styles['Normal'];s.font.name='Calibri';s.font.size=Pt(11);s.paragraph_format.space_after=Pt(0);s.paragraph_format.line_spacing=1
cases=[]
for size in [8,11,16,24]:
 for height in [4,8,16]:
  c=dict(name=f'f{size}-h{height}',size=size,height=height);cases.append(c)
  p=d.add_paragraph(c['name']);p.paragraph_format.page_break_before=len(cases)>1
  p=d.add_paragraph();r=p.add_run('Before');r.font.name='Verdana';r.font.size=Pt(8)
  p=d.add_paragraph();r=p.add_run();r.font.size=Pt(size);b=BytesIO();Image.new('RGB',(4,4),'#e03040').save(b,format='PNG');r.add_picture(BytesIO(b.getvalue()),width=Pt(height),height=Pt(height));r=p.add_run('Tail');r.font.name='Verdana';r.font.size=Pt(8)
  p=d.add_paragraph();r=p.add_run('After');r.font.name='Verdana';r.font.size=Pt(8)
d.save(out/'source.docx');(out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
for setting in d.settings.element.iter(qn('w:compatSetting')):
 if setting.get(qn('w:name'))=='compatibilityMode':setting.set(qn('w:val'),'15')
d.save(out/'modern.docx')
rules=[('multiple',.5),('multiple',1.5),('multiple',2),('exact',6),('exact',24),('atLeast',6),('atLeast',24)]
for i,(rule,value) in enumerate(rules):
 p=d.paragraphs[i*4+2]
 p.paragraph_format.line_spacing=value if rule=='multiple' else Pt(value)
 if rule=='atLeast':p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
 cases[i].update(rule=rule,spacing=value)
for p in list(d.paragraphs[28:]):p._p.getparent().remove(p._p)
d.save(out/'spacing.docx');(out/'spacing-cases.json').write_text(json.dumps(cases[:7],indent=2)+'\n')
