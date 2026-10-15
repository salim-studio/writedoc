from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
out=Path(__file__).resolve().parents[1]/'tests/fixtures/border-device-pixels';out.mkdir(exist_ok=True)
def prop(p,n,**a):
 e=OxmlElement('w:'+n)
 for k,v in a.items():e.set(qn('w:'+k),str(v))
 p.append(e);return e
d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.5);sec.left_margin=sec.right_margin=Inches(.5)
n=d.styles['Normal'];n.font.name='Arial';n.font.size=Pt(10);n.paragraph_format.line_spacing=Pt(12);n.paragraph_format.space_after=Pt(0)
for size in [2,4,6,8,12,16]:
 t=d.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Inches(3);t.cell(0,0).width=Inches(3);t.cell(0,0).text=f'Border {size}/8 pt';t.rows[0].height=Pt(20);t.rows[0].height_rule=2
 b=prop(t._tbl.tblPr,'tblBorders');t._tbl.tblPr.remove(b);t._tbl.tblPr.insert(1,b)
 for side in ['top','left','bottom','right']:prop(b,side,val='single',sz=size,color='6D7F91')
 p=d.add_paragraph('');p.paragraph_format.line_spacing=Pt(6)
d.save(out/'control-14.docx')
for e in d.settings.element.iter(qn('w:compatSetting')):
 if e.get(qn('w:name'))=='compatibilityMode':e.set(qn('w:val'),'15')
d.save(out/'control-15.docx')
