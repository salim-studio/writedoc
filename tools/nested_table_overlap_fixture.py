"""Independent nested floating-table collision controls, original OOXML only."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
out=Path(__file__).resolve().parents[1]/'tests/fixtures/nested-table-overlap';out.mkdir(exist_ok=True)
def prop(parent,name,**attrs):
 n=OxmlElement('w:'+name)
 for k,v in attrs.items():n.set(qn('w:'+k),str(v))
 parent.append(n);return n
cases=[]
for width in [108,252]:
 for modes in [('overlap','overlap'),('never','overlap'),('overlap','never'),('never','never')]:cases.append(dict(widths=[width,width],modes=modes))
for width,right,left,bottom,top in [(108,0,12,0,0),(108,12,0,0,0),(108,180,0,0,0),(108,0,180,0,0),(252,180,180,0,0),(252,0,0,180,0),(252,0,0,0,180),(252,180,180,180,180)]:cases.append(dict(widths=[width,width],right=right,left=left,bottom=bottom,top=top))
cases += [dict(widths=[108,108,108],heights=[45,15,15]),dict(widths=[252,108,108],heights=[30,15,15]),dict(widths=[108,108],border=16),dict(widths=[252,252],border=16),dict(widths=[180,180]),dict(widths=[198,198])]
d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for n in d.settings.element.iter(qn('w:compatSetting')):
 if n.get(qn('w:name'))=='compatibilityMode':n.set(qn('w:val'),'15')
normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12);normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=Pt(15)
for i,case in enumerate(cases,1):
 case['case']=i;p=d.add_paragraph(f'Case {i}');p.paragraph_format.page_break_before=i>1
 outer=d.add_table(rows=1,cols=1);outer.autofit=False;outer.columns[0].width=Inches(6);cell=outer.cell(0,0);cell.width=Inches(6);cell.paragraphs[0].text='Before'
 for j,width in enumerate(case['widths']):
  inner=cell.add_table(rows=1,cols=1);inner.autofit=False;inner.columns[0].width=Pt(width);inner.cell(0,0).width=Pt(width);tw=inner._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(width*20))
  pp=prop(inner._tbl.tblPr,'tblpPr',horzAnchor='page',vertAnchor='page',tblpX=720,tblpY=1440,leftFromText=case.get('left',0) if j else 0,rightFromText=case.get('right',0) if j==0 else 0,topFromText=case.get('top',0) if j else 0,bottomFromText=case.get('bottom',0) if j==0 else 0);inner._tbl.tblPr.remove(pp);inner._tbl.tblPr.insert(0,pp)
  ov=prop(inner._tbl.tblPr,'tblOverlap',val=case.get('modes',['never']*len(case['widths']))[j]);inner._tbl.tblPr.remove(ov);inner._tbl.tblPr.insert(1,ov)
  ic=inner.cell(0,0);ic.text=chr(65+j)+str(i);prop(ic._tc.get_or_add_tcPr(),'shd',val='clear',fill=['FFD0C0','C0D8FF','C0E8C0'][j])
  for k in range(case.get('heights',[15]*len(case['widths']))[j]//15-1):ic.add_paragraph(f'Line {j+1}.{k+1}')
  if case.get('border'):
   border=prop(ic._tc.get_or_add_tcPr(),'tcBorders')
   tcpr=ic._tc.get_or_add_tcPr();tcpr.remove(border);tcpr.insert(tcpr.index(tcpr.find(qn('w:shd'))),border)
   for side in ['top','left','bottom','right']:prop(border,side,val='single',sz=case['border'],color='222222')
  cell.paragraphs[-1].text='Anchor '+str(j+1)
 d.add_paragraph('After outer')
d.save(out/'control.docx');(out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
