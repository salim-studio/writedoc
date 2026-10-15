"""Original bookmark/internal-link fixture and independent python-docx checks."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/bookmarks.docx'

def mark(parent, id, name=None):
    node = OxmlElement('w:bookmarkStart' if name is not None else 'w:bookmarkEnd')
    node.set(qn('w:id'), str(id))
    if name is not None: node.set(qn('w:name'), name)
    parent.append(node)

def link(p, text, anchor=None, external=None):
    node = OxmlElement('w:hyperlink')
    if anchor is not None: node.set(qn('w:anchor'), anchor)
    if external: node.set(qn('r:id'), p.part.relate_to(external, RT.HYPERLINK, is_external=True))
    p._p.append(node); r = p.add_run(text); r.style = 'Hyperlink'; node.append(r._r)
    return node

def generate(factory):
    doc = factory(); sec = doc.sections[0]
    sec.page_width = Inches(6); sec.page_height = Inches(7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.5)
    doc.styles['Normal'].font.name = 'Noto Sans'; doc.styles['Normal'].font.size = Pt(12)
    style = doc.styles.add_style('Hyperlink', WD_STYLE_TYPE.CHARACTER); style.font.color.rgb = RGBColor.from_string('0563C1'); style.font.underline = True
    doc.add_paragraph('Bookmarks and internal links')
    p = doc.add_paragraph('Before '); mark(p._p, 1, 'Target'); p.add_run('target a\u0301😀'); mark(p._p, 1); p.add_run(' after.')
    p = doc.add_paragraph('Across '); mark(p._p, 20, 'Across'); p.add_run('two paragraphs with '); mark(p._p, 21, 'Overlap'); p.add_run('overlapping markers')
    p = doc.add_paragraph('and their '); mark(p._p, 20); p.add_run('last words'); mark(p._p, 21)
    p = doc.add_paragraph(); mark(p._p, 30, '_Point'); mark(p._p, 30); p.add_run('Point destination')
    p = doc.add_paragraph('Link to '); node = link(p, 'jump ', 'Target'); mark(node, 40, 'InsideLink'); r = p.add_run('middle'); r.style = style; r.bold = True; node.append(r._r); mark(node, 40); r = p.add_run(' end'); r.style = style; node.append(r._r)
    p = doc.add_paragraph(); link(p, 'External precedence', 'Target', 'https://example.com/external')
    p = doc.add_paragraph(); link(p, 'Missing target', 'NoSuch'); p.add_run(' / '); link(p, 'Top link')
    for id, text in [(50, 'First duplicate'), (51, 'Second duplicate')]:
        p = doc.add_paragraph(); mark(p._p, id, 'Duplicate'); p.add_run(text); mark(p._p, id)
    p = doc.add_paragraph(); mark(p._p, 52, '任意の名前 with spaces'); p.add_run('Unicode bookmark'); mark(p._p, 52)
    p = doc.add_paragraph(); mark(p._p, 53, ''); mark(p._p, 53); p.add_run('Unnamed bookmark')
    p = doc.add_paragraph(); link(p, 'Duplicate link', 'Duplicate'); p.add_run(' / '); link(p, 'Unicode link', '任意の名前 with spaces')
    p = doc.add_paragraph(); mark(p._p, 62, 'IntoTable'); p.add_run('Before the table')
    table = doc.add_table(rows=3, cols=2, style='Table Grid'); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p = table.cell(0,0).paragraphs[0]; mark(p._p, 60, 'HeaderTarget'); p.add_run('Header target'); mark(p._p, 60)
    link(table.cell(0,1).paragraphs[0], 'Header link', 'Target')
    p = table.cell(1,0).paragraphs[0]; mark(p._p, 61, 'CellTarget'); p.add_run('Cell target'); mark(p._p, 61); mark(p._p, 62)
    table.cell(1,1).text = 'Long table content. ' * 100
    table.cell(2,0).text = 'Last row'; table.cell(2,1).text = 'Last cell'
    p = doc.add_paragraph('Destination at document end'); mark(p._p, 70, 'EndTarget'); mark(p._p, 70)
    p = doc.add_paragraph(); link(p, 'End link', 'EndTarget')
    doc.save(SOURCE); print(f'Generated independent bookmark fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file); starts = {}; ends = set(); text = ''; spans = []
    # Traverse actual XML document order, including hyperlinks and table cells.
    for node in doc.element.iter():
        if node.tag == qn('w:t'): text += node.text or ''
        elif node.tag == qn('w:bookmarkStart'):
            id = int(node.get(qn('w:id'))); assert id not in starts; starts[id] = (node.get(qn('w:name')), len(text))
        elif node.tag == qn('w:bookmarkEnd'):
            id = int(node.get(qn('w:id'))); assert id in starts and id not in ends; ends.add(id); name, start = starts[id]; spans.append((name, text[start:]))
    assert len(starts) == len(ends)
    links = doc.element.findall('.//' + qn('w:hyperlink'))
    labels = [(''.join(t.text or '' for t in n.findall('.//' + qn('w:t'))), n) for n in links]
    if fresh:
        assert len(starts) == 2
        assert ('FreshTarget', 'Destination a\u0301😀') in spans and ('NewPoint', '') in spans
        jump = next(n for label,n in labels if label == 'Jump'); assert jump.get(qn('w:anchor')) == 'FreshTarget' and jump.get(qn('r:id')) is None
    else:
        assert len(starts) == 14
        assert ('RenamedTarget', 'edited target a\u0301😀') in spans and ('Created', 'Bookmarks') in spans
        jump = next(n for label,n in labels if label == 'jump middle end'); assert jump.get(qn('w:anchor')) == 'RenamedTarget' and jump.get(qn('r:id')) is None
        assert ('InsideLink','middle') in spans
        for expected in [('Across','two paragraphs with overlapping markersand their '), ('Overlap','overlapping markersand their last words'), ('HeaderTarget','Header target'), ('CellTarget','Cell target'), ('IntoTable','Before the tableHeader targetHeader linkCell target'), ('_Point',''), ('EndTarget',''), ('任意の名前 with spaces','Unicode bookmark'), ('','')]:
            assert expected in spans, expected
        assert [value for name,value in spans if name == 'Duplicate'] == ['First duplicate','Second duplicate']
        header = next(n for label,n in labels if label == 'Header link'); assert header.get(qn('w:anchor')) == 'RenamedTarget'
        missing = next(n for label,n in labels if label == 'Missing target'); assert missing.get(qn('w:anchor')) == 'NoSuch'
        top = next(n for label,n in labels if label == 'Top link'); assert top.get(qn('w:anchor')) is None and top.get(qn('r:id')) is None
        original_text = ''.join(n.text or '' for n in Document(SOURCE).element.iter(qn('w:t')))
        assert text == original_text.replace('target a\u0301😀', 'edited target a\u0301😀', 1)
        external = next(n for label,n in labels if label == 'External precedence'); assert external.get(qn('w:anchor')) == 'Target'; assert doc.part.rels[external.get(qn('r:id'))].target_ref == 'https://example.com/external'
        with ZipFile(SOURCE) as old, ZipFile(file) as new:
            for name in old.namelist():
                if name != 'word/document.xml': assert old.read(name) == new.read(name), name
    print(f'PASS {file}: independent bookmark pairing, ranges, internal targets and package preservation')
