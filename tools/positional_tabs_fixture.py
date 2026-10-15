"""Original positional-tab controls with marker images for alignment measurement."""
from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'tests/fixtures/positional-tabs';OUT.mkdir(exist_ok=True)
cases=[]
for relative in ['margin','indent']:
 for alignment in ['left','center','right']:
  for lead in [False,True]:cases.append(dict(name=f'{relative}-{alignment}-{lead}',relativeTo=relative,alignment=alignment,lead=lead,left=36,right=24,first=18,leader='none'))
for leader in ['dot','hyphen','underscore','middleDot']:
 cases.append(dict(name='leader-'+leader,relativeTo='margin',alignment='center',lead=True,left=36,right=24,first=0,leader=leader))
cases.extend([dict(name='past-center',relativeTo='margin',alignment='center',lead=True,left=0,right=0,first=0,leader='none',prefix='Long prefix '*8),dict(name='left-hanging',relativeTo='margin',alignment='left',lead=False,left=36,right=24,first=-18,leader='none')])
for left,right,first in [(0,0,0),(36,0,0),(0,24,0),(36,24,0),(36,24,-18),(36,24,36),(-18,0,0)]:
 for relative in ['margin','indent']:
  for alignment in ['left','center','right']:
   cases.append(dict(name=f'edges-{left}-{right}-{first}-{relative}-{alignment}',relativeTo=relative,alignment=alignment,lead=True,left=left,right=right,first=first,leader='none'))

def marker(run,color):
 b=BytesIO();Image.new('RGB',(4,4),color).save(b,format='PNG');run.add_picture(BytesIO(b.getvalue()),width=Pt(4),height=Pt(4))
d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.left_margin=s.right_margin=Inches(1)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.line_spacing=Pt(18);d.styles['Normal'].paragraph_format.space_after=Pt(0)
for i,c in enumerate(cases):
 label=d.add_paragraph(c['name']);label.paragraph_format.page_break_before=bool(i)
 p=d.add_paragraph();p.paragraph_format.left_indent=Pt(c['left']);p.paragraph_format.right_indent=Pt(c['right']);p.paragraph_format.first_line_indent=Pt(c['first'])
 p.paragraph_format.tab_stops.add_tab_stop(Pt(72));p.paragraph_format.tab_stops.add_tab_stop(Pt(108))
 if c['lead']:marker(p.add_run(),'#777777');p.add_run(c.get('prefix','Lead '))
 node=OxmlElement('w:ptab')
 for key in ['alignment','relativeTo','leader']:node.set(qn('w:'+key),c[key])
 p.add_run()._r.append(node);marker(p.add_run(),'#e03040');p.add_run('Tail');marker(p.add_run(),'#2050d0')
d.save(OUT/'source.docx');(OUT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n',encoding='utf-8');print(len(cases),'controls')
