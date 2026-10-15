"""Independent native controls for shaded paragraph spacing and adjacency."""
def generate(mode,next_before):
    from pathlib import Path
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.style import WD_STYLE_TYPE
    import json
    out=Path(__file__).resolve().parents[1]/'tests/fixtures/shading-spacing';out.mkdir(exist_ok=True)
    d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');n=d.styles['Normal'];n.font.name='Arial';n.font.size=Pt(12);n.paragraph_format.space_after=Pt(0);n.paragraph_format.line_spacing=1
    for name in ['Probe','Other']:s=d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);s.base_style=n
    cases=[]
    for before,after in [(0,0),(12,0),(0,18),(12,18)]:cases.append(dict(before=before,after=after,next='none',context=False,border=False,cell=False))
    for next in ['same','different','none']:
     for context in [False,True]:cases.append(dict(before=12,after=18,next=next,context=context,border=False,cell=False))
    for border,cell in [(True,False),(False,True),(True,True)]:cases.append(dict(before=12,after=18,next='none',context=False,border=border,cell=cell))
    cases += [dict(before=12,after=18,next='different',context=False,border=False,cell=False,pageBreak=True),dict(before=12,after=18,next='different',context=False,border=False,cell=False,indent=True)]
    for i,c in enumerate(cases):
     label=d.add_paragraph(f'Case {i+1}');label.paragraph_format.page_break_before=i>0
     container=d.add_table(rows=1,cols=1).cell(0,0) if c['cell'] else d
     container.add_paragraph('Before')
     p=container.add_paragraph('Shaded','Probe');p.paragraph_format.space_before=Pt(c['before']);p.paragraph_format.space_after=Pt(c['after']);pr=p._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:fill'),'00AABB');pr.insert(1,sh)
     if c['context']:pr.append(OxmlElement('w:contextualSpacing'))
     if c['border']:
      border=OxmlElement('w:pBdr')
      for side in ['top','left','bottom','right']:
       e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'8');e.set(qn('w:space'),'0');e.set(qn('w:color'),'000000');border.append(e)
      pr.insert(1,border)
     p=container.add_paragraph('Following','Probe' if c['context'] else 'Other');p.paragraph_format.space_before=Pt(next_before)
     if c.get('pageBreak'):p.paragraph_format.page_break_before=True
     if c.get('indent'):p.paragraph_format.left_indent=Pt(24)
     if c['next']!='none':
      sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:fill'),'00AABB' if c['next']=='same' else 'DDAA00');p._p.get_or_add_pPr().insert(1,sh)
     container.add_paragraph('After')
    for setting in d.settings.element.iter(qn('w:compatSetting')):
     if setting.get(qn('w:name'))=='compatibilityMode':setting.set(qn('w:val'),str(mode))
    order=['pStyle','pageBreakBefore','pBdr','shd','spacing','ind','contextualSpacing']
    for pr in d.element.iter(qn('w:pPr')):
     children=list(pr)
     for child in children:pr.remove(child)
     for child in sorted(children,key=lambda c:order.index(c.tag.split('}')[-1])):pr.append(child)
    name=f'mode{mode}-before{next_before}'
    d.save(out/(name+'.docx'));(out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')

for mode in [14,15]:
    for before in [6,30]:generate(mode,before)
