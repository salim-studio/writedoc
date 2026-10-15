from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
p=Path(__file__).resolve().parents[1]/'tests/fixtures/plain-text-controls'
d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for i,name in enumerate(['temp-prompt','persistent-prompt']):
 para=d.add_paragraph(name+': ');r=para.add_run('[Enter value]');r.font.color.rgb=RGBColor.from_string('777777')
 s=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');s.append(pr)
 rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'2050C0');rp.append(color);pr.append(rp)
 for tag,value in [('alias',name),('id',str(101+i))]:
  el=OxmlElement('w:'+tag);el.set(qn('w:val'),value);pr.append(el)
 if i==0:pr.append(OxmlElement('w:temporary'))
 pr.append(OxmlElement('w:showingPlcHdr'));pr.append(OxmlElement('w:text'))
 content=OxmlElement('w:sdtContent');s.append(content);para._p.replace(r._r,s);content.append(r._r)
d.save(p/'placeholders.docx')
