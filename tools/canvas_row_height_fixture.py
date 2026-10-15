"""Original row-height fixtures and independent checks using python-docx."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/row-heights.docx'

def height(parent, value=None, rule=None):
    node = OxmlElement('w:trHeight')
    if value is not None: node.set(qn('w:val'), str(value))
    if rule is not None: node.set(qn('w:hRule'), rule)
    parent.append(node)

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('Canvas Row Base', WD_STYLE_TYPE.TABLE); base.base_style = doc.styles['Table Grid']
    row_pr = OxmlElement('w:trPr'); height(row_pr, 960, 'atLeast'); base.element.append(row_pr)
    child = doc.styles.add_style('Canvas Row Child', WD_STYLE_TYPE.TABLE); child.base_style = base
    doc.add_paragraph('Row heights A\u0301😀')
    table = doc.add_table(rows=7, cols=2, style=child); table.autofit = False
    for col in table.columns: col.width = Inches(2.5)
    labels = ['Inherited minimum', 'Direct automatic', 'Omitted rule', 'Content grows', 'Zero minimum', 'Unit height', 'Missing value']
    for i, row in enumerate(table.rows):
        row.cells[0].text = labels[i]; row.cells[1].text = f'Neighbor {i + 1}'
        for cell in row.cells: cell.width = Inches(2.5); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    height(table.rows[1]._tr.get_or_add_trPr(), 7200, 'auto')
    height(table.rows[2]._tr.get_or_add_trPr(), 1440)
    height(table.rows[3]._tr.get_or_add_trPr(), 240, 'atLeast'); table.cell(3, 1).text = 'First\nSecond\nThird'
    height(table.rows[4]._tr.get_or_add_trPr(), 0, 'atLeast')
    height(table.rows[5]._tr.get_or_add_trPr(), '1in', 'atLeast')
    height(table.rows[6]._tr.get_or_add_trPr(), rule='atLeast')
    doc.add_paragraph('Before merged rows')
    merged = doc.add_table(rows=3, cols=2, style='Table Grid'); merged.autofit = False
    for i, row in enumerate(merged.rows):
        row.height = Pt(30 + 10 * i); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.cells[1].text = f'Merged neighbor {i + 1}'
    merged.cell(0, 0).merge(merged.cell(2, 0)).text = 'Merged owner A\u0301😀'
    merged.cell(0, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    doc.add_paragraph('Before repeated rows')
    flowing = doc.add_table(rows=3, cols=2, style='Table Grid'); flowing.autofit = True
    flowing.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    flowing.rows[0].height = Pt(36); flowing.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    flowing.cell(0, 0).text = 'Header height'; flowing.cell(0, 1).text = 'Other header'
    flowing.rows[1].height = Pt(900); flowing.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    flowing.cell(1, 0).text = 'Tall minimum'; flowing.cell(1, 1).text = 'Bottom on first fragment'
    flowing.cell(1, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    flowing.cell(2, 0).text = 'Flowing row text. ' * 60; flowing.cell(2, 1).text = 'Last neighbor'
    doc.add_paragraph('Tail'); doc.save(SOURCE); print(f'Generated independent row-height fixture: {SOURCE}')

def verify(file, mode):
    doc = Document(file)
    if mode == 'new':
        assert len(doc.tables) == 1
        table = doc.tables[0]; assert len(table.rows) == 2
        assert [cell.text for row in table.rows for cell in row.cells] == ['First A\u0301😀', 'Second', 'Third', 'Fourth']
        for row, rule, value in zip(table.rows, ['atLeast', 'auto'], ['1200', '0']):
            node = row._tr.trPr.find(qn('w:trHeight'))
            assert node.attrib == {qn('w:val'): value, qn('w:hRule'): rule}
    else:
        old = Document(SOURCE)
        assert len(doc.tables) == len(old.tables) == 3
        before = old.element.findall('.//' + qn('w:p')); after = doc.element.findall('.//' + qn('w:p'))
        def data(node): return [(n.tag, sorted(n.attrib.items()), n.text) for n in node.iter()]
        assert [data(p) for p in after] == [data(p) for p in before], 'Text or paragraph formatting changed'
        for i, (table, original) in enumerate(zip(doc.tables, old.tables)):
            assert data(table._tbl.tblPr) == data(original._tbl.tblPr)
            assert data(table._tbl.tblGrid) == data(original._tbl.tblGrid)
            assert len(table.rows) == len(original.rows)
            for j, (row, prior) in enumerate(zip(table.rows, original.rows)):
                node = row._tr.trPr.find(qn('w:trHeight')) if row._tr.trPr is not None else None
                expected = {(0, 0): ('atLeast', '1440'), (0, 1): None, (0, 2): ('auto', '0'), (1, 0): ('atLeast', '960'), (1, 1): ('atLeast', '960'), (1, 2): ('atLeast', '960'), (2, 0): ('atLeast', '840')}
                if (i, j) in expected:
                    setting = expected[i, j]
                    if setting is None: assert node is None
                    else: assert node.attrib == {qn('w:hRule'): setting[0], qn('w:val'): setting[1]}
                    props = lambda r: [data(n) for n in r._tr.trPr if n.tag != qn('w:trHeight')] if r._tr.trPr is not None else []
                    assert props(row) == props(prior)
                    cells = lambda r: [data(n) for n in r._tr if n.tag != qn('w:trPr')]
                    assert cells(row) == cells(prior), 'Cell structure or formatting changed'
                else: assert data(row._tr) == data(prior._tr)
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: row heights, inheritance declarations, exact text, formatting and unrelated package parts')
