"""Original Find/Replace fixture and independent export checks; python-docx tooling."""
import re
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/find-replace.docx'

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(6)
    base = doc.styles.add_style('Search Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal; base.font.bold = True
    child = doc.styles.add_style('Search Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    doc.add_paragraph('Find and replace sample')
    p = doc.add_paragraph(); p.add_run('nee').bold = True; p.add_run('dle').italic = True; p.add_run(' in rich text')
    doc.add_paragraph('Needle needle needles haystack')
    doc.add_paragraph('A\u0301😀 needle and A\u0301😀.', child)
    p = doc.add_paragraph(); link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), doc.part.relate_to('https://example.org/resource', RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); props = OxmlElement('w:rPr'); props.append(OxmlElement('w:b')); r.append(props)
    text = OxmlElement('w:t'); text.text = 'needle link'; r.append(text); link.append(r); p._p.append(link)
    p = doc.add_paragraph(); marker = OxmlElement('w:bookmarkStart'); marker.set(qn('w:id'), '0'); marker.set(qn('w:name'), 'CanvasNeedle'); p._p.append(marker)
    p.add_run('needle'); marker = OxmlElement('w:bookmarkEnd'); marker.set(qn('w:id'), '0'); p._p.append(marker); p.add_run(' bookmark')
    p = doc.add_paragraph('needle\tneedle\nneedle'); p.add_run().add_break(WD_BREAK.PAGE); p.add_run('needle after page break')
    doc.add_paragraph('needle numbered', 'List Number')
    doc.add_paragraph('Literal [a+b] $& . and café cafe\u0301')
    table = doc.add_table(rows=2, cols=2, style='Table Grid'); table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(0, 0).text = 'needle header'; table.cell(0, 1).text = 'Other header'
    table.cell(1, 0).text = 'needle ' + 'Flowing cell text. ' * 100
    table.cell(1, 1).text = 'Neighbor content. ' * 100
    doc.add_paragraph(' edge needle ')
    doc.save(SOURCE); print(f'Generated independent Find/Replace fixture: {SOURCE}')

def verify(file, mode):
    doc = Document(file)
    if mode == 'new':
        assert [p.text for p in doc.paragraphs] == ['$& $&', '$& boxes', 'A\u0301😀']
        assert doc.paragraphs[0].runs[0].bold is True
        assert all(r.bold is not True for p in doc.paragraphs[1:] for r in p.runs)
    else:
        before = Document(SOURCE)
        old = before.element.findall('.//' + qn('w:p')); new = doc.element.findall('.//' + qn('w:p'))
        assert len(old) == len(new)
        def text(node):
            return ''.join(n.text or '' if n.tag == qn('w:t') else '\t' if n.tag == qn('w:tab') else '\f' if n.get(qn('w:type')) == 'page' else '\n' for n in node.iter() if n.tag in [qn('w:t'), qn('w:tab'), qn('w:br')])
        def data(node): return None if node is None else [(n.tag, sorted(n.attrib.items()), n.text) for n in node.iter()]
        for a, b in zip(old, new):
            assert text(b) == re.sub(r'\bneedle\b', 'thread', text(a), flags=re.IGNORECASE)
            assert data(a.find(qn('w:pPr'))) == data(b.find(qn('w:pPr')))
            for tag in ['bookmarkStart', 'bookmarkEnd', 'br', 'tab']:
                assert [data(n) for n in a.iter(qn('w:' + tag))] == [data(n) for n in b.iter(qn('w:' + tag))]
        assert doc.paragraphs[1].runs[0].text == 'thread' and doc.paragraphs[1].runs[0].bold is True
        assert doc.paragraphs[1].runs[0].italic is not True
        assert doc.paragraphs[3].style.name == 'Search Child' and doc.styles['Search Base'].font.bold is True
        link = doc.element.find('.//' + qn('w:hyperlink')); assert text(link) == 'thread link'
        assert link.get(qn('r:id')) == before.element.find('.//' + qn('w:hyperlink')).get(qn('r:id'))
        for tag in ['tblPr', 'tblGrid', 'trPr', 'tcPr']:
            assert [data(n) for n in before.element.iter(qn('w:' + tag))] == [data(n) for n in doc.element.iter(qn('w:' + tag))]
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: literal replacement, formatting, inherited styles, hyperlinks, bookmarks, breaks, tables and unrelated parts')
