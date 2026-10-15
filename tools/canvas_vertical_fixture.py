"""Independent superscript/subscript fixture and export checks using python-docx."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/script-text.docx'

def prop(parent, name, value):
    node = OxmlElement('w:' + name); node.set(qn('w:val'), value); parent.append(node)
    return node

def align(parent, value):
    return prop(parent, 'vertAlign', value)

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = section.page_height = Inches(6)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    defaults = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    align(defaults, 'baseline')
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    base = doc.styles.add_style('ScriptParagraph', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    base.font.superscript = True
    child = doc.styles.add_style('ScriptChild', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    align(child.element.get_or_add_rPr(), 'baseline')
    char = doc.styles.add_style('ScriptCharacter', WD_STYLE_TYPE.CHARACTER); char.font.subscript = True
    grand = doc.styles.add_style('ScriptCharacterChild', WD_STYLE_TYPE.CHARACTER); grand.base_style = char
    table_style = doc.styles.add_style('ScriptTable', WD_STYLE_TYPE.TABLE); table_style.base_style = doc.styles['Table Grid']; table_style.font.superscript = True
    doc.add_paragraph('Script text A\u0301😀')
    p = doc.add_paragraph('H'); p.add_run('2').font.subscript = True; p.add_run('O and E = mc'); p.add_run('2').font.superscript = True
    for text, kind in [('Raised direct', 'superscript'), ('Lowered direct', 'subscript')]:
        run = doc.add_paragraph(text).runs[0]; run.font.size = Pt(16); run.italic = True; run.underline = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW; align(run._r.get_or_add_rPr(), kind)
    doc.add_paragraph('Paragraph raised', style=base)
    doc.add_paragraph('Character lowered', style=base).runs[0].style = grand
    align(doc.add_paragraph('Explicit baseline', style=base).runs[0]._r.get_or_add_rPr(), 'baseline')
    doc.add_paragraph('Child baseline', style=child)
    doc.add_paragraph('Styled character').runs[0].style = char
    p = doc.add_paragraph(); mark = OxmlElement('w:rPr'); prop(mark, 'sz', '40'); align(mark, 'superscript'); p._p.get_or_add_pPr().append(mark)
    run = doc.add_paragraph('Tab\tTail').runs[0]; run.font.subscript = True; run.underline = True; run.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
    p = doc.add_paragraph('Exact raised'); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; p.paragraph_format.line_spacing = Pt(6)
    p.runs[0].font.size = Pt(32); p.runs[0].font.superscript = True
    p = doc.add_paragraph('Script number'); numpr = p._p.get_or_add_pPr().get_or_add_numPr(); numpr.get_or_add_ilvl().val = 0; numpr.get_or_add_numId().val = 7
    root = doc.part.numbering_part.element
    instance = next(n for n in root.findall(qn('w:num')) if n.get(qn('w:numId')) == '7')
    aid = instance.find(qn('w:abstractNumId')).get(qn('w:val'))
    abstract = next(n for n in root.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == aid)
    level = abstract.find(qn('w:lvl')); rpr = OxmlElement('w:rPr'); align(rpr, 'superscript'); level.append(rpr)
    table = doc.add_table(rows=3, cols=2, style=table_style); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).text = 'Repeated script header'; table.cell(0, 1).text = 'Header baseline'
    align(table.cell(0, 1).paragraphs[0].runs[0]._r.get_or_add_rPr(), 'baseline')
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Autofitlongword'; table.cell(1, 1).text = 'Script body across pages. ' * 100
    table.cell(2, 0).text = 'Table inherited'; table.cell(2, 1).text = 'Tail cell'
    doc.add_paragraph('End script sample')
    doc.save(SOURCE); print(f'Generated independent script fixture: {SOURCE}')

def declared(parent):
    node = None if parent is None else parent.find(qn('w:vertAlign'))
    return None if node is None else node.get(qn('w:val'))

def verify(file, fresh):
    doc = Document(file)
    if fresh:
        assert [p.text for p in doc.paragraphs] == ['H2O A\u0301😀', 'E = mc2']
        runs = [r for p in doc.paragraphs for r in p.runs]
        assert [(r.text, declared(r._r.rPr)) for r in runs if declared(r._r.rPr) not in [None, 'baseline']] == [('2', 'subscript'), ('2', 'superscript')]
        assert all(r.font.size in [None, Pt(12)] for r in runs), 'Display transform changed saved font size'
    else:
        old = Document(SOURCE)
        assert declared(doc.paragraphs[2].runs[0]._r.rPr) == 'subscript'
        assert doc.paragraphs[2].runs[0].font.size == Pt(16)
        assert declared(doc.styles['ScriptParagraph'].element.rPr) == 'subscript'
        assert declared(doc.styles['ScriptCharacter'].element.rPr) == 'baseline'
        # Compare every XML node/attribute except the three intended values and
        # equivalent xml:space declarations introduced by text reconstruction.
        def signature(root, changed):
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n in changed and k == qn('w:val')) and not (n.tag == qn('w:t') and k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in root.iter()]
        def edited(d):
            return [d.paragraphs[2].runs[0]._r.rPr.find(qn('w:vertAlign'))]
        assert signature(old.element, edited(old)) == signature(doc.element, edited(doc)), 'Unrelated document content changed'
        def style_edits(d):
            return [d.styles[s].element.rPr.find(qn('w:vertAlign')) for s in ['ScriptParagraph', 'ScriptCharacter']]
        assert signature(old.styles.element, style_edits(old)) == signature(doc.styles.element, style_edits(doc)), 'Unrelated styles changed'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: script declarations, stored sizes, text and unrelated package content')
