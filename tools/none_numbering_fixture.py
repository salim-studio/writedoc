"""Independent OOXML literal/placeholder controls for numFmt=none."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
root=Path(__file__).resolve().parents[1]/'tests/fixtures/none-numbering'
root.mkdir(parents=True,exist_ok=True)
doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(11)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
def node(name,**attrs):
 n=OxmlElement('w:'+name)
 for k,v in attrs.items():n.set(qn('w:'+k),str(v))
 return n
numbering=doc.part.numbering_part.element
cases=[]
def definition(levels):
 aid=100+len(cases);num=node('num',numId=aid);num.append(node('abstractNumId',val=aid))
 abstract=node('abstractNum',abstractNumId=aid);abstract.append(node('nsid',val=f'{aid:08X}'));abstract.append(node('multiLevelType',val='multilevel'))
 for i,(fmt,text,legal) in enumerate(levels):
  level=node('lvl',ilvl=i)
  for n in [node('start',val=3),node('numFmt',val=fmt),node('suff',val='tab'),node('lvlText',val=text),node('lvlJc',val='left')]:level.append(n)
  if legal:level.insert(2,node('isLgl'))
  ppr=node('pPr');tabs=node('tabs');tabs.append(node('tab',val='num',pos=1440));ppr.append(tabs);ppr.append(node('ind',left=1440,hanging=1080));level.append(ppr)
  rpr=node('rPr');rpr.append(node('rFonts',ascii='Consolas',hAnsi='Consolas'));rpr.append(node('color',val='0070C0'));level.append(rpr)
  abstract.append(level)
 first_num=next((i for i,n in enumerate(numbering) if n.tag==qn('w:num')),len(numbering))
 numbering.insert(first_num,abstract);numbering.append(num);return aid
def paragraph(aid,index,label):
 p=doc.add_paragraph(label);p.paragraph_format.space_after=Pt(5)
 numpr=node('numPr');numpr.append(node('ilvl',val=index));numpr.append(node('numId',val=aid));p._p.get_or_add_pPr().insert(0,numpr)
 return p
doc.add_paragraph('Literal text with no numbering format')
for fmt,text,legal in [('none','',False),('none','>',False),('none','%1',False),('none','[%1]',False),('none','A%1B%2C',False),('decimal','[%1]',False),('none','[%1]',True)]:
 aid=definition([(fmt,text,legal)]);label=f'C{len(cases):02} {fmt} {text or "empty"} legal={legal}'
 cases.append({'label':label,'levels':[(fmt,text,legal)]});paragraph(aid,0,label)
for legal in [False,True]:
 levels=[('decimal','%1.',False),('none','>%1:%2',legal),('decimal','%1/%2/%3',legal)]
 aid=definition(levels);label=f'C{len(cases):02} nested legal={legal}';cases.append({'label':label,'levels':levels})
 for level in [0,1,1,2,0]:paragraph(aid,level,label+f' level={level}')
doc.add_page_break()
for fmt in ['none','decimal']:
 for defined in [False,True]:
  levels=[(fmt,'A%1B%2C',False)]+([('decimal','%2',False)] if defined else [])
  aid=definition(levels);label=f'C{len(cases):02} {fmt} future defined={defined}'
  cases.append({'label':label,'levels':levels});paragraph(aid,0,label)
levels=[('upperRoman','%1.',False),('lowerLetter','%1/%2/%3',False),('decimal','%1.%2.%3',True)]
aid=definition(levels);label=f'C{len(cases):02} nested future reference';cases.append({'label':label,'levels':levels})
for level in [0,1,2]:paragraph(aid,level,label+f' level={level}')
for text in ['%1','%2','A%1B']:
 levels=[('bullet',text,False)]
 aid=definition(levels);label=f'C{len(cases):02} bullet literal {text}';cases.append({'label':label,'levels':levels});paragraph(aid,0,label)
for legal in [False,True]:
 levels=[('bullet','*',False),('decimal','%1/%2',legal)]
 aid=definition(levels);label=f'C{len(cases):02} bullet parent legal={legal}';cases.append({'label':label,'levels':levels})
 for level in [0,1]:paragraph(aid,level,label+f' level={level}')
doc.add_page_break()
style=doc.styles.add_style('SuppressedIndent',1)
style.base_style=doc.styles['Normal']
style.paragraph_format.left_indent=Pt(36)
style.paragraph_format.first_line_indent=Pt(12)
for mode in ['direct','style','both','tabs']:
 levels=[('decimal','A%1B%2C',False)]
 aid=definition(levels);label=f'C{len(cases):02} suppressed geometry {mode}'
 cases.append({'label':label,'levels':levels});p=paragraph(aid,0,label)
 if mode in ['style','both']:p.style=style
 if mode in ['direct','both']:
  p.paragraph_format.left_indent=Pt(54)
  p.paragraph_format.first_line_indent=Pt(18)
 if mode=='tabs':
  p.paragraph_format.tab_stops.add_tab_stop(Pt(108))
for left,first in [(54,0),(54,-18),(0,18),(36,36),(36,12),(0,36),(18,0),(18,18),(0,-18),(0,12),(0,24),(0,35),(18,6)]:
 levels=[('decimal','A%1B%2C',False)]
 aid=definition(levels);label=f'C{len(cases):02} suppressed left={left} first={first}'
 cases.append({'label':label,'levels':levels});p=paragraph(aid,0,label)
 p.paragraph_format.left_indent=Pt(left);p.paragraph_format.first_line_indent=Pt(first)
for level_left,level_hang in [(1800,900),(1440,1440),(1440,360)]:
 levels=[('decimal','A%1B%2C',False)]
 aid=definition(levels);label=f'C{len(cases):02} changed level {level_left}/{level_hang}'
 cases.append({'label':label,'levels':levels});p=paragraph(aid,0,label)
 p.paragraph_format.left_indent=Pt(0);p.paragraph_format.first_line_indent=Pt(12)
 abstract=next(n for n in numbering if n.tag==qn('w:abstractNum') and n.get(qn('w:abstractNumId'))==str(aid))
 ind=abstract.find('.//'+qn('w:ind'));ind.set(qn('w:left'),str(level_left));ind.set(qn('w:hanging'),str(level_hang))
for tab in [9,18,24,108,15]:
 levels=[('decimal','A%1B%2C',False)]
 aid=definition(levels);label=f'C{len(cases):02} suppressed custom tab={tab}'
 cases.append({'label':label,'levels':levels});p=paragraph(aid,0,label)
 p.paragraph_format.left_indent=Pt(0);p.paragraph_format.first_line_indent=Pt(12)
 p.paragraph_format.tab_stops.add_tab_stop(Pt(tab))
for suffix in ['space','nothing']:
 levels=[('decimal','A%1B%2C',False)]
 aid=definition(levels);label=f'C{len(cases):02} suppressed suffix={suffix}'
 cases.append({'label':label,'levels':levels});p=paragraph(aid,0,label)
 p.paragraph_format.left_indent=Pt(0);p.paragraph_format.first_line_indent=Pt(12)
 abstract=next(n for n in numbering if n.tag==qn('w:abstractNum') and n.get(qn('w:abstractNumId'))==str(aid))
 abstract.find('.//'+qn('w:suff')).set(qn('w:val'),suffix)
doc.save(root/'control.docx');(root/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
