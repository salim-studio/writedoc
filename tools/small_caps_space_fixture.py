from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from pathlib import Path
import json
ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures/small-caps-spaces'
ROOT.mkdir(parents=True,exist_ok=True)
doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(8)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
cases=[]
def add(text, family='Constantia', size=21, **options):
 cases.append(dict(id='S%02d'%len(cases),text=text,family=family,size=size,**options))
for text in ['a A','A a','a a','A A','a  A',' a ','a! A','a. A','a9 A','a\u00a0A','a\u2002A','a\u2003A','a\u2009A','a\tA']:
 add(text)
for text in ['Alpha Beta','Alpha BETA','ALPHA beta','Alpha  Beta','Alpha! Beta']:
 add(text,'Calibri Light',42,bold=True,center=True,newPage=text=='Alpha Beta')
for text in ['a A','A a','a a','a  A','a\u00a0A']:
 add(text,'Arial',11,newPage=text=='a A')
for parts in [['a',' A'],['a ','A'],['a',' ','A'],['a','\u00a0A'],['a ',' A']]:
 add(''.join(parts),parts=parts)
add('a A',parts=['a',' A'],colorChange=True)
add('a A',parts=['a',' A'],italicChange=True)
add('a A',parts=['a',' A'],sizeChange=True)
add('a A',parts=['a',' A'],fontChange=True)
add('a A',parts=['a',' A'],capsChange=True)
add('alpha beta gamma delta epsilon zeta',width=100,newPage=True)
add('alpha beta gamma delta epsilon zeta',width=100,center=True)
add('alpha\n beta',center=True)
for c in cases:
 label=doc.add_paragraph(c['id']);label.paragraph_format.page_break_before=c.get('newPage',False)
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
 if c.get('center'):p.alignment=1
 if c.get('width'):p.paragraph_format.right_indent=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin-Pt(c['width'])
 for i,text in enumerate(c.get('parts',[c['text']])):
  r=p.add_run(text);r.font.name='Arial' if i and c.get('fontChange') else c['family'];r.font.size=Pt(30 if i and c.get('sizeChange') else c['size']);r.font.small_caps=not(i and c.get('capsChange'));r.font.bold=c.get('bold',False)
  if i and c.get('colorChange'):r.font.color.rgb=RGBColor(180,20,20)
  if i and c.get('italicChange'):r.font.italic=True
doc.save(ROOT/'control.docx')
(ROOT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
