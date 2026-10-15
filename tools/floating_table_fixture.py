"""Original floating-table positioning controls, with no corpus content."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'tests/fixtures/floating-tables'
OUT.mkdir(parents=True, exist_ok=True)
CASES = [
    ('Inline baseline', None),
    ('Page offsets', dict(horzAnchor='page', vertAnchor='page', tblpX='2880', tblpY='2880')),
    ('Margin offsets', dict(horzAnchor='margin', vertAnchor='margin', tblpX='1440', tblpY='1440')),
    ('Paragraph anchor', dict(horzAnchor='margin', vertAnchor='text', tblpX='0', tblpY='0')),
    ('Negative paragraph offset', dict(horzAnchor='margin', vertAnchor='text', tblpX='0', tblpY='-240')),
    ('Paragraph wrap distances', dict(horzAnchor='margin', vertAnchor='text', tblpX='0', tblpY='0', rightFromText='360', bottomFromText='240')),
    ('Page centered', dict(horzAnchor='page', vertAnchor='page', tblpXSpec='center', tblpYSpec='center', tblpX='0', tblpY='0')),
    ('Margin right bottom', dict(horzAnchor='margin', vertAnchor='margin', tblpXSpec='right', tblpYSpec='bottom')),
    ('Default anchors', dict(tblpX='2880', tblpY='2880')),
    ('Paragraph vertical alignment ignored', dict(horzAnchor='margin', vertAnchor='text', tblpX='0', tblpY='240', tblpYSpec='center')),
    ('Margin centered', dict(horzAnchor='margin', vertAnchor='margin', tblpXSpec='center', tblpYSpec='center')),
    ('Text horizontal anchor', dict(horzAnchor='text', vertAnchor='text', tblpX='720', tblpY='0')),
    ('Explicit margin default comparison', dict(horzAnchor='margin', vertAnchor='margin', tblpX='2880', tblpY='2880')),
    ('Omitted horizontal anchor only', dict(vertAnchor='page', tblpX='2880', tblpY='2880')),
    ('Omitted vertical anchor only', dict(horzAnchor='page', tblpX='2880', tblpY='2880')),
    ('Paragraph positive offset without alignment', dict(horzAnchor='margin', vertAnchor='text', tblpX='0', tblpY='240')),
    ('Narrow page centered', dict(horzAnchor='page', vertAnchor='page', tblpXSpec='center', tblpY='2880')),
    ('Wide margin right', dict(horzAnchor='margin', vertAnchor='page', tblpXSpec='right', tblpY='2880')),
    ('Text alignment larger font', dict(horzAnchor='margin', vertAnchor='text', tblpX='0', tblpY='480', tblpYSpec='center')),
    ('Small wrap distance', dict(horzAnchor='margin', vertAnchor='text', tblpX='0', tblpY='0', rightFromText='40')),
]

def element(parent, name, **attrs):
    node = OxmlElement('w:' + name)
    for key, value in attrs.items(): node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node

doc = Document()
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
for setting in doc.settings.element.iter(qn('w:compatSetting')):
    if setting.get(qn('w:name')) == 'compatibilityMode': setting.set(qn('w:val'), '15')
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
style = doc.styles['Normal']
style.font.name, style.font.size = 'Arial', Pt(12)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.line_spacing = Pt(15)
for index, (name, position) in enumerate(CASES):
    heading = doc.add_paragraph(name)
    heading.paragraph_format.page_break_before = bool(index)
    doc.add_paragraph('Before table. This paragraph precedes its logical position.')
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    control_width = {'Narrow page centered': 72, 'Wide margin right': 288}.get(name, 144)
    table.columns[0].width = Pt(control_width)
    table.cell(0, 0).width = Pt(control_width)
    pr = table._tbl.tblPr
    width = pr.find(qn('w:tblW')); width.set(qn('w:type'), 'dxa'); width.set(qn('w:w'), str(control_width * 20))
    if position is not None:
        node = OxmlElement('w:tblpPr')
        for key, value in position.items(): node.set(qn('w:' + key), value)
        pr.insert(0, node)
    margins = element(pr, 'tblCellMar')
    pr.remove(margins)
    pr.insert(list(pr).index(pr.find(qn('w:tblLook'))), margins)
    for side in ['top', 'left', 'bottom', 'right']: element(margins, side, w=0, type='dxa')
    element(table.rows[0]._tr.get_or_add_trPr(), 'trHeight', val=1440, hRule='exact')
    element(table.cell(0, 0)._tc.get_or_add_tcPr(), 'shd', val='clear', fill='E02030')
    table.cell(0, 0).text = 'Table control'
    anchor = doc.add_paragraph('Anchor paragraph. ' + 'The text should flow around the floating table and continue below it. ' * 8)
    if name == 'Text horizontal anchor': anchor.paragraph_format.left_indent = Pt(24)
    if name == 'Text alignment larger font':
        for run in anchor.runs: run.font.size = Pt(18)
    doc.add_paragraph('Following paragraph. ' + 'This second paragraph checks whether the table occupies normal flow. ' * 3)
target = OUT / 'floating-tables.docx'
temporary = OUT / 'generated.docx'
doc.save(temporary)
def parts(path):
    if not path.exists(): return {}
    with ZipFile(path) as z: return {name: z.read(name) for name in z.namelist()}
if parts(target) != parts(temporary): target.write_bytes(temporary.read_bytes())
temporary.unlink()
(OUT / 'cases.json').write_text(json.dumps([{'name': name, 'position': position} for name, position in CASES], indent=2) + '\n')
