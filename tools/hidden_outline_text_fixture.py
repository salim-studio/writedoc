"""Independent text-frame controls for explicit and inherited invisible line widths."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
import json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import canvas_group_fixture as g
out=Path(__file__).resolve().parents[1]/'tests/fixtures/hidden-outline-text';out.mkdir(exist_ok=True)
d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12);normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
cases=[]
for kind,anchor,height in [('rect','t',100),('rect','ctr',100),('rect','b',100),('ellipse','ctr',76.5),('rect','ctr',54),('rect','b',54)]:
 for mode,width in [('absent',0),('zero',0),('direct',2.25),('wide',6),('theme',2.25),('override',0)]:
  i=len(cases)+1;cases.append(dict(case=i,kind=kind,anchor=anchor,height=height,mode=mode,width=width));p=d.add_paragraph(f'Case {i}: {kind}, {anchor}, {mode}');p.paragraph_format.page_break_before=i>1
  p=d.add_paragraph();frame=g.add(g.add(p.add_run()._r,'w:drawing'),'wp:inline',distT=0,distB=0,distL=0,distR=0);w,h=int(100*12700),int(height*12700);g.add(frame,'wp:extent',cx=w,cy=h);g.add(frame,'wp:docPr',id=i,name=f'Text frame {i}');g.add(frame,'wp:cNvGraphicFramePr');data=g.add(g.add(frame,'a:graphic'),'a:graphicData',uri=g.NS['wps']);s=g.shape(data,i+100,w,h,0,0,'DDEEFF','&');s.find('.//'+g.tag('a:prstGeom')).set('prst',kind)
  body=s.find(g.tag('wps:bodyPr'));body.set('anchor',anchor)
  for key in ['lIns','rIns','tIns','bIns']:body.set(key,'0')
  rp=s.find('.//'+g.tag('w:rPr'));rp.find(g.tag('w:sz')).set(qn('w:val'),'96');g.add(rp,'w:rFonts',**{'w:ascii':'Arial','w:hAnsi':'Arial'})
  line=s.find('.//'+g.tag('a:ln'))
  if mode not in ['absent','theme']:line.set('w',str(int(width*12700)))
  if mode in ['theme','override']:
   style=g.add(s,'wps:style');s.remove(style);s.insert(3,style)
   for name,idx in [('lnRef',2),('fillRef',0),('effectRef',0)]:g.add(g.add(style,'a:'+name,idx=idx),'a:srgbClr',val='000000')
   g.add(g.add(style,'a:fontRef',idx='none'),'a:srgbClr',val='123456')
d.save(out/'control.docx')
with ZipFile(out/'control.docx') as z:parts={n:z.read(n) for n in z.namelist()}
theme=E.fromstring(parts['word/theme/theme1.xml']);theme.find('.//'+g.tag('a:lnStyleLst'))[1].set('w',str(int(2.25*12700)));parts['word/theme/theme1.xml']=E.tostring(theme,xml_declaration=True,encoding='utf-8',standalone=True)
with ZipFile(out/'control.docx','w',ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
(out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
