"""Original small-palette CMYK calibration and independent holdout controls."""
from pathlib import Path
from itertools import product
from random import Random
from io import BytesIO
from PIL import Image,ImageDraw
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import json,argparse
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output',type=Path,default=Path(__file__).resolve().parents[1]/'tests/fixtures/cmyk-color-grid')
parser.add_argument('--step',type=int,choices=[16,32,64],default=64)
parser.add_argument('--holdout',type=int,default=128)
parser.add_argument('--batch',type=int,choices=[16,128],default=16)
args=parser.parse_args()
if not 1<=args.holdout<=65536:parser.error('--holdout must be between 1 and 65536')
out=args.output;out.mkdir(parents=True,exist_ok=True)
cols,cellw,cellh=(4,64,32) if args.batch==16 else (16,32,32)
size=(cols*cellw,args.batch//cols*cellh)
levels=list(range(0,256,args.step))+[255];rng=Random(155)
sets=[('grid',list(product(levels,repeat=4))),('holdout',[tuple(rng.randrange(256) for _ in range(4)) for _ in range(args.holdout)])]
d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');cases=[];payloads=[]
for name,colors in sets:
 for start in range(0,len(colors),args.batch):
  batch=colors[start:start+args.batch];im=Image.new('CMYK',size);draw=ImageDraw.Draw(im)
  for i,c in enumerate(batch):draw.rectangle((i%cols*cellw,i//cols*cellh,i%cols*cellw+cellw-1,i//cols*cellh+cellh-1),fill=c)
  b=BytesIO();im.save(b,'JPEG',quality=100,subsampling=0);payload=b.getvalue();decoded=Image.open(BytesIO(payload))
  assert [decoded.getpixel((i%cols*cellw+cellw//2,i//cols*cellh+cellh//2)) for i in range(len(batch))]==batch
  # Add through a distinct RGB placeholder; replace its media part afterward.
  placeholder=BytesIO();identifier=len(cases)+1;Image.new('RGB',size,(identifier%256,identifier//256%256,identifier//65536)).save(placeholder,'PNG')
  if cases:d.add_page_break()
  d.add_paragraph(name+' '+str(start));d.add_picture(BytesIO(placeholder.getvalue()),width=Pt(288),height=Pt(144))
  cases.append({'set':name,'start':start,'colors':batch});payloads.append(payload)
buf=BytesIO();d.save(buf)
from zipfile import ZipFile,ZIP_DEFLATED,ZipInfo
with ZipFile(BytesIO(buf.getvalue())) as z:parts={n:z.read(n) for n in z.namelist()}
# Preserve PNG filenames but fix their declared MIME; encoded bytes are JPEG.
parts['[Content_Types].xml']=parts['[Content_Types].xml'].replace(b'Extension="png" ContentType="image/png"',b'Extension="png" ContentType="image/jpeg"')
for i in range(len(cases)):parts[f'word/media/image{i+1}.png']=payloads[i]
with ZipFile(out/'control.docx','w',ZIP_DEFLATED) as z:
 for name,data in parts.items():
  info=ZipInfo(name,(2026,9,27,0,0,0));info.compress_type=ZIP_DEFLATED;z.writestr(info,data)
(out/'cases.json').write_text(json.dumps(cases))
print('Original palette cases',len(cases))
