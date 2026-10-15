"""Independent group text geometry fixture. Original XML; python-docx MIT/lxml BSD.

Equal physical frames with unequal child coordinate systems expose accidental
scaling of text, insets, wrapping and paragraph geometry. No user artwork.
"""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt
from canvas_group_fixture import NS, tag, add, shape

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/group-text-scaling.docx'


def group(parent, element, width, height, sx, sy, x=0, y=0):
    g = add(parent, element)
    if element == 'wpg:grpSp': add(g, 'wpg:cNvPr', id=100, name='Nested group')
    add(g, 'wpg:cNvGrpSpPr')
    t = add(add(g, 'wpg:grpSpPr'), 'a:xfrm')
    add(t, 'a:off', x=x, y=y); add(t, 'a:ext', cx=width, cy=height)
    add(t, 'a:chOff', x=25400, y=-12700)
    add(t, 'a:chExt', cx=int(width / sx), cy=int(height / sy))
    return g


def generate():
    doc = Document()
    doc.settings.element.find(tag('w:zoom')).set(tag('w:percent'), '100')
    doc.sections[0].top_margin = doc.sections[0].bottom_margin = Inches(.5)
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)
    doc.add_paragraph('Group text scaling')
    cases = [('Identity', 1, 1, False, 't'), ('Double', 2, 2, False, 't'),
             ('Nonuniform', 4, .5, False, 'ctr'), ('Nested', 2, 2, True, 'b')]
    for ident, (label, sx, sy, nested, anchor) in enumerate(cases, 1):
        doc.add_paragraph(label)
        p = doc.add_paragraph()
        drawing = E.SubElement(p.add_run()._r, tag('w:drawing'), nsmap={k: NS[k] for k in ['wpg', 'wps', 'a']})
        inline = add(drawing, 'wp:inline')
        width, height = 2743200, 914400
        add(inline, 'wp:extent', cx=width, cy=height)
        add(inline, 'wp:docPr', id=ident, name=label)
        data = add(add(inline, 'a:graphic'), 'a:graphicData', uri=NS['wpg'])
        g = group(data, 'wpg:wgp', width, height, sx, sy)
        sw, sh = int(width / sx), int(height / sy)
        if nested:
            g = group(g, 'wpg:grpSp', sw, sh, 2, .5, 25400, -12700)
            sw, sh = sw // 2, sh * 2
        s = shape(g, 200 + ident, sw, sh, 25400, -12700, 'C2CB93', 'Fixed point size wraps inside the physical frame.')
        s.find('.//' + tag('w:sz')).set(tag('w:val'), '24')
        body = s.find(tag('wps:bodyPr'))
        for key, value in dict(lIns=152400, rIns=76200, tIns=50800, bIns=101600, anchor=anchor).items(): body.set(key, str(value))
    doc.add_paragraph('End of group text scaling.')
    doc.save(SOURCE)


def verify(path):
    def drawings(xml):
        root = E.fromstring(xml)
        return [[(x.tag, sorted(x.attrib.items()), x.text or '') for x in n.iter()]
                for n in root.iter(tag('w:drawing'))]
    with ZipFile(SOURCE) as original, ZipFile(path) as saved:
        assert set(original.namelist()) == set(saved.namelist())
        assert drawings(original.read('word/document.xml')) == drawings(saved.read('word/document.xml'))
        for name in original.namelist():
            if name != 'word/document.xml': assert original.read(name) == saved.read(name), name
        root = E.fromstring(saved.read('word/document.xml'))
        assert ''.join(root.itertext()).startswith('Edited Group text scaling')
    print('PASS retained group text geometry:', path)


if __name__ == '__main__':
    import sys
    if len(sys.argv) == 1: generate()
    else:
        for path in sys.argv[1:]: verify(path)
