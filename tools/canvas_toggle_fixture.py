"""Independent style-toggle fixtures and verification, authored with python-docx."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]

def source(enabled):
    return ROOT / ('public/fixtures/style-toggles-' + ('on' if enabled else 'off') + '.docx')

def boolean(parent, tag, enabled):
    node = OxmlElement('w:' + tag); node.set(qn('w:val'), '1' if enabled else '0'); parent.append(node)

def generate(factory):
    for enabled in [False, True]:
        doc = factory(); section = doc.sections[0]
        section.page_width, section.page_height = Inches(6), Inches(6)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
        defaults = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
        boolean(defaults, 'b', enabled); boolean(defaults, 'i', enabled)
        normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
        normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
        for kind, prefix in [(WD_STYLE_TYPE.PARAGRAPH, 'Toggle'), (WD_STYLE_TYPE.CHARACTER, 'Char'), (WD_STYLE_TYPE.TABLE, 'Table')]:
            for name, value in [('On', True), ('Off', False)]:
                style = doc.styles.add_style(prefix + name, kind)
                if kind == WD_STYLE_TYPE.PARAGRAPH: style.base_style = normal
                if kind == WD_STYLE_TYPE.TABLE: style.base_style = doc.styles['Table Grid']
                style.font.bold = value; style.font.italic = value
        child = doc.styles.add_style('ToggleChild', WD_STYLE_TYPE.PARAGRAPH); child.base_style = doc.styles['ToggleOn']; child.font.italic = False
        child = doc.styles.add_style('CharChild', WD_STYLE_TYPE.CHARACTER); child.base_style = doc.styles['CharOn']; child.font.bold = False
        cases = [('Default text', None, None), ('Paragraph on', 'ToggleOn', None), ('Paragraph off', 'ToggleOff', None),
                 ('Nearest paragraph child', 'ToggleChild', None), ('On with character on', 'ToggleOn', 'CharOn'),
                 ('On with character off', 'ToggleOn', 'CharOff'), ('Off with character on', 'ToggleOff', 'CharOn'),
                 ('Child with character child', 'ToggleChild', 'CharChild'), ('Unicode A\u0301😀', 'ToggleOff', 'CharOff'),
                 ('Direct off', 'ToggleOn', 'CharOff')]
        for text, pstyle, rstyle in cases:
            p = doc.add_paragraph(text, style=pstyle)
            if rstyle: p.runs[0].style = doc.styles[rstyle]
            if text == 'Direct off': p.runs[0].bold = False; p.runs[0].italic = False
        p = doc.add_paragraph(style='ToggleOff'); mark = OxmlElement('w:rPr'); ref = OxmlElement('w:rStyle'); ref.set(qn('w:val'), 'CharOff'); mark.append(ref); p._p.get_or_add_pPr().append(mark)
        p = doc.add_paragraph('Numbered toggle', style='ToggleChild'); numpr = OxmlElement('w:numPr')
        for tag, value in [('ilvl', '0'), ('numId', '7')]:
            n = OxmlElement('w:' + tag); n.set(qn('w:val'), value); numpr.append(n)
        p._p.get_or_add_pPr().append(numpr)
        table = doc.add_table(rows=3, cols=2, style='TableOn'); table.autofit = False
        for col in table.columns: col.width = Inches(2.5)
        for row in table.rows:
            for cell in row.cells: cell.width = Inches(2.5)
        for row, texts, styles in [(0, ['Repeated toggle header', 'Off header'], ['ToggleOn', 'ToggleOff']),
                                  (1, ['Wrapping WIDE bold italic text. ' * 55, 'Table only'], ['ToggleOn', 'Normal']),
                                  (2, ['Off in table', 'Child in table'], ['ToggleOff', 'ToggleChild'])]:
            for cell, text, style in zip(table.rows[row].cells, texts, styles):
                cell.text = text; cell.paragraphs[0].style = style
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        table = doc.add_table(rows=1, cols=2, style='TableOff'); table.autofit = False
        for col in table.columns: col.width = Inches(2.5)
        for cell, text, style in zip(table.rows[0].cells, ['Table off only', 'On in off table'], ['Normal', 'ToggleOn']):
            cell.width = Inches(2.5); cell.text = text; cell.paragraphs[0].style = style
        doc.add_paragraph('Tail')
        doc.save(source(enabled)); print(f'Generated independent toggle fixture: {source(enabled)}')

def verify(file, enabled):
    doc, old = Document(file), Document(source(enabled))
    def flag(parent, tag):
        node = None if parent is None else parent.find(qn('w:' + tag))
        return None if node is None else node.get(qn('w:val')) not in ['0', 'false', 'off']
    assert flag(doc.styles['ToggleOn'].element.rPr, 'b') is False
    assert flag(doc.styles['CharOff'].element.rPr, 'i') is True
    assert doc.paragraphs[9].runs[0].bold is True and doc.paragraphs[9].runs[0].italic is False
    assert [p.text for p in doc.paragraphs] == [p.text for p in old.paragraphs]
    assert [[c.text for row in t.rows for c in row.cells] for t in doc.tables] == [[c.text for row in t.rows for c in row.cells] for t in old.tables]
    # Compare all declarations independently, excluding only the three edited
    # booleans. xml:space on unchanged text without edge spaces is equivalent.
    def signature(root, omit):
        return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space') and n.text == (n.text or '').strip())), n.text)
                for n in root.iter() if n not in omit]
    replacements = {s.style_id: s for s in doc.styles}
    for style in old.styles:
        replacement = replacements[style.style_id]
        tag = 'b' if style.style_id == 'ToggleOn' else 'i' if style.style_id == 'CharOff' else None
        omitted = lambda s: [s.element.rPr.find(qn('w:' + tag))] if tag else []
        assert signature(style.element, omitted(style)) == signature(replacement.element, omitted(replacement)), f'Style changed: {style.style_id}'
    old_b = old.paragraphs[9].runs[0]._r.rPr.find(qn('w:b')); new_b = doc.paragraphs[9].runs[0]._r.rPr.find(qn('w:b'))
    assert signature(old.element, [old_b]) == signature(doc.element, [new_b]), 'Unrelated document content changed'
    with ZipFile(source(enabled)) as original, ZipFile(file) as exported:
        assert set(original.namelist()) == set(exported.namelist())
        for name in original.namelist():
            if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), f'Unrelated part changed: {name}'
    print(f'PASS: {file}: explicit style resets, character declarations, direct overrides, marks, tables and unchanged package parts')
