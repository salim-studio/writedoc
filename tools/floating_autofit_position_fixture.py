from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
out=Path(__file__).resolve().parents[1]/'tests/fixtures/floating-autofit-position';out.mkdir(exist_ok=True)
d=Document();n=d.styles['Normal'];n.font.name='Arial';n.font.size=Pt(10);n.paragraph_format.space_after=Pt(0);n.paragraph_format.line_spacing=Pt(12)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for e in d.settings.element.iter(qn('w:compatSetting')):
 if e.get(qn('w:name'))=='compatibilityMode':e.set(qn('w:val'),'15')
for i,(nested,x) in enumerate([(n,x) for n in [False,True] for x in [0,4320,7200]],1):
 p=d.add_paragraph(f'Case {i}');p.paragraph_format.page_break_before=i>1
 if nested:
  outer=d.add_table(rows=1,cols=1);outer.autofit=False;outer.columns[0].width=Inches(6);c=outer.cell(0,0);c.width=Inches(6);c.paragraphs[0].text='Before';t=c.add_table(rows=3,cols=2)
 else:t=d.add_table(rows=3,cols=2)
 t.autofit=True
 for j,width in enumerate([36,132]):
  t.columns[j].width=Pt(width)
  for row in t.rows:row.cells[j].width=Pt(width)
 for j,row in enumerate(t.rows):
  row.cells[0].text=str(j+1);row.cells[1].text=['Location: Assembly room','Date: September 22','Time: 2:00 PM'][j]
 pp=OxmlElement('w:tblpPr')
 for k,v in dict(horzAnchor='margin',vertAnchor='text',tblpX=x,tblpY=240).items():pp.set(qn('w:'+k),str(v))
 t._tbl.tblPr.insert(0,pp)
 if nested:c.paragraphs[-1].text='Anchor'
 else:d.add_paragraph('Anchor')
 d.add_paragraph('After')
d.save(out/'control.docx')
