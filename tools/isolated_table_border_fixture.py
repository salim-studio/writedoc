"""Original isolated table-border controls; no corpus content."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import json
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / 'tests/fixtures/isolated-table-border'
OUT.mkdir(parents=True, exist_ok=True)
CASES = [
    dict(name='Omitted layout empty', layout=None, text='', size=36),
    dict(name='Explicit AutoFit empty', layout='autofit', text='', size=36),
    dict(name='Fixed empty', layout='fixed', text='', size=36),
    dict(name='Omitted layout text', layout=None, text='Cell', size=36),
    dict(name='Omitted layout cell margins', layout=None, text='Cell', size=36, margin=115),
    dict(name='Top edge', layout=None, text='', size=36, side='top'),
    dict(name='Thin bottom edge', layout=None, text='', size=8),
    dict(name='Thick bottom edge', layout=None, text='', size=64),
    dict(name='Bottom with right corner', layout=None, text='', size=36, extra='right'),
    dict(name='Bottom with left corner', layout=None, text='', size=36, extra='left'),
    dict(name='Two columns bottom edge', layout=None, text='', size=36, columns=2),
    dict(name='Internal row edge', layout=None, text='', size=36, side='insideH', rows=2),
]
def add(parent,name,**attrs):
    node=OxmlElement('w:'+name)
    for key,value in attrs.items(): node.set(qn('w:'+key),str(value))
    parent.append(node)
    return node

doc=Document()
sec=doc.sections[0]
sec.page_width,sec.page_height=Inches(8.5),Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=Inches(.5)
doc.styles['Normal'].font.name='Arial'
doc.styles['Normal'].font.size=Pt(12)
doc.styles['Normal'].paragraph_format.space_after=Pt(0)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for node in doc.settings.element.iter(qn('w:compatSetting')):
    if node.get(qn('w:name'))=='compatibilityMode': node.set(qn('w:val'),'15')
for index,case in enumerate(CASES):
    p=doc.add_paragraph(case['name']);p.paragraph_format.page_break_before=bool(index)
    p.paragraph_format.line_spacing=Pt(18)
    table=doc.add_table(rows=case.get('rows',1),cols=case.get('columns',1))
    for column in table.columns: column.width=Pt(288/len(table.columns))
    for row in table.rows:
        for cell in row.cells: cell.width=Pt(288/len(table.columns))
    pr=table._tbl.tblPr
    width=pr.find(qn('w:tblW'));width.set(qn('w:type'),'dxa');width.set(qn('w:w'),'5760')
    look=pr.find(qn('w:tblLook'))
    borders=add(pr,'tblBorders');add(borders,case.get('side','bottom'),val='single',sz=case['size'],space=0,color='20AA60')
    if case.get('extra'):add(borders,case['extra'],val='single',sz=case['size'],space=0,color='20AA60')
    border_order=['top','left','bottom','right','insideH','insideV']
    for node in sorted(list(borders),key=lambda n:border_order.index(n.tag.split('}')[1])):borders.append(node)
    pr.remove(borders);pr.insert(list(pr).index(look),borders)
    if case['layout']:
        layout=add(pr,'tblLayout',type=case['layout']);pr.remove(layout);pr.insert(list(pr).index(look),layout)
    margins=add(pr,'tblCellMar')
    for side in ['top','left','bottom','right']:add(margins,side,w=case.get('margin',0) if side in ['left','right'] else 0,type='dxa')
    pr.remove(margins);pr.insert(list(pr).index(look),margins)
    for row in table.rows:
        add(row._tr.get_or_add_trPr(),'trHeight',val=259,hRule='exact')
        for cell in row.cells:
            add(cell._tc.get_or_add_tcPr(),'shd',val='clear',fill='F2F2F2')
            cell.text=case['text']
    if index==0:
        header=sec.header._element
        header.insert(0,deepcopy(table._tbl))
    doc.add_paragraph('After table')
temporary=OUT/'generated.docx';target=OUT/'isolated-table-border.docx'
doc.save(temporary)
def parts(p):
    if not p.exists():return {}
    with ZipFile(p) as z:return {n:z.read(n) for n in z.namelist()}
if parts(temporary)!=parts(target):target.write_bytes(temporary.read_bytes())
temporary.unlink()
(OUT/'cases.json').write_text(json.dumps(CASES,indent=2)+'\n')
