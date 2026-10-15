"""Independent baseline-shift fixture and export verification using python-docx."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/baseline-shifts.docx'

def prop(parent, name, value):
    node = OxmlElement('w:' + name); node.set(qn('w:val'), str(value)); parent.append(node)
    return node

def position(parent, value):
    # Insert before size/decorations according to CT_RPr's sequence.
    node = OxmlElement('w:position'); node.set(qn('w:val'), str(value))
    after = {'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang'}
    anchor = next((n for n in parent if n.tag.split('}')[-1] in after), None)
    if anchor is None: parent.append(node)
    else: anchor.addprevious(node)
    return node

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = section.page_height = Inches(6)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    defaults = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    position(defaults, '0')
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    base = doc.styles.add_style('ShiftParagraph', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal; position(base.element.get_or_add_rPr(), 8)
    child = doc.styles.add_style('ShiftChild', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base; position(child.element.get_or_add_rPr(), 0)
    char = doc.styles.add_style('ShiftCharacter', WD_STYLE_TYPE.CHARACTER); position(char.element.get_or_add_rPr(), -4)
    table_style = doc.styles.add_style('ShiftTable', WD_STYLE_TYPE.TABLE); table_style.base_style = doc.styles['Table Grid']; position(table_style.element.get_or_add_rPr(), 8)
    doc.add_paragraph('Baseline shifts A\u0301😀')
    p = doc.add_paragraph('Base ')
    r = p.add_run('Raised'); r.underline = True; r.font.highlight_color = WD_COLOR_INDEX.YELLOW; position(r._r.get_or_add_rPr(), 12)
    p.add_run(' plain ')
    r = p.add_run('Lowered'); r.underline = True; r.font.highlight_color = WD_COLOR_INDEX.TURQUOISE; position(r._r.get_or_add_rPr(), -8)
    p.add_run(' end')
    doc.add_paragraph('Paragraph raised', style=base)
    doc.add_paragraph('Character lowered', style=base).runs[0].style = char
    position(doc.add_paragraph('Explicit zero', style=base).runs[0]._r.get_or_add_rPr(), 0)
    doc.add_paragraph('Child reset', style=child)
    for text, sub, shift in [('Raised super', False, 6), ('Lowered sub', True, -6)]:
        r = doc.add_paragraph(text).runs[0]; r.font.size = Pt(16); r.italic = True; r.underline = True
        if sub: r.font.subscript = True
        else: r.font.superscript = True
        position(r._r.get_or_add_rPr(), shift)
    position(doc.add_paragraph('Universal units').runs[0]._r.get_or_add_rPr(), '6pt')
    r = doc.add_paragraph('Shift\tTab').runs[0]; r.underline = True; r.font.highlight_color = WD_COLOR_INDEX.YELLOW; position(r._r.get_or_add_rPr(), 10)
    p = doc.add_paragraph('Exact lowered'); p.paragraph_format.line_spacing = Pt(6); r = p.runs[0]; r.font.size = Pt(32); position(r._r.get_or_add_rPr(), -60)
    p = doc.add_paragraph(); rpr = OxmlElement('w:rPr'); position(rpr, 10); p._p.get_or_add_pPr().append(rpr)
    p = doc.add_paragraph('Shifted number'); numpr = p._p.get_or_add_pPr().get_or_add_numPr(); numpr.get_or_add_ilvl().val = 0; numpr.get_or_add_numId().val = 7
    root = doc.part.numbering_part.element; instance = next(n for n in root.findall(qn('w:num')) if n.get(qn('w:numId')) == '7'); aid = instance.find(qn('w:abstractNumId')).get(qn('w:val'))
    abstract = next(n for n in root.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == aid)
    level = abstract.find(qn('w:lvl')); rpr = OxmlElement('w:rPr'); position(rpr, 4); level.append(rpr)
    table = doc.add_table(rows=3, cols=2, style=table_style); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).text = 'Shifted header'; table.cell(0, 1).text = 'Zero header'; position(table.cell(0, 1).paragraphs[0].runs[0]._r.get_or_add_rPr(), 0)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Autofitlongword'; table.cell(1, 1).text = 'Shifted body across pages. ' * 70
    table.cell(2, 0).text = 'Table inherited'; table.cell(2, 1).text = 'Tail cell'
    doc.add_paragraph('End position sample')
    unit_style = doc.styles.add_style('UniversalShift', WD_STYLE_TYPE.PARAGRAPH); unit_style.base_style = normal; position(unit_style.element.get_or_add_rPr(), '9.525mm')
    doc.add_paragraph('Universal style', style=unit_style)
    doc.save(SOURCE); print(f'Generated independent baseline-shift fixture: {SOURCE}')

def declared(parent):
    node = None if parent is None else parent.find(qn('w:position'))
    return None if node is None else node.get(qn('w:val'))

def verify(file, fresh):
    doc = Document(file)
    if fresh:
        assert [p.text for p in doc.paragraphs] == ['Raised text', 'Lowered text', 'Plain text', 'Shifted typing']
        for p, expected in zip(doc.paragraphs, ['8', '-6', '0', '5']):
            assert all(declared(r._r.rPr) == expected for r in p.runs), p.text
            assert all(r.font.size in [None, Pt(12)] for r in p.runs), 'Saved size changed'
    else:
        old = Document(SOURCE)
        assert declared(doc.paragraphs[1].runs[1]._r.rPr) == '17'
        assert declared(doc.styles['ShiftParagraph'].element.rPr) == '-12'
        assert declared(doc.styles['ShiftCharacter'].element.rPr) is None
        def signature(root, omit):
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in root.iter() if n not in omit]
        def edits(d): return [d.paragraphs[1].runs[1]._r.rPr.find(qn('w:position'))]
        assert signature(old.element, edits(old)) == signature(doc.element, edits(doc)), 'Unrelated document content changed'
        def edits(d): return [d.styles[s].element.rPr.find(qn('w:position')) for s in ['ShiftParagraph', 'ShiftCharacter']]
        assert signature(old.styles.element, edits(old)) == signature(doc.styles.element, edits(doc)), 'Unrelated style content changed'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: baseline shifts, source text, stored sizes and unchanged package content')
