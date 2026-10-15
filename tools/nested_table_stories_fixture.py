"""Original nested body table with changing first/even/default story clearance."""
from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];d=Document();s=d.sections[0]
s.top_margin=s.bottom_margin=Inches(.5);s.left_margin=s.right_margin=Inches(1);s.header_distance=s.footer_distance=Inches(.25)
s.different_first_page_header_footer=True;d.settings.odd_and_even_pages_header_footer=True
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
n=d.styles['Normal'];n.font.name='Arial';n.font.size=Pt(10);n.paragraph_format.space_after=Pt(0);n.paragraph_format.line_spacing=Pt(12)
for story,label,count in [(s.first_page_header,'First header',4),(s.even_page_header,'Even header',3),(s.header,'Default header',2)]:
 story.paragraphs[0].text='\n'.join([label]*count)
for story,label in [(s.first_page_footer,'First footer'),(s.even_page_footer,'Even footer'),(s.footer,'Default footer')]:story.paragraphs[0].text=label
b=BytesIO();Image.new('RGB',(24,6),'#1255aa').save(b,format='PNG')
for i in range(3):
 if i:d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
 d.add_paragraph(f'Body page {i+1}')
 outer=d.add_table(rows=1,cols=1);outer.autofit=False;outer.columns[0].width=Inches(5)
 outer.cell(0,0).paragraphs[0].text='Outer cell'
 inner=outer.cell(0,0).add_table(rows=2,cols=1);inner.autofit=False;inner.columns[0].width=Inches(4)
 for j,cell in enumerate(inner.column_cells(0)):
  cell.paragraphs[0].add_run().add_picture(BytesIO(b.getvalue()),width=Pt(24));cell.paragraphs[0].add_run(f' Nested {i+1}.{j+1}')
folder=root/'tests/fixtures/nested-table-stories';folder.mkdir(exist_ok=True);d.save(folder/'control.docx')
