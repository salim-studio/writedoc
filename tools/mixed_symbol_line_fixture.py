"""Original mixed-font line controls, independent of corpus XML."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1] / 'tests/fixtures/mixed-symbol-line-metrics'

def generate():
    ROOT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(8)
    normal.paragraph_format.space_before = Pt(12)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    for setting in doc.settings.element.iter(qn('w:compatSetting')):
        if setting.get(qn('w:name')) == 'compatibilityMode':
            setting.set(qn('w:val'), '15')
    cases = []
    for table in [False, True]:
        for size in [9, 11]:
            for family, glyph in [('Calibri', 'X'), ('MS Gothic', 'X'), ('MS Gothic', '\u2610')]:
                cases.append(dict(table=table, size=size, family=family, glyph=glyph))
                title = doc.add_paragraph('Case ' + str(len(cases)))
                title.paragraph_format.page_break_before = len(cases) > 1
                doc.add_paragraph('Before')
                doc.add_paragraph()
                p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0] if table else doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(12)
                mark = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(size * 2)); mark.append(sz); p._p.get_or_add_pPr().append(mark)
                r = p.add_run(glyph); r.font.name = family; r.font.size = Pt(size)
                r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), family)
                r = p.add_run(' Options'); r.font.name = 'Calibri'; r.font.size = Pt(size)
                doc.add_paragraph('After')
    pure_cases = [dict(pure=True, size=size, family=family)
                  for family in ['Calibri', 'MS Gothic', 'MS PGothic', 'MS UI Gothic']
                  for size in [6, 9, 11, 20, 48, 96, 144]]
    pure_cases += [dict(pure=True, size=size, family=family, bold=bold, italic=italic)
                   for family in ['MS Gothic', 'MS PGothic', 'MS UI Gothic']
                   for size in [9, 20] for bold,italic in [(True,False),(False,True),(True,True)]]
    for case in pure_cases:
            family, size = case['family'], case['size']
            cases.append(case)
            title = doc.add_paragraph('Case ' + str(len(cases)))
            title.paragraph_format.page_break_before = True
            for text in ['A', 'B', 'C']:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(text); r.font.name = family; r.font.size = Pt(size)
                r.bold = case.get('bold', False); r.italic = case.get('italic', False)
                mark = OxmlElement('w:rPr')
                fonts = OxmlElement('w:rFonts')
                for slot in ['ascii', 'hAnsi', 'eastAsia']:
                    fonts.set(qn('w:' + slot), family)
                mark.append(fonts)
                sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(size * 2)); mark.append(sz)
                p._p.get_or_add_pPr().append(mark)
    doc.save(ROOT / 'controls.docx')
    (ROOT / 'cases.json').write_text(json.dumps(cases, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

if __name__ == '__main__':
    generate()
