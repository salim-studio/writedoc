"""Original relative-frame drawing controls; no corpus XML."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree as E
import canvas_group_fixture as g

ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures/relative-shape-frames'
WP14='http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing'

def generate():
 ROOT.mkdir(exist_ok=True)
 cases=[dict(angle=a,effects=e) for e in [False,True] for a in [0,30,45,90,135,180,270,315]]
 cases += [dict(angle=0,text=True),dict(angle=0,text=True,effects=True),dict(angle=30,pattern=True),dict(angle=90,nested=True,effects=True)]
 cases += [dict(angle=90,nested=True,nonuniform=True),dict(angle=90,intrinsic=True),dict(angle=90,intrinsic=True,fixed=True)]
 doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(12)
 doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
 for s in doc.settings.element.iter(qn('w:compatSetting')):
  if s.get(qn('w:name'))=='compatibilityMode':s.set(qn('w:val'),'15')
 for i,case in enumerate(cases,1):
  p=doc.add_paragraph('Relative shape '+str(i));p.paragraph_format.page_break_before=i>1
  a=g.add(g.add(p.add_run()._r,'w:drawing'),'wp:anchor',distT=0,distB=0,distL=0,distR=0,simplePos=0,relativeHeight=i,behindDoc=1,locked=0,layoutInCell=1,allowOverlap=1)
  g.add(a,'wp:simplePos',x=0,y=0)
  for axis,pos in [('H',144),('V',180)]:g.add(g.add(a,'wp:position'+axis,relativeFrom='page'),'wp:posOffset',text=str(pos*12700))
  w,h=120*12700,240*12700;g.add(a,'wp:extent',cx=w,cy=h)
  if case.get('effects'):g.add(a,'wp:effectExtent',l=5*12700,t=7*12700,r=11*12700,b=13*12700)
  g.add(a,'wp:wrapNone');g.add(a,'wp:docPr',id=i,name='Independent relative group');g.add(a,'wp:cNvGraphicFramePr')
  group=g.add(g.add(g.add(a,'a:graphic'),'a:graphicData',uri=g.NS['wpg']),'wpg:wgp');g.add(group,'wpg:cNvGrpSpPr')
  gp=g.add(group,'wpg:grpSpPr');g.transform(gp,w,h,group=True)
  gp[0].find(g.tag('a:chOff')).attrib.update({'x':'0','y':'0'});gp[0].find(g.tag('a:chExt')).attrib.update({'cx':str(w),'cy':str(h)})
  if case.get('intrinsic'):
   gp[0].find(g.tag('a:ext')).set('cx',str(w//2));gp[0].find(g.tag('a:chExt')).set('cx',str(110*12700))
  if case.get('nested'):
   group=g.add(group,'wpg:grpSp');g.add(group,'wpg:cNvGrpSpPr');gp=g.add(group,'wpg:grpSpPr');g.transform(gp,w,h,group=True)
   gp[0].find(g.tag('a:chOff')).attrib.update({'x':'0','y':'0'})
   if case.get('nonuniform'):gp[0].find(g.tag('a:chExt')).set('cy',str(h))
  shape=g.shape(group,100+i,80*12700,140*12700,20*12700,30*12700,'00A060','Control text' if case.get('text') else None,custom=not case.get('text'))
  props=shape.find(g.tag('wps:spPr'));props[0].set('rot',str(case['angle']*60000))
  if case.get('text'):
   rp=shape.find('.//'+g.tag('w:rPr'));rp.find(g.tag('w:sz')).set(g.tag('w:val'),'28');g.add(rp,'w:rFonts',**{'w:ascii':'Arial','w:hAnsi':'Arial'})
  if case.get('effects'):
   line=props.find(g.tag('a:ln'));line.clear();line.set('w',str(2*12700));g.add(g.add(line,'a:solidFill'),'a:srgbClr',val='00A060')
  if case.get('pattern'):
   solid=props.find(g.tag('a:solidFill'));index=props.index(solid);props.remove(solid)
   fill=E.Element(g.tag('a:pattFill'),prst='dkUpDiag');props.insert(index,fill)
   for role,color in [('fgClr','00A060'),('bgClr','FFFFFF')]:g.add(g.add(fill,'a:'+role),'a:srgbClr',val=color)
  for axis,pct,tag in ([] if case.get('fixed') else [('H',40000,'pctWidth'),('V',30000,'pctHeight')]):
   n=E.SubElement(a,'{'+WP14+'}sizeRel'+axis,relativeFrom='page');E.SubElement(n,'{'+WP14+'}'+tag).text=str(pct)
 doc.save(ROOT/'controls.docx');(ROOT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
 print('Generated',len(cases),'relative frame controls')

if __name__=='__main__':generate()
