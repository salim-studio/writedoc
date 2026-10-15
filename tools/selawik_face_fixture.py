"""Original exact-face control: Selawik Semibold at two sizes and four styles."""
from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
out=root/'tests/fixtures/selawik-face';out.mkdir(exist_ok=True)
marker=BytesIO();Image.new('RGB',(4,4),'#1255aa').save(marker,format='PNG')
d=Document();d.sections[0].left_margin=d.sections[0].right_margin=Inches(.75)
d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(9)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for size in (16,24):
 for bold,italic in ((False,False),(True,False),(False,True),(True,True)):
  d.add_paragraph(f'Selawik Semibold {size}pt bold={bold} italic={italic}')
  p=d.add_paragraph();r=p.add_run('Measured typography: AV fi 0123.');r.font.name='Selawik Semibold';r.font.size=Pt(size);r.bold=bold;r.italic=italic
  p.add_run().add_picture(BytesIO(marker.getvalue()),width=Pt(4))
d.save(out/'control.docx')
