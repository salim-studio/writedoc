"""Original picture-control behavior cases built from an authored PNG."""
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
import json
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'tests/fixtures/picture-controls';OUT.mkdir(exist_ok=True)
with ZipFile(ROOT/'public/fixtures/picture-alpha.docx') as z:png=z.read('word/media/main.png')
(OUT/'replacement.png').write_bytes(png)
d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');cases=[]
for temporary in [False,True]:
 for action in ['resize','type','delete','replace']:
  name=('temporary' if temporary else 'persistent')+'-'+action;cases.append(dict(name=name,temporary=temporary,action=action))
  p=d.add_paragraph(name+': ');run=p.add_run();run.add_picture(BytesIO(png),width=Inches(1))
  s=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');s.append(pr)
  for tag,value in [('alias',name),('id',str(len(cases)))]:
   n=OxmlElement('w:'+tag);n.set(qn('w:val'),value);pr.append(n)
  if temporary:pr.append(OxmlElement('w:temporary'))
  pr.append(OxmlElement('w:picture'));content=OxmlElement('w:sdtContent');s.append(content)
  p._p.replace(run._r,s);content.append(run._r)
d.save(OUT/'control.docx');(OUT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
