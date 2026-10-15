"""Independent positional-tab controls in page stories, cells and columns."""
from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(__file__).resolve().parents[1] / 'tests/fixtures/positional-tabs'
d = Document()
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
s = d.sections[0]
s.page_width = Inches(8.5); s.page_height = Inches(11)
s.left_margin = s.right_margin = Inches(1)
normal = d.styles['Normal']
normal.font.name = 'Arial'; normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = Pt(18)
normal.paragraph_format.space_after = Pt(0)

def marker(p, color):
    b = BytesIO(); Image.new('RGB', (4, 4), color).save(b, format='PNG')
    p.add_run().add_picture(BytesIO(b.getvalue()), width=Pt(4), height=Pt(4))

def control(p, alignment, relative='margin'):
    p.paragraph_format.left_indent = Pt(12)
    p.paragraph_format.right_indent = Pt(6)
    marker(p, '#777777'); p.add_run('Lead ')
    tab = OxmlElement('w:ptab')
    for key, value in dict(alignment=alignment, relativeTo=relative, leader='none').items():
        tab.set(qn('w:'+key), value)
    p.add_run()._r.append(tab)
    marker(p, '#e03040'); p.add_run('Tail'); marker(p, '#2050d0')

control(s.header.paragraphs[0], 'center')
control(s.footer.paragraphs[0], 'right')
d.add_paragraph('Body and repeated positional tabs')
for relative in ['margin', 'indent']:
    for alignment in ['left', 'center', 'right']:
        control(d.add_paragraph(), alignment, relative)
p = d.add_paragraph()
control(p, 'center'); control(p, 'right')
d.add_section(WD_SECTION_START.NEW_PAGE)
d.add_paragraph('Fixed table cells')
table = d.add_table(rows=3, cols=2); table.autofit = False
for row, alignment in zip(table.rows, ['left', 'center', 'right']):
    for cell, relative in zip(row.cells, ['margin', 'indent']):
        control(cell.paragraphs[0], alignment, relative)
s = d.add_section(WD_SECTION_START.NEW_PAGE)
cols = s._sectPr.find(qn('w:cols')); cols.set(qn('w:num'), '2'); cols.set(qn('w:space'), '720')
d.add_paragraph('Two columns')
for alignment in ['left', 'center', 'right']:
    control(d.add_paragraph(), alignment)
d.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
for alignment in ['left', 'center', 'right']:
    control(d.add_paragraph(), alignment, 'indent')
d.save(out / 'contexts.docx')
table.autofit = True
d.save(out / 'contexts-autofit.docx')
