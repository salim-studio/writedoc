"""Original proofing-annotation controls; python-docx package, authored OOXML."""
from pathlib import Path
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt
from canvas_content_control_fixture import control, child, W, W15

ROOT = Path(__file__).resolve().parents[1]
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def generate():
    doc = Document()
    doc.settings.element.find(f'{{{W}}}zoom').set(f'{{{W}}}percent', '100')
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.75)
    section.top_margin = section.bottom_margin = Inches(.75)
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal'].paragraph_format.space_after = Pt(10)
    doc._element.set(f'{{{W15}}}probe', 'unused')
    prefix = next(k for k, v in doc._element.nsmap.items() if v == W15)
    del doc._element.attrib[f'{{{W15}}}probe']
    mc = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
    doc._element.set(f'{{{mc}}}Ignorable', doc._element.get(f'{{{mc}}}Ignorable', '') + ' ' + prefix)
    doc.add_paragraph('Content control proofing markers')
    cases = [('[Your name]', True), ('Existing company', False), ('[Plain value]', True), ('Linked company', False), ('Persistent company', False)]
    for i, (label, placeholder) in enumerate(cases, 1):
        paragraph = doc.add_paragraph(f'Case {i}: ')
        node = control(paragraph, i, [label], placeholder)
        properties = node.find(f'{{{W}}}sdtPr')
        if i == 3:
            properties.remove(properties.find(f'{{{W}}}richText'))
            child(properties, 'text')
        if i == 5:
            properties.remove(properties.find(f'{{{W}}}temporary'))
        content = node.find(f'{{{W}}}sdtContent')
        if i == 4:
            run = content[0]
            content.remove(run)
            link = child(content, 'hyperlink')
            link.set(f'{{{R}}}id', doc.part.relate_to('https://example.com/', R + '/hyperlink', is_external=True))
            link.append(run)
            content = link
        for index, marker in [(0, 'spellStart'), (1, 'gramStart')]:
            proof = E.Element(f'{{{W}}}proofErr', {f'{{{W}}}type': marker})
            content.insert(index, proof)
        child(content, 'proofErr', type='spellEnd')
        child(content, 'proofErr', type='gramEnd')
        paragraph.add_run(' tail')
    doc.add_paragraph('End of proofing controls')
    directory = ROOT / 'tests/fixtures/control-proofing'
    directory.mkdir(parents=True, exist_ok=True)
    doc.save(directory / 'source.docx')


if __name__ == '__main__':
    generate()
