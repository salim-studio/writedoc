"""Independent ordinary-space fixtures and export checks using python-docx (MIT)."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1] / 'public/fixtures'

def generate(factory):
    for wrapped in [False, True]:
        doc = factory(); section = doc.sections[0]
        section.page_width = section.page_height = Inches(6)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
        normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
        normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
        style = doc.styles.add_style('SpaceStyle', WD_STYLE_TYPE.PARAGRAPH); style.base_style = normal; style.font.underline = True
        char = doc.styles.add_style('SpaceCharacter', WD_STYLE_TYPE.CHARACTER); char.font.italic = True
        compat = doc.settings.element.find(qn('w:compat'))
        # CT_Compat order: ulTrailSpace precedes wrapTrailSpaces; both precede
        # the template's doNotAutofitConstrainedTables and compatSetting nodes.
        for tag in ['ulTrailSpace', 'wrapTrailSpaces']:
            node = OxmlElement('w:' + tag); node.set(qn('w:val'), 'on' if wrapped else 'off'); compat.insert(0 if tag == 'ulTrailSpace' else 1, node)
        doc.add_paragraph('Spaces A\u0301😀')
        doc.add_paragraph('Separated' + ' ' * 160 + 'Next word', style=style)
        for title, alignment in [('Left', WD_ALIGN_PARAGRAPH.LEFT), ('Centered', WD_ALIGN_PARAGRAPH.CENTER), ('Right', WD_ALIGN_PARAGRAPH.RIGHT)]:
            p = doc.add_paragraph(title, style=style); p.alignment = alignment
            r = p.add_run(' ' * 160); r.font.highlight_color = WD_COLOR_INDEX.YELLOW; r.style = char
        doc.add_paragraph(' ' * 160, style=style)
        doc.add_paragraph(' ' * 160 + 'Leading spaces')
        doc.add_paragraph('A\u00a0B A\u202fB A\u2007B A \u0301B')
        p = doc.add_paragraph('Manual   ', style=style); p.add_run().add_break(WD_BREAK.LINE); p.add_run('Following   ')
        doc.add_paragraph('Tab\tAnchor   ', style=style)
        for fixed in [False, True]:
            table = doc.add_table(rows=3, cols=2, style='Table Grid'); table.autofit = not fixed
            for col in table.columns: col.width = Inches(2.5)
            for row in table.rows:
                for cell in row.cells: cell.width = Inches(2.5)
            table.cell(0, 0).text = 'Header' + ' ' * 100; table.cell(0, 1).text = 'Other header'
            table.cell(0, 0).paragraphs[0].style = style
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            table.cell(1, 0).text = 'Cell' + ' ' * 100 + 'content'
            table.cell(1, 1).text = 'Table content with separators. ' * 70
            table.cell(2, 0).text = 'End   '; table.cell(2, 1).text = 'Tail   '
        doc.add_paragraph('End spaces   ')
        path = ROOT / ('wrapped-spaces.docx' if wrapped else 'trailing-spaces.docx'); doc.save(path); print(f'Generated independent spaces fixture: {path}')

def verify(file, kind):
    doc = Document(file)
    if kind == 'new':
        assert [p.text for p in doc.paragraphs] == ['Fresh A\u0301😀' + ' ' * 200 + 'Tail', 'Right   ', '   ']
        assert doc.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
        assert all(r.underline for p in doc.paragraphs for r in p.runs)
    else:
        source = ROOT / (kind + '-spaces.docx'); old = Document(source)
        assert doc.paragraphs[0].text == 'Edited ' + old.paragraphs[0].text
        assert [p.text for p in doc.paragraphs[1:]] == [p.text for p in old.paragraphs[1:]]
        assert doc.paragraphs[3].runs[0].text == 'Centered' and doc.paragraphs[3].runs[0].bold
        assert doc.styles['SpaceStyle'].font.size.pt == 14
        assert [[c.text for row in t.rows for c in row.cells] for t in doc.tables] == [[c.text for row in t.rows for c in row.cells] for t in old.tables]
        for key in ['wrapTrailSpaces', 'ulTrailSpace']:
            assert doc.settings.element.find(qn('w:compat')).find(qn('w:' + key)).get(qn('w:val')) == ('on' if kind == 'wrapped' else 'off')
        with ZipFile(source) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), name
    for node in doc.element.iter(qn('w:t')):
        if node.text and node.text != node.text.strip(): assert node.get(qn('xml:space')) == 'preserve', repr(node.text)
    print(f'PASS: {file}: independent space counts, Unicode, formatting, compatibility flags and package preservation')
