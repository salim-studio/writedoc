"""Original justified-paragraph fixtures and independent python-docx checks."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
def source(soft=False): return ROOT / 'public/fixtures' / ('justification-soft-breaks.docx' if soft else 'justification.docx')

def generate(factory):
    for soft in [False, True]:
        doc = factory(); section = doc.sections[0]
        section.page_width = Inches(6); section.page_height = Inches(7)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
        normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
        normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1
        base = doc.styles.add_style('Canvas Justified', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal; base.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        child = doc.styles.add_style('Canvas Derived Justified', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
        if soft: doc.settings.element.find(qn('w:compat')).insert(0, OxmlElement('w:doNotExpandShiftReturn'))
        doc.add_paragraph('Justified paragraphs A\u0301😀')
        phrase = 'Measured words keep their shapes while spaces fill the line. '
        p = doc.add_paragraph('Inherited words ', child); r = p.add_run(phrase * 3); r.underline = True; r.font.color.rgb = RGBColor.from_string('224466')
        p = doc.add_paragraph('Direct left ' + phrase * 2, child); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p = doc.add_paragraph('Direct both ' + phrase * 2, child); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p = doc.add_paragraph('Soft line words A\u0301😀', child); p.add_run().add_break(WD_BREAK.LINE); p.add_run('Following soft line ' + phrase)
        p = doc.add_paragraph('Spaces: Alpha\u00a0Beta narrow\u202fjoin and regular words. ' + phrase + '    ', child)
        p = doc.add_paragraph('Mixed ', child); r = p.add_run('small caps words '); r.font.small_caps = True; r.font.color.rgb = RGBColor.from_string('8855CC')
        r = p.add_run('raised 2 words '); r.font.superscript = True; r.font.strike = True
        r = p.add_run(phrase); r.bold = True; r.font.highlight_color = 7
        p = doc.add_paragraph('Hanging ' + phrase * 2, child); p.paragraph_format.left_indent = Pt(18); p.paragraph_format.first_line_indent = Pt(-12); p.paragraph_format.right_indent = Pt(12)
        p = doc.add_paragraph('First label\t' + phrase * 2, child); p.paragraph_format.tab_stops.add_tab_stop(Inches(1.25), WD_TAB_ALIGNMENT.LEFT)
        p = doc.add_paragraph('Numbered ' + phrase * 2, 'List Number'); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p = doc.add_paragraph('Page line words', child); p.add_run().add_break(WD_BREAK.PAGE); p.add_run('After page break ' + phrase)
        doc.add_paragraph('Before justified table')
        table = doc.add_table(rows=2, cols=2, style='Table Grid'); table.autofit = False
        for col in table.columns: col.width = Inches(2.5)
        for row in table.rows:
            for cell in row.cells:
                cell.width = Inches(2.5); cell.paragraphs[0].style = child
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        table.cell(0,0).paragraphs[0].add_run('Header justified words wrap here now')
        table.cell(0,1).paragraphs[0].add_run('Other header words fit here')
        table.cell(1,0).paragraphs[0].add_run('Flowing justified words. ' * 80)
        table.cell(1,1).paragraphs[0].add_run('Neighbor words remain editable')
        doc.add_paragraph('Tail'); doc.save(source(soft)); print(f'Generated independent justification fixture: {source(soft)}')

def verify(file, mode):
    doc = Document(file)
    if mode == 'new':
        assert [p.text for p in doc.paragraphs] == ['New justified words ' * 15, 'Final left paragraph']
        assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert doc.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.LEFT
    else:
        original = Document(source(mode == 'soft'))
        before = original.element.findall('.//' + qn('w:p')); after = doc.element.findall('.//' + qn('w:p'))
        text = lambda p: ''.join(n.text or '' for n in p.findall('.//' + qn('w:t')))
        assert [text(p) for p in after] == [text(p) for p in before]
        for p, old in zip(after, before):
            changed = text(p).startswith('Direct left') or mode == 'main' and text(p).startswith('Direct both')
            if changed:
                expected = 'both' if text(p).startswith('Direct left') else 'right'
                assert p.find(qn('w:pPr')).find(qn('w:jc')).get(qn('w:val')) == expected
            def data(element): return [(n.tag, sorted(n.attrib.items()), n.text) for n in element.iter() if not (changed and n.tag == qn('w:jc'))]
            assert data(p) == data(old), f'Unrelated paragraph data changed: {text(p)}'
        with ZipFile(source(mode == 'soft')) as original_zip, ZipFile(file) as exported:
            assert set(original_zip.namelist()) == set(exported.namelist())
            for name in original_zip.namelist():
                if name != 'word/document.xml': assert original_zip.read(name) == exported.read(name), name
    print(f'PASS: {file}: justified alignment, exact text, styles, breaks, formatting and unrelated package parts')
