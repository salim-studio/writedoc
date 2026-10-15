"""Independent WPS autofit controls, without corpus content or renderer calls."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree as E
import canvas_group_fixture as g

ROOT = Path(__file__).resolve().parents[1] / 'tests/fixtures/shape-autofit'
TEXT = 'Alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu nu xi omicron pi'

def generate():
    ROOT.mkdir(exist_ok=True)
    doc = Document()
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    ident = 100
    for grouped in [False, True]:
        for height in [18, 120]:
            for anchor in ['t', 'ctr', 'b']:
                ident += 1
                label = doc.add_paragraph(f'Case {ident-100}: {"group" if grouped else "shape"}, {height}pt, {anchor}')
                label.paragraph_format.page_break_before = ident > 101
                doc.add_paragraph('Before')
                frame = g.add(g.add(doc.add_paragraph().add_run()._r, 'w:drawing'), 'wp:inline', distT=0, distB=0, distL=0, distR=0)
                width, high = 120*12700, height*12700
                g.add(frame, 'wp:extent', cx=width, cy=high)
                g.add(frame, 'wp:docPr', id=ident, name=f'Control {ident}')
                g.add(frame, 'wp:cNvGraphicFramePr')
                data = g.add(g.add(frame, 'a:graphic'), 'a:graphicData', uri=g.NS['wpg' if grouped else 'wps'])
                parent = data
                if grouped:
                    parent = g.add(data, 'wpg:wgp'); g.add(parent, 'wpg:cNvGrpSpPr')
                    props = g.add(parent, 'wpg:grpSpPr'); g.transform(props, width, high, group=True)
                    props[0].find(g.tag('a:chOff')).attrib.update({'x': '0', 'y': '0'})
                shape = g.shape(parent, ident+100, width//(2 if grouped else 1), high//(2 if grouped else 1), 0, 0, 'DDEEFF', TEXT)
                rp = shape.find('.//' + g.tag('w:rPr'))
                rp.find(g.tag('w:sz')).set(g.tag('w:val'), '24')
                rp.insert(0, E.Element(g.tag('w:rFonts'), {g.tag('w:ascii'): 'Arial', g.tag('w:hAnsi'): 'Arial'}))
                body = shape.find(g.tag('wps:bodyPr')); body.set('anchor', anchor)
                for side in ['lIns', 'rIns', 'tIns', 'bIns']: body.set(side, str(3*12700))
                doc.add_paragraph('After')
    base = ROOT / 'noAutofit.docx'; doc.save(base)
    with ZipFile(base) as z: parts = {n:z.read(n) for n in z.namelist()}
    root = E.fromstring(parts['word/document.xml'])
    for mode in ['noAutofit', 'spAutoFit', 'omitted']:
        variant = deepcopy(root)
        for body in variant.iter(g.tag('wps:bodyPr')):
            for child in list(body): body.remove(child)
            if mode != 'omitted': g.add(body, 'a:'+mode)
        parts['word/document.xml'] = E.tostring(variant, encoding='UTF-8', xml_declaration=True, standalone=True)
        with ZipFile(ROOT/(mode+'.docx'), 'w', ZIP_DEFLATED) as z:
            for name, data in parts.items(): z.writestr(name, data)
        settings = E.fromstring(parts['word/settings.xml'])
        for setting in settings.iter(g.tag('w:compatSetting')):
            if setting.get(g.tag('w:name')) == 'compatibilityMode':
                setting.set(g.tag('w:val'), '15')
        mode15 = {**parts, 'word/settings.xml': E.tostring(settings, encoding='UTF-8', xml_declaration=True, standalone=True)}
        with ZipFile(ROOT/(mode+'-mode15.docx'), 'w', ZIP_DEFLATED) as z:
            for name, data in mode15.items(): z.writestr(name, data)

if __name__ == '__main__': generate()
