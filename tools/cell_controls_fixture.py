"""Original temporary cell-control fixture; python-docx boilerplate is MIT."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add(parent,name,**attrs):
 n=OxmlElement('w:'+name)
 for k,v in attrs.items():n.set(qn('w:'+k),str(v))
 parent.append(n);return n

def wrap(cell,ident,placeholder=False):
 tc=cell._tc;parent=tc.getparent();s=OxmlElement('w:sdt');parent.replace(tc,s)
 pr=add(s,'sdtPr');add(pr,'alias',val='Original cell '+str(ident));add(pr,'id',val=ident);add(pr,'temporary')
 if placeholder:add(pr,'showingPlcHdr')
 add(pr,'richText');add(s,'sdtContent').append(tc)
 return s

doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(12)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
style.paragraph_format.space_after=Pt(0)
s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.left_margin=s.right_margin=Inches(1)
doc.add_paragraph('Temporary cell controls')
t=doc.add_table(rows=3,cols=3);t.style='Table Grid';t.autofit=False
for row in t.rows:
 for cell in row.cells:cell.width=Inches(2)
t.cell(0,0).text='Cell alpha';t.cell(0,0).add_paragraph('Second paragraph');t.cell(0,1).text='Cell beta';t.cell(0,2).text='Ordinary gamma'
t.cell(1,0).merge(t.cell(1,1)).text='Spanning delta';t.cell(1,2).text=''
t.cell(2,0).text='Cell epsilon';t.cell(2,1).text='Ordinary zeta';t.cell(2,2).text='Cell eta'
for cell,i,p in [(t.cell(0,0),101,False),(t.cell(0,1),102,True),(t.cell(1,0),103,False),(t.cell(1,2),104,True),(t.cell(2,0),105,False),(t.cell(2,2),106,False)]:wrap(cell,i,p)
doc.add_paragraph('After table')
doc.save(Path(__file__).resolve().parents[1]/'public/fixtures/cell-controls.docx')
for control in list(doc.element.iter(qn('w:sdt'))):
 parent=control.getparent();cell=control.find(qn('w:sdtContent'))[0]
 parent.replace(control,cell)
doc.save(Path(__file__).resolve().parents[1]/'public/fixtures/cell-controls-baseline.docx')
