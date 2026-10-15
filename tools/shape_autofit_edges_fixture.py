"""Original autofit controls for group bounds, text rectangles and spacing."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree as E
import canvas_group_fixture as g
from shape_autofit_fixture import TEXT

ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures/shape-autofit-edges'
def generate():
 ROOT.mkdir(exist_ok=True)
 for fit in ['noAutofit','spAutoFit']:
  d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
  for c in d.settings.element.iter(qn('w:compatSetting')):
   if c.get(qn('w:name'))=='compatibilityMode':c.set(qn('w:val'),'15')
  n=d.styles['Normal'];n.font.name='Arial';n.font.size=Pt(12);n.paragraph_format.space_after=Pt(0);n.paragraph_format.line_spacing=1
  def transform(parent,w,h):
   props=g.add(parent,'wpg:grpSpPr');g.transform(props,w,h,group=True)
   props[0].find(g.tag('a:chOff')).attrib.update({'x':'0','y':'0'})
   props[0].find(g.tag('a:chExt')).attrib.update({'cx':str(w),'cy':str(h)})
  for i,kind in enumerate(['mixed','nested','ellipse','outline','spacing','empty']):
   p=d.add_paragraph(kind);p.paragraph_format.page_break_before=i>0;d.add_paragraph('Before')
   grouped=kind in ['mixed','nested'];w,h=(360*12700,240*12700) if grouped else (120*12700,120*12700)
   frame=g.add(g.add(d.add_paragraph().add_run()._r,'w:drawing'),'wp:inline',distT=0,distB=0,distL=0,distR=0)
   g.add(frame,'wp:extent',cx=w,cy=h);g.add(frame,'wp:docPr',id=i+1,name=kind);g.add(frame,'wp:cNvGraphicFramePr')
   data=g.add(g.add(frame,'a:graphic'),'a:graphicData',uri=g.NS['wpg' if grouped else 'wps']);parent=data
   if grouped:
    parent=g.add(data,'wpg:wgp');g.add(parent,'wpg:cNvGrpSpPr');transform(parent,w,h)
    if kind=='nested':
     parent=g.add(parent,'wpg:grpSp');g.add(parent,'wpg:cNvPr',id=100+i,name='Nested');g.add(parent,'wpg:cNvGrpSpPr');transform(parent,w,h)
   sh=g.shape(parent,200+i,120*12700,(18 if grouped else 120)*12700,(20 if grouped else 0)*12700,(20 if grouped else 0)*12700,'DDEEFF',TEXT)
   if grouped:g.shape(parent,300+i,60*12700,20*12700,180*12700,180*12700,'11AABB')
   body=sh.find(g.tag('wps:bodyPr'));body[0].tag=g.tag('a:'+fit)
   for side in ['lIns','rIns','tIns','bIns']:body.set(side,str(3*12700))
   rp=sh.find('.//'+g.tag('w:rPr'));rp.find(g.tag('w:sz')).set(g.tag('w:val'),'24');rp.insert(0,E.Element(g.tag('w:rFonts'),{g.tag('w:ascii'):'Arial',g.tag('w:hAnsi'):'Arial'}))
   if kind=='ellipse':sh.find('.//'+g.tag('a:prstGeom')).set('prst','ellipse')
   if kind=='outline':
    ln=sh.find('.//'+g.tag('a:ln'));ln.set('w',str(4*12700));ln.remove(ln[0]);g.add(g.add(ln,'a:solidFill'),'a:srgbClr',val='000000')
   if kind=='spacing':sh.find('.//'+g.tag('w:spacing')).attrib.update({g.tag('w:before'):'240',g.tag('w:after'):'360'})
   if kind=='empty':sh.find('.//'+g.tag('w:t')).text=''
   d.add_paragraph('After')
  d.save(ROOT/(fit+'.docx'))
if __name__=='__main__':generate()
