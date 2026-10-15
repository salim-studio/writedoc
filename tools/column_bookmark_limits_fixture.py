"""Independent column-index bounds controls, including a merged first row."""
import json
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

out = Path(__file__).resolve().parents[1] / 'tests/fixtures/column-bookmarks'
cases = [
    ('PlainLastBeyond', 1, 3, False),
    ('MergedAll', 0, 2, True),
    ('MergedLastBeyond', 1, 2, True),
    ('MergedFirstBeyond', 2, 2, True),
    ('MergedBothBeyond', 3, 3, True),
    ('Reversed', 2, 1, False),
    ('FarBeyond', 99, 99, False),
    ('NegativeFirst', -1, 1, False),
    ('NegativeBoth', -1, -1, False),
    ('PlainFirstBeyond', 3, 3, False),
]
doc = Document()
doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(0)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
for index, (name, first, last, merged) in enumerate(cases):
    if index:
        doc.add_page_break()
    doc.add_paragraph(name)
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f'R{r+1}C{c+1}'
    if merged:
        table.cell(0, 0).merge(table.cell(0, 1))
    start = OxmlElement('w:bookmarkStart')
    for key, value in dict(id=index, name=name, colFirst=first, colLast=last).items():
        start.set(qn('w:' + key), str(value))
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(index))
    table.cell(0, 0)._tc.insert(1, start)
    table._tbl.append(end)
    doc.add_paragraph('After ' + name)
doc.save(out / 'limits-control.docx')
(out / 'limits-cases.json').write_text(json.dumps(cases, indent=2) + '\n')
