"""Independent controls for applying OOXML bold to installed heavy font faces."""
from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml.ns import qn
import json
root=Path(__file__).resolve().parents[1]/'tests/fixtures/heavy-face-bold'
root.mkdir(parents=True,exist_ok=True)
d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
d.sections[0].top_margin=d.sections[0].bottom_margin=Inches(.5)
cases=[]
for i,family in enumerate(['Franklin Gothic Demi','Franklin Gothic Book','Franklin Gothic Medium','Arial Black','Arial Bold','Arial']):
 if i:d.add_page_break()
 d.add_paragraph(family)
 for size in [12,24,85]:
  for italic,bold in [(False,False),(False,True),(True,False),(True,True)]:
   p=d.add_paragraph();p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
   r=p.add_run('AgMW');r.font.name=family;r.font.size=Pt(size);r.bold=bold;r.italic=italic
   cases.append({'family':family,'size':size,'italic':italic,'bold':bold,'text':'AgMW'})
d.save(root/'control.docx');(root/'cases.json').write_text(json.dumps(cases,indent=2)+'\n',encoding='utf-8')
