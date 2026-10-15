"""Original annual-report feature fixture, independent of the JavaScript engine.

python-docx (MIT) builds the package; lxml (BSD-3-Clause) supplies original XML.
All PNG pixels are generated below; no user template assets are redistributed.
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import struct
import zlib
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/annual-report.docx'
NS = {'w': qn('w:p').split('}')[0][1:], 'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
      'wp14': 'http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing',
      'w14': 'http://schemas.microsoft.com/office/word/2010/wordml',
      'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006'}

def tag(name):
    prefix, local = name.split(':')
    return '{' + NS[prefix] + '}' + local

def add(parent, name, text=None, **attrs):
    n = E.SubElement(parent, tag(name), {tag(k) if ':' in k else k: str(v) for k, v in attrs.items()})
    n.text = text
    return n

def png():
    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xffffffff)
    pixels = b''.join(b'\0' + bytes([40, 130, 180] * 16 + [210, 60, 120] * 16) for _ in range(32))
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 32, 32, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(pixels)) + chunk(b'IEND', b'')

def picture(p, floating=False, band=False):
    run = p.add_run(); inline = run.add_picture(BytesIO(png()), width=Inches(2 if band else .4))._inline
    if not band: add(run._r.get_or_add_rPr(), 'w:position', **{'w:val': -23})
    if not floating: return
    anchor = E.Element(tag('wp:anchor'), nsmap={'wp14': NS['wp14']}, distT='0', distB='0', distL='114300', distR='114300', simplePos='0', relativeHeight='12', behindDoc='0', locked='0', layoutInCell='1', allowOverlap='1')
    anchor.set(tag('mc:Ignorable'), 'wp14')
    add(anchor, 'wp:simplePos', x=0, y=0)
    add(add(anchor, 'wp:positionH', relativeFrom='column'), 'wp:posOffset', text='-91440' if band else '95250')
    add(add(anchor, 'wp:positionV', relativeFrom='paragraph'), 'wp:posOffset', text='285750' if band else '19050')
    anchor.append(deepcopy(inline.find(tag('wp:extent'))))
    add(anchor, 'wp:effectExtent', l=0, t=0, r=0 if band else 3810, b=0 if band else 3810)
    add(anchor, 'wp:wrapTopAndBottom' if band else 'wp:wrapNone')
    for child in inline:
        if child.tag != tag('wp:extent'): anchor.append(deepcopy(child))
    if not band:
        add(add(anchor, 'wp14:sizeRelH', relativeFrom='page'), 'wp14:pctWidth', text='0')
        add(add(anchor, 'wp14:sizeRelV', relativeFrom='page'), 'wp14:pctHeight', text='0')
    inline.getparent().replace(inline, anchor)

def generate():
    doc = Document(); sec = doc.sections[0]
    doc.settings.element.find(tag('w:zoom')).set(tag('w:percent'), '100')
    sec.left_margin = sec.right_margin = Inches(.75)
    sec.top_margin = sec.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(6)
    doc.add_paragraph('Annual report features Á😀')
    p = doc.add_paragraph(); add(p._p.get_or_add_pPr(), 'w:spacing', **{'w:line': 0, 'w:lineRule': 'auto', 'w:before': 120, 'w:after': 120})
    p = doc.add_paragraph('Lowered inline icon '); picture(p); p.add_run(' after the icon.')
    table = doc.add_table(rows=3, cols=4); table.autofit = False
    for column in table.columns: column.width = Inches(1.5)
    pr = table._tbl.tblPr
    add(pr, 'w:tblInd', **{'w:w': -360, 'w:type': 'dxa'})
    pr.find(tag('w:tblW')).set(tag('w:type'), 'dxa'); pr.find(tag('w:tblW')).set(tag('w:w'), '8640')
    # Put the property in schema order (tblInd precedes tblLayout).
    pr.insert(list(pr).index(pr.find(tag('w:tblLayout'))), pr.find(tag('w:tblInd')))
    for i, row in enumerate(table.rows):
        row._tr.set(tag('w14:paraId'), f'{i + 100:08X}')
        ex = E.Element(tag('w:tblPrEx')); margins = add(ex, 'w:tblCellMar')
        for side, value in [('top', 120), ('left', 180), ('bottom', 120), ('right', 180)]: add(margins, 'w:' + side, **{'w:w': value, 'w:type': 'dxa'})
        row._tr.insert(0, ex)
        for j, cell in enumerate(row.cells):
            cell.width = Inches(1.5); cell.text = f'Cell {i + 1}.{j + 1}'
    cell = table.cell(0, 0)
    add(cell.paragraphs[0]._p.get_or_add_pPr(), 'w:ind', **{'w:left': -240, 'w:right': -240})
    margins = add(cell._tc.get_or_add_tcPr(), 'w:tcMar'); add(margins, 'w:top', **{'w:w': 30, 'w:type': 'dxa'})
    row = table.rows[1]._tr
    row.remove(row.findall(tag('w:tc'))[-1]); row.remove(row.findall(tag('w:tc'))[0])
    rp = row.get_or_add_trPr()
    add(rp, 'w:gridBefore', **{'w:val': 1}); add(rp, 'w:gridAfter', **{'w:val': 1})
    add(rp, 'w:wBefore', **{'w:w': 2160, 'w:type': 'dxa'}); add(rp, 'w:wAfter', **{'w:w': 2160, 'w:type': 'dxa'})
    # Work with physical XML after omitting cells (python-docx's grid proxy is rectangular).
    p = doc.add_paragraph('Floating cell'); picture(p, floating=True)
    tc = table.rows[2]._tr.findall(tag('w:tc'))[1]
    tc.remove(tc.find(tag('w:p'))); tc.append(p._p)
    cells = table.rows[2]._tr.findall(tag('w:tc'))
    add(cells[2].find(tag('w:tcPr')), 'w:gridSpan', **{'w:val': 2})
    cells[2].find('.//' + tag('w:tcW')).set(tag('w:w'), '4320')
    table.rows[2]._tr.remove(cells[3])
    doc.add_paragraph('After the table.')
    p = doc.add_paragraph('Before the wrapping band.'); picture(p, floating=True, band=True)
    doc.add_paragraph('After the wrapping band.')
    doc.save(SOURCE); print(f'Generated {SOURCE}')

def verify(path):
    def signature(node): return [(n.tag, sorted(n.attrib.items()), n.text or '') for n in node.iter()]
    with ZipFile(SOURCE) as old, ZipFile(path) as new:
        a, b = [E.fromstring(z.read('word/document.xml')) for z in [old, new]]
        text = lambda root: ''.join(root.xpath('//w:t/text()', namespaces=NS))
        formatted = Path(path).name == 'annual-table-format.docx'
        assert text(b) == ('' if formatted else 'Edited ') + text(a)
        if formatted:
            table = b.find('.//' + tag('w:tbl'))
            assert table.find('.//' + tag('w:tblInd')).get(tag('w:w')) == '-180'
            assert [n.get(tag('w:w')) for n in table.findall('./w:tblGrid/w:gridCol', NS)] == ['1440', '2160', '2160', '2160']
            assert table.find('.//' + tag('w:wBefore')).get(tag('w:w')) == '1440'
            assert table.find('.//' + tag('w:wAfter')).get(tag('w:w')) == '2160'
            assert table.find('.//' + tag('w:trHeight')).attrib == {tag('w:val'): '1440', tag('w:hRule'): 'atLeast'}
            assert [n.tag for n in table.find(tag('w:tr'))][:2] == [tag('w:tblPrEx'), tag('w:trPr')]
        for name in [*(['w:tblPrEx'] if formatted else ['w:tbl']), 'w:drawing', 'w:sectPr']:
            assert [signature(n) for n in a.findall('.//' + tag(name))] == [signature(n) for n in b.findall('.//' + tag(name))], name
        assert set(old.namelist()) == set(new.namelist())
        for name in old.namelist():
            if name != 'word/document.xml': assert old.read(name) == new.read(name), name
    print(f'PASS {path}: independent annual-report text, tables, drawings and package retention')

if __name__ == '__main__':
    import sys
    if len(sys.argv) == 1: generate()
    else:
        for path in sys.argv[1:]: verify(path)
