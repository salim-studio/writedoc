from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor
p=Path('tests/fixtures/textbox-page-break-before/oversized-body.docx')
d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');s=d.sections[0];s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(.5)
for i,c in enumerate('AB'):
 para=d.add_paragraph();para.paragraph_format.page_break_before=bool(i);r=para.add_run(c);r.font.name='Arial';r.font.size=Pt(650);r.font.color.rgb=RGBColor.from_string('D02030')
d.save(p)
