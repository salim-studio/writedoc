"""Original light-face selection control with regular, bold and italic variants."""
from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
marker=BytesIO();Image.new('RGB',(4,4),'#1255aa').save(marker,format='PNG')
d=Document();d.sections[0].left_margin=d.sections[0].right_margin=Inches(1)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
for family in ['Source sans pro','Source sans pro light']:
 for bold,italic in [(False,False),(True,False),(False,True),(True,True)]:
  d.add_paragraph(f'{family} bold={bold} italic={italic}')
  p=d.add_paragraph();r=p.add_run('Measured typography uses the declared font face.');r.font.name=family;r.font.size=Pt(16);r.bold=bold;r.italic=italic; p.add_run().add_picture(BytesIO(marker.getvalue()),width=Pt(4))
(root/'tests/fixtures/source-sans-light').mkdir(exist_ok=True)
d.save(root/'tests/fixtures/source-sans-light/control.docx')
