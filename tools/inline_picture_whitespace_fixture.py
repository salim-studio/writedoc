"""Independent reference controls for whitespace on inline-picture lines."""
from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

out = Path(__file__).resolve().parents[1] / 'tests/fixtures/inline-picture-whitespace'
out.mkdir(exist_ok=True)
d = Document(); s = d.sections[0]
s.page_width = Inches(8.5); s.page_height = Inches(11)
s.left_margin = s.right_margin = Inches(1)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
for setting in d.settings.element.iter(qn('w:compatSetting')):
    if setting.get(qn('w:name')) == 'compatibilityMode': setting.set(qn('w:val'), '15')
normal = d.styles['Normal']; normal.font.name = 'Verdana'; normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = 1
cases = []
for mark_size in [9, 20]:
    for name, text in [('empty',''),('space',' '),('nbsp','\u00a0'),('en','\u2002'),('em','\u2003'),('thin','\u2009'),('narrow-nbsp','\u202f'),('tab','\t'),('zero-width','\u200b'),('letter','x'),
                       *[(f'U{code:04X}', chr(code)) for code in [0x1680,0x2000,0x2001,0x2004,0x2005,0x2006,0x2007,0x2008,0x200a,0x205f,0x3000,0xfeff,0x2060]]]:
        case = dict(name=f'{name}-mark{mark_size}', text=text, markSize=mark_size); cases.append(case)
        label = d.add_paragraph(case['name']); label.paragraph_format.page_break_before = len(cases) > 1
        d.add_paragraph('Before')
        p = d.add_paragraph(); rpr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(mark_size*2)); rpr.append(sz); p._p.get_or_add_pPr().append(rpr)
        b = BytesIO(); Image.new('RGB',(4,4),'#e03040').save(b,format='PNG')
        p.add_run().add_picture(BytesIO(b.getvalue()), width=Pt(14.25), height=Pt(14.25))
        p.add_run(text)
        d.add_paragraph('After')
d.save(out/'source.docx')
(out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n',encoding='utf-8')
for drawing in list(d.element.iter(qn('w:drawing'))): drawing.getparent().remove(drawing)
d.save(out/'text-only.docx')
# The reference PDFs show glyph substitution for these characters in Verdana. Keep
# the original controls and add explicit-face variants to isolate line metrics
# from the separate fallback-font problem. These names never enter engine code.
fallbacks = {0x1680:'Segoe UI Historic',0x205f:'Cambria Math',0x2060:'Segoe UI Symbol',
             **{code:'Arial' for code in [0x2000,0x2001,0x2004,0x2006,0x2007,0x2008,0x2009,0x200a]}}
for name in ['source','text-only']:
    control = Document(out/(name+'.docx'))
    for i, case in enumerate(cases):
        family = fallbacks.get(ord(case['text'])) if len(case['text']) == 1 else None
        if family: control.paragraphs[i*4+2].runs[-1].font.name = family
    control.save(out/(name+'-explicit.docx'))
