"""Original oversized-inline probe: independent WordprocessingML and PNG artwork."""
from pathlib import Path
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]

def generate():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8); section.page_height = Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)
    doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    artwork = Image.new('RGB', (260, 40), '#1255aa'); draw = ImageDraw.Draw(artwork)
    draw.rectangle((130, 0, 259, 39), fill='#ee6622')
    data = BytesIO(); artwork.save(data, format='PNG')
    for i, (where, alignment) in enumerate((where, alignment) for where in ['body', 'cell'] for alignment in ['left', 'center', 'right']):
        if i: doc.add_page_break()
        doc.add_paragraph(f'{where} {alignment}')
        if where == 'cell':
            table = doc.add_table(rows=1, cols=2); table.autofit = False
            for col in table.columns: col.width = Inches(3)
            for cell in table.rows[0].cells: cell.width = Inches(3)
            margins = OxmlElement('w:tblCellMar')
            for side in ['left', 'right']:
                n = OxmlElement('w:' + side); n.set(qn('w:w'), '0'); n.set(qn('w:type'), 'dxa'); margins.append(n)
            table._tbl.tblPr.insert_element_before(margins, 'w:tblLook')
            p = table.cell(0, 0).paragraphs[0]; table.cell(0, 1).text = 'Neighbor'
        else: p = doc.add_paragraph()
        p.alignment = getattr(WD_ALIGN_PARAGRAPH, alignment.upper())
        p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(.5)
        p.add_run('Before ')
        p.add_run().add_picture(BytesIO(data.getvalue()), width=Inches(3.25 if where == 'cell' else 6.5))
        p.add_run(' after')
        doc.add_paragraph('End')
    path = ROOT / 'public/fixtures/oversized-inline.docx'; doc.save(path); print(path)

if __name__ == '__main__': generate()
