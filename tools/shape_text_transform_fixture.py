"""Independent WPS text-transform controls; no corpus content."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import canvas_group_fixture as g

ROOT = Path(__file__).resolve().parents[1] / 'tests/fixtures/shape-text-transform'
CASES = [
    ('superscript', [('vertAlign', 'superscript')], 'Test'),
    ('subscript', [('vertAlign', 'subscript')], 'Test'),
    ('raised', [('position', '14')], 'Test'),
    ('lowered', [('position', '-14')], 'Test'),
    ('raised superscript', [('vertAlign', 'superscript'), ('position', '14')], 'Test'),
    ('raised superscript space', [('vertAlign', 'superscript'), ('position', '14'), ('spacing', '120')], ' '),
    ('condensed', [('w', '50')], 'Test'),
    ('expanded', [('w', '150')], 'Test'),
]

def generate():
    ROOT.mkdir(exist_ok=True)
    for fit in ['noAutofit', 'spAutoFit']:
        doc = Document()
        doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
        normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(12)
        normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = 1
        for setting in doc.settings.element.iter(qn('w:compatSetting')):
            if setting.get(qn('w:name')) == 'compatibilityMode': setting.set(qn('w:val'), '15')
        ident = 0
        for grouped in [False, True]:
            for label, properties, text in CASES:
                ident += 1
                p = doc.add_paragraph(f'Case {ident}: {label}, group={grouped}')
                p.paragraph_format.page_break_before = ident > 1
                doc.add_paragraph('Before')
                frame = g.add(g.add(doc.add_paragraph().add_run()._r, 'w:drawing'), 'wp:inline', distT=0, distB=0, distL=0, distR=0)
                width, height = 240*12700, 90*12700
                g.add(frame, 'wp:extent', cx=width, cy=height)
                g.add(frame, 'wp:docPr', id=ident, name=f'Control {ident}')
                g.add(frame, 'wp:cNvGraphicFramePr')
                data = g.add(g.add(frame, 'a:graphic'), 'a:graphicData', uri=g.NS['wpg' if grouped else 'wps'])
                parent = data
                if grouped:
                    parent = g.add(data, 'wpg:wgp'); g.add(parent, 'wpg:cNvGrpSpPr')
                    props = g.add(parent, 'wpg:grpSpPr'); g.transform(props, width, height, group=True)
                    props[0].find(g.tag('a:chOff')).attrib.update({'x':'0','y':'0'})
                shape = g.shape(parent, ident+100, width//(2 if grouped else 1), height//(2 if grouped else 1), 0, 0, 'DDEEFF', 'placeholder')
                p = shape.find('.//' + g.tag('w:p'))
                for r in list(p.findall(g.tag('w:r'))): p.remove(r)
                for content, attrs in [('Start ', []), (text, properties), (' end', [])]:
                    r = g.add(p, 'w:r'); rp = g.add(r, 'w:rPr')
                    g.add(rp, 'w:rFonts', **{'w:ascii':'Arial','w:hAnsi':'Arial'})
                    g.add(rp, 'w:sz', **{'w:val':40})
                    g.add(rp, 'w:color', **{'w:val':'123456'})
                    for tag, value in attrs: g.add(rp, 'w:'+tag, **{'w:val':value})
                    t = g.add(r, 'w:t', text=content); t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                body = shape.find(g.tag('wps:bodyPr'))
                for child in list(body): body.remove(child)
                g.add(body, 'a:'+fit)
                for side in ['lIns','rIns','tIns','bIns']: body.set(side,str(3*12700))
                doc.add_paragraph('After')
        doc.save(ROOT/(fit+'.docx'))
        print(f'Generated {fit}: {ident} pages')

if __name__ == '__main__': generate()
