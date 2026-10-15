"""Original WordprocessingML text-fill controls; no corpus XML or artwork."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
import canvas_group_fixture as g

ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures/text-fill'
W14='http://schemas.microsoft.com/office/word/2010/wordml'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def effect(parent, body, name='textFill', attrs=''):
    parent.append(parse_xml(f'<w14:{name} xmlns:w14="{W14}" {attrs}>{body}</w14:{name}>'))

def fill(parent, value):
    body='<w14:noFill/>' if value=='none' else '' if value=='empty' else f'<w14:solidFill><w14:srgbClr w14:val="{value}"/></w14:solidFill>'
    effect(parent,body)

def outline(parent):
    effect(parent,'<w14:solidFill><w14:srgbClr w14:val="008080"/></w14:solidFill><w14:bevel/>','textOutline','w14:w="9525"')

def generate():
    ROOT.mkdir(parents=True,exist_ok=True)
    doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(18)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    for name,color in [('NoFill','none'),('SolidFill','208020')]:
        style=doc.styles.add_style(name,WD_STYLE_TYPE.CHARACTER);fill(style.element.get_or_add_rPr(),color)
    names=['normal','no-fill','solid-green','empty-fill','underline-auto','underline-red','strike','outline','bold','small-caps','style-no-fill','style-no-fill-color','style-no-fill-solid','style-solid-no-fill']
    cases=[]
    for i,name in enumerate(names):
        p=doc.add_paragraph();p.add_run(f'{i+1:02d} | ')
        r=p.add_run('Abcdef');pr=r._r.get_or_add_rPr();r.font.color.rgb=RGBColor.from_string('C02020')
        if name=='solid-green':fill(pr,'208020')
        elif name=='empty-fill':fill(pr,'empty')
        elif name.startswith('style-'):
            r.style=doc.styles['SolidFill' if name=='style-solid-no-fill' else 'NoFill']
            if name=='style-no-fill-solid':fill(pr,'2040C0')
            elif name=='style-solid-no-fill':fill(pr,'none')
            elif name=='style-no-fill-color':r.font.color.rgb=RGBColor.from_string('2040C0')
        elif name!='normal':
            fill(pr,'none')
            if name.startswith('underline-'):
                r.underline=True
                if name=='underline-red':pr.find(f'{{{W}}}u').set(f'{{{W}}}color','C02020')
            elif name=='strike':r.font.strike=True
            elif name=='outline':outline(pr)
            elif name=='bold':r.bold=True
            elif name=='small-caps':r.font.small_caps=True
        p.add_run(' | Tail')
        cases.append(dict(name=name,text='Abcdef'))
    table=doc.add_table(rows=1,cols=2)
    for i,cell in enumerate(table.rows[0].cells):
        p=cell.paragraphs[0];p.add_run(f'T{i+1} | ');r=p.add_run('Abcdef');fill(r._r.get_or_add_rPr(),'none' if i==0 else '208020');p.add_run(' | Tail')
    doc.add_paragraph('After table')
    doc.save(ROOT/'controls.docx');(ROOT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')

def extra_controls():
    doc=Document(); normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(18)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    order=['empty','2040C0','empty',None,'empty','none','empty','208020','empty']
    for i,value in enumerate(order):
        p=doc.add_paragraph();p.add_run(f'{i+1:02d} | ');r=p.add_run('Abcdef');r.font.color.rgb=RGBColor.from_string('C02020')
        if value is not None:fill(r._r.get_or_add_rPr(),value)
        p.add_run(' | Tail')
    doc.save(ROOT/'empty-order.docx')
    doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(18)
    normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    for ident,value in enumerate(['none','208020','empty'],1):
        doc.add_paragraph(f'Shape {ident}: {value}')
        frame=g.add(g.add(doc.add_paragraph().add_run()._r,'w:drawing'),'wp:inline',distT=0,distB=0,distL=0,distR=0)
        width,height=360*12700,50*12700
        g.add(frame,'wp:extent',cx=width,cy=height);g.add(frame,'wp:docPr',id=ident,name=f'Fill {ident}');g.add(frame,'wp:cNvGraphicFramePr')
        data=g.add(g.add(frame,'a:graphic'),'a:graphicData',uri=g.NS['wps'])
        shape=g.shape(data,ident+100,width,height,0,0,'FFFFFF','placeholder')
        p=shape.find('.//'+g.tag('w:p'))
        for r in list(p.findall(g.tag('w:r'))):p.remove(r)
        for content,fill_value in [('Before ',None),('Abcdef',value),(' after',None)]:
            r=g.add(p,'w:r');rp=g.add(r,'w:rPr');g.add(rp,'w:rFonts',**{'w:ascii':'Arial','w:hAnsi':'Arial'});g.add(rp,'w:sz',**{'w:val':36})
            g.add(rp,'w:color',**{'w:val':'C02020'})
            if fill_value:fill(rp,fill_value)
            t=g.add(r,'w:t',text=content);t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
        body=shape.find(g.tag('wps:bodyPr'))
        for side in ['lIns','rIns','tIns','bIns']:body.set(side,str(3*12700))
    doc.save(ROOT/'shapes.docx')
    doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(18)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    for i,(value,color) in enumerate([(None,'auto'),(None,'C02020'),('none','C02020'),('208020','auto'),('208020','C02020'),('empty','C02020')]):
        p=doc.add_paragraph();p.add_run(f'{i+1:02d} | ');r=p.add_run('Abcdef');r.font.color.rgb=RGBColor.from_string('C02020');r.underline=True
        r._r.get_or_add_rPr().find(qn('w:u')).set(qn('w:color'),color)
        if value:fill(r._r.get_or_add_rPr(),value)
        p.add_run(' | Tail')
    doc.save(ROOT/'decorations.docx')

if __name__=='__main__':generate();extra_controls()
