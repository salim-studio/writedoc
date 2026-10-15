"""Original table-column bookmark controls, independent of corpus documents."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

out=Path(__file__).resolve().parents[1]/'tests/fixtures/column-bookmarks'
out.mkdir(parents=True,exist_ok=True)
doc=Document()
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(0)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
cases=[dict(name='MiddleColumn',first=1,last=1,start=[0,1],end=[2,1]),
       dict(name='TwoColumns',first=0,last=1,start=[0,0],end=[2,1]),
       dict(name='OneCell',first=1,last=1,start=[1,1],end=[1,1]),
       dict(name='MergedTop',first=1,last=1,start=[0,0],end=[2,1],merged=True),
       dict(name='CollapsedCell',first=0,last=0,start=[0,0],end=[0,0],collapsed=True),
       dict(name='CollapsedInline',first=1,last=1,start=[0,1],end=[0,1],collapsed=True,inline=True)]
for index,case in enumerate(cases):
    doc.add_paragraph(case['name'])
    table=doc.add_table(rows=3,cols=3);table.style='Table Grid'
    for row in range(3):
        for col in range(3):table.cell(row,col).text=f'T{index+1}R{row+1}C{col+1}'
    if case.get('merged'):table.cell(0,0).merge(table.cell(0,1))
    start=OxmlElement('w:bookmarkStart')
    for key,value in dict(id=index,name=case['name'],colFirst=case['first'],colLast=case['last']).items():start.set(qn('w:'+key),str(value))
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(index))
    tc=table.cell(*case['start'])._tc
    if case.get('inline'):
        p=table.cell(*case['start']).paragraphs[0]._p;p.insert(0,start);p.insert(1,end)
    else:
        tc.insert(1,start)
        if case.get('collapsed'):tc.insert(2,end)
        else:table.cell(*case['end'])._tc.append(end)
    doc.add_paragraph('After '+case['name'])
doc.save(out/'control.docx')
for n in list(doc.element.iter()):
    if n.tag in [qn('w:bookmarkStart'),qn('w:bookmarkEnd')] and int(n.get(qn('w:id'))) < 4:n.getparent().remove(n)
doc.save(out/'collapsed.docx')
(out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
