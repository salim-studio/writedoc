"""Original exact-row fixture and independent export checks.

Generated using python-docx (MIT), lxml (BSD), and original XML/PNG bytes.
No user documents or third-party artwork are redistributed.
"""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
import struct
import zlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/exact-row-heights.docx'

def child(parent, name, **attrs):
    node = OxmlElement('w:' + name)
    for key, value in attrs.items(): node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node

def row_height(row, points, rule=WD_ROW_HEIGHT_RULE.EXACTLY):
    row.height = Pt(points); row.height_rule = rule

def margins(cell, top=0, bottom=0):
    node = child(cell._tc.get_or_add_tcPr(), 'tcMar')
    for side, value in [('top', top), ('left', 120), ('bottom', bottom), ('right', 120)]:
        child(node, side, w=value, type='dxa')

def generate():
    doc = Document(); doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('Canvas Exact Rows', WD_STYLE_TYPE.TABLE); base.base_style = doc.styles['Table Grid']
    child(child(base.element, 'trPr'), 'trHeight', val=480, hRule='exact')
    doc.add_paragraph('Exact row heights Á😀')
    def table(rows, style='Table Grid'):
        result = doc.add_table(rows=rows, cols=2, style=style); result.autofit = False
        for col in result.columns: col.width = Inches(2.5)
        for row in result.rows:
            for cell in row.cells: cell.width = Inches(2.5); margins(cell)
        return result
    first = table(6, base)
    first.cell(0, 0).text = 'Exact visible Á😀\nHIDDEN ONE\nHIDDEN TWO'
    first.cell(0, 1).text = 'Padded neighbor\nHIDDEN NEIGHBOR'
    for cell, bottom in zip(first.rows[0].cells, [90, 150]):
        cell._tc.tcPr.remove(cell._tc.tcPr.find(qn('w:tcMar'))); margins(cell, top=60, bottom=bottom)
    for index, text in [(1, 'Following automatic row'), (3, 'Below clipped picture'), (5, 'Below clipped decorations')]:
        row_height(first.rows[index], 0, WD_ROW_HEIGHT_RULE.AUTO)
        first.cell(index, 0).text = text; first.cell(index, 1).text = 'Unchanged neighbor'
    def chunk(t, data): return struct.pack('>I', len(data)) + t + data + struct.pack('>I', zlib.crc32(t + data) & 0xffffffff)
    png = b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 2, 2, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(b'\0\xe0\x20\x70\xe0\x20\x70' * 2)) + chunk(b'IEND', b'')
    row_height(first.rows[2], 36)
    first.cell(2, 0).paragraphs[0].add_run().add_picture(BytesIO(png), width=Inches(1), height=Inches(1))
    first.cell(2, 1).text = 'Picture clipped here'
    row_height(first.rows[4], 24)
    p = first.cell(4, 0).paragraphs[0]
    p.add_run('Decorations\nCLIPPED DECORATIONS\nMORE OVERFLOW').font.color.rgb = RGBColor.from_string('104080')
    p.runs[0].underline = True; p.runs[0].font.strike = True
    borders = child(p._p.get_or_add_pPr(), 'pBdr')
    for side in ['top', 'left', 'bottom', 'right']: child(borders, side, val='single', sz=4, space=1, color='B02080')
    child(p._p.get_or_add_pPr(), 'shd', val='clear', fill='F4D260')
    first.cell(4, 1).text = 'Centered short text'; first.cell(4, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph('All exact merged rows')
    merged = table(3)
    for i, row in enumerate(merged.rows): row_height(row, 18); row.cells[1].text = f'Fixed neighbor {i + 1}'
    merged.cell(0, 0).merge(merged.cell(2, 0)).text = '\n'.join(['Fixed merged owner'] + [f'MERGED HIDDEN {i}' for i in range(1, 8)])
    merged.cell(0, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    doc.add_paragraph('Mixed exact and automatic merged rows')
    mixed = table(3)
    for i, row in enumerate(mixed.rows):
        if i != 1: row_height(row, 12)
        row.cells[1].text = f'Mixed neighbor {i + 1}'
    mixed.cell(0, 0).merge(mixed.cell(2, 0)).text = '\n'.join(['Mixed merged owner'] + [f'Flexible line {i}' for i in range(1, 9)])
    doc.add_paragraph('Repeated exact header and a tall exact row')
    flowing = table(3)
    child(flowing.rows[0]._tr.get_or_add_trPr(), 'tblHeader'); row_height(flowing.rows[0], 18)
    flowing.cell(0, 0).text = 'Exact header\nHIDDEN HEADER'; flowing.cell(0, 1).text = 'Repeated neighbor'
    row_height(flowing.rows[1], 900)
    flowing.cell(1, 0).text = '\n'.join(f'Tall fixed line {i}' for i in range(100))
    flowing.cell(1, 1).text = 'Bottom aligned'; flowing.cell(1, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    flowing.cell(2, 0).text = 'After tall fixed row'; flowing.cell(2, 1).text = 'Final neighbor'
    doc.add_paragraph('Exact-row tail'); doc.save(SOURCE)
    print(f'Generated {SOURCE}')

def verify(path, edited=False):
    original = Document(SOURCE); exported = Document(path)
    # xml:space="preserve" on text nodes is an equivalent explicit spelling for
    # this fixture's text, including its internal spaces and line-break nodes.
    def data(node): return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space'))), n.text) for n in node.iter()]
    if edited:
        row_height(original.tables[0].rows[0], 36)
        # Browser editing replaces one complete plain-text run fragment.
        node = next(n for n in original.element.iter(qn('w:t')) if n.text == 'HIDDEN TWO')
        node.text = 'REVISED HIDDEN Á😀'
        assert data(exported.element) == data(original.element), 'Unexpected document content, formatting or structure change'
    with ZipFile(SOURCE) as a, ZipFile(path) as b:
        assert set(a.namelist()) == set(b.namelist())
        for name in a.namelist():
            if not edited or name != 'word/document.xml': assert a.read(name) == b.read(name), name
    print(f'PASS {path}: exact rows, retained overflow, merged cells, headers, formatting and package parts')

if __name__ == '__main__':
    import sys
    if len(sys.argv) == 1: generate()
    else: verify(sys.argv[1], '--edited' in sys.argv[2:])
