"""Original native underline controls; no corpus content or font binaries."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures/native-underlines'

def generate():
    ROOT.mkdir(parents=True,exist_ok=True)
    doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(10);normal.paragraph_format.line_spacing=1
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    cases=[]
    for family in ['Arial','Calibri','Times New Roman','Cambria']:
        for bold,italic in [(False,False),(True,False),(False,True),(True,True)]:
            page=len(cases)//6
            title=doc.add_paragraph(f'{family}: bold={bold}, italic={italic}')
            title.paragraph_format.page_break_before=page>0
            for size in [8,10,12,18,24,36]:
                ident=len(cases)+1;p=doc.add_paragraph();p.add_run(f'{ident:03d} | ')
                r=p.add_run('Underline sample');r.font.name=family;r.font.size=Pt(size);r.bold=bold;r.italic=italic;r.underline=True
                r._r.get_or_add_rPr().find(qn('w:u')).set(qn('w:color'),'008080')
                p.add_run(' | end')
                cases.append(dict(id=ident,page=page,font=family,bold=bold,italic=italic,size=size))
    doc.save(ROOT/'controls.docx');(ROOT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')

def patterns():
    names=['single','words','double','thick','dotted','dottedHeavy','dash','dashedHeavy','dashLong','dashLongHeavy','dotDash','dashDotHeavy','dotDotDash','dashDotDotHeavy','wave','wavyHeavy','wavyDouble']
    doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    for bold in [False,True]:
        p=doc.add_paragraph(f'Bold={bold}');p.paragraph_format.page_break_before=bold
        for name in names:
            p=doc.add_paragraph(name+' | ');r=p.add_run('Underline sample');r.font.size=Pt(18);r.bold=bold;r.underline=True
            u=r._r.get_or_add_rPr().find(qn('w:u'));u.set(qn('w:val'),name);u.set(qn('w:color'),'008080')
    doc.save(ROOT/'patterns.docx');(ROOT/'patterns.json').write_text(json.dumps(names,indent=2)+'\n')

if __name__=='__main__':generate();patterns()
