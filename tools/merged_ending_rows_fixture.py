from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE as H,WD_CELL_VERTICAL_ALIGNMENT as V
import json,sys
from docx.oxml.ns import qn
mode=int(sys.argv[1]) if len(sys.argv)>1 else 14
folder=Path(__file__).resolve().parents[1]/'tests/fixtures/merged-ending-rows'
folder.mkdir(exist_ok=True)
out=folder/f'controls-{mode}.docx'
d=Document();s=d.sections[0];s.page_width=Inches(6);s.page_height=Inches(6);s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(.5)
d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.space_after=Pt(0)
cases=[]
for nested in [False,True]:
 for rules in ['AA','AE','EA','EE','AEA','AAE','EAA','EAE']:
  i=len(cases)
  if i:d.add_page_break()
  d.add_paragraph(f'Case {i+1} {rules} '+('nested' if nested else 'text'))
  t=d.add_table(rows=len(rules),cols=2);t.autofit=False
  for j,(r,rule) in enumerate(zip(t.rows,rules)):
   r.height=Pt(40);r.height_rule=H.EXACTLY if rule=='E' else H.AT_LEAST
   t.cell(j,0).text=f'ROW{j+1}';t.cell(j,0).vertical_alignment=V.BOTTOM
  c=t.cell(0,1).merge(t.cell(len(rules)-1,1))
  if nested:
   c.text='';n=c.add_table(rows=1,cols=1);n.rows[0].height=Pt(170);n.rows[0].height_rule=H.AT_LEAST;n.cell(0,0).text='MERGED';n.cell(0,0).vertical_alignment=V.BOTTOM
  else:
   c.text='MERGED';c.paragraphs[0].paragraph_format.line_spacing=Pt(170)
  d.add_paragraph('AFTER');cases.append({'nested':nested,'rules':rules})
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
compat=d.settings.element.find('.//' + qn('w:compatSetting'))
compat.set(qn('w:val'),str(mode))
d.save(out);out.with_suffix('.json').write_text(json.dumps(cases))


