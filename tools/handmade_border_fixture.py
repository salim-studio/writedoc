"""Independent page-art border sizing/placement controls; no corpus artwork."""
from pathlib import Path
import json
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1] / 'tests/fixtures/handmade-borders'
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(12)
cases = []
for index, (size, origin) in enumerate([(8, 'text'), (21, 'text'), (40, 'text'), (21, 'page')]):
    section = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    for old in list(section._sectPr.findall(qn('w:pgBorders'))):
        section._sectPr.remove(old)
    borders = OxmlElement('w:pgBorders')
    borders.set(qn('w:offsetFrom'), origin)
    for side in ['top', 'left', 'bottom', 'right']:
        node = OxmlElement('w:' + side)
        for key, value in dict(val='handmade1', sz=size, space=12, color='0070C0').items():
            node.set(qn('w:' + key), str(value))
        borders.append(node)
    section._sectPr.insert_element_before(borders, 'w:lnNumType', 'w:pgNumType', 'w:cols', 'w:formProt', 'w:vAlign', 'w:noEndnote', 'w:titlePg', 'w:textDirection', 'w:bidi', 'w:rtlGutter', 'w:docGrid')
    label = f'C{index:02} handmade1 size={size}, origin={origin}'
    doc.add_paragraph(label)
    doc.add_paragraph('Independent page-border control. No template art is included.')
    cases.append(dict(label=label, size=size, origin=origin, space=12, color='0070C0'))
doc.save(root / 'control.docx')
(root / 'cases.json').write_text(json.dumps(cases, indent=2) + '\n')
