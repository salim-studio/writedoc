"""Original vertical-merge, omitted-grid and nested-table bookmark controls."""
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

out = Path(__file__).resolve().parents[1] / 'tests/fixtures/column-bookmarks'
doc = Document()
doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(0)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
for index, name in enumerate(['VerticalMiddle', 'StartInContinuation', 'VerticalFirst', 'NestedMiddle', 'OmittedFirst', 'UnequalGrid']):
    if index:
        doc.add_page_break()
    doc.add_paragraph(name)
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f'R{r+1}C{c+1}'
    start_row = 0
    if name in ['VerticalMiddle', 'StartInContinuation']:
        table.cell(0, 1).merge(table.cell(1, 1))
        if name == 'StartInContinuation':
            start_row = 1
    if name == 'VerticalFirst':
        table.cell(0, 0).merge(table.cell(1, 0))
    if name == 'NestedMiddle':
        nested = table.cell(1, 1).add_table(rows=1, cols=2)
        nested.cell(0, 0).text = 'InnerA'
        nested.cell(0, 1).text = 'InnerB'
    if name == 'OmittedFirst':
        row = table.rows[0]._tr
        row.remove(row.findall(qn('w:tc'))[0])
        before = OxmlElement('w:gridBefore'); before.set(qn('w:val'), '1')
        row.get_or_add_trPr().append(before)
    if name == 'UnequalGrid':
        for i, width in enumerate([1200, 3000, 4440]):
            table._tbl.tblGrid.gridCol_lst[i].set(qn('w:w'), str(width))
            for row in table.rows:
                row.cells[i]._tc.tcPr.tcW.set(qn('w:w'), str(width))
    start = OxmlElement('w:bookmarkStart')
    for key, value in dict(id=index, name=name, colFirst=1, colLast=1).items():
        start.set(qn('w:' + key), str(value))
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(index))
    table.rows[start_row]._tr.findall(qn('w:tc'))[0].insert(1, start)
    table._tbl.append(end)
    doc.add_paragraph('After ' + name)
doc.save(out / 'geometry-control.docx')
