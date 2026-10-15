"""Original OOXML plain-text controls for native editing behavior checks."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];folder=root/'tests/fixtures/plain-text-controls';folder.mkdir(exist_ok=True)
cases=[('temp-format',True,False,False),('persistent-format',False,False,False),('temp-caret-format',True,False,False),('persistent-caret-format',False,False,False),('temp-insert',True,False,False),('persistent-insert',False,False,False),('temp-enter',True,False,False),('persistent-enter',False,False,False),('multi-enter',False,True,False),('block-temp-format',True,True,True),('block-temp-insert',True,True,True),('existing-break',True,False,False)]
d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for i,(name,temporary,multi,block) in enumerate(cases):
 p=d.add_paragraph(name+': ');r=p.add_run('Alpha Bravo');r.font.color.rgb=None
 if name=='existing-break':r.add_break();r.add_text('Charlie')
 s=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');s.append(pr)
 for tag,val in [('alias',name),('id',str(i+1))]:
  n=OxmlElement('w:'+tag);n.set(qn('w:val'),val);pr.append(n)
 if temporary:pr.append(OxmlElement('w:temporary'))
 t=OxmlElement('w:text');t.set(qn('w:multiLine'),'1' if multi else '0');pr.append(t)
 content=OxmlElement('w:sdtContent');s.append(content)
 if block:
  body=p._p.getparent();body.replace(p._p,s);content.append(p._p)
 else:
  p._p.replace(r._r,s);content.append(r._r)
d.save(folder/'control.docx')
(folder/'cases.json').write_text(json.dumps([dict(name=n,temporary=t,multiLine=m,block=b) for n,t,m,b in cases],indent=2)+'\n')
