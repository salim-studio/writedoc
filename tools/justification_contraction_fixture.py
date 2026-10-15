"""Regenerate original contraction controls from explicit case parameters."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json
from lxml import etree as E
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / 'tests/fixtures/justification-contraction'

def build(name, spec, folder=FOLDER):
    d = Document()
    s = d.sections[0]
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
    d.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    for setting in d.settings.element.iter(qn('w:compatSetting')):
        if setting.get(qn('w:name')) == 'compatibilityMode':
            setting.set(qn('w:val'), str(spec.get('mode', 15)))
    n = d.styles['Normal']
    n.font.name = 'Arial'
    n.font.size = Pt(10)
    n.paragraph_format.space_after = Pt(0)
    n.paragraph_format.line_spacing = 1
    for i, case in enumerate(spec['cases']):
        p = d.add_paragraph(case['name'])
        p.paragraph_format.page_break_before = i > 0
        p = d.add_paragraph()
        p.paragraph_format.right_indent = Pt(468 - case['width'])
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if case.get('align') == 'left' else WD_ALIGN_PARAGRAPH.JUSTIFY
        runs = case.get('runs') or [dict(text=case['text'], family=case['family'], size=10)]
        if spec.get('joined'):
            runs = [dict(runs[0], text=''.join(r['text'] for r in runs))]
        for part in runs:
            r = p.add_run(part['text'])
            r.font.name = part['family']
            r.font.size = Pt(part['size'])
            if spec.get('joined'):
                r._r.find(qn('w:t')).set(qn('xml:space'), 'preserve')
    if spec.get('noExpand'):
        d.settings.element.find(qn('w:compat')).insert(0, OxmlElement('w:doNotExpandShiftReturn'))
    folder.mkdir(parents=True, exist_ok=True)
    destination = folder / (name + '.docx')
    d.save(destination)
    # Match the XML declaration choices of the independently exported packages.
    if spec.get('stripDeclarations'):
        with ZipFile(destination) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        for part in spec['stripDeclarations']:
            parts[part] = E.tostring(E.fromstring(parts[part]))
        with ZipFile(destination, 'w', ZIP_DEFLATED) as z:
            for part, data in parts.items():
                z.writestr(part, data)

if __name__ == '__main__':
    for name, spec in json.loads((FOLDER / 'cases.json').read_text()).items():
        build(name, spec)
