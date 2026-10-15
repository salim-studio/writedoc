"""Original tab-leader fixture and independent checks using python-docx."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/tab-leaders.docx'

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('Canvas Leader Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    base.paragraph_format.tab_stops.add_tab_stop(Inches(4.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    child = doc.styles.add_style('Canvas Leader Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    doc.add_paragraph('Tab leaders A\u0301😀')
    doc.add_paragraph('Inherited contents\t12', child)
    p = doc.add_paragraph('Direct none\t34', child)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(4.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)
    p._p.pPr.tabs[0].set(qn('w:leader'), 'none')
    for name, leader in [('dot', WD_TAB_LEADER.DOTS), ('hyphen', WD_TAB_LEADER.DASHES), ('underscore', WD_TAB_LEADER.LINES), ('heavy', WD_TAB_LEADER.HEAVY), ('middleDot', WD_TAB_LEADER.MIDDLE_DOT)]:
        p = doc.add_paragraph(f'{name} price\t123.45')
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.7), WD_TAB_ALIGNMENT.DECIMAL, leader)
    p = doc.add_paragraph('Two stops\tcenter\tend')
    p.paragraph_format.tab_stops.add_tab_stop(Inches(2), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.DOTS)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(4.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DASHES)
    p = doc.add_paragraph('Styled tab')
    r = p.add_run('\t'); r.bold = True; r.italic = True; r.font.size = Pt(16); r.font.color.rgb = RGBColor.from_string('993366')
    p.add_run('Unicode A\u0301😀'); p.paragraph_format.tab_stops.add_tab_stop(Inches(4.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.MIDDLE_DOT)
    doc.add_paragraph('Default\ttab')
    p = doc.add_paragraph('Clear inherited\tstop', child)
    stop = OxmlElement('w:tab'); stop.set(qn('w:val'), 'clear'); stop.set(qn('w:pos'), '6480'); p._p.get_or_add_pPr().get_or_add_tabs().append(stop)
    p = doc.add_paragraph('Numbered leader', 'List Number')
    p.paragraph_format.tab_stops.add_tab_stop(Inches(.25), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS)
    p = doc.add_paragraph('Wrapped\t' + 'text follows the tab and wraps naturally. ' * 8)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(1.6), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS)
    table = doc.add_table(rows=2, cols=2, style='Table Grid'); table.autofit = False
    for col in table.columns: col.width = Inches(2.5)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(2.5)
            cell.paragraphs[0].paragraph_format.tab_stops.add_tab_stop(Inches(2.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    table.cell(0, 0).text = 'Header\t1'; table.cell(0, 1).text = 'Price\t2'
    # Assigning cell.text replaces its paragraph properties.
    for cell in table.rows[0].cells: cell.paragraphs[0].paragraph_format.tab_stops.add_tab_stop(Inches(2.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    table.cell(1, 0).paragraphs[0].add_run('Flowing\t3\n' * 30 + 'End\t4')
    table.cell(1, 1).paragraphs[0].add_run('Neighbor\t5')
    doc.add_paragraph('Tail'); doc.save(SOURCE); print(f'Generated independent tab-leader fixture: {SOURCE}')

def verify(file, mode):
    doc = Document(file)
    if mode == 'new':
        assert [p.text for p in doc.paragraphs] == ['New contents\t42', 'Second\t99']
        for p in doc.paragraphs:
            stops = p.paragraph_format.tab_stops
            assert len(stops) == 1 and stops[0].position.twips == 4320
            assert stops[0].alignment == WD_TAB_ALIGNMENT.RIGHT and stops[0].leader == WD_TAB_LEADER.DOTS
    else:
        before = Document(SOURCE).element.findall('.//' + qn('w:p')); after = doc.element.findall('.//' + qn('w:p'))
        text = lambda p: ''.join(n.text or '' for n in p.findall('.//' + qn('w:t')))
        assert [text(p) for p in after] == [text(p) for p in before]
        def data(element):
            # Rewritten text receives xml:space="preserve"; text and every run
            # property must remain identical, including tabs and break kinds.
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space'))), n.text) for n in element.iter()]
        for p, old in zip(after, before):
            expected = 'hyphen' if text(p).startswith('Inherited contents') else 'middleDot' if text(p).startswith('Direct none') else None
            if expected:
                tabs = p.find(qn('w:pPr')).find(qn('w:tabs'))
                assert len(tabs) == 1 and tabs[0].attrib == {qn('w:pos'): '6480', qn('w:val'): 'right', qn('w:leader'): expected}
                assert [data(n) for n in p.find(qn('w:pPr')) if n.tag != qn('w:tabs')] == [data(n) for n in old.find(qn('w:pPr')) if n.tag != qn('w:tabs')]
                assert [data(r) for r in p.findall(qn('w:r'))] == [data(r) for r in old.findall(qn('w:r'))]
            else: assert data(p) == data(old), text(p)
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: leaders, tab alignment, source text, rich runs and unrelated package parts')
