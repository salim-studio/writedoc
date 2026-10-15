"""Original TIFF controls; keep the TIFFs embedded rather than converting them."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import json
from PIL import Image, ImageDraw, TiffImagePlugin
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from lxml import etree as E

root=Path(__file__).resolve().parents[1];folder=root/'tests/fixtures/tiff-pictures';folder.mkdir(exist_ok=True)
base=Image.new('RGB',(90,60));draw=ImageDraw.Draw(base)
for i,color in enumerate(['#ff0000','#00ff00','#0000ff','#ffff00','#ffffff','#000000']):
 draw.rectangle((i%3*30,i//3*30,i%3*30+29,i//3*30+29),fill=color)
cases=[{'name':f'orientation-{i}','orientation':i,'mode':'RGB','compression':'raw'} for i in range(1,9)]
cases += [{'name':name,'orientation':1,'mode':mode,'compression':compression} for name,mode,compression in [('lzw','RGB','tiff_lzw'),('deflate','RGB','tiff_adobe_deflate'),('packbits','RGB','packbits'),('gray','L','raw'),('bilevel','1','group4'),('palette','P','raw'),('alpha','RGBA','raw'),('cmyk','CMYK','raw'),('multipage','RGB','raw')]]
d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Inches(.5)
n=d.styles['Normal'];n.font.name='Arial';n.font.size=Pt(10);n.paragraph_format.space_after=Pt(0)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
payloads={}
for i,c in enumerate(cases):
 if i and i%4==0:d.add_page_break()
 d.add_paragraph(c['name'])
 # Unique PNG placeholders prevent python-docx deduplicating the picture parts.
 placeholder=base.copy();placeholder.putpixel((0,0),(i,0,0));png=BytesIO();placeholder.save(png,format='PNG')
 d.add_paragraph().add_run().add_picture(BytesIO(png.getvalue()),width=Pt(144),height=Pt(72))
 pic=base.convert(c['mode'],dither=Image.Dither.NONE);info=TiffImagePlugin.ImageFileDirectory_v2();info[274]=c['orientation']
 if c['mode']=='RGBA':pic.putalpha(128)
 data=BytesIO();opts={'format':'TIFF','compression':c['compression'],'tiffinfo':info}
 if c['name']=='multipage':opts.update(save_all=True,append_images=[Image.new('RGB',base.size,'#ff00ff')])
 pic.save(data,**opts);name=f'image{i+1}.tiff';payloads[name]=data.getvalue();(folder/name).write_bytes(data.getvalue())
buffer=BytesIO();d.save(buffer)
with ZipFile(BytesIO(buffer.getvalue())) as z:parts={n:z.read(n) for n in z.namelist()}
r=E.fromstring(parts['word/_rels/document.xml.rels'])
for relation in r:
 target=relation.get('Target')
 if target and target.startswith('media/image'):
  old='word/'+target;new=target.rsplit('.',1)[0]+'.tiff';relation.set('Target',new);parts.pop(old);parts['word/'+new]=payloads[Path(new).name]
parts['word/_rels/document.xml.rels']=E.tostring(r,xml_declaration=True,encoding='UTF-8')
t=E.fromstring(parts['[Content_Types].xml']);E.SubElement(t,'{'+E.QName(t).namespace+'}Default',Extension='tiff',ContentType='image/tiff');parts['[Content_Types].xml']=E.tostring(t,xml_declaration=True,encoding='UTF-8')
with ZipFile(folder/'control.docx','w',ZIP_DEFLATED) as z:
 for n,data in parts.items():z.writestr(n,data)
(folder/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
