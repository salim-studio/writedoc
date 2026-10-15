"""Original controls: numbering cancellation within a paragraph style chain."""
from pathlib import Path
from io import BytesIO
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
d=Document();d.sections[0].left_margin=d.sections[0].right_margin=Inches(1)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
b=BytesIO();Image.new('RGB',(40,10),'#1255aa').save(b,format='PNG')
cases=[]
for numbered in [False,True]:
 base=d.styles.add_style('Base'+str(numbered),WD_STYLE_TYPE.PARAGRAPH);base.base_style=d.styles['List Bullet' if numbered else 'Normal']
 base.paragraph_format.left_indent=Pt(18);base.paragraph_format.right_indent=Pt(12);base.paragraph_format.first_line_indent=Pt(-18)
 if numbered:base.element.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val=1
 for variant in ['none','sameLeft','sameFirst','childLeft','childFirst','directLeft']:
  cancel=d.styles.add_style('Cancel'+str(numbered)+variant,WD_STYLE_TYPE.PARAGRAPH);cancel.base_style=base
  cancel.element.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val=0
  if variant=='sameLeft':cancel.paragraph_format.left_indent=Pt(30)
  if variant=='sameFirst':cancel.paragraph_format.first_line_indent=Pt(6)
  child=d.styles.add_style('Child'+str(numbered)+variant,WD_STYLE_TYPE.PARAGRAPH);child.base_style=cancel
  if variant=='childLeft':child.paragraph_format.left_indent=Pt(24)
  if variant=='childFirst':child.paragraph_format.first_line_indent=Pt(9)
  if cases:d.add_page_break()
  d.add_paragraph(f'numbered={numbered} {variant}')
  p=d.add_paragraph(style=child)
  if variant=='directLeft':p.paragraph_format.left_indent=Pt(36)
  for i in range(2):
   if i:p.add_run().add_break()
   p.add_run().add_picture(BytesIO(b.getvalue()),width=Pt(40))
  cases.append(variant)
d.save(ROOT/'public/fixtures/numbering-cancel-style.docx')
