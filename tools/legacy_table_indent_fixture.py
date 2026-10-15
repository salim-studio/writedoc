"""Original table-indent controls; no application layout code is used."""
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
(ROOT/'tests/fixtures/legacy-table-indent').mkdir(parents=True,exist_ok=True)
(ROOT/'tests/fixtures/legacy-table-indent-12').mkdir(parents=True,exist_ok=True)
(ROOT/'tests/fixtures/legacy-table-indent-15').mkdir(parents=True,exist_ok=True)
(ROOT/'tests/fixtures/legacy-table-origin').mkdir(parents=True,exist_ok=True)
(ROOT/'tests/fixtures/legacy-table-origin-no-styles').mkdir(parents=True,exist_ok=True)
(ROOT/'tests/fixtures/legacy-table-origin-custom-style').mkdir(parents=True,exist_ok=True)
(ROOT/'tests/fixtures/table-style-origin-properties').mkdir(parents=True,exist_ok=True)

# make-legacy-table-indent.py
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from io import BytesIO
from PIL import Image
from pathlib import Path
import json
d=Document(); d.sections[0].left_margin=d.sections[0].right_margin=Inches(1)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
b=BytesIO();Image.new('RGB',(30,15),'#1255aa').save(b,format='PNG')
cases=[]
for width in ['auto','dxa']:
 for indent in [None,0,360]:
  for margin in [108,216]:
   i=len(cases);cases.append(dict(width=width,indent=indent,margin=margin))
   if i:d.add_page_break()
   d.add_paragraph(str(cases[-1]));t=d.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=t.cell(0,0).width=Pt(237)
   pr=t._tbl.tblPr;tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),width);tw.set(qn('w:w'),'0' if width=='auto' else '4740')
   if indent is not None:
    n=OxmlElement('w:tblInd');n.set(qn('w:w'),str(indent));n.set(qn('w:type'),'dxa');pr.insert_element_before(n,'w:tblBorders','w:shd','w:tblLayout','w:tblCellMar','w:tblLook')
   m=OxmlElement('w:tblCellMar')
   for side in ['left','right']:
    n=OxmlElement('w:'+side);n.set(qn('w:w'),str(margin));n.set(qn('w:type'),'dxa');m.append(n)
   pr.insert_element_before(m,'w:tblLook');t.cell(0,0).paragraphs[0].add_run().add_picture(BytesIO(b.getvalue()),width=Pt(30))
d.save(ROOT/'public/fixtures/legacy-table-indent.docx');(ROOT/'tests/fixtures/legacy-table-indent/cases.json').write_text(json.dumps(cases))

# make-legacy-table-origin.py
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy
from pathlib import Path
import json
d=Document(ROOT/'public/fixtures/legacy-table-indent.docx');table=deepcopy(d.tables[0]._tbl)
body=d._element.body
for n in list(body):
 if n.tag!=qn('w:sectPr'):body.remove(n)
cases=[]
for border in [0,32]:
 for align in [None,'left','center','right']:
  if cases:d.add_page_break()
  cases.append(dict(border=border,align=align));d.add_paragraph(str(cases[-1]));t=deepcopy(table);pr=t.find(qn('w:tblPr'))
  if align:
   n=OxmlElement('w:jc');n.set(qn('w:val'),align);pr.insert_element_before(n,'w:tblInd','w:tblBorders','w:tblLayout','w:tblCellMar','w:tblLook')
  if border:
   b=OxmlElement('w:tblBorders')
   for side in ['top','left','bottom','right']:
    n=OxmlElement('w:'+side);n.set(qn('w:val'),'single');n.set(qn('w:sz'),str(border));n.set(qn('w:color'),'009900');b.append(n)
   pr.insert_element_before(b,'w:tblLayout','w:tblCellMar','w:tblLook')
  t.xpath('.//wp:docPr')[0].set('id',str(len(cases)))
  body.insert(-1,t)
d.save(ROOT/'public/fixtures/legacy-table-origin.docx');(ROOT/'tests/fixtures/legacy-table-origin/cases.json').write_text(json.dumps(cases))

# make-table-indent-modes.py
from docx import Document
from docx.oxml.ns import qn
for mode in [12,15]:
 d=Document(ROOT/'public/fixtures/legacy-table-indent.docx')
 for c in d.settings.element.xpath('.//w:compatSetting'):
  if c.get(qn('w:name'))=='compatibilityMode':c.set(qn('w:val'),str(mode))
 d.save(ROOT/f'public/fixtures/legacy-table-indent-{mode}.docx')

# make-no-table-styles.py
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
with ZipFile(ROOT/'public/fixtures/legacy-table-origin.docx') as z:parts={n:z.read(n) for n in z.namelist() if n not in ['word/styles.xml','word/stylesWithEffects.xml']}
for path,attr in [('word/_rels/document.xml.rels','Type'),('[Content_Types].xml','ContentType')]:
 r=E.fromstring(parts[path])
 for n in list(r):
  if 'styles' in n.get(attr,''):r.remove(n)
 parts[path]=E.tostring(r,encoding='UTF-8',xml_declaration=True)
with ZipFile(ROOT/'public/fixtures/legacy-table-origin-no-styles.docx','w',ZIP_DEFLATED) as z:
 for n,v in parts.items():z.writestr(n,v)

# make-custom-table-style.py
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
with ZipFile(ROOT/'public/fixtures/legacy-table-origin-no-styles.docx') as z:parts={n:z.read(n) for n in z.namelist()}
parts['word/styles.xml']=f'<w:styles xmlns:w="{W}"><w:style w:type="table" w:default="1" w:customStyle="1" w:styleId="ProbeTable"><w:name w:val="Probe table"/></w:style></w:styles>'.encode()
r=E.fromstring(parts['word/_rels/document.xml.rels']);E.SubElement(r,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='probeStyles',Type=R+'/styles',Target='styles.xml');parts['word/_rels/document.xml.rels']=E.tostring(r)
r=E.fromstring(parts['[Content_Types].xml']);E.SubElement(r,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',PartName='/word/styles.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml');parts['[Content_Types].xml']=E.tostring(r)
with ZipFile(ROOT/'public/fixtures/legacy-table-origin-custom-style.docx','w',ZIP_DEFLATED) as z:
 for n,v in parts.items():z.writestr(n,v)

# make-style-origin-properties.py
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy
from pathlib import Path
import json
d=Document(ROOT/'public/fixtures/legacy-table-origin.docx');tbase=deepcopy(d.tables[0]._tbl);body=d._element.body
for n in list(body):
 if n.tag!=qn('w:sectPr'):body.remove(n)
for i,kind in enumerate(['empty','indent','zero-margins','margins','based-on']):
 if i:d.add_page_break()
 d.add_paragraph(kind);style=d.styles.add_style('Origin '+kind,WD_STYLE_TYPE.TABLE)
 if kind=='based-on':style.base_style=d.styles['Normal Table']
 if kind in ['indent','zero-margins','margins']:
  pr=OxmlElement('w:tblPr');style.element.append(pr)
  if kind=='indent':
   n=OxmlElement('w:tblInd');n.set(qn('w:w'),'0');n.set(qn('w:type'),'dxa');pr.append(n)
  else:
   m=OxmlElement('w:tblCellMar');pr.append(m)
   for side in ['left','right']:
    n=OxmlElement('w:'+side);n.set(qn('w:w'),'0' if kind=='zero-margins' else '108');n.set(qn('w:type'),'dxa');m.append(n)
 t=deepcopy(tbase);pr=t.find(qn('w:tblPr'));n=OxmlElement('w:tblStyle');n.set(qn('w:val'),style.style_id);pr.insert(0,n);t.xpath('.//wp:docPr')[0].set('id',str(i+1));body.insert(-1,t)
d.save(ROOT/'public/fixtures/table-style-origin-properties.docx')

(ROOT/'tests/fixtures/legacy-table-indent-no-styles').mkdir(parents=True,exist_ok=True)

(ROOT/'tests/fixtures/legacy-table-origin-asymmetric').mkdir(parents=True,exist_ok=True)

(ROOT/'tests/fixtures/legacy-table-origin-no-styles-asymmetric').mkdir(parents=True,exist_ok=True)

(ROOT/'tests/fixtures/legacy-table-origin-left-asymmetric').mkdir(parents=True,exist_ok=True)

# make-no-indent-styles.py
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
with ZipFile(ROOT/'public/fixtures/legacy-table-indent.docx') as z:parts={n:z.read(n) for n in z.namelist() if n not in ['word/styles.xml','word/stylesWithEffects.xml']}
for path,attr in [('word/_rels/document.xml.rels','Type'),('[Content_Types].xml','ContentType')]:
 r=E.fromstring(parts[path])
 for n in list(r):
  if 'styles' in n.get(attr,''):r.remove(n)
 parts[path]=E.tostring(r,encoding='UTF-8',xml_declaration=True)
with ZipFile(ROOT/'public/fixtures/legacy-table-indent-no-styles.docx','w',ZIP_DEFLATED) as z:
 for n,v in parts.items():z.writestr(n,v)


# make-asymmetric-origin.py
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
from pathlib import Path
repo=ROOT;W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
for suffix in ['','-no-styles']:
 name='legacy-table-origin'+suffix
 with ZipFile(repo/'public/fixtures'/f'{name}.docx') as z:parts={n:z.read(n) for n in z.namelist()}
 r=E.fromstring(parts['word/document.xml'])
 for n in r.iter(W+'tblCellMar'):n.find(W+'right').set(W+'w','216')
 parts['word/document.xml']=E.tostring(r,encoding='UTF-8',xml_declaration=True)
 with ZipFile(ROOT/f'public/fixtures/{name}-asymmetric.docx','w',ZIP_DEFLATED) as z:
  for n,v in parts.items():z.writestr(n,v)

# make-left-asymmetric-origin.py
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
from pathlib import Path
repo=ROOT;W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
for suffix in ['']:
 name='legacy-table-origin'+suffix
 with ZipFile(repo/'public/fixtures'/f'{name}.docx') as z:parts={n:z.read(n) for n in z.namelist()}
 r=E.fromstring(parts['word/document.xml'])
 for n in r.iter(W+'tblCellMar'):n.find(W+'left').set(W+'w','216')
 parts['word/document.xml']=E.tostring(r,encoding='UTF-8',xml_declaration=True)
 with ZipFile(ROOT/f'public/fixtures/{name}-left-asymmetric.docx','w',ZIP_DEFLATED) as z:
  for n,v in parts.items():z.writestr(n,v)

