from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

d=Document(); s=d.sections[0]
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.left_margin=s.right_margin=s.top_margin=s.bottom_margin=Inches(1)
n=d.styles['Normal']; n.font.name='Arial'; n.font.size=Pt(10)
n.paragraph_format.space_after=Pt(0); n.paragraph_format.line_spacing=1
cases=[]
for floating in [False, True]:
 for name,text in [('empty',''),('space',' '),('spaces','  '),('tab','\t'),('nbsp','\u00a0'),('thin','\u2009'),('letter','x'),('trailing','x ')]:
  label=f'{"floating" if floating else "plain"}-{name}'
  p=d.add_paragraph(label); p.paragraph_format.page_break_before=bool(cases)
  d.add_paragraph('Before')
  p=d.add_paragraph(); pr=OxmlElement('w:rPr'); sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz);p._p.get_or_add_pPr().append(pr)
  if floating:
   b=BytesIO();Image.new('RGB',(4,4),'#e03040').save(b,format='PNG')
   inline=p.add_run().add_picture(BytesIO(b.getvalue()),width=Pt(14),height=Pt(14))._inline
   inline.tag=qn('wp:anchor')
   for k,v in dict(distT='0',distB='0',distL='0',distR='0',simplePos='0',relativeHeight='0',behindDoc='1',locked='0',layoutInCell='1',allowOverlap='1').items(): inline.set(k,v)
   simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');inline.insert(0,simple)
   for at,axis in [(1,'H'),(2,'V')]:
    pos=OxmlElement('wp:position'+axis);pos.set('relativeFrom','page');off=OxmlElement('wp:posOffset');off.text='254000';pos.append(off);inline.insert(at,pos)
   inline.insert(4,OxmlElement('wp:wrapNone'))
  p.add_run(text).font.size=Pt(24)
  d.add_paragraph('After')
  cases.append(label)
out=Path(__file__).resolve().parents[1]/'tests/fixtures/floating-blank-lines'
out.mkdir(exist_ok=True)
d.save(out/'source.docx')
(out/'cases.json').write_text(json.dumps(cases))

