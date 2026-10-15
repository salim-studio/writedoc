"""Original large-font controls; no user template content or external artwork."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree as E
from canvas_group_fixture import NS,tag,add,shape
root=Path(__file__).resolve().parents[1]/'tests/fixtures/large-fonts';root.mkdir(parents=True,exist_ok=True)
doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for setting in doc.settings.element.findall('.//'+qn('w:compatSetting')):
 if setting.get(qn('w:name'))=='compatibilityMode':setting.set(qn('w:val'),'15')
s=doc.sections[0];s.page_width=s.page_height=Inches(22);s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(.5)
n=doc.styles['Normal'];n.font.name='Arial';n.font.size=Pt(12);n.paragraph_format.space_after=Pt(0)
style=doc.styles.add_style('LargeGlyph',WD_STYLE_TYPE.CHARACTER);style.font.name='Arial';style.font.size=Pt(650)
cases=[('minimum',1),('old-limit',200),('half-step',200.5),('large',650),('bold',650),('inherited',650),('complex-size-inactive',650),('frame-near-max',1637.5),('frame-max',1638)]
for i,(name,size) in enumerate(cases):
 if i:doc.add_page_break()
 doc.add_paragraph('Large font: '+name)
 p=doc.add_paragraph()
 if name.startswith('frame'):
  drawing=E.SubElement(p.add_run()._r,tag('w:drawing'),nsmap={k:NS[k] for k in ['wps','a']})
  inline=add(drawing,'wp:inline');w,h=1400*12700,1350*12700
  add(inline,'wp:extent',cx=w,cy=h);add(inline,'wp:docPr',id=100+i,name=name)
  data=add(add(inline,'a:graphic'),'a:graphicData',uri=NS['wps'])
  node=shape(data,200+i,w,h,0,0,'FFFFFF','A')
  rp=node.find('.//'+tag('w:rPr'));rp.clear();add(rp,'w:rFonts',**{'w:ascii':'Arial','w:hAnsi':'Arial','w:cs':'Arial'});add(rp,'w:color',**{'w:val':'D02030'});add(rp,'w:sz',**{'w:val':int(size*2)});add(rp,'w:szCs',**{'w:val':int(size*2)})
  body=node.find(tag('wps:bodyPr'))
  for key in ['lIns','rIns','tIns','bIns']:body.set(key,'0')
 else:
  r=p.add_run('A');r.font.color.rgb=__import__('docx').shared.RGBColor.from_string('D02030')
  if name=='inherited':r.style=style
  else:r.font.size=Pt(size)
  if name=='bold':r.bold=True
  if name=='complex-size-inactive':
   rpr=r._r.get_or_add_rPr();rpr.get_or_add_rFonts().set(qn('w:cs'),'Arial');sz=OxmlElement('w:szCs');sz.set(qn('w:val'),str(int(size*2)));rpr.append(sz);r.font.size=Pt(12)
 doc.add_paragraph('After')
doc.save(root/'control.docx')
print(root/'control.docx')

# Same seven body controls without retained frames, for named-style editing.
# Keep this derivation separate so the native nine-page control is unchanged.
import zipfile
from io import BytesIO
parts={}
with zipfile.ZipFile(root/'control.docx') as z:
 parts={n:z.read(n) for n in z.namelist()}
bodydoc=E.fromstring(parts['word/document.xml']);body=bodydoc.find(qn('w:body'));nodes=list(body)
start=next(i for i,n in enumerate(nodes) if ''.join(n.itertext())=='Large font: frame-near-max')-1
for n in nodes[start:]:
 if n.tag!=qn('w:sectPr'):body.remove(n)
parts['word/document.xml']=E.tostring(bodydoc,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(root/'body-control.docx','w',zipfile.ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
