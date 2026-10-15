"""Original small-text controls; no corpus text or artwork."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
d=Document();d.sections[0].left_margin=d.sections[0].right_margin=Inches(1)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
d.add_paragraph('Small text: native font rasterization control')
for family in ['Book Antiqua','Arial']:
 for size in [7,10,14]:
  d.add_paragraph(f'{family} {size} pt: regular / bold')
  table=d.add_table(rows=1,cols=2);table.autofit=False
  for i,cell in enumerate(table.rows[0].cells):
   shading=OxmlElement('w:shd');shading.set(qn('w:val'),'clear');shading.set(qn('w:fill'),'55826B');cell._tc.get_or_add_tcPr().append(shading)
   run=cell.paragraphs[0].add_run('Small type 333');run.font.name=family;run.font.size=Pt(size);run.bold=bool(i)
   color=OxmlElement('w:color');color.set(qn('w:val'),'FFFFFF');run._r.get_or_add_rPr().append(color)
(root/'tests/fixtures/small-text-raster').mkdir(exist_ok=True)
d.save(root/'tests/fixtures/small-text-raster/control.docx')
