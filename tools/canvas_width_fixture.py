"""Independent text-width fixture and export verification using python-docx."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/text-width.docx'

def scale(parent, value):
    # Original fixture authoring: width precedes kern/position/size in CT_RPr.
    node = OxmlElement('w:w'); node.set(qn('w:val'), str(value))
    after = {'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang'}
    anchor = next((n for n in parent if n.tag.split('}')[-1] in after), None)
    if anchor is None: parent.append(node)
    else: anchor.addprevious(node)
    return node

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = section.page_height = Inches(6)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    defaults = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    scale(defaults, '100')
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    base = doc.styles.add_style('WidthParagraph', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal; scale(base.element.get_or_add_rPr(), 125)
    child = doc.styles.add_style('WidthChild', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base; scale(child.element.get_or_add_rPr(), 100)
    char = doc.styles.add_style('WidthCharacter', WD_STYLE_TYPE.CHARACTER); scale(char.element.get_or_add_rPr(), 75)
    table_style = doc.styles.add_style('WidthTable', WD_STYLE_TYPE.TABLE); table_style.base_style = doc.styles['Table Grid']; scale(table_style.element.get_or_add_rPr(), 150)
    doc.add_paragraph('Text width A\u0301😀')
    p = doc.add_paragraph('Normal ')
    r = p.add_run('Wide'); r.underline = True; r.font.highlight_color = WD_COLOR_INDEX.YELLOW; scale(r._r.get_or_add_rPr(), 200)
    p.add_run(' plain ')
    r = p.add_run('Compressed'); r.underline = True; r.font.highlight_color = WD_COLOR_INDEX.TURQUOISE; scale(r._r.get_or_add_rPr(), 50)
    p.add_run(' end')
    doc.add_paragraph('Paragraph wide', style=base)
    doc.add_paragraph('Character compressed', style=base).runs[0].style = char
    scale(doc.add_paragraph('Explicit normal', style=base).runs[0]._r.get_or_add_rPr(), 100)
    doc.add_paragraph('Child reset', style=child)
    for text, sub, percent in [('Wide super', False, 150), ('Narrow sub', True, 75)]:
        r = doc.add_paragraph(text).runs[0]; r.font.size = Pt(16); r.italic = True; r.underline = True
        if sub: r.font.subscript = True
        else: r.font.superscript = True
        scale(r._r.get_or_add_rPr(), percent)
        position = OxmlElement('w:position'); position.set(qn('w:val'), '4'); r._r.rPr.find(qn('w:w')).addnext(position); pitch = OxmlElement('w:spacing'); pitch.set(qn('w:val'), '20'); r._r.rPr.find(qn('w:w')).addprevious(pitch)
    scale(doc.add_paragraph('Percent width').runs[0]._r.get_or_add_rPr(), '075%')
    r = doc.add_paragraph('Aligned\tField').runs[0]; r.underline = True; r.font.highlight_color = WD_COLOR_INDEX.YELLOW; scale(r._r.get_or_add_rPr(), 125)
    p = doc.add_paragraph('Exact wide'); p.paragraph_format.line_spacing = Pt(6); r = p.runs[0]; r.font.size = Pt(32); scale(r._r.get_or_add_rPr(), 180)
    p = doc.add_paragraph(); rpr = OxmlElement('w:rPr'); scale(rpr, 150); p._p.get_or_add_pPr().append(rpr)
    p = doc.add_paragraph('Wider number'); numpr = p._p.get_or_add_pPr().get_or_add_numPr(); numpr.get_or_add_ilvl().val = 0; numpr.get_or_add_numId().val = 7
    root = doc.part.numbering_part.element; instance = next(n for n in root.findall(qn('w:num')) if n.get(qn('w:numId')) == '7'); aid = instance.find(qn('w:abstractNumId')).get(qn('w:val'))
    abstract = next(n for n in root.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == aid)
    level = abstract.find(qn('w:lvl')); rpr = OxmlElement('w:rPr'); scale(rpr, 160); level.append(rpr)
    r = doc.add_paragraph('Missing value reset', style=base).runs[0]; r._r.get_or_add_rPr().append(OxmlElement('w:w'))
    table = doc.add_table(rows=3, cols=2, style=table_style); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).text = 'Wide header'; table.cell(0, 1).text = 'Normal header'; scale(table.cell(0, 1).paragraphs[0].runs[0]._r.get_or_add_rPr(), 100)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Autofitlongword'; table.cell(1, 1).text = 'Wide body across pages. ' * 70
    table.cell(2, 0).text = 'Table inherited'; table.cell(2, 1).text = 'Tail cell'
    doc.add_paragraph('End width sample')
    unit_style = doc.styles.add_style('PercentWidth', WD_STYLE_TYPE.PARAGRAPH); unit_style.base_style = normal; scale(unit_style.element.get_or_add_rPr(), '125%')
    doc.add_paragraph('Percent style', style=unit_style)
    doc.save(SOURCE); print(f'Generated independent text-width fixture: {SOURCE}')

def declared(parent):
    node = None if parent is None else parent.find(qn('w:w'))
    return None if node is None else node.get(qn('w:val'))

def verify(file, fresh):
    doc = Document(file)
    if fresh:
        assert [p.text for p in doc.paragraphs] == ['Wide text', 'Narrow text', 'Normal text', 'Width typing']
        for p, expected in zip(doc.paragraphs, ['150', '75', '100', '125']):
            assert all(declared(r._r.rPr) == expected for r in p.runs), p.text
            assert all(r.font.size in [None, Pt(12)] for r in p.runs), 'Saved size changed'
    else:
        old = Document(SOURCE)
        assert declared(doc.paragraphs[1].runs[1]._r.rPr) == '175'
        assert declared(doc.styles['WidthParagraph'].element.rPr) == '150'
        assert declared(doc.styles['WidthCharacter'].element.rPr) is None
        def signature(root, omit):
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in root.iter() if n not in omit]
        def edits(d): return [d.paragraphs[1].runs[1]._r.rPr.find(qn('w:w'))]
        assert signature(old.element, edits(old)) == signature(doc.element, edits(doc)), 'Unrelated document content changed'
        def edits(d): return [d.styles[s].element.rPr.find(qn('w:w')) for s in ['WidthParagraph', 'WidthCharacter']]
        assert signature(old.styles.element, edits(old)) == signature(doc.styles.element, edits(doc)), 'Unrelated style content changed'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: text width, source text, stored sizes and unchanged package content')
