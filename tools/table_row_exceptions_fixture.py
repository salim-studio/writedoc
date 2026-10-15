"""Original row-level table property controls; no corpus XML."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_ROW_HEIGHT_RULE

ROOT = Path(__file__).resolve().parents[1] / 'tests/fixtures/table-row-exceptions'
SIDES = ['top','left','bottom','right','insideH','insideV']
COLORS = ['D02020','208020','2040D0','9020A0','008090','D07000']

def add(parent, tag, **attrs):
    node = OxmlElement('w:'+tag)
    for k,v in attrs.items(): node.set(qn('w:'+k),str(v))
    parent.append(node)
    return node

def borders(parent, name, sides, size=24, style='single'):
    b=add(parent,name)
    for side,color in sides.items(): add(b,side,val=style,sz=size,color=color)

def generate():
    ROOT.mkdir(parents=True,exist_ok=True)
    doc=Document(); normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    for s in doc.settings.element.iter(qn('w:compatSetting')):
        if s.get(qn('w:name'))=='compatibilityMode':s.set(qn('w:val'),'15')
    style=add(doc.styles.element,'style',type='table',styleId='RowExceptionStyle',customStyle=1)
    add(style,'name',val='Row Exception Style');sp=add(style,'tblPr');add(sp,'tblStyleRowBandSize',val=1)
    for region,color in [('firstRow','E8B0B0'),('firstCol','B0E8B0'),('band1Horz','B0D0E8'),('band2Horz','E8D0B0')]:
        part=add(style,'tblStylePr',type=region);add(add(part,'tcPr'),'shd',val='clear',fill=color)
    names=['baseline','all-middle','inside-middle','partial-middle','none-middle','nil-middle','direct-cells','adjacent-rows','vertical-merge','horizontal-merge','outer-rows','look-first-column','look-disable-first','look-bands','paged-merge','repeated-header']
    cases=[]
    for name in names:
        cases.append(dict(name=name));p=doc.add_paragraph(name);p.paragraph_format.page_break_before=len(cases)>1
        paged=name in ['paged-merge','repeated-header']
        table=doc.add_table(rows=8 if paged else 3,cols=2);table.autofit=False
        for col in table.columns:col.width=Inches(2)
        pr=table._tbl.tblPr;edge=add(pr,'tblBorders');pr.remove(edge);pr.insert(list(pr).index(pr.find(qn('w:tblLayout'))),edge)
        for side in SIDES:add(pr.find(qn('w:tblBorders')),side,val='single',sz=8,color='303030')
        for ri,row in enumerate(table.rows):
            row.height=Pt(100 if paged else 36);row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
            for ci,cell in enumerate(row.cells):cell.width=Inches(2);cell.text=f'Row {ri+1} Col {ci+1}'
        if name.startswith('look-'):
            st=OxmlElement('w:tblStyle');st.set(qn('w:val'),'RowExceptionStyle');pr.insert(0,st)
            look=pr.find(qn('w:tblLook'));look.attrib.clear();look.set(qn('w:val'),'06A0' if name=='look-disable-first' else '0600')
            ri=0 if name=='look-disable-first' else 1
            ex=OxmlElement('w:tblPrEx');table.rows[ri]._tr.insert(0,ex)
            add(ex,'tblLook',val={'look-disable-first':'0600','look-first-column':'0680','look-bands':'0400'}[name])
        if name=='vertical-merge':table.cell(0,0).merge(table.cell(2,0))
        if paged:table.cell(1 if name=='repeated-header' else 0,0).merge(table.cell(7,0))
        if name=='repeated-header':add(table.rows[0]._tr.get_or_add_trPr(),'tblHeader')
        if name=='horizontal-merge':table.cell(1,0).merge(table.cell(1,1))
        if name!='baseline' and not name.startswith('look-'):
            rows=[1,5,6] if paged else [0,2] if name=='outer-rows' else [1,2] if name=='adjacent-rows' else [1]
            for ri in rows:
                ex=OxmlElement('w:tblPrEx');table.rows[ri]._tr.insert(0,ex)
                sides=dict(zip(SIDES,COLORS))
                if name=='inside-middle':sides={s:sides[s] for s in ['insideH','insideV']}
                if name=='partial-middle':sides={'left':'208020'}
                borders(ex,'tblBorders',sides,style='none' if name=='none-middle' else 'nil' if name=='nil-middle' else 'single')
                if name=='direct-cells':
                    for cell in table.rows[ri].cells:borders(cell._tc.get_or_add_tcPr(),'tcBorders',dict.fromkeys(['top','left','bottom','right'],'B08000'),size=4)
        doc.add_paragraph('After')
    doc.save(ROOT/'controls.docx');(ROOT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')

if __name__=='__main__':generate()
