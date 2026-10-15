"""Independent kerning fixture and export verification using python-docx."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/kerning.docx'

def kern(parent, value):
    node = OxmlElement('w:kern'); node.set(qn('w:val'), str(value))
    after = {'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang'}
    anchor = next((n for n in parent if n.tag.split('}')[-1] in after), None)
    if anchor is None: parent.append(node)
    else: anchor.addprevious(node)
    return node

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = section.page_height = Inches(6)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    base = doc.styles.add_style('KernParagraph', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal; kern(base.element.get_or_add_rPr(), 24)
    child = doc.styles.add_style('KernChild', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base; kern(child.element.get_or_add_rPr(), 0)
    char = doc.styles.add_style('KernCharacter', WD_STYLE_TYPE.CHARACTER); kern(char.element.get_or_add_rPr(), 28)
    table_style = doc.styles.add_style('KernTable', WD_STYLE_TYPE.TABLE); table_style.base_style = doc.styles['Table Grid']; kern(table_style.element.get_or_add_rPr(), 24)
    doc.add_paragraph('Kerning A\u0301😀')
    for text, value, size in [('Enabled AVATAR', 24, 12), ('Disabled AVATAR', 0, 12), ('Below AVATAR', 28, 12), ('At AVATAR', 28, 14), ('Above AVATAR', 28, 16)]:
        r = doc.add_paragraph(text).runs[0]; r.font.size = Pt(size); r.underline = True; r.font.highlight_color = WD_COLOR_INDEX.YELLOW; kern(r._r.get_or_add_rPr(), value)
    doc.add_paragraph('Paragraph AVATAR', style=base)
    doc.add_paragraph('Character AVATAR', style=base).runs[0].style = char
    doc.add_paragraph('Child AVATAR', style=child)
    kern(doc.add_paragraph('Direct zero AVATAR', style=base).runs[0]._r.get_or_add_rPr(), 0)
    kern(doc.add_paragraph('Universal AVATAR').runs[0]._r.get_or_add_rPr(), '12pt')
    for text, sub in [('Super AVATAR', False), ('Sub AVATAR', True)]:
        r = doc.add_paragraph(text).runs[0]; r.font.size = Pt(20); r.italic = True; r.underline = True
        if sub: r.font.subscript = True
        else: r.font.superscript = True
        rpr = r._r.get_or_add_rPr(); k = kern(rpr, 40)
        for tag, value in [('spacing', 20), ('w', 150)]:
            n = OxmlElement('w:' + tag); n.set(qn('w:val'), str(value)); k.addprevious(n)
    p = doc.add_paragraph(); mark = OxmlElement('w:rPr'); kern(mark, 24); p._p.get_or_add_pPr().append(mark)
    p = doc.add_paragraph('AVATAR\tTo WA', style=base)
    p = doc.add_paragraph('Number AVATAR'); p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 7
    root = doc.part.numbering_part.element; instance = next(n for n in root.findall(qn('w:num')) if n.get(qn('w:numId')) == '7')
    abstract = next(n for n in root.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == instance.find(qn('w:abstractNumId')).get(qn('w:val')))
    rpr = OxmlElement('w:rPr'); kern(rpr, 24); abstract.find(qn('w:lvl')).append(rpr)
    table = doc.add_table(rows=3, cols=2, style=table_style); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).text = 'Header AVATAR'; table.cell(0, 1).text = 'Plain header'; kern(table.cell(0, 1).paragraphs[0].runs[0]._r.get_or_add_rPr(), 0)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'AVATARAVATARAVATAR'; table.cell(1, 1).text = 'AVATAR To WA table content. ' * 80
    table.cell(2, 0).text = 'Inherited table'; table.cell(2, 1).text = 'End table'
    doc.add_paragraph('End kerning sample')
    doc.save(SOURCE); print(f'Generated independent kerning fixture: {SOURCE}')

def declared(parent):
    node = None if parent is None else parent.find(qn('w:kern'))
    return None if node is None else node.get(qn('w:val'))

def verify(file, fresh):
    doc = Document(file)
    if fresh:
        assert [p.text for p in doc.paragraphs] == ['Enabled AVATAR', 'Disabled AVATAR', 'Future AVATAR']
        for p, expected in zip(doc.paragraphs, ['24', '0', '28']):
            assert all(declared(r._r.rPr) == expected for r in p.runs), p.text
    else:
        old = Document(SOURCE)
        assert declared(doc.paragraphs[1].runs[0]._r.rPr) == '28'
        assert declared(doc.styles['KernParagraph'].element.rPr) == '20'
        assert declared(doc.styles['KernCharacter'].element.rPr) is None
        def signature(root, omit):
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in root.iter() if n not in omit]
        def edits(d): return [d.paragraphs[1].runs[0]._r.rPr.find(qn('w:kern'))]
        assert signature(old.element, edits(old)) == signature(doc.element, edits(doc)), 'Unrelated document XML changed'
        def edits(d): return [d.styles[s].element.rPr.find(qn('w:kern')) for s in ['KernParagraph', 'KernCharacter']]
        assert signature(old.styles.element, edits(old)) == signature(doc.styles.element, edits(doc)), 'Unrelated style XML changed'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: independent kerning thresholds, source text and unchanged package content')
