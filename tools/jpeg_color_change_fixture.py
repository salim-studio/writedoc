"""Independent JPEG/PNG color-key comparison using identical decoded RGB."""
from pathlib import Path
from io import BytesIO
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
import json

out=Path(__file__).resolve().parents[1]/'tests/fixtures/jpeg-color-change'
out.mkdir(parents=True,exist_ok=True)
im=Image.new('RGB',(256,96),'white');draw=ImageDraw.Draw(im)
for i in range(16):
 v=255-i;draw.rectangle((i*16,0,i*16+15,23),fill=(v,v,v))
draw.ellipse((35,30,120,91),fill='#2060c0');draw.polygon([(150,88),(180,29),(237,88)],fill='#c04080')
jpeg=BytesIO();im.save(jpeg,format='JPEG',quality=85,subsampling=2)
decoded=Image.open(BytesIO(jpeg.getvalue())).convert('RGB');png=BytesIO();decoded.save(png,format='PNG')
doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
doc.add_paragraph('JPEG and decoded PNG color-key controls')
for index,(fmt,effect) in enumerate([('JPEG',False),('PNG',False),('JPEG',True),('PNG',True)]):
 doc.add_paragraph(f'C{index} {fmt} '+('white transparent' if effect else 'original'))
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
 p._p.get_or_add_pPr().insert(0,parse_xml('<w:shd '+nsdecls('w')+' w:val="clear" w:fill="6699AA"/>'))
 run=p.add_run();pic=run.add_picture(BytesIO(jpeg.getvalue() if fmt=='JPEG' else png.getvalue()),width=Inches(3))
 if effect:
  blip=pic._inline.xpath('.//a:blip')[0]
  blip.append(parse_xml('<a:clrChange '+nsdecls('a')+'><a:clrFrom><a:srgbClr val="FFFFFF"/></a:clrFrom><a:clrTo><a:srgbClr val="FFFFFF"><a:alpha val="0"/></a:srgbClr></a:clrTo></a:clrChange>'))
doc.save(out/'control.docx')

# Large constant interiors separate color matching from JPEG boundary artifacts.
colors=[(v,v,v) for v in range(255,-1,-1)]
for channel in range(3):
 for delta in range(64):
  rgb=[255,255,255];rgb[channel]-=delta;colors.append(tuple(rgb))
key=(120,100,80)
for channel in range(3):
 for delta in range(-40,41):
  rgb=list(key);rgb[channel]+=delta;colors.append(tuple(rgb))
for delta in [-17,-16,-15,15,16,17]:
 colors.append(tuple(c+delta for c in key))
columns=32;size=16;rows=(len(colors)+columns-1)//columns
grid=Image.new('RGB',(columns*size,rows*size),'black');draw=ImageDraw.Draw(grid)
for i,rgb in enumerate(colors):
 x=(i%columns)*size;y=(i//columns)*size;draw.rectangle((x,y,x+size-1,y+size-1),fill=rgb)
jpeg=BytesIO();grid.save(jpeg,format='JPEG',quality=95,subsampling=0)
decoded=Image.open(BytesIO(jpeg.getvalue())).convert('RGB');png=BytesIO();decoded.save(png,format='PNG')
doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
cases=[('JPEG','FFFFFF',None,100000,0),('PNG','FFFFFF',None,100000,0),('JPEG','FFFFFF','0',100000,0),('JPEG','786450',None,100000,0),('PNG','786450',None,100000,0),('JPEG','786450','0',100000,0),('JPEG','FFFFFF',None,50000,0),('JPEG','FFFFFF',None,100000,100000)]
cases += [('JPEG','FFFFFF',None,round(alpha/255*100000),0) for alpha in [239,238,254,253]]
for i,(fmt,key,use,alpha,replacement) in enumerate(cases):
 if i:doc.add_page_break()
 doc.add_paragraph(f'C{i} {fmt} key {key} useA {use} alpha {alpha} replacement {replacement}')
 p=doc.add_paragraph();p._p.get_or_add_pPr().append(parse_xml('<w:shd '+nsdecls('w')+' w:val="clear" w:fill="6699AA"/>'))
 pic=p.add_run().add_picture(BytesIO(jpeg.getvalue() if fmt=='JPEG' else png.getvalue()),width=Inches(6))
 use_attr='' if use is None else f' useA="{use}"'
 effect=f'<a:clrChange {nsdecls("a")}{use_attr}><a:clrFrom><a:srgbClr val="{key}"><a:alpha val="{alpha}"/></a:srgbClr></a:clrFrom><a:clrTo><a:srgbClr val="1177DD"><a:alpha val="{replacement}"/></a:srgbClr></a:clrTo></a:clrChange>'
 pic._inline.xpath('.//a:blip')[0].append(parse_xml(effect))
doc.save(out/'thresholds.docx')
(out/'thresholds.json').write_text(json.dumps({'columns':columns,'size':size,'colors':colors,'decoded':[decoded.getpixel(((i%columns)*size+size//2,(i//columns)*size+size//2)) for i in range(len(colors))],'cases':cases},indent=2)+'\n')
