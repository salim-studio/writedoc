"""Original paragraph-border fixture, authored independently with python-docx."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/paragraph-borders.docx'
SIDES = ['top', 'left', 'bottom', 'right', 'between', 'bar']

def borders(ppr, sides, **attributes):
    parent = ppr.find(qn('w:pBdr'))
    if parent is None:
        parent = OxmlElement('w:pBdr')
        ppr.insert_element_before(parent, 'w:shd', 'w:tabs', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:jc', 'w:rPr', 'w:sectPr')
    for side in SIDES:
        if side not in sides: continue
        node = OxmlElement('w:' + side)
        for key, value in {'val': 'single', 'sz': '8', 'space': '2', 'color': '224466', **attributes}.items(): node.set(qn('w:' + key), value)
        parent.append(node)

def shade(ppr, color):
    node = OxmlElement('w:shd'); node.set(qn('w:val'), 'clear'); node.set(qn('w:fill'), color)
    ppr.insert_element_before(node, 'w:tabs', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:jc', 'w:rPr', 'w:sectPr')

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('Canvas Border Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    borders(base.element.get_or_add_pPr(), ['top', 'left', 'bottom', 'right'])
    borders(base.element.get_or_add_pPr(), ['between'], sz='4', space='12', color='CC3300')
    child = doc.styles.add_style('Canvas Border Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    borders(child.element.get_or_add_pPr(), ['left'], val='dotted', sz='12', space='3')
    doc.add_paragraph('Paragraph borders A\u0301😀')
    doc.add_paragraph('Joined first', child); doc.add_paragraph('Joined second', child)
    p = doc.add_paragraph('Different spacing breaks the group', child); borders(p._p.get_or_add_pPr(), ['bottom'], space='4')
    p = doc.add_paragraph('Partial direct border', child); borders(p._p.get_or_add_pPr(), ['right'], val='nil')
    p = doc.add_paragraph('No borders', child); borders(p._p.get_or_add_pPr(), SIDES, val='none')
    p = doc.add_paragraph('Themed bottom'); borders(p._p.get_or_add_pPr(), ['bottom'], val='double', sz='16', space='3', color='010203', themeColor='accent3', themeTint='99', themeShade='00')
    for pattern in ['single', 'double', 'triple', 'dotted', 'dashed', 'dashSmallGap', 'dotDash', 'dotDotDash']:
        p = doc.add_paragraph('Pattern ' + pattern); borders(p._p.get_or_add_pPr(), ['bottom'], val=pattern, sz='12', space='2', color='663399')
    p = doc.add_paragraph('Hanging box across wrapped lines. ' * 5, child)
    p.paragraph_format.left_indent = Inches(.5); p.paragraph_format.right_indent = Inches(.25); p.paragraph_format.first_line_indent = Inches(-.25)
    shade(p._p.get_or_add_pPr(), 'DDEEFF')
    p = doc.add_paragraph('Centered box', child); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(.25); p.paragraph_format.right_indent = Inches(.5)
    p = doc.add_paragraph('Flowing bordered paragraph A\u0301😀. ' * 90, child); shade(p._p.get_or_add_pPr(), 'FFE4CC')
    table = doc.add_table(rows=3, cols=2, style='Table Grid'); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(2.5)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                n = OxmlElement('w:' + side); n.set(qn('w:w'), '120'); n.set(qn('w:type'), 'dxa'); margins.append(n)
            cell._tc.get_or_add_tcPr().append(margins)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(0,0).text = 'Bordered header'; table.cell(0,1).text = 'Other header'
    p = table.cell(0,0).paragraphs[0]; borders(p._p.get_or_add_pPr(), ['top', 'left', 'bottom', 'right'], space='1', color='005588'); shade(p._p.get_or_add_pPr(), 'DDEEFF')
    table.cell(1,0).text = 'Cell first'; p = table.cell(1,0).paragraphs[0]; p.style = base
    p = table.cell(1,0).add_paragraph('Cell second', base)
    p = table.cell(1,1).paragraphs[0]; p.add_run('Flowing bordered cell. ' * 95)
    borders(p._p.get_or_add_pPr(), ['top', 'left', 'bottom', 'right'], space='1', color='227744'); shade(p._p.get_or_add_pPr(), 'DDFFEE')
    table.cell(2,0).text = 'Last row'; table.cell(2,1).text = 'Last cell'
    doc.add_paragraph('Tail')
    doc.save(SOURCE); print(f'Generated independent paragraph-border fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file)
    paragraphs = doc.element.findall('.//' + qn('w:p'))
    text = lambda p: ''.join(n.text or '' for n in p.findall('.//' + qn('w:t')))
    by_text = {text(p): p for p in paragraphs}
    def edge(label, side, style, size=None, space=None, color=None):
        parent = by_text[label].find(qn('w:pPr')).find(qn('w:pBdr'))
        node = parent.find(qn('w:' + side))
        assert node is not None and node.get(qn('w:val')) == style, (label, side)
        for key, value in [('sz', size), ('space', space), ('color', color)]:
            if value is not None: assert node.get(qn('w:' + key)) == str(value), (label, side, key)
        assert not any('theme' in key for key in node.attrib), (label, side)
    if fresh:
        assert [text(p) for p in paragraphs] == ['Box paragraph', 'Bottom paragraph', 'Plain paragraph']
        for side in SIDES[:4]: edge('Box paragraph', side, 'single', 8, 2, '224466')
        edge('Bottom paragraph', 'bottom', 'dashed', 12, 3, 'CC3300')
        for side in SIDES:
            if side != 'bottom': edge('Bottom paragraph', side, 'none')
            edge('Plain paragraph', side, 'none')
    else:
        original_doc = Document(SOURCE)
        original_paragraphs = original_doc.element.findall('.//' + qn('w:p'))
        assert [text(p) for p in paragraphs] == [text(p) for p in original_paragraphs]
        edge('Joined first', 'bottom', 'double', 16, 4, 'CC3300')
        for side in SIDES:
            if side != 'bottom': edge('Joined first', side, 'none')
            edge('Themed bottom', side, 'none')
        for side in SIDES[:4]: edge('Bordered header', side, 'single', 8, 2, 'AA3377')
        changed = {'Joined first', 'Themed bottom', 'Bordered header'}
        for p, before in zip(paragraphs, original_paragraphs):
            def data(element):
                # Compare every unedited node/attribute, tolerating only the
                # writer's redundant xml:space="preserve" on plain text.
                skip = set()
                if text(p) in changed:
                    for border in element.findall('./' + qn('w:pPr') + '/' + qn('w:pBdr')): skip.update(border.iter())
                return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in element.iter() if n not in skip]
            assert data(p) == data(before), f'Unrelated paragraph data changed: {text(p)}'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), f'Unrelated part changed: {name}'
    print(f'PASS: {file}: paragraph borders, explicit resets, text, retained formatting and unrelated package parts')
