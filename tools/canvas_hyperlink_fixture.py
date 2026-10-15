"""Independent external-link fixtures/export checks using python-docx (MIT)."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/hyperlinks.docx'
TARGET = 'https://example.com/guide?q=one%20two&lang=en#section'

def add_link(p, target, pieces, **attrs):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    for key, value in attrs.items(): link.set(qn('w:' + key), value)
    p._p.append(link)
    for text, bold in pieces:
        run = p.add_run(text); run.style = 'Hyperlink'; run.bold = bold
        link.append(run._r)
    return link

def generate(factory):
    doc = factory(); sec = doc.sections[0]
    sec.page_width = Inches(6); sec.page_height = Inches(7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.5)
    doc.styles['Normal'].font.name = 'Noto Sans'; doc.styles['Normal'].font.size = Pt(12)
    style = doc.styles.add_style('Hyperlink', WD_STYLE_TYPE.CHARACTER)
    style.font.color.rgb = RGBColor.from_string('0563C1'); style.font.underline = True
    doc.add_paragraph('External hyperlinks a\u0301😀')
    p = doc.add_paragraph('Before ')
    add_link(p, TARGET, [('Read ', False), ('the guide a\u0301😀', True)], tooltip='Guide tip', history='1', anchor='ignored', tgtFrame='_top')
    p.add_run(' after.')
    p = doc.add_paragraph('Shared: '); add_link(p, TARGET, [('Same destination', False)])
    p = doc.add_paragraph('Adjacent: ')
    add_link(p, TARGET, [('First', False)]); add_link(p, TARGET, [('Second', False)])
    p = doc.add_paragraph(); add_link(p, 'mailto:hello@example.com?subject=Hello', [('Email us', False)])
    p.add_run(' / '); add_link(p, 'tel:+1555010100', [('Call us', False)])
    p = doc.add_paragraph(); add_link(p, 'https://example.com/wrapped', [('Wrapped link a\u0301😀 across many words. ' * 10, False)])
    p = doc.add_paragraph(); add_link(p, 'https://example.com/breaks', [('Tab\tthen line\nnext', False)])
    p = doc.add_paragraph('Picture: '); link = add_link(p, 'https://example.com/picture', [])
    with ZipFile(SOURCE.with_name('inline-images.docx')) as z:
        image = z.read(next(n for n in z.namelist() if n.startswith('word/media/') and n.endswith('.png')))
    run = p.add_run(); run.add_picture(BytesIO(image), width=Inches(.4)); link.append(run._r)
    table = doc.add_table(rows=3, cols=2, style='Table Grid'); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    add_link(table.cell(0,0).paragraphs[0], 'https://example.com/header', [('Linked header', False)])
    table.cell(0,1).text = 'Header two'; table.cell(1,0).text = 'Cell text'
    add_link(table.cell(1,1).paragraphs[0], 'https://example.com/cell', [('Long linked cell content. ' * 85, False)])
    table.cell(2,0).text = 'Last row'; table.cell(2,1).text = 'Last cell'
    doc.add_paragraph('End hyperlinks')
    doc.save(SOURCE); print(f'Generated independent hyperlink fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file)
    links = doc.element.findall('.//' + qn('w:hyperlink'))
    with ZipFile(file) as package:
        relations = etree.fromstring(package.read('word/_rels/document.xml.rels'))
        ids = [r.get('Id') for r in relations]
        assert len(ids) == len(set(ids)) and all(ids), 'Relationship IDs must be unique'
        for link in links:
            matches = [r for r in relations if r.get('Id') == link.get(qn('r:id'))]
            assert len(matches) == 1 and matches[0].get('Type') == RT.HYPERLINK and matches[0].get('TargetMode') == 'External'
    def target(link):
        rel = doc.part.rels[link.get(qn('r:id'))]
        assert rel.reltype == RT.HYPERLINK and rel.is_external
        return rel.target_ref
    def label(link): return ''.join(n.text or '' for n in link.findall('.//' + qn('w:t')))
    if fresh:
        assert [(label(l), target(l)) for l in links] == [('New link a\u0301😀', 'https://example.com/new'), ('Mail', 'mailto:new@example.com')]
    else:
        guide = next(l for l in links if label(l) == 'Read the guide a\u0301😀')
        assert target(guide) == 'https://example.com/edited#new'
        assert guide.get(qn('w:tooltip')) == 'Guide tip' and guide.get(qn('w:anchor')) == 'ignored'
        assert guide.get(qn('w:history')) == '1' and guide.get(qn('w:tgtFrame')) == '_top'
        shared = next(l for l in links if label(l) == 'Same destination'); assert target(shared) == TARGET
        assert not any(label(l) == 'Email us' for l in links)
        assert any(l.find('.//' + qn('w:drawing')) is not None for l in links)
        assert any(label(l) == 'Linked header' for l in links)
        with ZipFile(SOURCE) as original, ZipFile(file) as result:
            for name in original.namelist():
                if name not in ['word/document.xml', 'word/_rels/document.xml.rels']:
                    assert original.read(name) == result.read(name), name
        assert len([r for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK and r.target_ref == TARGET]) == 1
    for link in links: target(link)
    print(f'PASS {file}: independent hyperlink text, wrappers, relationship targets and preservation')
