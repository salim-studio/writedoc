"""Original temporary-control fixture, using python-docx (MIT) and lxml (BSD).

The package boilerplate/thumbnail is python-docx's MIT default template. All
additional XML and content are original; no user documents are redistributed.
"""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/content-controls.docx'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
W15 = 'http://schemas.microsoft.com/office/word/2012/wordml'
NS = {'w': W}

def child(parent, name, **attrs):
    node = OxmlElement('w:' + name)
    for key, value in attrs.items(): node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node

def control(p, number, texts, placeholder=True):
    sdt = child(p._p, 'sdt'); pr = child(sdt, 'sdtPr')
    if number == 1:
        rp = child(pr, 'rPr'); child(rp, 'b'); child(rp, 'color', val='2050C0')
    child(pr, 'alias', val=f'Prompt {number}'); child(pr, 'tag', val=f'original-{number}')
    child(pr, 'id', val=-number); child(pr, 'lock', val='unlocked')
    child(child(pr, 'placeholder'), 'docPart', val=f'ExamplePrompt{number}')
    child(pr, 'temporary')
    if placeholder: child(pr, 'showingPlcHdr')
    child(pr, 'richText')
    appearance = etree.SubElement(pr, '{%s}appearance' % W15, nsmap={'w15': W15})
    appearance.set('{%s}val' % W15, 'hidden')
    child(sdt, 'sdtEndPr'); content = child(sdt, 'sdtContent')
    for i, text in enumerate(texts):
        r = child(content, 'r'); rp = child(r, 'rPr')
        if i: child(rp, 'i')
        child(rp, 'color', val='777777' if placeholder else '208050')
        t = child(r, 't'); t.text = text; t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return sdt

def generate():
    doc = Document(); section = doc.sections[0]
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    # Keep the optional w15 appearance extension in its MCE scope.
    doc._element.set('{%s}appearanceScope' % W15, 'unused')
    prefix = next(k for k, v in doc._element.nsmap.items() if v == W15)
    del doc._element.attrib['{%s}appearanceScope' % W15]
    doc._element.set('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable', doc._element.get('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable', '') + ' ' + prefix)
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(10)
    doc.add_paragraph('Temporary content controls Á😀', 'Heading 1')
    p = doc.add_paragraph('Dear '); control(p, 1, ['[Your ', 'name]']); p.add_run(', welcome.')
    p = doc.add_paragraph('Code: '); control(p, 2, ['EXISTING ', 'VALUE'], False); p.add_run(' — retained rich text')
    p = doc.add_paragraph(); control(p, 3, ['[First]']); control(p, 4, ['[Second]'])
    table = doc.add_table(rows=2, cols=2, style='Table Grid'); table.autofit = False
    for column in table.columns: column.width = Inches(2.5)
    table.cell(0, 0).text = 'Recipient'; control(table.cell(0, 1).paragraphs[0], 5, ['[Recipient Á😀]'])
    table.cell(1, 0).text = 'Message'; control(table.cell(1, 1).paragraphs[0], 6, ['[Message]'])
    doc.add_paragraph('Controls tail')
    doc.save(SOURCE)

def verify(path, edited=False):
    with ZipFile(SOURCE) as z: before = {n: z.read(n) for n in z.namelist()}
    with ZipFile(path) as z: after = {n: z.read(n) for n in z.namelist()}
    assert set(before) == set(after)
    for name in before:
        if name != 'word/document.xml': assert before[name] == after[name], name
    old = etree.fromstring(before['word/document.xml']); new = etree.fromstring(after['word/document.xml'])
    def canonical(el):
        # Prefixes and redundant namespace declarations are not document semantics.
        return (el.tag, sorted(el.attrib.items()), el.text or '', [canonical(c) for c in el])
    controls = old.findall('.//w:sdt', NS); output = new.findall('.//w:sdt', NS)
    assert len(output) == (5 if edited else 6)
    assert [canonical(c) for c in output] == [canonical(c) for c in (controls[1:] if edited else controls)]
    def text(root): return ''.join(root.xpath('//w:body//w:t/text()', namespaces=NS))
    assert text(new) == (text(old).replace('[Your name]', 'Ada Á😀') if edited else text(old))
    if edited:
        run = new.xpath('//w:r[w:t="Ada Á😀"]', namespaces=NS)[0]
        assert run.find('w:rPr/w:b', NS) is not None
        assert run.find('w:rPr/w:color', NS).get(qn('w:val')) == '2050C0'
    else: assert before == after
    print('Independent temporary-control checks passed:', path)

if __name__ == '__main__':
    import sys
    if len(sys.argv) == 1: generate()
    else: verify(sys.argv[1], '--edited' in sys.argv[2:])
