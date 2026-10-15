"""Original controls for border-dependent nested floating table alignment."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def prop(parent, name, **attrs):
    node = OxmlElement('w:' + name)
    for k, v in attrs.items(): node.set(qn('w:' + k), str(v))
    parent.append(node)
    return node

doc = Document()
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(1)
normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = Pt(15)
cases = [dict(side=side, size=size, row=row) for side, size, row in [
    ('right',0,0), ('right',8,0), ('right',24,0), ('right',48,0),
    ('left',24,0), ('both',24,0), ('right',24,1), ('left',24,1), ('both',24,1)]]
cases += [dict(side=side, size=24, row=0, alignment=alignment)
          for alignment in ['left','right','offset'] for side in ['left','right','both']]
cases += [dict(side='right', size=size, row=0) for size in [12,16]]
for index, case in enumerate(cases):
    heading = doc.add_paragraph('Nested border case ' + str(index))
    heading.paragraph_format.page_break_before = bool(index)
    outer = doc.add_table(rows=1, cols=2); outer.autofit = False
    for col in outer.columns: col.width = Inches(3)
    for i, cell in enumerate(outer.rows[0].cells):
        cell.width = Inches(3); cell.paragraphs[0].text = 'Before nested ' + str(i)
        inner = cell.add_table(rows=2, cols=1); inner.autofit = False
        inner.columns[0].width = Inches(1)
        width = inner._tbl.tblPr.find(qn('w:tblW'))
        width.set(qn('w:type'), 'dxa'); width.set(qn('w:w'), '1440')
        if i:
            alignment = case.get('alignment', 'center')
            horizontal = {'tblpX':'720'} if alignment == 'offset' else {'tblpXSpec':alignment}
            p = prop(inner._tbl.tblPr, 'tblpPr', horzAnchor='margin', vertAnchor='text', tblpY='720', **horizontal)
            inner._tbl.tblPr.remove(p); inner._tbl.tblPr.insert(0, p)
        for ri, row in enumerate(inner.rows):
            child = row.cells[0]; child.width = Inches(1); child.text = 'Nested ' + str(ri + 1)
            prop(child._tc.get_or_add_tcPr(), 'shd', val='clear', fill='D0E8C0')
            if i and case['size'] and ri == case['row']:
                borders = OxmlElement('w:tcBorders')
                child._tc.get_or_add_tcPr().insert(1, borders)
                for side in ['left','right'] if case['side'] == 'both' else [case['side']]:
                    prop(borders, side, val='single', sz=case['size'], color='000000')
        cell.paragraphs[-1].text = 'After nested ' + str(i)
    doc.add_paragraph('Following outer table')
root = Path(__file__).resolve().parents[1]
target = root / 'public/fixtures/nested-floating-borders.docx'
doc.save(target)
folder = root / 'tests/fixtures/nested-floating-borders'; folder.mkdir(exist_ok=True)
(folder / 'cases.json').write_text(json.dumps(cases, indent=2) + '\n')
print(target)
