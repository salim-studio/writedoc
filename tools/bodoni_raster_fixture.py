"""Original native/browser raster controls; no font files are bundled."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'tests/fixtures/bodoni-raster';OUT.mkdir(exist_ok=True)
d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.left_margin=s.right_margin=Inches(.75);s.top_margin=s.bottom_margin=Inches(.5)
d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(12);cases=[]
for family,size,bold in [('Bodoni MT',12,False),('Bodoni MT',24,False),('Bodoni MT',12,True),('Georgia',12,False),('Arial',12,False)]:
 for dark in [False,True]:
  p=d.add_paragraph();p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=Pt(size+12)
  shade=OxmlElement('w:shd');shade.set(qn('w:val'),'clear');shade.set(qn('w:fill'),'222C37' if dark else 'FFFFFF');p._p.get_or_add_pPr().insert(0,shade)
  r=p.add_run('See you soon / HAMBURG 012345');r.font.name=family;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string('FFFFFF' if dark else '000000')
  cases.append(dict(index=len(cases),family=family,size=size,bold=bold,dark=dark))
d.save(OUT/'control.docx');(OUT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n',encoding='utf-8')
