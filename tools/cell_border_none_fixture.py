"""Original, independently authored OOXML controls for cell border none semantics.

Generate with Python/python-docx/lxml; no corpus document content is used.
"""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
from lxml import etree as E
from canvas_group_fixture import Document, Inches, Pt, tag, add
from docx.oxml import OxmlElement

root=Path(__file__).resolve().parents[1]/'tests/fixtures/cell-border-none'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
q=lambda s:'{'+W+'}'+s

def matrix():
    cases=[(size,color) for size in [0,8,16] for color in ['auto','A030A0']]
    doc=Document();doc.settings.element.find(tag('w:zoom')).set(tag('w:percent'),'100')
    section=doc.sections[0];section.page_width=section.page_height=Inches(6)
    section.left_margin=section.right_margin=section.top_margin=section.bottom_margin=Inches(.5)
    doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(0)

    def borders(parent,kind,value,color,size,sides=('top','left','bottom','right')):
        b=add(parent,kind)
        for side in sides:add(b,'w:'+side,**{'w:val':value,'w:color':color,'w:sz':size})

    for i,(size,color) in enumerate(cases):
        if i:doc.add_page_break()
        doc.add_paragraph(f'matrix-{size}-{color}')
        sty=add(doc.styles.element,'w:style',**{'w:type':'table','w:styleId':f'Probe{i}'})
        add(sty,'w:name',**{'w:val':f'Probe {i}'})
        holder=add(sty,'w:tblStylePr',**{'w:type':'firstRow'})
        borders(add(holder,'w:tcPr'),'w:tcBorders','single','2050D0',8)
        table=doc.add_table(rows=1,cols=2)
        table.autofit=False
        pr=table._tbl.tblPr
        style=OxmlElement('w:tblStyle');style.set(tag('w:val'),f'Probe{i}');pr.insert(0,style)
        look=pr.find(tag('w:tblLook'));look.set(tag('w:val'),'0020')
        for attr,value in [('firstRow','1'),('firstColumn','0'),('lastRow','0'),('lastColumn','0'),('noHBand','0'),('noVBand','0')]:look.set(tag('w:'+attr),value)
        borders(pr,'w:tblBorders','single','E02030',48,('top','left','bottom','right','insideH','insideV'))
        # CT_TblPr places borders before layout/look.
        pr.find(tag('w:tblLayout')).addprevious(pr.find(tag('w:tblBorders')))
        for row in table.rows:
            rp=row._tr.get_or_add_trPr();add(rp,'w:trHeight',**{'w:val':'1200','w:hRule':'atLeast'})
            for j,cell in enumerate(row.cells):
                cell.text=f'Cell {j+1}'
                borders(cell._tc.get_or_add_tcPr(),'w:tcBorders','none',color,size)
        doc.add_paragraph('Following paragraph')
    doc.save(root/'matrix.docx')

def edges():
    variants = [
     ('omitted', {}), ('space', {'sz':'0','color':'auto','space':'1'}),
     ('frame', {'sz':'0','color':'auto','frame':'1'}),
     ('shadow', {'sz':'0','color':'auto','shadow':'1'}),
     ('theme', {'sz':'0','color':'auto','themeColor':'accent1'}),
     ('explicit-defaults', {'sz':'0','color':'auto','space':'0','frame':'0','shadow':'0'}),
    ]
    with ZipFile(root/'matrix.docx') as z:
     parts={n:z.read(n) for n in z.namelist()}
    xml=E.fromstring(parts['word/document.xml']);body=xml.find(q('body'))
    tables=body.findall(q('tbl'))
    for table,(label,attrs) in zip(tables,variants):
     p=table.getprevious()
     p.find('.//'+q('t')).text=label
     for border in table.findall('.//'+q('tcBorders')+'/*'):
      border.attrib.clear();border.set(q('val'),'none')
      for k,v in attrs.items():border.set(q(k),v)
    parts['word/document.xml']=E.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    target=root/'edges.docx'
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
     for name,data in parts.items():z.writestr(name,data)
    print(target)

def style():
    with ZipFile(root/'matrix.docx') as z:parts={n:z.read(n) for n in z.namelist()}
    xml=E.fromstring(parts['word/document.xml']);styles=E.fromstring(parts['word/styles.xml'])
    for table in xml.find(q('body')).findall(q('tbl')):
     pr=table.find(q('tblPr'));sid=pr.find(q('tblStyle')).get(q('val'))
     pr.find(q('tblLook')).set(q('lastRow'),'1');pr.find(q('tblLook')).set(q('val'),'0060')
     style=next(s for s in styles if s.get(q('styleId'))==sid)
     group=E.SubElement(style,q('tblStylePr'),{q('type'):'lastRow'});cellpr=E.SubElement(group,q('tcPr'))
     for index,cp in enumerate(table.findall('.//'+q('tcPr'))):
      border=cp.find(q('tcBorders'))
      if border is not None:
       cp.remove(border)
       if index==0:cellpr.append(border)
    for name,tree in [('word/document.xml',xml),('word/styles.xml',styles)]:
     parts[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    target=root/'style.docx'
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
     for name,data in parts.items():z.writestr(name,data)
    print(target)

def spacing():
    cases=[('none',1,0),('none',1,60),('none',6,60),('nil',1,0),('nil',1,60),('nil',6,60)]
    with ZipFile(root/'matrix.docx') as z:parts={n:z.read(n) for n in z.namelist()}
    xml=E.fromstring(parts['word/document.xml'])
    for table,(style,space,margin) in zip(xml.find(q('body')).findall(q('tbl')),cases):
     table.getprevious().find('.//'+q('t')).text=f'{style}-space{space}-margin{margin}'
     for cp in table.findall('.//'+q('tcPr')):
      for border in cp.find(q('tcBorders')):
       border.attrib.clear()
       for k,v in {'val':style,'sz':'0','color':'auto','space':str(space)}.items():border.set(q(k),v)
      margins=E.SubElement(cp,q('tcMar'))
      for side in ['top','left','bottom','right']:E.SubElement(margins,q(side),{q('w'):str(margin),q('type'):'dxa'})
    parts['word/document.xml']=E.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    target=root/'spacing.docx'
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
     for name,data in parts.items():z.writestr(name,data)
    print(target)

def spacing_rules():
    with ZipFile(root/'spacing.docx') as z:parts={n:z.read(n) for n in z.namelist()}
    xml=E.fromstring(parts['word/document.xml']);styles=E.fromstring(parts['word/styles.xml'])
    tables=xml.find(q('body')).findall(q('tbl'))
    prototype=deepcopy(tables[2])
    for index,(old,label) in enumerate(zip(tables,['exact','mixed','shared','nested','style-space','bottom-only'])):
     table=deepcopy(prototype);old.getparent().replace(old,table)
     table.getprevious().find('.//'+q('t')).text=label
     # Point each replacement at the original distinct blue first-row style.
     table.find(q('tblPr')).find(q('tblStyle')).set(q('val'),f'Probe{index}')
     if label=='exact':table.find('.//'+q('trHeight')).set(q('hRule'),'exact')
     if label=='mixed':
      for ci,cp in enumerate(table.findall('.//'+q('tcPr'))):
       for b in cp.find(q('tcBorders')):b.set(q('space'),'1' if ci==0 else '6')
       for m in cp.find(q('tcMar')):m.set(q('w'),'120' if ci==0 else '0')
     if label=='shared':
      row=deepcopy(table.find(q('tr')));table.append(row)
      for t in row.findall('.//'+q('t')):t.text=t.text.replace('Cell','Next')
     if label=='nested':
      outer=deepcopy(table);row=outer.find(q('tr'));row.remove(row.findall(q('tc'))[1]);cell=row.find(q('tc'))
      for child in list(cell):
       if child.tag!=q('tcPr'):cell.remove(child)
      cp=cell.find(q('tcPr'));cp.remove(cp.find(q('tcBorders')))
      cp.find(q('tcW')).set(q('w'),'7200');cp.insert(1,E.Element(q('gridSpan'),{q('val'):'2'}))
      for m in cp.find(q('tcMar')):m.set(q('w'),'0')
      outer.find(q('tblPr')).remove(outer.find(q('tblPr')).find(q('tblStyle')))
      outer.find(q('tblPr')).remove(outer.find(q('tblPr')).find(q('tblBorders')))
      table.getparent().replace(table,outer);cell.append(table);E.SubElement(cell,q('p'))
     if label=='style-space':
      style=next(s for s in styles if s.get(q('styleId'))==f'Probe{index}')
      group=style.find(q('tblStylePr'));bp=group.find(q('tcPr'));bp.remove(bp.find(q('tcBorders')))
      for ci,cp in enumerate(table.findall('.//'+q('tcPr'))):
       b=cp.find(q('tcBorders'));cp.remove(b)
       if ci==0:bp.append(b)
     if label=='bottom-only':
      for cp in table.findall('.//'+q('tcPr')):
       for b in cp.find(q('tcBorders')):
        if b.tag!=q('bottom'):b.set(q('space'),'0');b.set(q('val'),'nil')
    for name,tree in [('word/document.xml',xml),('word/styles.xml',styles)]:parts[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    target=root/'spacing-rules.docx'
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
     for name,data in parts.items():z.writestr(name,data)
    print(target)

def style_only():
    with ZipFile(root/'matrix.docx') as z:parts={n:z.read(n) for n in z.namelist()}
    xml=E.fromstring(parts['word/document.xml']);styles=E.fromstring(parts['word/styles.xml'])
    for table in xml.find(q('body')).findall(q('tbl')):
     sid=table.find(q('tblPr')).find(q('tblStyle')).get(q('val'))
     style=next(s for s in styles if s.get(q('styleId'))==sid)
     cp=style.find(q('tblStylePr')).find(q('tcPr'));cp.remove(cp.find(q('tcBorders')))
     for index,cellpr in enumerate(table.findall('.//'+q('tcPr'))):
      borders=cellpr.find(q('tcBorders'));cellpr.remove(borders)
      if index==0:cp.append(borders)
    for name,tree in [('word/document.xml',xml),('word/styles.xml',styles)]:parts[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    target=root/'style-only.docx'
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
     for name,data in parts.items():z.writestr(name,data)
    print(target)

def priority():
    with ZipFile(root/'matrix.docx') as z:parts={n:z.read(n) for n in z.namelist()}
    xml=E.fromstring(parts['word/document.xml']);styles=E.fromstring(parts['word/styles.xml'])
    cases=[('firstRow','single','lastRow','none'),('firstRow','none','lastRow','single'),
           ('firstCol','single','firstRow','none'),('firstCol','none','firstRow','single'),
           ('firstRow','single','lastRow','single'),('firstCol','single','firstRow','single')]
    for index,(table,case) in enumerate(zip(xml.find(q('body')).findall(q('tbl')),cases)):
     pr=table.find(q('tblPr'));sid=pr.find(q('tblStyle')).get(q('val'))
     look=pr.find(q('tblLook'));look.set(q('val'),'00E0')
     for key in ['firstRow','lastRow','firstColumn']:look.set(q(key),'1')
     table.getprevious().find('.//'+q('t')).text='/'.join(case)
     style=next(s for s in styles if s.get(q('styleId'))==sid)
     for group in list(style.findall(q('tblStylePr'))):style.remove(group)
     for region,value,color in [(case[0],case[1],'2050D0'),(case[2],case[3],'A030A0')]:
      group=E.SubElement(style,q('tblStylePr'),{q('type'):region});cp=E.SubElement(group,q('tcPr'));bp=E.SubElement(cp,q('tcBorders'))
      for side in ['top','left','bottom','right']:
       E.SubElement(bp,q(side),{q('val'):value,q('sz'):'16',q('color'):color})
     for cp in table.findall('.//'+q('tcPr')):cp.remove(cp.find(q('tcBorders')))
    for name,tree in [('word/document.xml',xml),('word/styles.xml',styles)]:parts[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    target=root/'priority.docx'
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
     for name,data in parts.items():z.writestr(name,data)
    print(target)

def properties():
    with ZipFile(root/'priority.docx') as z:parts={n:z.read(n) for n in z.namelist()}
    styles=E.fromstring(parts['word/styles.xml'])
    for style in styles:
     if not (style.get(q('styleId')) or '').startswith('Probe'):continue
     for index,group in enumerate(style.findall(q('tblStylePr'))):
      pp=E.Element(q('pPr'));E.SubElement(pp,q('ind'),{q('left'):str(240*(index+1))});group.insert(0,pp)
      rp=E.Element(q('rPr'));E.SubElement(rp,q('b' if index==0 else 'i'));E.SubElement(rp,q('sz'),{q('val'):'28' if index==0 else '36'});group.insert(1,rp)
      E.SubElement(group.find(q('tcPr')),q('shd'),{q('val'):'clear',q('color'):'auto',q('fill'):'D0E0F0' if index==0 else 'F0D0E0'})
    parts['word/styles.xml']=E.tostring(styles,xml_declaration=True,encoding='UTF-8',standalone=True)
    target=root/'properties.docx'
    with ZipFile(target,'w',ZIP_DEFLATED) as z:
     for name,data in parts.items():z.writestr(name,data)
    print(target)

if __name__=='__main__':
    root.mkdir(parents=True,exist_ok=True)
    matrix(); edges(); style(); spacing(); spacing_rules(); style_only(); priority(); properties()
