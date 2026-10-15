"""Original DrawingML pattern controls, independent of corpus documents."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree as E
import canvas_group_fixture as g

ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures/pattern-fills'
PRESETS='pct5 pct10 pct20 pct25 pct30 pct40 pct50 pct60 pct70 pct75 pct80 pct90 horz vert ltHorz ltVert dkHorz dkVert narHorz narVert dashHorz dashVert cross dnDiag upDiag ltDnDiag ltUpDiag dkDnDiag dkUpDiag wdDnDiag wdUpDiag dashDnDiag dashUpDiag diagCross smCheck lgCheck smGrid lgGrid dotGrid smConfetti lgConfetti horzBrick diagBrick solidDmnd openDmnd dotDmnd plaid sphere weave divot shingle wave trellis zigZag'.split()

def generate():
 ROOT.mkdir(exist_ok=True)
 cases=[dict(name=p,preset=p) for p in PRESETS]
 cases += [dict(name='omitted preset',preset=None),dict(name='omitted foreground',preset='dkUpDiag',omit='fgClr'),dict(name='omitted background',preset='dkUpDiag',omit='bgClr'),dict(name='omitted colors',preset='dkUpDiag',omit='both')]
 cases += [dict(name='group scale',preset='dkUpDiag',group=True),dict(name='shape rotation',preset='dkUpDiag',rotation=30),dict(name='horizontal flip',preset='dkUpDiag',flip=True),dict(name='translucent',preset='dkUpDiag',alpha=True)]
 cases += [dict(name='shifted origin',preset='dkUpDiag',shift=True),dict(name='uneven frame',preset='dkUpDiag',uneven=True),dict(name='nonuniform group',preset='dkUpDiag',group=True,nonuniform=True),dict(name='group rotation',preset='dkUpDiag',group=True,groupRotation=30),dict(name='ellipse clipping',preset='dkUpDiag',ellipse=True),dict(name='group inherited fill',preset='dkUpDiag',group=True,inherited=True)]
 for preset in ['dnDiag','upDiag','diagCross']:
  cases.append(dict(name=preset+' monochrome',preset=preset,colors=['000000','FFFFFF']))
  cases.append(dict(name=preset+' reversed colors',preset=preset,colors=['F0D080','1248A0']))
 doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12);normal.paragraph_format.space_after=Pt(0)
 doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
 for s in doc.settings.element.iter(qn('w:compatSetting')):
  if s.get(qn('w:name'))=='compatibilityMode':s.set(qn('w:val'),'15')
 for i,case in enumerate(cases,1):
  p=doc.add_paragraph(case['name']);p.paragraph_format.page_break_before=i>1
  paragraph=doc.add_paragraph()
  if case.get('shift'):paragraph.paragraph_format.left_indent=Pt(1.25);paragraph.paragraph_format.space_before=Pt(2.25)
  frame=g.add(g.add(paragraph.add_run()._r,'w:drawing'),'wp:inline',distT=0,distB=0,distL=0,distR=0)
  w,h=180*12700,120*12700
  if case.get('uneven'):w+=15875;h+=28575
  g.add(frame,'wp:extent',cx=w,cy=h);g.add(frame,'wp:docPr',id=i,name=case['name']);g.add(frame,'wp:cNvGraphicFramePr')
  parent=g.add(g.add(frame,'a:graphic'),'a:graphicData',uri=g.NS['wpg' if case.get('group') else 'wps'])
  if case.get('group'):
   parent=g.add(parent,'wpg:wgp');g.add(parent,'wpg:cNvGrpSpPr');props=g.add(parent,'wpg:grpSpPr');g.transform(props,w,h,group=True);props[0].find(g.tag('a:chOff')).attrib.update({'x':'0','y':'0'});w//=2;h//=2
   if case.get('nonuniform'):h*=2;props[0].find(g.tag('a:chExt')).set('cy',str(h))
   if case.get('groupRotation'):props[0].set('rot',str(case['groupRotation']*60000))
  shape=g.shape(parent,100+i,w,h,0,0,'FFFFFF');props=shape.find(g.tag('wps:spPr'))
  if case.get('ellipse'):props.find(g.tag('a:prstGeom')).set('prst','ellipse')
  if case.get('rotation'):props[0].set('rot',str(case['rotation']*60000))
  if case.get('flip'):props[0].set('flipH','1')
  solid=props.find(g.tag('a:solidFill'));index=props.index(solid);props.remove(solid)
  pattern=g.add(props,'a:pattFill',**({'prst':case['preset']} if case['preset'] else {}));props.remove(pattern);props.insert(index,pattern)
  for role,color in zip(['fgClr','bgClr'],case.get('colors',['1248A0','F0D080'])):
   if case.get('omit') in [role,'both']:continue
   c=g.add(g.add(pattern,'a:'+role),'a:srgbClr',val=color)
   if case.get('alpha'):g.add(c,'a:alpha',val=50000 if role=='fgClr' else 25000)
  if case.get('inherited'):
   props.remove(pattern);parent.find(g.tag('wpg:grpSpPr')).append(pattern);props.insert(index,E.Element(g.tag('a:grpFill')))
  doc.add_paragraph('After')
 doc.save(ROOT/'controls.docx');(ROOT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
 print(f'Generated {len(cases)} independent pattern pages')

if __name__=='__main__':generate()
