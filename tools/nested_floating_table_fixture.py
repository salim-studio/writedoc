"""Original nested table positioning probes; no corpus content."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def prop(parent, name, **attrs):
    node = OxmlElement('w:' + name)
    for k,v in attrs.items(): node.set(qn('w:' + k), str(v))
    parent.append(node)
    return node

doc = Document()
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(1)
normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = Pt(15)
cases = [None,
    dict(horzAnchor='page',vertAnchor='page',tblpX='4320',tblpY='4320'),
    dict(horzAnchor='margin',vertAnchor='text',tblpX='720',tblpY='-240',rightFromText='720',bottomFromText='720'),
    dict(horzAnchor='margin',vertAnchor='margin',tblpXSpec='right',tblpYSpec='bottom')]
# Equal offsets distinguish cell-relative placement from physical-page anchors.
cases += [dict(horzAnchor=h, vertAnchor=v, tblpX='720', tblpY='720')
          for h, v in [('page','page'), ('margin','margin'), ('text','text')]]
cases += [dict(horzAnchor='margin', vertAnchor='text', tblpX=x, tblpY=y)
          for x,y in [('0','0'), ('-720','-720'), ('4320','0'), ('720','-240')]]
cases += [dict(horzAnchor='margin', vertAnchor='margin', tblpXSpec=x, tblpYSpec=y)
          for x,y in [('left','top'), ('center','center'), ('right','bottom')]]
cases += [dict(horzAnchor='margin', vertAnchor='text', tblpX='720', tblpY='-240',
               **{distance:'720'}) for distance in ['leftFromText','rightFromText','topFromText','bottomFromText']]
for index, position in enumerate(cases):
    heading = doc.add_paragraph('Nested positioning case ' + str(index))
    heading.paragraph_format.page_break_before = bool(index)
    outer = doc.add_table(rows=1, cols=2); outer.autofit = False
    for col in outer.columns: col.width = Inches(3)
    for i, cell in enumerate(outer.rows[0].cells):
        cell.width = Inches(3); cell.paragraphs[0].text = 'Before nested ' + str(i)
        inner = cell.add_table(rows=1, cols=1); inner.autofit = False
        inner.columns[0].width = Inches(1); inner.cell(0,0).width = Inches(1)
        width=inner._tbl.tblPr.find(qn('w:tblW')); width.set(qn('w:type'),'dxa'); width.set(qn('w:w'),'1440')
        if position and i == 1:
            p = prop(inner._tbl.tblPr,'tblpPr',**position); inner._tbl.tblPr.remove(p);inner._tbl.tblPr.insert(0,p)
        prop(inner.cell(0,0)._tc.get_or_add_tcPr(),'shd',val='clear',fill='D0E8C0')
        inner.cell(0,0).text = 'Nested text'
        cell.paragraphs[-1].text = 'After nested ' + str(i)
    doc.add_paragraph('Following outer table')
target=Path(__file__).resolve().parents[1]/'public/fixtures/nested-floating-tables.docx'
doc.save(target)
(target.parents[2] / 'tests/fixtures/nested-floating-tables/cases.json').write_text(
    json.dumps(cases, indent=2) + '\n')
print(target)
