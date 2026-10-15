"""Original DrawingML ellipse probe; no user artwork or editor serialization."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from lxml import etree as E
from canvas_group_fixture import NS, tag, add, shape
from canvas_group_text_fixture import group

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/ellipses.docx'

def generate():
    doc = Document()
    doc.settings.element.find(tag('w:zoom')).set(tag('w:percent'), '100')
    doc.sections[0].top_margin = doc.sections[0].bottom_margin = Inches(.5)
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)
    doc.add_paragraph('Ellipse geometry')
    for ident, (sx, sy, text, anchor) in enumerate([(1, 1, None, 't'), (2, .5, 'Ellipse text', 't'), (1, 1, 'Ellipse text', 'ctr')], 1):
        p = doc.add_paragraph()
        drawing = E.SubElement(p.add_run()._r, tag('w:drawing'), nsmap={k: NS[k] for k in ['wpg', 'wps', 'a']})
        inline = add(drawing, 'wp:inline')
        width, height = 2743200, 914400
        add(inline, 'wp:extent', cx=width, cy=height)
        add(inline, 'wp:docPr', id=ident, name=f'Ellipse {ident}')
        data = add(add(inline, 'a:graphic'), 'a:graphicData', uri=NS['wpg'])
        g = group(data, 'wpg:wgp', width, height, sx, sy)
        s = shape(g, 100 + ident, int(width / sx), int(height / sy), 25400, -12700, '29A36A', text)
        props = s.find(tag('wps:spPr'))
        props.find(tag('a:prstGeom')).set('prst', 'ellipse')
        line = props.find(tag('a:ln')); line.clear(); line.set('w', '25400')
        add(add(line, 'a:solidFill'), 'a:srgbClr', val='2040D0')
        body = s.find(tag('wps:bodyPr'))
        for key, value in dict(lIns=0, rIns=0, tIns=0, bIns=0, anchor=anchor).items(): body.set(key, str(value))
    doc.add_paragraph('End ellipse geometry.')
    doc.save(SOURCE)

if __name__ == '__main__': generate()
