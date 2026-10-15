"""Original controls isolating automatic line spacing in installed font faces."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures/native-heading-line-metrics'
def generate():
 ROOT.mkdir(parents=True,exist_ok=True)
 d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(8)
 normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
 for s in d.settings.element.iter(qn('w:compatSetting')):
  if s.get(qn('w:name'))=='compatibilityMode':s.set(qn('w:val'),'15')
 cases=[dict(family=f,size=s,bold=b,spacing=a) for f in ['Arial','Microsoft YaHei'] for s in [9,20,44,96] for b in [False,True] for a in [.9,1.2]]
 cases += [dict(family='Microsoft YaHei',size=s,bold=b,italic=True,spacing=a) for s in [20,44] for b in [False,True] for a in [.9,1.2]]
 cases += [dict(family=f,size=s,bold=True,spacing='atLeast24') for f in ['Arial','Microsoft YaHei'] for s in [20,44]]
 cases += [dict(family=f,size=20,bold=True,spacing='exact24') for f in ['Arial','Microsoft YaHei']]
 for i,c in enumerate(cases):
  title=d.add_paragraph('Case '+str(i+1));title.paragraph_format.page_break_before=i>0
  p=d.add_paragraph();p.paragraph_format.line_spacing=Pt(24) if isinstance(c['spacing'],str) else c['spacing']
  if c['spacing']=='atLeast24':p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
  r=p.add_run('Alpha\nBeta\nGamma');r.font.name=c['family'];r.font.size=Pt(c['size']);r.bold=c['bold'];r.italic=c.get('italic',False)
  mark=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
  for slot in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+slot),c['family'])
  mark.append(fonts);b=OxmlElement('w:b');b.set(qn('w:val'),'1' if c['bold'] else '0');mark.append(b)
  sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(c['size']*2));mark.append(sz);p._p.get_or_add_pPr().append(mark)
  d.add_paragraph('After')
 d.save(ROOT/'control.docx');(ROOT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n',encoding='utf-8')
if __name__=='__main__':generate()
