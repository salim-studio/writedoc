"""Independent underline fixtures and export checks using python-docx (MIT)."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_UNDERLINE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/underlines.docx'
PATTERNS = ['single', 'words', 'double', 'thick', 'dotted', 'dottedHeavy', 'dash', 'dashedHeavy', 'dashLong', 'dashLongHeavy', 'dotDash', 'dashDotHeavy', 'dotDotDash', 'dashDotDotHeavy', 'wave', 'wavyHeavy', 'wavyDouble', 'none']

def set_u(parent, pattern=None, color=None):
    old = parent.find(qn('w:u'))
    if old is not None: parent.remove(old)
    u = OxmlElement('w:u')
    if pattern is not None: u.set(qn('w:val'), pattern)
    if color is not None: u.set(qn('w:color'), color)
    parent.append(u)
    return u

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(7); section.page_height = Inches(8)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('UnderlineBase', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal; set_u(base.element.get_or_add_rPr(), 'double', 'A00000')
    child = doc.styles.add_style('UnderlineChild', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base; set_u(child.element.get_or_add_rPr(), color='0000A0')
    char = doc.styles.add_style('UnderlineCharacter', WD_STYLE_TYPE.CHARACTER); set_u(char.element.get_or_add_rPr(), 'wave')
    off = doc.styles.add_style('UnderlineOff', WD_STYLE_TYPE.CHARACTER); off.font.underline = False
    table_style = doc.styles.add_style('UnderlineTable', WD_STYLE_TYPE.TABLE); table_style.base_style = doc.styles['Table Grid']; set_u(table_style.element.get_or_add_rPr(), 'dottedHeavy', '008800')
    doc.add_paragraph('Underline patterns a\u0301😀')
    for pattern in PATTERNS:
        r = doc.add_paragraph(f'{pattern}: Alpha beta 123').runs[0]
        r.underline = next(v for v in WD_UNDERLINE if v.xml_value == pattern)
    doc.add_paragraph('Inherited double red', style=base)
    doc.add_paragraph('Child double blue', style=child)
    doc.add_paragraph('Character wave blue', style=child).runs[0].style = char
    p = doc.add_paragraph('Direct pattern inherits blue', style=child); set_u(p.runs[0]._r.get_or_add_rPr(), 'dash')
    p = doc.add_paragraph('Direct color inherits double', style=base); set_u(p.runs[0]._r.get_or_add_rPr(), color='008800')
    p = doc.add_paragraph('Empty property inherits double', style=base); set_u(p.runs[0]._r.get_or_add_rPr())
    p = doc.add_paragraph('Empty property has no underline'); set_u(p.runs[0]._r.get_or_add_rPr())
    doc.add_paragraph('Character none', style=base).runs[0].style = off
    p = doc.add_paragraph('Words a b\u00a0c\u202fd\u2007e\tem a\u0301😀  '); p.runs[0].underline = WD_UNDERLINE.WORDS
    p = doc.add_paragraph();
    for text, color in [('ava', 'FF0000'), ('tar', '0000FF')]:
        r = p.add_run(text); r.underline = WD_UNDERLINE.DASH; r.font.color.rgb = RGBColor.from_string(color)
    r = doc.add_paragraph('Script wavy double').runs[0]; r.font.underline = WD_UNDERLINE.WAVY_DOUBLE; r.font.superscript = True; r.font.size = Pt(24); r.italic = True
    for tag, value in [('spacing', '10'), ('w', '150'), ('position', '6')]:
        n = OxmlElement('w:' + tag); n.set(qn('w:val'), value); r._r.get_or_add_rPr().insert(0, n)
    # Put properties in schema order independently of the editor serializer.
    rpr = r._r.get_or_add_rPr(); order = ['i', 'spacing', 'w', 'position', 'sz', 'u', 'vertAlign']; rpr[:] = sorted(rpr, key=lambda n: order.index(n.tag.split('}')[-1]))
    p = doc.add_paragraph('Tail' + ' ' * 80); p.runs[0].underline = WD_UNDERLINE.DOUBLE
    p = doc.add_paragraph(); mark = OxmlElement('w:rPr'); set_u(mark, 'wavyHeavy', 'A00000'); p._p.get_or_add_pPr().append(mark)
    p = doc.add_paragraph('Numbered words'); p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 7
    root = doc.part.numbering_part.element; instance = next(n for n in root.findall(qn('w:num')) if n.get(qn('w:numId')) == '7'); abstract = next(n for n in root.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == instance.find(qn('w:abstractNumId')).get(qn('w:val')))
    level = abstract.find(qn('w:lvl')); level.find(qn('w:lvlText')).set(qn('w:val'), 'item %1.'); rpr = OxmlElement('w:rPr'); set_u(rpr, 'words', '008800'); level.append(rpr)
    table = doc.add_table(rows=3, cols=2, style=table_style); table.autofit = True
    for col in table.columns: col.width = Inches(3)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(3)
    table.cell(0,0).text = 'Pattern header'; table.cell(0,1).text = 'Wave header'; table.cell(0,1).paragraphs[0].runs[0].underline = WD_UNDERLINE.WAVY_DOUBLE
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1,0).text = 'Table inherited'; table.cell(1,1).text = 'Underlines across continued rows. ' * 70
    table.cell(2,0).text = 'End cell'; table.cell(2,1).text = 'Other end'
    doc.add_paragraph('End underlines'); doc.save(SOURCE); print(f'Generated independent underline fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file)
    if fresh:
        assert [p.text for p in doc.paragraphs] == ['Double a\u0301😀', 'Words a b\u00a0c', 'Wave text', 'Future dashed']
        for p, pattern in zip(doc.paragraphs, ['double','words','wavyDouble','dash']):
            assert all(r._r.rPr.u.val.xml_value == pattern for r in p.runs)
        assert doc.paragraphs[0].runs[0]._r.rPr.u.get(qn('w:color')) == 'A00000'
    else:
        old = Document(SOURCE)
        assert [p.text for p in doc.paragraphs] == [p.text for p in old.paragraphs]
        assert doc.paragraphs[3].runs[0]._r.rPr.u.val.xml_value == 'wavyDouble'
        assert doc.paragraphs[3].runs[0]._r.rPr.u.get(qn('w:color')) == '008800'
        assert doc.styles['UnderlineBase'].element.rPr.u.val.xml_value == 'dash'
        assert doc.styles['UnderlineChild'].element.rPr.u is None or doc.styles['UnderlineChild'].element.rPr.u.get(qn('w:color')) is None
        def signature(root, omitted):
            return [(n.tag, sorted((k,v) for k,v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in root.iter() if n not in omitted]
        assert signature(old.element, [old.paragraphs[3].runs[0]._r.rPr.u]) == signature(doc.element, [doc.paragraphs[3].runs[0]._r.rPr.u])
        def omit(d): return [d.styles[name].element.rPr.u for name in ['UnderlineBase', 'UnderlineChild'] if d.styles[name].element.rPr.u is not None]
        assert signature(old.styles.element, omit(old)) == signature(doc.styles.element, omit(doc))
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/document.xml','word/styles.xml']: assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: independent underline patterns, color inheritance, text and package preservation')
