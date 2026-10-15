"""Original flat cell-shading fixture and independent python-docx export checks."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/cell-shading.docx'

def shade(pr, **attrs):
    node = OxmlElement('w:shd')
    for key, value in attrs.items(): node.set(qn('w:' + key), value)
    pr.append(node)

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('Canvas Shaded Table', WD_STYLE_TYPE.TABLE); base.base_style = doc.styles['Table Grid']
    pr = OxmlElement('w:tblPr'); base.element.append(pr); shade(pr, val='clear', fill='FFE4CC')
    pr = OxmlElement('w:tcPr'); base.element.append(pr)
    shade(pr, val='clear', fill='112233', themeFill='accent3', themeFillTint='99', themeFillShade='10')
    child = doc.styles.add_style('Canvas Shading Child', WD_STYLE_TYPE.TABLE); child.base_style = base
    reset = doc.styles.add_style('Canvas Shading Reset', WD_STYLE_TYPE.TABLE); reset.base_style = child
    pr = OxmlElement('w:tcPr'); reset.element.append(pr); shade(pr, val='clear', color='auto')
    def table(rows, cols, style=child, height=48):
        result = doc.add_table(rows=rows, cols=cols, style=style); result.autofit = False
        for col in result.columns: col.width = Inches(5 / cols)
        for row in result.rows:
            row.height = Pt(height); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                cell.width = Inches(5 / cols)
                margins = OxmlElement('w:tcMar')
                for side, value in [('top', 90), ('left', 120), ('bottom', 150), ('right', 120)]:
                    n = OxmlElement('w:' + side); n.set(qn('w:w'), str(value)); n.set(qn('w:type'), 'dxa'); margins.append(n)
                cell._tc.get_or_add_tcPr().append(margins)
        return result
    def fill(cell, **attrs):
        # Shading precedes margins in CT_TcPr.
        pr = cell._tc.get_or_add_tcPr(); shade(pr, **attrs); pr.find(qn('w:tcMar')).addprevious(pr[-1])
    doc.add_paragraph('Cell shading A\u0301😀')
    first = table(4, 3)
    for r, labels in enumerate([['Inherited theme', 'Direct blue', 'No fill'], ['Solid foreground', 'Automatic foreground', 'Atomic clear'], ['Horizontal shade', '', 'Nil inherited'], ['Paragraph layer', 'Unchanged theme', '']]):
        for c, label in enumerate(labels): first.cell(r,c).text = label
    fill(first.cell(0,1), val='clear', color='auto', fill='DDEEFF')
    fill(first.cell(0,2), val='clear', color='auto', fill='auto')
    fill(first.cell(1,0), val='solid', color='224466', fill='FFFF00')
    fill(first.cell(1,1), val='solid', color='auto', fill='FFFF00')
    fill(first.cell(1,2), val='clear', color='FF0000')
    first.cell(2,0).merge(first.cell(2,1)); fill(first.cell(2,0), val='clear', fill='CCAAEE')
    fill(first.cell(2,2), val='nil', fill='FF0000')
    fill(first.cell(3,0), val='clear', fill='224466')
    shade(first.cell(3,0).paragraphs[0]._p.get_or_add_pPr(), val='clear', fill='FFFF00')
    fill(first.cell(3,1), val='solid', color='FF0000', themeColor='accent4', themeShade='BF')
    fill(first.cell(3,2), val='clear', fill='55AA77')
    doc.add_paragraph('Style resets')
    clear = table(1, 2, reset); clear.cell(0,0).text = 'Style no fill'; clear.cell(0,1).text = 'Still no fill'
    doc.add_paragraph('Merged colors')
    merged = table(3, 2); merged.cell(0,0).merge(merged.cell(2,0)); merged.cell(0,0).text = 'Vertical theme'
    for r, label in enumerate(['First neighbor', 'Second neighbor', 'Last neighbor']): merged.cell(r,1).text = label
    doc.add_paragraph('Before repeated table')
    repeated = table(3, 2)
    repeated.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    repeated.cell(0,0).text = 'Header theme'; repeated.cell(0,1).text = 'Header blue'
    fill(repeated.cell(0,1), val='clear', fill='DDEEFF')
    repeated.cell(1,0).text = 'Flowing theme. ' * 150; repeated.cell(1,1).text = 'Short no fill'
    fill(repeated.cell(1,1), val='clear', fill='auto')
    repeated.rows[2].height = Inches(7)
    repeated.cell(2,0).text = 'Tall theme'; repeated.cell(2,1).text = 'Tall automatic'
    fill(repeated.cell(2,1), val='solid', color='auto')
    doc.add_paragraph('Tail')
    doc.save(SOURCE); print(f'Generated independent cell-shading fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file); cells = doc.element.findall('.//' + qn('w:tc'))
    text = lambda node: ''.join(n.text or '' for n in node.findall('.//' + qn('w:t')))
    by_text = {text(cell): cell for cell in cells}
    def shading(cell, value):
        node = cell.find(qn('w:tcPr')).find(qn('w:shd'))
        if value is None: assert node is None; return
        assert node is not None
        expected = {'val': 'solid' if value == 'foreground' else 'clear', 'color': 'auto', 'fill': 'auto' if value == 'foreground' else value}
        assert dict(node.attrib) == {qn('w:' + k): v for k,v in expected.items()}, text(cell)
    if fresh:
        assert [text(cell) for cell in cells] == ['New blue', 'New clear', 'New automatic']
        for label, value in [('New blue', '224466'), ('New clear', 'auto'), ('New automatic', 'foreground')]: shading(by_text[label], value)
    else:
        original_doc = Document(SOURCE); before = original_doc.element.findall('.//' + qn('w:tc'))
        assert [text(cell) for cell in cells] == [text(cell) for cell in before]
        changes = {'Inherited theme': '224466', 'Direct blue': None, 'No fill': '55AA77', 'Vertical theme': 'CCAAEE', 'Header theme': '55AA77'}
        # This fixture's sole vertical merge has two physical continuation cells.
        for cell, original in zip(cells, before):
            continuation = cell.find(qn('w:tcPr')).find(qn('w:vMerge'))
            changed = text(cell) in changes or continuation is not None
            if changed: shading(cell, changes.get(text(cell), 'CCAAEE'))
            def data(element):
                return [(n.tag, sorted(n.attrib.items()), n.text) for n in element.iter() if not (changed and n.tag == qn('w:shd') and n.getparent().tag == qn('w:tcPr'))]
            assert data(cell) == data(original), f'Unrelated cell properties or content changed: {text(cell)}'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: cell shading, inheritance, merged segments, source themes and unrelated package parts')
