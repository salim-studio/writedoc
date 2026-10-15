"""Independent drawing/section fixture and export checks.

Original XML and PNG pixels, constructed with python-docx (MIT), lxml (BSD),
and the Python standard library. No user template or font is redistributed.
"""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
from copy import deepcopy
import struct
import zlib
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/grouped-drawings.docx'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
      'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
      'wpg': 'http://schemas.microsoft.com/office/word/2010/wordprocessingGroup',
      'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'}
def tag(name):
    prefix, local = name.split(':'); return '{' + NS[prefix] + '}' + local
def add(parent, element, text=None, **attrs):
    node = E.SubElement(parent, tag(element), {(tag(k) if ':' in k else k): str(v) for k, v in attrs.items()})
    if text is not None: node.text = text
    return node

def transform(parent, width, height, x=0, y=0, group=False):
    node = add(parent, 'a:xfrm'); add(node, 'a:off', x=x, y=y); add(node, 'a:ext', cx=width, cy=height)
    if group: add(node, 'a:chOff', x=1000, y=2000); add(node, 'a:chExt', cx=width // 2, cy=height // 2)

def shape(parent, ident, width, height, x, y, color, text=None, custom=False):
    node = add(parent, 'wps:wsp'); add(node, 'wps:cNvPr', id=ident, name=f'Shape {ident}'); add(node, 'wps:cNvSpPr')
    props = add(node, 'wps:spPr'); transform(props, width, height, x, y)
    if custom:
        geom = add(props, 'a:custGeom'); add(geom, 'a:avLst'); add(geom, 'a:gdLst'); add(geom, 'a:ahLst'); add(geom, 'a:cxnLst'); add(geom, 'a:rect', l='l', t='t', r='r', b='b')
        path = add(add(geom, 'a:pathLst'), 'a:path', w=100, h=100)
        for op, px, py in [('moveTo', 0, 0), ('lnTo', 100, 0), ('lnTo', 90, 100), ('lnTo', 0, 100)]: add(add(path, 'a:' + op), 'a:pt', x=px, y=py)
        add(path, 'a:close')
    else: add(add(props, 'a:prstGeom', prst='rect'), 'a:avLst')
    add(add(props, 'a:solidFill'), 'a:srgbClr', val=color); add(add(props, 'a:ln'), 'a:noFill')
    if text:
        content = add(add(node, 'wps:txbx'), 'w:txbxContent'); p = add(content, 'w:p'); pr = add(p, 'w:pPr'); add(pr, 'w:spacing', **{'w:after': 0})
        r = add(p, 'w:r'); rp = add(r, 'w:rPr'); add(rp, 'w:sz', **{'w:val': 16}); add(rp, 'w:color', **{'w:val': '123456'}); add(r, 'w:t', text=text)
    body = add(node, 'wps:bodyPr', lIns=10000, rIns=10000, tIns=10000, bIns=10000, anchor='t', wrap='square'); add(body, 'a:noAutofit')
    return node

def generate():
    doc = Document(); doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    section = doc.sections[0]; section.page_width = Inches(7); section.page_height = Inches(8)
    section.left_margin = section.right_margin = Inches(.5); section.top_margin = Inches(.25); section.bottom_margin = Inches(.5)
    doc.styles['Normal'].font.size = Pt(11); doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.add_paragraph('Grouped drawings Á😀')
    def chunk(t, b): return struct.pack('>I', len(b)) + t + b + struct.pack('>I', zlib.crc32(t + b) & 0xffffffff)
    png = b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 2, 2, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(b'\0\xff\x20\x10\x20\x50\xc0' * 2)) + chunk(b'IEND', b'')
    temp = doc.add_paragraph(); pic = temp.add_run().add_picture(BytesIO(png), width=Inches(.5))._inline
    fallback = deepcopy(pic.getparent()); picture = deepcopy(pic.find('.//' + tag('pic:pic'))); temp._p.getparent().remove(temp._p)
    for ident, wrapping in [(21, 'wrapNone'), (22, 'wrapTopAndBottom'), (23, None)]:
        p = doc.add_paragraph('Behind the artwork' if ident == 21 else '')
        r = p.add_run()._r; alt = E.SubElement(r, tag('mc:AlternateContent'), nsmap={k: NS[k] for k in ['mc', 'wpg', 'wps', 'a', 'pic', 'r']})
        choice = add(alt, 'mc:Choice', Requires='wpg'); drawing = add(choice, 'w:drawing')
        frame = add(drawing, 'wp:anchor' if wrapping else 'wp:inline', distT=0, distB=0, distL=0, distR=0)
        if wrapping:
            for k,v in {'simplePos': 0, 'relativeHeight': ident, 'behindDoc': int(ident == 21), 'locked': 0, 'layoutInCell': 1, 'allowOverlap': 1}.items(): frame.set(k, str(v))
            add(frame, 'wp:simplePos', x=0, y=0); add(add(frame, 'wp:positionH', relativeFrom='column'), 'wp:posOffset', text='95250'); add(add(frame, 'wp:positionV', relativeFrom='paragraph'), 'wp:posOffset', text='0' if ident == 21 else '95250')
        width, height = 2743200, 685800
        add(frame, 'wp:extent', cx=width, cy=height); add(frame, 'wp:effectExtent', l=0, t=0, r=0, b=0)
        if wrapping: add(frame, 'wp:' + wrapping)
        add(frame, 'wp:docPr', id=ident, name=f'Group {ident}'); add(frame, 'wp:cNvGraphicFramePr')
        data = add(add(frame, 'a:graphic'), 'a:graphicData', uri=NS['wpg']); group = add(data, 'wpg:wgp'); add(group, 'wpg:cNvGrpSpPr'); transform(add(group, 'wpg:grpSpPr'), width, height, group=True)
        nested = add(group, 'wpg:grpSp'); add(nested, 'wpg:cNvPr', id=100+ident, name='Nested'); add(nested, 'wpg:cNvGrpSpPr'); transform(add(nested, 'wpg:grpSpPr'), width // 2, height // 2, 1000, 2000, True)
        shape(nested, 200+ident, width // 4, height // 4, 1000, 2000, 'C2CB93', 'Group text', custom=True)
        nested.append(deepcopy(picture)); t = nested[-1].find('.//' + tag('a:xfrm')); t.find(tag('a:off')).set('x', '400000'); t.find(tag('a:off')).set('y', '2000'); t.find(tag('a:ext')).set('cx', '120000'); t.find(tag('a:ext')).set('cy', '120000')
        fb = add(alt, 'mc:Fallback'); drawing = deepcopy(fallback); drawing.find('.//' + tag('wp:docPr')).set('id', str(ident)); fb.append(drawing)
        if ident == 22: doc.add_paragraph('Text after the wrapping band.')
    doc.add_paragraph('End of the first section.')
    section = doc.add_section(WD_SECTION_START.NEW_PAGE); section.top_margin = Inches(1)
    doc.add_paragraph('Second section starts lower.'); doc.add_table(rows=1, cols=2).cell(0, 0).text = 'Table retained'
    doc.add_paragraph('End of report.')
    doc.save(SOURCE); print(f'Generated {SOURCE}')

def verify(path):
    def signature(n): return [(x.tag, sorted(x.attrib.items()), x.text or '') for x in n.iter()]
    with ZipFile(SOURCE) as old, ZipFile(path) as new:
        before, after = [E.fromstring(z.read('word/document.xml')) for z in [old, new]]
        texts = lambda root: root.xpath('//w:body/w:p/w:r/w:t/text()', namespaces=NS)
        assert texts(after) == ['Edited ' + texts(before)[0], *texts(before)[1:]] or ''.join(texts(after)) == 'Edited ' + ''.join(texts(before))
        for name in ['mc:AlternateContent', 'w:sectPr', 'w:tbl']:
            assert [signature(n) for n in before.findall('.//' + tag(name))] == [signature(n) for n in after.findall('.//' + tag(name))], f'Changed {name}'
        assert set(old.namelist()) == set(new.namelist())
        for name in old.namelist():
            if name != 'word/document.xml': assert old.read(name) == new.read(name), name
    print(f'PASS {path}: independently verified both drawing branches, section properties, table and untouched package parts')

if __name__ == '__main__':
    import sys
    if len(sys.argv) == 1: generate()
    else:
        for path in sys.argv[1:]: verify(path)
