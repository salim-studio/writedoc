"""Original multi-paragraph plain-control cache probes; no corpus content."""
from pathlib import Path
import json
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1] / 'tests/fixtures/plain-text-block-placeholders'
CASES = [
    ('body-placeholder', True, True, True, False, False),
    ('cell-placeholder', True, True, True, True, False),
    ('body-existing', True, False, True, False, False),
    ('cell-existing', True, False, True, True, False),
    ('singleline-placeholder', True, True, False, False, False),
    ('mixed-placeholder', True, True, True, False, True),
    ('persistent-placeholder', False, True, True, False, False),
    ('persistent-existing', False, False, True, False, False),
]


def child(parent, name, value=None):
    node = OxmlElement('w:' + name)
    if value is not None:
        node.set(qn('w:val'), str(value))
    parent.append(node)
    return node


def generate():
    doc = Document()
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.space_after = Pt(4)
    for index, (name, temporary, placeholder, multiline, cell, mixed) in enumerate(CASES, 1):
        doc.add_paragraph(name)
        host = doc.add_table(rows=1, cols=1).cell(0, 0) if cell else doc
        first = host.paragraphs[0] if cell else host.add_paragraph()
        first.add_run('Alpha Bravo')
        second = host.add_paragraph('Second line')
        third = host.add_paragraph('Third line')
        if mixed:
            second.runs[0].bold = True
        sdt = OxmlElement('w:sdt')
        properties = child(sdt, 'sdtPr')
        child(properties, 'alias', name); child(properties, 'id', index)
        if temporary:
            child(properties, 'temporary')
        if placeholder:
            child(properties, 'showingPlcHdr')
        text = child(properties, 'text'); text.set(qn('w:multiLine'), '1' if multiline else '0')
        content = child(sdt, 'sdtContent')
        first._p.addprevious(sdt)
        for paragraph in [first, second, third]:
            content.append(paragraph._p)
        if cell:
            host.add_paragraph('Cell tail')
        doc.add_paragraph('Outside ' + name)
    ROOT.mkdir(parents=True, exist_ok=True)
    doc.save(ROOT / 'source.docx')
    # Isolate cases: the reference application may reject a semantically invalid control even though
    # the package passes schemas. One rejected case must not hide the others.
    with ZipFile(ROOT / 'source.docx') as package:
        parts = {name: package.read(name) for name in package.namelist()}
    for name, *_ in CASES:
        root = E.fromstring(parts['word/document.xml'])
        for control in root.iter(qn('w:sdt')):
            if control.find(qn('w:sdtPr')).find(qn('w:alias')).get(qn('w:val')) != name:
                control.getparent().remove(control)
        variant = dict(parts)
        variant['word/document.xml'] = E.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        with ZipFile(ROOT / (name + '-source.docx'), 'w', ZIP_DEFLATED) as package:
            for part, content in variant.items():
                package.writestr(part, content)
    (ROOT / 'cases.json').write_text(json.dumps([
        dict(name=n, temporary=t, placeholder=p, multiLine=m, cell=c, mixed=f)
        for n, t, p, m, c, f in CASES], indent=2) + '\n')


if __name__ == '__main__':
    generate()
