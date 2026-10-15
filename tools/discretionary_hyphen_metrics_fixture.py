from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
root=Path(__file__).resolve().parents[1]/'tests/fixtures/discretionary-hyphen-metrics';root.mkdir(parents=True,exist_ok=True);doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for s in doc.settings.element.findall('.//'+qn('w:compatSetting')):
 if s.get(qn('w:name'))=='compatibilityMode':s.set(qn('w:val'),'15')
doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].paragraph_format.space_after=Pt(0);doc.styles['Normal'].paragraph_format.space_before=Pt(12)
buf=BytesIO();Image.new('RGB',(120,40),'#D02030').save(buf,format='PNG');cases=[]
for size in [10,15,20]:
 for height in [18,36]:
  for kind,tail in [('soft','\u00adXYZ'),('literal','-XYZ'),('space',' XYZ'),('letter','a XYZ')]:
   if cases:doc.add_page_break()
   doc.styles['Normal'].font.size=Pt(15)
   doc.add_paragraph(f'{size} {height} {kind}')
   t=doc.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Inches(1);c=t.cell(0,0);c.width=Inches(1)
   m=OxmlElement('w:tblCellMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'0');e.set(qn('w:type'),'dxa');m.append(e)
   t._tbl.tblPr.insert_element_before(m,'w:tblLook')
   p=c.paragraphs[0];r=p.add_run();r.font.size=Pt(size);r.add_picture(BytesIO(buf.getvalue()),width=Pt(54),height=Pt(height))
   for char in tail:
    r=p.add_run();r.font.size=Pt(size)
    if char=='\u00ad':r._r.append(OxmlElement('w:softHyphen'))
    else:r.add_text(char)
   c.add_paragraph('Next');doc.add_paragraph('After');cases.append(dict(size=size,height=height,kind=kind))
doc.save(root/'control.docx');(root/'cases.json').write_text(json.dumps(cases,indent=2));print(len(cases))
