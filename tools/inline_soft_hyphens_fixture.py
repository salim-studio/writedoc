from pathlib import Path
import argparse
parser=argparse.ArgumentParser();parser.add_argument("--mode",type=int,choices=[14,15],default=15);args=parser.parse_args()
from io import BytesIO
from PIL import Image
from canvas_group_fixture import NS, tag, add, shape
from canvas_group_text_fixture import group
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path('tests/fixtures/inline-soft-hyphens');root.mkdir(parents=True,exist_ok=True)
doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for setting in doc.settings.element.findall('.//'+qn('w:compatSetting')):
 if setting.get(qn('w:name'))=='compatibilityMode':setting.set(qn('w:val'),str(args.mode))
doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(15)
doc.styles['Normal'].paragraph_format.space_after=Pt(0)
doc.styles['Normal'].paragraph_format.space_before=Pt(12)
im=Image.new('RGB',(120,40),'#D02030');buf=BytesIO();im.save(buf,format='PNG')
cases=[('plain','',90),('terminal','\u00ad',90),('repeated','\u00ad\u00ad',90),('word','\u00adXYZ',90),('space',' ',90),('break','\u00ad\n',90),('fitting','\u00ad',54),('fitting-word','\u00adXYZ',54),('effects','',54),('effects-terminal','\u00ad',54),('effects-word','\u00adXYZ',54),('effects-space',' ',54),('group','',90),('group-terminal','\u00ad',90),('group-effects','',54),('group-effects-terminal','\u00ad',54)]
for i,(name,tail,width) in enumerate(cases):
 if i:doc.add_page_break()
 doc.add_paragraph('Inline soft hyphen: '+name)
 t=doc.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Inches(1);c=t.cell(0,0);c.width=Inches(1)
 pr=t._tbl.tblPr;m=OxmlElement('w:tblCellMar')
 for edge in ['top','left','bottom','right']:
  e=OxmlElement('w:'+edge);e.set(qn('w:w'),'0');e.set(qn('w:type'),'dxa');m.append(e)
 pr.insert_element_before(m,'w:tblLook','w:tblCaption','w:tblDescription','w:tblPrChange')
 p=c.paragraphs[0];picture=p.add_run().add_picture(BytesIO(buf.getvalue()),width=Pt(width))
 if name.startswith('group'):
  data=picture._inline.find('.//'+tag('a:graphicData'));data.clear();data.set('uri',NS['wpg'])
  g=group(data,'wpg:wgp',width*12700,int(width/3*12700),1,1)
  shape(g,200+i,width*12700,int(width/3*12700),25400,-12700,'D02030')
 if 'effects' in name:
  effect=OxmlElement('wp:effectExtent')
  for edge in ['l','r']:effect.set(edge,str(18*12700))
  for edge in ['t','b']:effect.set(edge,'0')
  picture._inline.insert(1,effect)
 for ch in tail:
  r=p.add_run()
  if ch=='\u00ad':r._r.append(OxmlElement('w:softHyphen'))
  elif ch=='\n':r.add_break()
  else:r.add_text(ch)
 c.add_paragraph('Next')
 doc.add_paragraph('After')
doc.save(root/f'control{args.mode}.docx');print(root/f'control{args.mode}.docx')
