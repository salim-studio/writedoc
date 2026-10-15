"""Independent controls for blank-character fonts on text lines."""
def make_text(mode):
    from pathlib import Path
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    import json
    out=Path(__file__).resolve().parents[1]/'tests/fixtures/text-space-metrics';out.mkdir(exist_ok=True)
    d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    for s in d.settings.element.iter(qn('w:compatSetting')):
     if s.get(qn('w:name'))=='compatibilityMode':s.set(qn('w:val'),str(mode))
    s=d.styles['Normal'];s.font.name='Calibri';s.font.size=Pt(12);s.paragraph_format.space_after=Pt(0);s.paragraph_format.line_spacing=1
    cases=[]
    for kind,text in [('space',' '),('nbsp','\u00a0'),('em','\u2003'),('tab','\t')]:
     for position in ['leading','middle','trailing']:
      for size in [9,24]:
       c=dict(kind=kind,position=position,size=size);cases.append(c)
       p=d.add_paragraph(f'Case {len(cases)}');p.paragraph_format.page_break_before=len(cases)>1
       d.add_paragraph('Before');p=d.add_paragraph()
       if position!='leading':p.add_run('First')
       r=p.add_run(text);r.font.size=Pt(size)
       if position!='trailing':p.add_run('Last')
       d.add_paragraph('After')
    d.save(out/f'mode{mode}.docx');(out/'text-cases.json').write_text(json.dumps(cases,indent=2))

def make_drawings():
    from pathlib import Path
    from io import BytesIO
    from PIL import Image
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    import json
    out=Path(__file__).resolve().parents[1]/'tests/fixtures/text-space-metrics';out.mkdir(exist_ok=True)
    d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    for s in d.settings.element.iter(qn('w:compatSetting')):
     if s.get(qn('w:name'))=='compatibilityMode':s.set(qn('w:val'),'15')
    s=d.styles['Normal'];s.font.name='Trebuchet MS';s.font.size=Pt(8);s.paragraph_format.space_after=Pt(0);s.paragraph_format.line_spacing=1
    cases=[]
    for kind in ['text','picture','shape']:
     for height in ([0] if kind=='text' else [.1,.45,.75,1,2,4]):
      for space in [False,True]:
       c=dict(kind=kind,height=height,space=space);cases.append(c)
       p=d.add_paragraph(f'Case {len(cases)}');p.paragraph_format.page_break_before=len(cases)>1
       d.add_paragraph('Before');p=d.add_paragraph('Website')
       if space:r=p.add_run(' ');r.font.size=Pt(9)
       if kind!='text':
        r=p.add_run();r.font.size=Pt(9);b=BytesIO();Image.new('RGB',(4,4),'red').save(b,format='PNG');r.add_picture(BytesIO(b.getvalue()),width=Pt(80),height=Pt(height))
        if kind=='shape':
         gd=r._r.find('.//'+qn('a:graphicData'));gd.clear();gd.set('uri','http://schemas.microsoft.com/office/word/2010/wordprocessingShape')
         gd.append(parse_xml(f'<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><wps:cNvCnPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1016000" cy="{int(height*12700)}"/></a:xfrm><a:prstGeom prst="line"><a:avLst/></a:prstGeom><a:noFill/><a:ln w="6350"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill></a:ln></wps:spPr><wps:bodyPr/></wps:wsp>'))
       d.add_paragraph('After')
    d.save(out/'mixed-drawings.docx');(out/'mixed-drawings-cases.json').write_text(json.dumps(cases,indent=2))

def make_decorations():
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_TAB_LEADER
    from pathlib import Path
    from docx.oxml.ns import qn
    out=Path(__file__).resolve().parents[1]/'tests/fixtures/text-space-metrics';out.mkdir(exist_ok=True)
    d=Document();d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
    s=d.styles['Normal'];s.font.name='Calibri';s.font.size=Pt(12);s.paragraph_format.space_after=Pt(0);s.paragraph_format.line_spacing=1
    for i,kind in enumerate(['space','underline','strike','highlight','tab','dots','line']):
     p=d.add_paragraph(kind);p.paragraph_format.page_break_before=i>0
     d.add_paragraph('Before');p=d.add_paragraph('First');r=p.add_run('\t' if kind in ['tab','dots','line'] else ' ');r.font.size=Pt(24)
     if kind=='underline':r.underline=True
     if kind=='strike':r.font.strike=True
     if kind=='highlight':r.font.highlight_color=7
     if kind in ['dots','line']:p.paragraph_format.tab_stops.add_tab_stop(Pt(100),leader=WD_TAB_LEADER.DOTS if kind=='dots' else WD_TAB_LEADER.LINES)
     p.add_run('Last');d.add_paragraph('After')
    d.save(out/'decorations.docx')

make_text(14)
make_text(15)
make_drawings()
make_decorations()
