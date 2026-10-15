"""Original inline pictures inside retained/grouped WordprocessingML text boxes."""
from pathlib import Path
from io import BytesIO
from copy import deepcopy
import argparse
from PIL import Image, ImageDraw
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt
from canvas_group_fixture import NS, tag, add, shape
from canvas_group_text_fixture import group

parser=argparse.ArgumentParser()
parser.add_argument('--reflections',action='store_true')
args=parser.parse_args()
root = Path(__file__).resolve().parents[1] / ('tests/fixtures/reflected-textboxes' if args.reflections else 'tests/fixtures/textbox-pictures')
root.mkdir(parents=True, exist_ok=True)
im = Image.new('RGB', (80, 40), '#D02030')
d = ImageDraw.Draw(im); d.rectangle((40, 0, 79, 39), fill='#2050D0'); d.rectangle((32, 12, 47, 27), fill='#F0C020')
buf = BytesIO(); im.save(buf, format='PNG')
doc = Document()
doc.settings.element.find(tag('w:zoom')).set(tag('w:percent'), '100')
doc.styles['Normal'].font.name = 'Arial'; doc.styles['Normal'].font.size = Pt(12)
doc.styles['Normal'].paragraph_format.space_after = Pt(0)
cases = [('identity',1,1,False,'t',0),('scaled',2,2,False,'t',0),
 ('nonuniform',2,.5,False,'t',0),('nested',2,2,True,'t',0),
 ('rotated',1,1,False,'t',15),('centered',1,1,False,'ctr',0),
 ('clipped',1,1,False,'t',0),('crop-flip',1,1,False,'t',0)]
if args.reflections:
 cases=[(name,sx,sy,nested,'t',rotation) for name,sx,sy,nested,rotation in [
  ('H',1,1,False,0),('V',1,1,False,0),('B',1,1,False,0),
  ('H-rotated',1,1,False,30),('V-rotated',1,1,False,30),('B-rotated',1,1,False,30),
  ('H-nested-H',2,2,True,0),('H-nested-V',2,2,True,0),('V-nested-H',2,2,True,30),
  ('H-nonuniform',2,.5,False,30),('V-nonuniform',2,.5,False,90),('H-shape-V',1,1,False,0)]]
for i,(name,sx,sy,nested,anchor,rotation) in enumerate(cases):
 if i: doc.add_page_break()
 doc.add_paragraph('Textbox pictures: '+name)
 temp = doc.add_paragraph()
 pic = temp.add_run().add_picture(BytesIO(buf.getvalue()),width=Inches(.9))._inline
 pic.find(tag('wp:docPr')).set('id',str(100+i))
 if name=='crop-flip':
  x=pic.find('.//'+tag('a:xfrm')); x.set('flipH','1')
  fill=pic.find('.//'+tag('pic:blipFill')); crop=E.Element(tag('a:srcRect'),l='25000',r='0',t='0',b='0');fill.insert(1,crop)
 picture = deepcopy(pic.getparent());temp._p.getparent().remove(temp._p)
 p=doc.add_paragraph()
 drawing=E.SubElement(p.add_run()._r,tag('w:drawing'),nsmap={k:NS[k] for k in ['wpg','wps','a']})
 inline=add(drawing,'wp:inline');width=2743200;height=274320 if name=='clipped' else 1280160
 add(inline,'wp:extent',cx=width,cy=height);add(inline,'wp:docPr',id=200+i,name=name)
 data=add(add(inline,'a:graphic'),'a:graphicData',uri=NS['wpg'])
 g=group(data,'wpg:wgp',width,height,sx,sy);sw,sh=int(width/sx),int(height/sy)
 if args.reflections:
  x=g.find(tag('wpg:grpSpPr')+'/'+tag('a:xfrm'))
  x.set('flipH',str(int(name[0] in ['H','B'])));x.set('flipV',str(int(name[0] in ['V','B'])))
 if nested:
  g=group(g,'wpg:grpSp',sw,sh,2,.5,25400,-12700);sw,sh=sw//2,sh*2
  if args.reflections:
   x=g.find(tag('wpg:grpSpPr')+'/'+tag('a:xfrm'));x.set('flip'+name[-1],'1')
 s=shape(g,300+i,sw,sh,25400,-12700,'E4E8EC','placeholder')
 if rotation:s.find('.//'+tag('a:xfrm')).set('rot',str(rotation*60000))
 if name=='H-shape-V':s.find('.//'+tag('a:xfrm')).set('flipV','1')
 content=s.find('.//'+tag('w:txbxContent'));content.clear()
 p=add(content,'w:p');pr=add(p,'w:pPr');add(pr,'w:spacing',**{'w:after':0})
 for text in ['Before ',None,' After']:
  r=add(p,'w:r')
  if text is None:r.append(picture)
  else:
   t=add(r,'w:t',text=text);t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
 add(add(add(content,'w:p'),'w:r'),'w:t',text='Next line')
 body=s.find(tag('wps:bodyPr'))
 for key,value in dict(lIns=45720,rIns=45720,tIns=45720,bIns=45720,anchor=anchor).items():body.set(key,str(value))
 doc.add_paragraph('After drawing')
doc.save(root/'control.docx')
print(root/'control.docx')
