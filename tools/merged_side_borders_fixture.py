"""Original controls for physical side borders of vertically merged cells."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

folder = Path(__file__).resolve().parents[1] / 'tests/fixtures/merged-side-borders'
folder.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10)
normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = Pt(12)
cases = [
    ('restart-only', ['single', None, None], False, False),
    ('middle-gap', ['single', None, 'single'], False, False),
    ('continuation-only', [None, 'single', 'single'], False, False),
    ('nil-over-table', ['single', 'nil', 'single'], True, False),
    ('none-inherits-table', ['single', 'none', 'single'], True, False),
    ('opposing-neighbor', ['single', None, 'single'], False, True),
]

def edge(parent, side, style, color='D02020', size=16):
    node = OxmlElement('w:' + side)
    for key, value in [('val', style), ('sz', size), ('color', color)]:
        node.set(qn('w:' + key), str(value))
    parent.append(node)

for index, (name, sides, inherited, opposing) in enumerate(cases):
    title = doc.add_paragraph(name)
    title.paragraph_format.page_break_before = bool(index)
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    for column in table.columns:
        column.width = Inches(2)
    if inherited:
        borders = OxmlElement('w:tblBorders')
        edge(borders, 'insideV', 'single', '2060D0')
        props = table._tbl.tblPr
        props.insert(list(props).index(props.find(qn('w:tblLayout'))), borders)
    # Populate before adding merges, so python-docx cannot redirect a physical
    # continuation cell to its logical owner.
    for ri, row in enumerate(table.rows):
        row.height = Pt(36)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(2)
            cell.text = f'R{ri}' if ci else ('Merged' if ri == 0 else '')
    for ri, row in enumerate(table._tbl.tr_lst):
        for ci, cell in enumerate(row.tc_lst):
            props = cell.get_or_add_tcPr()
            if ci == 0:
                merge = OxmlElement('w:vMerge')
                if ri == 0:
                    merge.set(qn('w:val'), 'restart')
                props.append(merge)
            if ci == 0 and sides[ri] or ci == 1 and opposing and ri == 1:
                borders = OxmlElement('w:tcBorders')
                edge(borders, 'right' if ci == 0 else 'left', sides[ri] if ci == 0 else 'single', 'D02020' if ci == 0 else '208040')
                props.append(borders)
    doc.add_paragraph('After')
doc.save(folder / 'controls.docx')
(folder / 'cases.json').write_text(json.dumps([c[0] for c in cases], indent=2) + '\n')
