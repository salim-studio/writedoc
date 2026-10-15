"""Original cell-margin fixture and independent python-docx export checks."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/cell-margins.docx'
ORDER = ['top', 'start', 'left', 'bottom', 'end', 'right']

def margins(pr, values, tag='tcMar'):
    parent = pr.find(qn('w:' + tag))
    if parent is None:
        parent = OxmlElement('w:' + tag)
        pr.insert_element_before(parent, 'w:vAlign', 'w:hideMark', 'w:tblLook', 'w:tblCaption', 'w:tblDescription')
    for side in ORDER:
        if side not in values: continue
        old = parent.find(qn('w:' + side))
        if old is not None: parent.remove(old)
        value = values[side]; node = OxmlElement('w:' + side)
        for key, val in (value if isinstance(value, dict) else {'w': str(value), 'type': 'dxa'}).items(): node.set(qn('w:' + key), val)
        before = next((n for n in parent if ORDER.index(n.tag.split('}')[1]) > ORDER.index(side)), None)
        if before is None: parent.append(node)
        else: before.addprevious(node)

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('Canvas Margin Base', WD_STYLE_TYPE.TABLE); base.base_style = doc.styles['Table Grid']
    pr = OxmlElement('w:tblPr'); base.element.append(pr); margins(pr, {'top': 80, 'left': 160, 'bottom': 120, 'right': 200}, 'tblCellMar')
    pr = OxmlElement('w:tcPr'); base.element.append(pr); margins(pr, {'top': 60, 'left': 240})
    child = doc.styles.add_style('Canvas Margin Child', WD_STYLE_TYPE.TABLE); child.base_style = base
    pr = OxmlElement('w:tcPr'); child.element.append(pr); margins(pr, {'left': {'w': '999', 'type': 'nil'}, 'right': 180})
    def table(rows, cols, style=child, height=42):
        result = doc.add_table(rows=rows, cols=cols, style=style); result.autofit = False
        for col in result.columns: col.width = Inches(5 / cols)
        for row in result.rows:
            row.height = Pt(height); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells: cell.width = Inches(5 / cols)
        return result
    doc.add_paragraph('Cell margins A\u0301😀')
    first = table(3, 3)
    for r, labels in enumerate([['Inherited padding', 'Logical padding', 'Ignored padding'], ['No padding', 'Text wraps ' * 16, 'Bottom aligned'], ['Horizontal margin', '', 'Unicode A\u0301😀']]):
        for c, label in enumerate(labels): first.cell(r,c).text = label
    margins(first._tbl.tblPr, {'top': 100, 'left': 360}, 'tblCellMar')
    margins(first.cell(0,1)._tc.get_or_add_tcPr(), {'top': {'w': '12pt', 'type': 'pct'}, 'start': 540, 'left': 20, 'bottom': 0, 'right': {'w': '999', 'type': 'nil'}})
    margins(first.cell(0,2)._tc.get_or_add_tcPr(), {'top': {'w': '100', 'type': 'auto'}, 'left': {'w': '20%', 'type': 'pct'}, 'bottom': {'w': '999', 'type': 'nil'}, 'right': {'w': '123', 'type': 'auto'}})
    margins(first.cell(1,0)._tc.get_or_add_tcPr(), dict.fromkeys(['top','left','bottom','right'], 0))
    first.cell(1,2).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    first.cell(2,0).merge(first.cell(2,1)); first.cell(2,0).text = 'Horizontal margin'
    margins(first.cell(2,0)._tc.get_or_add_tcPr(), {'top': 90, 'bottom': 210})
    from canvas_paragraph_border_fixture import shade
    shade(first.cell(2,0).paragraphs[0]._p.get_or_add_pPr(), 'DDEEFF')
    doc.add_paragraph('Merged margins')
    merged = table(3, 2); merged.cell(0,0).merge(merged.cell(2,0)); merged.cell(0,0).text = 'Vertical margin'
    for r, label in enumerate(['First neighbor', 'Second neighbor', 'Last neighbor']):
        merged.cell(r,1).text = label
        margins(merged.rows[r]._tr.tc_lst[0].get_or_add_tcPr(), {'left': 180 if r == 0 else 60, 'top': 120, 'bottom': 180})
    doc.add_paragraph('Automatic widths')
    automatic = table(1, 2); automatic.autofit = True
    automatic.cell(0,0).text = 'AutoFit first'; automatic.cell(0,1).text = 'Longer words adapt to wider padding'
    bare = doc.styles.add_style('Canvas Bare Margins', WD_STYLE_TYPE.TABLE)
    defaults = table(1, 2, bare); defaults.cell(0,0).text = 'Default padding'; defaults.cell(0,1).text = 'Default neighbor'
    doc.add_paragraph('Before repeated table')
    repeated = table(3, 2)
    repeated.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    repeated.cell(0,0).text = 'Header padding'; repeated.cell(0,1).text = 'Other header'
    repeated.cell(1,0).text = 'Flowing margins. ' * 120; repeated.cell(1,1).text = 'Short neighbor'
    repeated.rows[2].height = Inches(7); repeated.cell(2,0).text = 'Tall margin'; repeated.cell(2,1).text = 'Tall neighbor'
    doc.add_paragraph('Tail'); doc.save(SOURCE); print(f'Generated independent cell-margin fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file); cells = doc.element.findall('.//' + qn('w:tc'))
    text = lambda cell: ''.join(n.text or '' for n in cell.findall('.//' + qn('w:t')))
    def declaration(cell, side, value):
        parent = cell.find(qn('w:tcPr')).find(qn('w:tcMar'))
        node = parent.find(qn('w:' + side)) if parent is not None else None
        if value is None: assert node is None; return
        assert node is not None and dict(node.attrib) == {qn('w:type'): 'dxa', qn('w:w'): str(value)}, (text(cell), side)
    if fresh:
        assert [text(cell) for cell in cells] == ['New padding', 'New zero']
        for cell in cells:
            values = {'top': 120, 'left': 240, 'bottom': 180, 'right': 300} if text(cell) == 'New padding' else dict.fromkeys(['top','left','bottom','right'], 0)
            for side, value in values.items(): declaration(cell, side, value)
    else:
        original_doc = Document(SOURCE); before = original_doc.element.findall('.//' + qn('w:tc'))
        assert [text(cell) for cell in cells] == [text(cell) for cell in before]
        changes = {'Inherited padding': {'top': 180, 'left': 300}, 'Logical padding': dict.fromkeys(ORDER), 'Ignored padding': {'left': 0}, 'Header padding': {'top': 240, 'bottom': 180}}
        for cell, original in zip(cells, before):
            edits = changes.get(text(cell), {})
            if cell.find(qn('w:tcPr')).find(qn('w:vMerge')) is not None: edits = {'left': 360}
            for side, value in edits.items(): declaration(cell, side, value)
            def data(element):
                return [(n.tag, sorted(n.attrib.items()), n.text) for n in element.iter() if n.tag != qn('w:tcMar') and not (n.getparent() is not None and n.getparent().tag == qn('w:tcMar') and n.tag.split('}')[1] in edits)]
            assert data(cell) == data(original), f'Unrelated properties or text changed: {text(cell)}'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: cell margins, aliases, inheritance, merge segments, text and unrelated package parts')
