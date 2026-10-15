"""Independent original controls for automatic WordprocessingML shape text color."""
from pathlib import Path
import sys,json
import canvas_group_fixture as g
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
out=Path(__file__).resolve().parents[1]/'tests/fixtures/shape-auto-color';out.mkdir(exist_ok=True)
doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12)
for name,typ in [('AutoParagraph',WD_STYLE_TYPE.PARAGRAPH),('AutoCharacter',WD_STYLE_TYPE.CHARACTER)]:
 s=doc.styles.add_style(name,typ);g.add(s.element.get_or_add_rPr(),'w:color',**{'w:val':'auto'})
cases=[]
for fill in ['FFFFFF','000000','548235']:
 for ref in ['FFFFFF','FF0000',None]:
  for level in (['omitted','paragraph','character','direct','explicit'] if ref else ['direct','explicit']):
   i=len(cases)+1;cases.append(dict(case=i,fill=fill,ref=ref,level=level))
   p=doc.add_paragraph(f'Case {i}: {fill}, ref {ref}, {level}');p.paragraph_format.page_break_before=i>1
   frame=g.add(g.add(doc.add_paragraph().add_run()._r,'w:drawing'),'wp:inline',distT=0,distB=0,distL=0,distR=0)
   w,h=240*12700,70*12700;g.add(frame,'wp:extent',cx=w,cy=h);g.add(frame,'wp:docPr',id=i,name=f'Color {i}');g.add(frame,'wp:cNvGraphicFramePr')
   data=g.add(g.add(frame,'a:graphic'),'a:graphicData',uri=g.NS['wps']);shape=g.shape(data,i+100,w,h,0,0,fill,'Color sample')
   if ref:
    style=g.add(shape,'wps:style');shape.remove(style);shape.insert(3,style)
    for name in ['lnRef','fillRef','effectRef']:g.add(g.add(style,'a:'+name,idx=0),'a:srgbClr',val='000000')
    g.add(g.add(style,'a:fontRef',idx='minor'),'a:srgbClr',val=ref)
   p=shape.find('.//'+g.tag('w:p'));r=p.find(g.tag('w:r'));rp=r.find(g.tag('w:rPr'));rp.remove(rp.find(g.tag('w:color')))
   if level=='explicit':g.add(rp,'w:color',**{'w:val':'0000FF'})
   if level=='direct':g.add(rp,'w:color',**{'w:val':'auto'})
   if level=='character':rp.insert(0,g.add(rp,'w:rStyle',**{'w:val':'AutoCharacter'}))
   if level=='paragraph':pr=p.find(g.tag('w:pPr'));pr.insert(0,g.add(pr,'w:pStyle',**{'w:val':'AutoParagraph'}))
doc.save(out/'control.docx');(out/'cases.json').write_text(json.dumps(cases,indent=2))
