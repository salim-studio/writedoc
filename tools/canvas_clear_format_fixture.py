"""Original direct-format reset fixture and independent python-docx checks."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import struct
import zlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/clear-text-formatting.docx'

def element(parent, tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items(): node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(6)
    base = doc.styles.add_style('Reset Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    base.font.bold = True; base.font.size = Pt(14); base.font.color.rgb = RGBColor.from_string('224488')
    child = doc.styles.add_style('Reset Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    char = doc.styles.add_style('Reset Character', WD_STYLE_TYPE.CHARACTER); char.font.italic = True; char.font.underline = True
    char_child = doc.styles.add_style('Reset Character Child', WD_STYLE_TYPE.CHARACTER); char_child.base_style = char
    huge = doc.styles.add_style('Reset Huge', WD_STYLE_TYPE.CHARACTER); huge.font.size = Pt(200)
    def direct(run, size=24):
        run.font.size = Pt(size); run.bold = False; run.italic = False; run.font.color.rgb = RGBColor.from_string('BB2200')
        props = run._r.get_or_add_rPr(); element(props, 'highlight', val='yellow'); element(props, 'lang', val='fr-CA')
        return run
    def mark(p):
        props = element(p._p.get_or_add_pPr(), 'rPr'); element(props, 'rStyle', val=char_child.style_id)
        element(props, 'sz', val=52); element(props, 'lang', val='fr-CA')
    doc.add_paragraph('Clear direct text formatting')
    p = doc.add_paragraph(style=child); p.add_run('Left ')
    r = direct(p.add_run('reset A\u0301😀', char_child)); r.font.name = 'Courier New'
    element(r._r.rPr, 'u', val='dotted', color='BB2200'); element(r._r.rPr, 'spacing', val=20)
    p.add_run(' right'); mark(p)
    p = doc.add_paragraph(style=child); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.left_indent = Inches(.25)
    direct(p.add_run('Paragraph overrides'))
    p = doc.add_paragraph(style=child); mark(p)
    p = doc.add_paragraph('Language only'); props = p.runs[0]._r.get_or_add_rPr()
    element(props, 'noProof'); element(props, 'lang', val='fr-CA')
    p = doc.add_paragraph(style=child); p.add_run('Style only', char_child)
    p = doc.add_paragraph(style=child); element(p._p, 'bookmarkStart', id=0, name='ResetLink')
    link = element(p._p, 'hyperlink'); link.set(qn('r:id'), doc.part.relate_to('https://example.org/reset', RT.HYPERLINK, is_external=True))
    r = element(link, 'r'); props = element(r, 'rPr'); element(props, 'rStyle', val=char_child.style_id); element(props, 'b', val=0); element(props, 'sz', val=48)
    element(r, 't').text = 'Linked reset'; element(p._p, 'bookmarkEnd', id=0)
    p = doc.add_paragraph(style='List Number'); direct(p.add_run('Numbered reset'), 20)
    p = doc.add_paragraph(); direct(p.add_run('Tabs\tand\nsoft breaks'), 18); p.add_run().add_break(WD_BREAK.PAGE); p.add_run('After page break')
    # Original two-color test pixels, encoded using only the Python standard library.
    def chunk(kind, data): return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xffffffff)
    rows = b''.join(b'\0' + bytes([30, 90, 180] * 8 + [245, 150, 35] * 8) for _ in range(16))
    png = b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 16, 16, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(rows)) + chunk(b'IEND', b'')
    p = doc.add_paragraph('Picture: '); direct(p.add_run(), 16).add_picture(BytesIO(png), width=Inches(.35))
    table = doc.add_table(rows=2, cols=2, style='Table Grid'); table.autofit = True
    element(table.rows[0]._tr.get_or_add_trPr(), 'tblHeader')
    for i, text in enumerate(['Reset header label', 'Other header']):
        p = table.cell(0, i).paragraphs[0]; p.style = child; direct(p.add_run(text), 8)
    for i in range(2):
        p = table.cell(1, i).paragraphs[0]; p.style = child; direct(p.add_run(('Flowing reset text. ' if not i else 'Neighbor content. ') * 65), 20)
    doc.add_paragraph('Tail')
    doc.save(SOURCE); print(f'Generated independent direct-format reset fixture: {SOURCE}')

def verify(file, mode):
    doc = Document(file)
    if mode == 'new':
        assert [p.text for p in doc.paragraphs] == ['First A\u0301😀', 'Second']
        assert all(p.style.name == 'Heading 1' for p in doc.paragraphs)
        assert all(r.style.name == 'Emphasis' for p in doc.paragraphs for r in p.runs)
    else:
        before = Document(SOURCE)
        def data(node): return None if node is None else [(n.tag, sorted(n.attrib.items()), n.text) for n in node.iter()]
        def content(p):
            return [(n.tag, n.text if n.tag == qn('w:t') else sorted(n.attrib.items())) for n in p.iter() if n.tag in [qn('w:t'), qn('w:tab'), qn('w:br'), qn('w:bookmarkStart'), qn('w:bookmarkEnd')]]
        old = before.element.findall('.//' + qn('w:p')); new = doc.element.findall('.//' + qn('w:p')); assert len(old) == len(new)
        def character_styles(p):
            result = []
            for run in p.iter(qn('w:r')):
                ref = run.find('./' + qn('w:rPr') + '/' + qn('w:rStyle')); style = None if ref is None else ref.get(qn('w:val'))
                for n in run:
                    if n.tag == qn('w:t'): result.extend((c, style) for c in n.text or '')
                    elif n.tag in [qn('w:tab'), qn('w:br'), qn('w:drawing')]: result.append((n.tag, style))
            return result
        for a, b in zip(old, new):
            # Run splitting/merging is allowed; text, control order and marker positions stay fixed.
            def stream(p):
                out = []
                for tag, value in content(p): out.extend([(tag, c) for c in value] if tag == qn('w:t') else [(tag, value)])
                return out
            assert stream(a) == stream(b)
            assert character_styles(a) == character_styles(b), 'A character-style reference changed'
            def mark_style(p):
                ref = p.find('./' + qn('w:pPr') + '/' + qn('w:rPr') + '/' + qn('w:rStyle'))
                return None if ref is None else ref.get(qn('w:val'))
            assert mark_style(a) == mark_style(b), 'A paragraph-mark style changed'
            for p in [a, b]:
                props = p.find('./' + qn('w:pPr') + '/' + qn('w:rPr'))
                if props is not None: props.getparent().remove(props)
            assert data(a.find(qn('w:pPr'))) == data(b.find(qn('w:pPr')))
        for tag in ['tblPr', 'tblGrid', 'trPr', 'tcPr', 'drawing']:
            assert [data(n) for n in before.element.iter(qn('w:' + tag))] == [data(n) for n in doc.element.iter(qn('w:' + tag))]
        assert doc.paragraphs[1].runs[1].style.name == 'Reset Character Child'
        links = doc.element.findall('.//' + qn('w:hyperlink')); assert len(links) == 1
        assert links[0].get(qn('r:id')) == before.element.find('.//' + qn('w:hyperlink')).get(qn('r:id'))
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    # Inspect the unmodified export again, including paragraph-mark properties.
    doc = Document(file)
    for props in doc.element.iter(qn('w:rPr')):
        assert all(n.tag == qn('w:rStyle') for n in props), 'A direct run declaration survived the reset'
    print(f'PASS: {file}: direct declarations removed, styles retained, source text and package structures preserved')
