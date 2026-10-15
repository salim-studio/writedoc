"""Independent OOXML controls for manual-break font metrics; no renderer imports."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
out=Path(__file__).resolve().parents[1]/'tests/fixtures/manual-break-metrics'
out.mkdir(exist_ok=True)
cases=[dict(kind=kind,size=size,family='Calibri') for kind in ['text','empty','spaces','trailing','repeated','nbsp','cell'] for size in [10,24]]
cases += [dict(kind='text',size=size,family=family) for size in [None,12,24] for family in ['Verdana','Calibri']]
for mode in [14,15]:
 d=Document();normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(12);normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
 d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
 for setting in d.settings.element.iter(qn('w:compatSetting')):
  if setting.get(qn('w:name'))=='compatibilityMode':setting.set(qn('w:val'),str(mode))
 for i,case in enumerate(cases):
  kind,size,family=case['kind'],case['size'],case['family']
  label=d.add_paragraph(f'Case {i+1}');label.paragraph_format.page_break_before=i>0
  p=d.add_table(rows=1,cols=1).cell(0,0).paragraphs[0] if kind=='cell' else d.add_paragraph()
  rpr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'28');rpr.append(sz);p._p.get_or_add_pPr().append(rpr)
  text='   ' if kind=='spaces' else '\u00a0' if kind=='nbsp' else '' if kind=='empty' else 'First'
  r=p.add_run(text);r.font.name='Verdana';r.font.size=Pt(10);r.bold=True
  r=p.add_run();r.bold=True
  if size is not None:r.font.name=family;r.font.size=Pt(size)
  r.add_break()
  if kind=='repeated':r.add_break()
  if kind!='trailing':
   r=p.add_run('Second');r.font.name='Verdana';r.font.size=Pt(10);r.bold=True
  d.add_paragraph('After')
 d.save(out/f'mode{mode}.docx')
(out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n',encoding='utf-8')

# A page-boundary control distinguishes ignored spaces from a large blank-line
# height: 36pt spaces cannot fit if their metrics contribute to the final break.
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
with ZipFile(out.parents[2]/'public/fixtures/break-metrics-spaces.docx') as z:
 parts={n:z.read(n) for n in z.namelist()}
xml=etree.fromstring(parts['word/document.xml']);ns={'w':xml.nsmap['w']}
p=xml.find('w:body',ns).findall('w:p',ns)[1]
p.find('w:r/w:rPr/w:sz',ns).set(qn('w:val'),'72')
parts['word/document.xml']=etree.tostring(xml)
with ZipFile(out/'page-break-large-space.docx','w',ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
