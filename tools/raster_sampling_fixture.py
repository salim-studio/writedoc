"""Original high-frequency raster controls; no corpus artwork is redistributed."""
from pathlib import Path
from io import BytesIO
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt
from lxml import etree as E
from canvas_group_fixture import NS, tag, add

root=Path(__file__).resolve().parents[1]
image=Image.new('RGB',(1000,500),'white');draw=ImageDraw.Draw(image)
for x in range(0,1000,2):draw.line((x,0,x,249),fill='black')
draw.polygon([(0,500),(1000,250),(1000,500)],fill='#1865B0')
b=BytesIO();image.save(b,format='PNG');b.seek(0)
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.5)
style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(11)
style.paragraph_format.space_after=Pt(6)
doc.add_paragraph('Raster sampling controls')
for name in ['inline','cropped','grouped','shape-fill']:
 doc.add_paragraph(name)
 doc.add_paragraph().add_run().add_picture(BytesIO(b.getvalue()),width=Inches(111/96),height=Inches(55.5/96))
out=BytesIO();doc.save(out)
with ZipFile(out) as z:parts={n:z.read(n) for n in z.namelist()}
settings=E.fromstring(parts['word/settings.xml'])
settings.find('w:zoom',NS).set(tag('w:percent'),'100')
parts['word/settings.xml']=E.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
xml=E.fromstring(parts['word/document.xml'])
pictures=xml.findall('.//pic:pic',NS)
fill=pictures[1].find('pic:blipFill',NS);fill.insert(1,E.Element(tag('a:srcRect'),l='100',r='300',t='0',b='0'))
data=pictures[2].getparent();pic=pictures[2];data.remove(pic);data.set('uri',NS['wpg'])
group=add(data,'wpg:wgp');add(group,'wpg:cNvGrpSpPr');xf=add(add(group,'wpg:grpSpPr'),'a:xfrm')
width=round(111*9525);height=round(55.5*9525)
for t,attrs in [('off',dict(x=0,y=0)),('ext',dict(cx=width,cy=height)),('chOff',dict(x=0,y=0)),('chExt',dict(cx=width,cy=height))]:add(xf,'a:'+t,**attrs)
group.append(pic)
pic=pictures[3];data=pic.getparent();data.remove(pic);data.set('uri',NS['wps'])
shape=add(data,'wps:wsp');add(shape,'wps:cNvSpPr');props=add(shape,'wps:spPr')
props.append(deepcopy(pic.find('pic:spPr/a:xfrm',NS)));add(add(props,'a:prstGeom',prst='rect'),'a:avLst')
fill=deepcopy(pic.find('pic:blipFill',NS));fill.tag=tag('a:blipFill');props.append(fill);add(shape,'wps:bodyPr')
parts['word/document.xml']=E.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
target=root/'public/fixtures/raster-sampling.docx'
with ZipFile(target,'w',ZIP_DEFLATED) as z:
 for n,data in parts.items():z.writestr(n,data)
print(target)
