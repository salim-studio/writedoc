"""Original cell-border fixture and independent python-docx export checks."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/cell-borders.docx'
ORDER = ['top', 'start', 'left', 'bottom', 'end', 'right', 'insideH', 'insideV', 'tl2br', 'tr2bl']

def borders(pr, sides, tag='tcBorders', **attrs):
    parent = pr.find(qn('w:' + tag))
    if parent is None:
        parent = OxmlElement('w:' + tag)
        pr.insert_element_before(parent, 'w:shd', 'w:noWrap', 'w:tcMar', 'w:vAlign', 'w:tblLayout', 'w:tblCellMar')
    for side in ORDER:
        if side not in sides: continue
        old = parent.find(qn('w:' + side))
        if old is not None: parent.remove(old)
        node = OxmlElement('w:' + side)
        for key, value in {'val': 'single', 'sz': '8', 'color': '224466', **attrs}.items(): node.set(qn('w:' + key), value)
        before = next((n for n in parent if ORDER.index(n.tag.split('}')[1]) > ORDER.index(side)), None)
        if before is None: parent.append(node)
        else: before.addprevious(node)

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('Canvas Cell Border Base', WD_STYLE_TYPE.TABLE); base.base_style = doc.styles['Table Grid']
    pr = OxmlElement('w:tblPr'); base.element.append(pr)
    borders(pr, ['top', 'left', 'bottom', 'right'], 'tblBorders', val='double', sz='12', color='112233', themeColor='accent4', themeTint='99')
    pr = OxmlElement('w:tcPr'); base.element.append(pr)
    borders(pr, ['top'], val='single', color='668899')
    borders(pr, ['left'], val='dotted', color='008844')
    borders(pr, ['right'], val='dashed', color='008844')
    child = doc.styles.add_style('Canvas Cell Border Child', WD_STYLE_TYPE.TABLE); child.base_style = base
    pr = OxmlElement('w:tcPr'); child.element.append(pr)
    borders(pr, ['bottom'], val='dashSmallGap', sz='8', color='006699')
    borders(pr, ['right'], val='none')
    def table(rows, cols, height=48):
        result = doc.add_table(rows=rows, cols=cols, style=child); result.autofit = False
        for col in result.columns: col.width = Inches(5 / cols)
        for row in result.rows:
            row.height = Pt(height); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                cell.width = Inches(5 / cols)
                margins = OxmlElement('w:tcMar')
                for side, value in [('top', 150), ('left', 180), ('bottom', 150), ('right', 180)]:
                    n = OxmlElement('w:' + side); n.set(qn('w:w'), str(value)); n.set(qn('w:type'), 'dxa'); margins.append(n)
                cell._tc.get_or_add_tcPr().append(margins)
        return result
    doc.add_paragraph('Cell borders A\u0301😀')
    first = table(4, 3)
    for r, labels in enumerate([['Inherited edges', 'Direct edges', 'None inherits'], ['Nil edge', 'Heavy single', 'Double wins'], ['Dotted loses', 'Thin single', 'Theme A\u0301😀'], ['Horizontal edges', '', 'Corner']]):
        for c, label in enumerate(labels): first.cell(r,c).text = label
    pr = first.cell(0,1)._tc.get_or_add_tcPr()
    borders(pr, ['top'], val='double', sz='24', color='FF0000', themeColor='accent6')
    borders(pr, ['start'], val='single', sz='16', color='993366'); borders(pr, ['left'], val='dotted', sz='8', color='FF0000')
    borders(pr, ['insideH'], val='single', color='FF0000'); borders(pr, ['tl2br'], val='nil')
    borders(first.cell(0,2)._tc.get_or_add_tcPr(), ['left'], val='none', color='FF0000')
    borders(first.cell(1,0)._tc.get_or_add_tcPr(), ['right'], val='nil')
    borders(first.cell(1,1)._tc.get_or_add_tcPr(), ['left'], val='single', sz='32', color='CC3300')
    borders(first.cell(1,1)._tc.get_or_add_tcPr(), ['right'], val='single', sz='12', color='00AA00')
    borders(first.cell(1,2)._tc.get_or_add_tcPr(), ['left'], val='double', sz='8', color='224466')
    borders(first.cell(2,0)._tc.get_or_add_tcPr(), ['right'], val='dotted', sz='96', color='FF0000')
    borders(first.cell(2,1)._tc.get_or_add_tcPr(), ['left'], val='single', sz='2', color='008844')
    borders(first.cell(2,2)._tc.get_or_add_tcPr(), ['right'], val='dotDash', sz='12', color='FF0000', themeColor='accent3', themeShade='BF')
    first.cell(3,0).merge(first.cell(3,1))
    borders(first.cell(3,0)._tc.get_or_add_tcPr(), ['bottom'], val='triple', sz='16', color='AA2244')
    borders(first.cell(3,2)._tc.get_or_add_tcPr(), ['right'], val='thick', sz='12', color='993366', frame='1')
    from canvas_cell_shading_fixture import shade
    pr = first.cell(3,0)._tc.get_or_add_tcPr(); shade(pr, val='clear', fill='DDEEFF'); pr.find(qn('w:tcMar')).addprevious(pr[-1])
    doc.add_paragraph('Merged borders')
    merged = table(3, 2, 42); merged.cell(0,0).merge(merged.cell(2,0)); merged.cell(0,0).text = 'Vertical edges'
    for r, label in enumerate(['First neighbor', 'Second neighbor', 'Last neighbor']): merged.cell(r,1).text = label
    doc.add_paragraph('Before repeated table')
    repeated = table(3, 2)
    repeated.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    repeated.cell(0,0).text = 'Header edges'; repeated.cell(0,1).text = 'Other header'
    borders(repeated.cell(0,0)._tc.get_or_add_tcPr(), ['top'], val='dotDotDash', sz='12', color='AA2244')
    repeated.cell(1,0).text = 'Flowing borders. ' * 140; repeated.cell(1,1).text = 'Short neighbor'
    repeated.rows[2].height = Inches(7)
    repeated.cell(2,0).text = 'Tall border'; repeated.cell(2,1).text = 'Tall neighbor'
    doc.add_paragraph('Tail')
    doc.save(SOURCE); print(f'Generated independent cell-border fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file); cells = doc.element.findall('.//' + qn('w:tc'))
    text = lambda cell: ''.join(n.text or '' for n in cell.findall('.//' + qn('w:t')))
    def declaration(cell, side, value):
        parent = cell.find(qn('w:tcPr')).find(qn('w:tcBorders'))
        node = parent.find(qn('w:' + side)) if parent is not None else None
        if value is None: assert node is None; return
        assert node is not None, (text(cell), side)
        assert dict(node.attrib) == {qn('w:' + k): v for k,v in zip(['val', 'sz', 'color'], value)}, (text(cell), side)
    if fresh:
        assert [text(cell) for cell in cells] == ['New box', 'New clear']
        for cell in cells:
            for side in ['top','left','bottom','right']: declaration(cell, side, ('single','8','auto') if text(cell) == 'New box' else ('nil','0','auto'))
    else:
        original_doc = Document(SOURCE); before = original_doc.element.findall('.//' + qn('w:tc'))
        assert [text(cell) for cell in cells] == [text(cell) for cell in before]
        changes = {'Inherited edges': {'top': ('double','16','AABBCC')}, 'Direct edges': dict.fromkeys(['top','left','start','bottom','right','end']), 'Heavy single': {'right': ('nil','0','auto')}, 'Header edges': {'bottom': ('triple','16','8855CC')}}
        for cell, original in zip(cells, before):
            edits = changes.get(text(cell), {})
            if cell.find(qn('w:tcPr')).find(qn('w:vMerge')) is not None: edits = {'bottom': ('double','16','CC3300')}
            for side, value in edits.items(): declaration(cell, side, value)
            def data(element):
                return [(n.tag, sorted(n.attrib.items()), n.text) for n in element.iter() if n.tag != qn('w:tcBorders') and not (n.getparent() is not None and n.getparent().tag == qn('w:tcBorders') and n.tag.split('}')[1] in edits)]
            assert data(cell) == data(original), f'Unrelated cell properties or text changed: {text(cell)}'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: cell borders, logical aliases, inheritance, merge segments, text and unrelated package parts')
