"""Original line-spacing fixture and independent export checks using python-docx."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/line-spacing.docx'

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(6)
    base = doc.styles.add_style('Canvas Line Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    base.paragraph_format.line_spacing = Pt(24)
    child = doc.styles.add_style('Canvas Line Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    reset = doc.styles.add_style('Canvas Line Reset', WD_STYLE_TYPE.PARAGRAPH); reset.base_style = base
    reset.paragraph_format.line_spacing = 1.5; reset.element.pPr.spacing.attrib.pop(qn('w:lineRule'))
    minimum = doc.styles.add_style('Canvas Line Minimum', WD_STYLE_TYPE.PARAGRAPH); minimum.base_style = normal
    minimum.paragraph_format.line_spacing = Pt(18); minimum.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    rule = doc.styles.add_style('Canvas Line Rule', WD_STYLE_TYPE.PARAGRAPH); rule.base_style = minimum
    spacing = OxmlElement('w:spacing'); spacing.set(qn('w:lineRule'), 'exact'); rule.element.get_or_add_pPr().append(spacing)
    doc.add_paragraph('Line spacing A\u0301😀')
    doc.add_paragraph('Inherited exact A\u0301😀\nSecond soft line', style=child)
    p = doc.add_paragraph('Direct minimum\nContent grows', style=child)
    p.paragraph_format.line_spacing = Pt(12); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
    doc.add_paragraph('Omitted rule\nUses a multiple', style=reset)
    doc.add_paragraph('Rule only\nInherits the line value', style=rule)
    p = doc.add_paragraph('Clipped exact gyp'); p.runs[0].font.size = Pt(26); p.runs[0].font.color.rgb = RGBColor.from_string('CC0000'); p.paragraph_format.line_spacing = Pt(8)
    p = doc.add_paragraph('Mixed typography '); p.add_run('TALL').font.size = Pt(30); p.add_run(' small').font.italic = True
    p.paragraph_format.line_spacing = Pt(6); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    p = doc.add_paragraph('Fractional height\nSecond line'); p.paragraph_format.line_spacing = Pt(18.05)
    p = doc.add_paragraph('Numbered line', style='List Number'); p.paragraph_format.line_spacing = Pt(20)
    doc.add_paragraph('Before flowing table')
    table = doc.add_table(rows=2, cols=2, style='Table Grid'); table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i, text in enumerate(['Header spacing', 'Other header']):
        p = table.cell(0, i).paragraphs[0]; p.style = child; p.add_run(text)
    for i in range(2):
        p = table.cell(1, i).paragraphs[0]; p.style = minimum; p.add_run(('Flowing text. ' if not i else 'Neighbor text. ') * 65)
    doc.add_paragraph('Tail'); doc.save(SOURCE); print(f'Generated independent line-spacing fixture: {SOURCE}')

def verify(file, mode):
    doc = Document(file)
    if mode == 'new':
        assert [p.text for p in doc.paragraphs] == ['First A\u0301😀', 'Second', 'Third']
        for p, rule, value in zip(doc.paragraphs, ['exact', 'atLeast', 'auto'], ['481', '360', '330']):
            assert p._p.pPr.spacing.get(qn('w:lineRule')) == rule
            assert p._p.pPr.spacing.get(qn('w:line')) == value
    else:
        old = Document(SOURCE)
        before = old.element.findall('.//' + qn('w:p')); after = doc.element.findall('.//' + qn('w:p'))
        assert len(before) == len(after)
        expected = {'Inherited exact': ('exact', '600'), 'Direct minimum': None, 'Omitted rule': ('auto', '300'), 'Clipped exact': ('atLeast', '200'), 'Header spacing': ('atLeast', '560')}
        def text(p): return ''.join(n.text or '' for n in p.iter(qn('w:t')))
        def data(node):
            # Serialization may add xml:space="preserve" to text with no edge
            # whitespace. It has no semantic effect there; keep checking the
            # attribute wherever it controls leading or trailing whitespace.
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (k == qn('xml:space') and n.tag == qn('w:t') and n.text == (n.text or '').strip())), n.text) for n in node.iter()]
        for p, prior in zip(after, before):
            key = next((label for label in expected if text(prior).startswith(label)), None)
            assert text(p) == text(prior)
            if key:
                node = p.find('./' + qn('w:pPr') + '/' + qn('w:spacing')); setting = expected[key]
                if setting is None: assert node is None or all(node.get(qn('w:' + a)) is None for a in ['line', 'lineRule'])
                else: assert (node.get(qn('w:lineRule')), node.get(qn('w:line'))) == setting
                for para in [p, prior]:
                    spacing = para.find('./' + qn('w:pPr') + '/' + qn('w:spacing'))
                    if spacing is not None:
                        for attr in ['line', 'lineRule']: spacing.attrib.pop(qn('w:' + attr), None)
                        if not spacing.attrib: spacing.getparent().remove(spacing)
            assert data(p) == data(prior), text(prior)
        assert data(doc.element) == data(old.element), 'Other document or table properties changed'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: line spacing, inherited resets, source text, formatting and unrelated package parts')
