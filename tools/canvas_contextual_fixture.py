"""Original contextual-spacing fixtures and independent python-docx checks."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1] / 'public/fixtures'

def flag(parent, value=True, name='contextualSpacing'):
    node = OxmlElement('w:' + name)
    if value is not None:
        node.set(qn('w:val'), '1' if value else '0')
    parent.append(node)
    return node

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = section.page_height = Inches(6)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('ContextBody', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    base.paragraph_format.space_before = base.paragraph_format.space_after = Pt(8)
    flag(base.element.get_or_add_pPr())
    child = doc.styles.add_style('ContextChild', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    reset = doc.styles.add_style('ContextReset', WD_STYLE_TYPE.PARAGRAPH); reset.base_style = base
    flag(reset.element.get_or_add_pPr(), False)
    def para(text, before=0, after=0, contextual=None, style=None):
        p = doc.add_paragraph(text, style=style)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if contextual is not None: flag(p._p.get_or_add_pPr(), contextual)
        return p
    para('Spec first A\u0301😀', after=10, contextual=True)
    para('Spec second', before=12, contextual=False)
    para('Both first', after=10, contextual=True)
    para('Both second', before=12, contextual=True)
    para('Reverse first', after=12, contextual=False)
    para('Reverse second', before=10, contextual=True)
    for text, style in [('Inherited first', base), ('Inherited second', base), ('Different child', child), ('Child same', child), ('Style reset', reset)]:
        doc.add_paragraph(text, style=style)
    p = doc.add_paragraph('Direct reset', style=base); flag(p._p.get_or_add_pPr(), False)
    para('Implicit default', after=10, contextual=True)
    p = para('Explicit default', before=12, contextual=True)
    p._p.get_or_add_pPr().get_or_add_pStyle().val = 'Normal'
    doc.add_paragraph('', style=base)
    list_style = doc.styles.add_style('ContextList', WD_STYLE_TYPE.PARAGRAPH); list_style.base_style = base
    num = list_style.element.get_or_add_pPr().get_or_add_numPr(); num.get_or_add_numId().val = 7
    numbering = doc.part.numbering_part.element
    instance = next(n for n in numbering.findall(qn('w:num')) if n.get(qn('w:numId')) == '7')
    abstract = next(n for n in numbering.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == instance.find(qn('w:abstractNumId')).get(qn('w:val')))
    level = abstract.find(qn('w:lvl')); flag(level.find(qn('w:pPr')), False)
    doc.add_paragraph('Inherited list', style=list_style)
    p = doc.add_paragraph('Direct list', style=list_style); p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 7
    p = doc.add_paragraph('Direct list reset', style=list_style); p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 7; flag(p._p.pPr)
    table_style = doc.styles.add_style('ContextTable', WD_STYLE_TYPE.TABLE); table_style.base_style = doc.styles['Table Grid']
    table_style.paragraph_format.space_before = table_style.paragraph_format.space_after = Pt(6)
    flag(table_style.element.get_or_add_pPr())
    table = doc.add_table(rows=3, cols=2, style=table_style); table.autofit = False
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    for cell, prefix in [(table.cell(0, 0), 'Header'), (table.cell(0, 1), 'Other cell')]:
        cell.text = prefix + ' first'; cell.add_paragraph(prefix + ' second')
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Cell before'; table.cell(1, 0).add_paragraph('Cell after')
    table.cell(1, 1).text = 'Long table content across pages. ' * 65
    table.cell(2, 0).text = 'Last row'; table.cell(2, 1).text = 'Last cell'
    para('After table', before=12, contextual=True)
    doc.save(ROOT / 'contextual-spacing.docx')
    compat = doc.settings.element.find(qn('w:compat'))
    node = flag(compat, True, 'allowSpaceOfSameStyleInTable')
    # CT_Compat requires this before the existing compatSetting declarations.
    compat.remove(node); compat.find(qn('w:compatSetting')).addprevious(node)
    doc.save(ROOT / 'contextual-table-spacing.docx')
    print('Generated independent contextual-spacing fixtures')

def declared(ppr):
    node = None if ppr is None else ppr.find(qn('w:contextualSpacing'))
    return None if node is None else node.get(qn('w:val'), '1')

def verify(file, mode):
    doc = Document(file)
    if mode == 'new':
        assert [p.text for p in doc.paragraphs] == ['Context first', 'Context second', 'Ordinary third']
        assert [declared(p._p.pPr) for p in doc.paragraphs] == ['1', '1', None]
    else:
        source = ROOT / ('contextual-table-spacing.docx' if mode == 'table' else 'contextual-spacing.docx')
        old = Document(source)
        assert [p.text for p in doc.paragraphs] == [p.text for p in old.paragraphs]
        assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [[c.text for c in r.cells] for r in old.tables[0].rows]
        for i, (p, original) in enumerate(zip(doc.paragraphs, old.paragraphs)):
            expected = '0' if mode == 'body' and i == 0 else declared(original._p.pPr)
            assert declared(p._p.pPr) == expected, p.text
            assert p.paragraph_format.space_before == original.paragraph_format.space_before
            assert p.paragraph_format.space_after == original.paragraph_format.space_after
            assert p.style.style_id == original.style.style_id
        assert declared(doc.styles['ContextBody'].element.find(qn('w:pPr'))) == ('0' if mode == 'body' else '1')
        header = doc.tables[0].cell(0, 0).paragraphs
        assert [declared(p._p.pPr) for p in header] == (['0', '0'] if mode == 'body' else ['0', None])
        def signature(root, parents):
            omitted = {n for p in parents if p is not None for n in p.findall(qn('w:contextualSpacing'))}
            # Direct cell edits may create a previously absent pPr wrapper.
            omitted.update(p for p in parents if p is not None and not p.attrib and all(n in omitted for n in p))
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in root.iter() if n not in omitted]
        def changed_parents(d):
            return ([d.paragraphs[0]._p.pPr] if mode == 'body' else []) + [p._p.pPr for p in d.tables[0].cell(0, 0).paragraphs[:2 if mode == 'body' else 1]]
        assert signature(old.element, changed_parents(old)) == signature(doc.element, changed_parents(doc)), 'Unrelated document XML changed'
        def changed_styles(d): return [d.styles['ContextBody'].element.find(qn('w:pPr'))] if mode == 'body' else []
        assert signature(old.styles.element, changed_styles(old)) == signature(doc.styles.element, changed_styles(doc)), 'Unrelated style XML changed'
        with ZipFile(source) as original, ZipFile(file) as saved:
            assert set(saved.namelist()) == set(original.namelist())
            changed = {'word/document.xml'} | ({'word/styles.xml'} if mode == 'body' else set())
            for name in original.namelist():
                if name not in changed: assert saved.read(name) == original.read(name), name
    print(f'PASS {file}: independent contextual spacing ({mode})')
