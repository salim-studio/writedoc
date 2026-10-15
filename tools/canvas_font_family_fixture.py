"""Independent font-family fixture: original XML, python-docx (MIT), lxml (BSD).

Only font names are referenced. No external font binary is distributed.
"""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as E

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/font-families.docx'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def fonts(parent, **attributes):
    old = parent.find(qn('w:rFonts'))
    if old is not None: parent.remove(old)
    node = OxmlElement('w:rFonts')
    for key, value in attributes.items(): node.set(qn('w:' + key), value)
    parent.insert(0, node)


def generate():
    doc = Document()
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(7), Inches(8)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    default = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    fonts(default, asciiTheme='minorHAnsi', hAnsiTheme='minorHAnsi', eastAsiaTheme='minorEastAsia', cstheme='minorBidi')
    base = doc.styles.add_style('Family Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    fonts(base.element.get_or_add_rPr(), asciiTheme='majorHAnsi', hAnsiTheme='majorHAnsi')
    child = doc.styles.add_style('Family Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    fonts(child.element.get_or_add_rPr(), hAnsi='Times New Roman')
    char = doc.styles.add_style('Family Character', WD_STYLE_TYPE.CHARACTER)
    fonts(char.element.get_or_add_rPr(), ascii='Arial')
    doc.add_paragraph('Font families')
    doc.add_paragraph('Theme default WWW iii')
    doc.add_paragraph('Paragraph theme WWW iii', style=base)
    doc.add_paragraph('Partial override AéΩ A\u0301😀', style=child)
    p = doc.add_paragraph(style=child); p.add_run('Character override é').style = char
    p = doc.add_paragraph(style=base)
    fonts(p.add_run('Direct override WWW iii')._r.get_or_add_rPr(), ascii='Times New Roman', hAnsi='Arial')
    p = doc.add_paragraph()
    fonts(p.add_run('Same declaration theme WWW iii')._r.get_or_add_rPr(), ascii='Arial', asciiTheme='majorHAnsi', hAnsiTheme='majorHAnsi')
    p = doc.add_paragraph()
    fonts(p.add_run('Unavailable WWW iii')._r.get_or_add_rPr(), ascii='Unavailable Fixture Typeface', hAnsi='Unavailable Fixture Typeface')
    p = doc.add_paragraph()
    mark = OxmlElement('w:rPr'); fonts(mark, ascii='Courier New', hAnsi='Courier New'); p._p.get_or_add_pPr().append(mark)
    table = doc.add_table(rows=1, cols=2); table.autofit = True
    for cell, name in zip(table.rows[0].cells, ['Arial', 'Courier New']):
        fonts(cell.paragraphs[0].add_run('WWWW iiii WWWW iiii')._r.get_or_add_rPr(), ascii=name, hAnsi=name)
    p = doc.add_paragraph('Numbering uses its own fonts')
    np = OxmlElement('w:numPr')
    for tag, value in [('ilvl', '0'), ('numId', '1')]:
        n = OxmlElement('w:' + tag); n.set(qn('w:val'), value); np.append(n)
    p._p.get_or_add_pPr().append(np)
    doc.add_paragraph('End of font families.')
    doc.save(SOURCE)
    with ZipFile(SOURCE) as z: parts = {n: z.read(n) for n in z.namelist()}
    theme = E.fromstring(parts['word/theme/theme1.xml'])
    for kind, family in [('majorFont', 'Courier New'), ('minorFont', 'Arial')]:
        collection = theme.find('.//{' + A + '}' + kind)
        collection.find('{' + A + '}latin').set('typeface', family)
        for name in ['ea', 'cs']: collection.find('{' + A + '}' + name).set('typeface', '')
        for node in collection.findall('{' + A + '}font'):
            if node.get('script') == 'Jpan': node.set('typeface', 'Noto Sans')
    settings = E.fromstring(parts['word/settings.xml'])
    settings.find(qn('w:themeFontLang')).set(qn('w:val'), 'en-US')
    settings.find(qn('w:themeFontLang')).set(qn('w:eastAsia'), 'ja-JP')
    numbering = E.fromstring(parts['word/numbering.xml'])
    abstract_id = numbering.find(qn('w:num')).find(qn('w:abstractNumId')).get(qn('w:val'))
    abstract = next(n for n in numbering.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == abstract_id)
    level = abstract.find(qn('w:lvl'))
    level.find(qn('w:numFmt')).set(qn('w:val'), 'decimal')
    level.find(qn('w:lvlText')).set(qn('w:val'), '§%1.')
    rp = level.find(qn('w:rPr'))
    if rp is None: rp = E.SubElement(level, qn('w:rPr'))
    fonts(rp, ascii='Courier New', hAnsi='Arial')
    for path, root in [('word/theme/theme1.xml', theme), ('word/settings.xml', settings), ('word/numbering.xml', numbering)]:
        parts[path] = E.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(SOURCE, 'w', ZIP_DEFLATED) as z:
        for name, data in parts.items(): z.writestr(name, data)


def verify(path):
    with ZipFile(SOURCE) as original, ZipFile(path) as saved:
        assert set(original.namelist()) == set(saved.namelist())
        for name in original.namelist():
            if name != 'word/document.xml': assert original.read(name) == saved.read(name), name
        before, after = [E.fromstring(z.read('word/document.xml')) for z in [original, saved]]
        assert [dict(n.attrib) for n in before.iter(qn('w:rFonts'))] == [dict(n.attrib) for n in after.iter(qn('w:rFonts'))]
        assert ''.join(after.itertext()).startswith('Edited Font families')
    print('PASS retained font declarations, theme and style parts:', path)


if __name__ == '__main__':
    import sys
    if len(sys.argv) == 1: generate()
    else:
        for path in sys.argv[1:]: verify(path)
