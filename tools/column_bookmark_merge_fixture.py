"""Original controls isolating column-bookmark indices from merged table grids."""
import json
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

out = Path(__file__).resolve().parents[1] / 'tests/fixtures/column-bookmarks'
cases = [
    dict(name='PlainMiddle', first=1, last=1),
    dict(name='TopMergedFirst', first=0, last=0, merge_row=0),
    dict(name='TopMergedSecond', first=1, last=1, merge_row=0),
    dict(name='TopMergedBoth', first=0, last=1, merge_row=0),
    dict(name='MiddleMergedFirst', first=0, last=0, merge_row=1),
    dict(name='MiddleMergedSecond', first=1, last=1, merge_row=1),
    dict(name='BottomMergedSecond', first=1, last=1, merge_row=2),
    dict(name='StartAfterMerge', first=1, last=1, merge_row=0, start_row=1),
    dict(name='StartAtMerge', first=1, last=1, merge_row=1, start_row=1),
    dict(name='EndBeforeMerge', first=1, last=1, merge_row=2, end_row=1),
    dict(name='InlineEndpoints', first=1, last=1, inline=True),
    dict(name='TableEnd', first=0, last=2, table_end=True),
]
doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(0)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
for index, case in enumerate(cases):
    if index:
        doc.add_page_break()
    doc.add_paragraph(case['name'])
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    for row in range(3):
        for col in range(3):
            table.cell(row, col).text = f'R{row+1}C{col+1}'
    if 'merge_row' in case:
        row = case['merge_row']
        table.cell(row, 0).merge(table.cell(row, 1))
    start = OxmlElement('w:bookmarkStart')
    for key, value in dict(id=index, name=case['name'], colFirst=case['first'], colLast=case['last']).items():
        start.set(qn('w:' + key), str(value))
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(index))
    first = table.cell(case.get('start_row', 0), 0)
    last = table.cell(case.get('end_row', 2), 2)
    if case.get('inline'):
        first.paragraphs[0]._p.insert(0, start)
        last.paragraphs[0]._p.append(end)
    else:
        first._tc.insert(1, start)
        (table._tbl if case.get('table_end') else last._tc).append(end)
    doc.add_paragraph('After ' + case['name'])
doc.save(out / 'merge-control.docx')
(out / 'merge-cases.json').write_text(json.dumps(cases, indent=2) + '\n')
