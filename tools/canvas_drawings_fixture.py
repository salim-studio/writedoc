"""Original chart/anchor fixture and independent checks; no engine-generated XML.

python-docx (MIT) constructs the package; lxml (BSD-3-Clause) builds DrawingML.
The PNG is original two-color test data encoded with the Python standard library.
"""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
import struct
import zlib
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/charts-and-anchors.docx'
C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
NS = {'c': C, 'a': A, 'wp': WP, 'w': qn('w:p').split('}')[0][1:], 'r': qn('r:id').split('}')[0][1:]}
def add(parent, tag, text=None, **attrs):
    prefix, local = tag.split(':'); node = E.SubElement(parent, '{' + NS[prefix] + '}' + local, {k: str(v) for k, v in attrs.items()})
    node.text = text
    return node

def chart_xml():
    root = E.Element('{' + C + '}chartSpace', nsmap={'c': C, 'a': A})
    add(root, 'c:lang', val='en-US'); chart = add(root, 'c:chart'); plot = add(chart, 'c:plotArea'); add(plot, 'c:layout')
    bars = add(plot, 'c:barChart'); add(bars, 'c:barDir', val='col'); add(bars, 'c:grouping', val='clustered')
    def cache(parent, kind, values):
        ref = add(parent, f'c:{kind}Ref'); add(ref, 'c:f', text='Sheet1!$A$1:$A$4')
        cached = add(ref, f'c:{kind}Cache')
        if kind == 'num': add(cached, 'c:formatCode', text='General')
        add(cached, 'c:ptCount', val=len(values))
        # XML order is not the data index; exercise indexed-cache semantics.
        for index in reversed(range(len(values))): add(add(cached, 'c:pt', idx=index), 'c:v', text=str(values[index]))
    def shape(parent, color=None, line=None):
        node = add(parent, 'c:spPr')
        if color: add(add(node, 'a:solidFill'), 'a:srgbClr', val=color)
        else: add(node, 'a:noFill')
        ln = add(node, 'a:ln')
        if line: add(add(ln, 'a:solidFill'), 'a:srgbClr', val=line)
        else: add(ln, 'a:noFill')
    for index, (name, color, values) in enumerate([('North', '004586', [6, 4, 8, 2]), ('South', 'ff420e', [3, -2, 5, 7]), ('West', 'ffd320', [2, 3, 4, 6])]):
        series = add(bars, 'c:ser'); add(series, 'c:idx', val=index); add(series, 'c:order', val=index)
        cache(add(series, 'c:tx'), 'str', [name]); shape(series, color)
        cache(add(series, 'c:cat'), 'str', ['Q1', 'Q2', 'Q3', 'Q4']); cache(add(series, 'c:val'), 'num', values)
    add(bars, 'c:gapWidth', val=100); add(bars, 'c:axId', val=10); add(bars, 'c:axId', val=20)
    for name, ident, cross, pos in [('catAx', 10, 20, 'b'), ('valAx', 20, 10, 'l')]:
        axis = add(plot, 'c:' + name); add(axis, 'c:axId', val=ident); add(add(axis, 'c:scaling'), 'c:orientation', val='minMax')
        add(axis, 'c:delete', val=0); add(axis, 'c:axPos', val=pos)
        if name == 'valAx': shape(add(axis, 'c:majorGridlines'), line='b3b3b3')
        add(axis, 'c:majorTickMark', val='out'); add(axis, 'c:minorTickMark', val='none'); add(axis, 'c:tickLblPos', val='nextTo')
        shape(axis, line='b3b3b3'); add(axis, 'c:crossAx', val=cross); add(axis, 'c:crossesAt', val=0)
        if name == 'catAx': add(axis, 'c:auto', val=1); add(axis, 'c:lblAlgn', val='ctr'); add(axis, 'c:lblOffset', val=100)
    shape(plot, line='b3b3b3'); legend = add(chart, 'c:legend'); add(legend, 'c:legendPos', val='r'); add(legend, 'c:overlay', val=0); shape(legend)
    add(chart, 'c:plotVisOnly', val=1); shape(root)
    return E.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

def generate(factory):
    doc = factory(); section = doc.sections[0]; section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.size = Pt(11); normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph('Charts and anchors Á😀'); p.runs[0].font.all_caps = False; p.runs[0].font.small_caps = False
    doc.add_paragraph('Before the chart.'); p = doc.add_paragraph(); drawing = add(p.add_run()._r, 'w:drawing')
    inline = add(drawing, 'wp:inline', distT=0, distB=0, distL=0, distR=0)
    add(inline, 'wp:extent', cx=Inches(4.5), cy=Inches(2.2)); add(inline, 'wp:effectExtent', l=0, t=0, r=0, b=0)
    add(inline, 'wp:docPr', id=7, name='Quarterly results', descr='Cached column chart'); add(inline, 'wp:cNvGraphicFramePr')
    graphic = add(inline, 'a:graphic'); data = add(graphic, 'a:graphicData', uri=C)
    part = Part(PackURI('/word/charts/chart1.xml'), 'application/vnd.openxmlformats-officedocument.drawingml.chart+xml', chart_xml(), doc.part.package)
    ref = add(data, 'c:chart'); ref.set(qn('r:id'), doc.part.relate_to(part, RT.CHART))
    doc.add_paragraph('After the chart.'); doc.add_paragraph('An ordinary list item', style='List Bullet')
    table = doc.add_table(rows=2, cols=2, style='Table Grid')
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells): cell.text = f'Cell {i + 1}.{j + 1}'
    doc.add_paragraph('Before the floating picture.'); p = doc.add_paragraph()
    def chunk(kind, data): return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data) & 0xffffffff)
    rows = b''.join(b'\0' + bytes([30, 90, 180] * 32 + [245, 150, 35] * 32) for _ in range(16))
    png = b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 64, 16, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(rows)) + chunk(b'IEND', b'')
    picture = p.add_run().add_picture(BytesIO(png), width=Inches(5)); inline = picture._inline
    inline.tag = '{' + WP + '}anchor'
    for name, value in dict(simplePos=0, relativeHeight=1, behindDoc=0, locked=0, layoutInCell=1, allowOverlap=1, distT=0, distB=0, distL=0, distR=0).items(): inline.set(name, str(value))
    prefix = E.Element('prefix'); add(prefix, 'wp:simplePos', x=0, y=0)
    add(add(prefix, 'wp:positionH', relativeFrom='column'), 'wp:align', text='center')
    add(add(prefix, 'wp:positionV', relativeFrom='paragraph'), 'wp:align', text='top')
    for index, node in enumerate(list(prefix)): inline.insert(index, node)
    wrap = E.Element('{' + WP + '}wrapSquare', wrapText='largest'); inline.insert(4, wrap)
    inline.find('{' + WP + '}docPr').set('descr', 'Floating test panels')
    locks = inline.find('.//{' + A + '}graphicFrameLocks'); locks.set('noChangeAspect', '1')
    pic_pr = inline.find('.//' + qn('pic:cNvPicPr')); add(pic_pr, 'a:picLocks', noChangeArrowheads=1)
    shape = inline.find('.//' + qn('pic:spPr')); add(shape, 'a:noFill'); line = add(shape, 'a:ln', w=9525); add(line, 'a:noFill'); add(line, 'a:miter', lim=800000); add(line, 'a:headEnd'); add(line, 'a:tailEnd')
    doc.add_paragraph('After the floating picture.'); doc.save(SOURCE); print(f'Generated {SOURCE}')

def verify(path):
    doc = Document(path); old = Document(SOURCE)
    assert doc.paragraphs[0].text == 'Edited ' + old.paragraphs[0].text
    assert [p.text for p in doc.paragraphs[1:]] == [p.text for p in old.paragraphs[1:]]
    assert [c.text for t in doc.tables for r in t.rows for c in r.cells] == [c.text for t in old.tables for r in t.rows for c in r.cells]
    root = doc.element
    drawings = root.findall('.//' + qn('w:drawing')); assert len(drawings) == 2
    inline = drawings[0].find('{' + WP + '}inline'); anchor = drawings[1].find('{' + WP + '}anchor')
    assert inline.find('{' + WP + '}extent').get('cx') == str(Inches(4))
    assert inline.find('{' + WP + '}docPr').get('descr') == 'Chart retained after editing'
    assert anchor.find('{' + WP + '}docPr').get('descr') == 'Picture retained after editing'
    def signature(node): return [(n.tag, sorted(n.attrib.items()), n.text) for n in node.iter()]
    original = old.element.find('.//{' + WP + '}anchor'); original.find('{' + WP + '}docPr').set('descr', 'Picture retained after editing')
    assert signature(anchor) == signature(original), 'Floating positioning or picture metadata changed'
    with ZipFile(path) as output, ZipFile(SOURCE) as source:
        assert set(output.namelist()) == set(source.namelist())
        for name in source.namelist():
            if name != 'word/document.xml': assert output.read(name) == source.read(name), f'Changed unrelated part: {name}'
    print(f'PASS {path}: independently verified chart caches, text, table, anchor and unrelated package parts')

if __name__ == '__main__':
    import sys
    if len(sys.argv) > 1:
        for path in sys.argv[1:]: verify(path)
    else:
        def factory():
            doc = Document(); doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100'); return doc
        generate(factory)
