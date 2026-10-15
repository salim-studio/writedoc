"""Original missing-font metadata and changed-theme controls."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
import json
from docx import Document
from docx.shared import Pt
from lxml import etree as E
repo=Path(__file__).resolve().parents[1]
folder=repo/'tests/fixtures/font-family-fallback';folder.mkdir(exist_ok=True)
doc=Document();doc.settings.element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}zoom').set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}percent','100');doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(12)
cases=[('no-entry',[]),('swiss',[('family','swiss')]),('panose',[('panose1','020B0604020202020204'),('family','swiss')]),('notTrueType',[('panose1','020B0604020202020204'),('family','swiss'),('notTrueType',None)]),('roman',[('family','roman')]),('modern',[('family','modern')]),('altName',[('altName','Calibri'),('family','swiss')])]
for i,(name,props) in enumerate(cases):
 run=doc.add_paragraph().add_run(f'{i+1} Hamburgefontsiv WWW iii 0123456789');run.font.name=f'Codex Missing Face {i+1:03}'
for name in ['Arial','Calibri','Times New Roman','Courier New']:
 run=doc.add_paragraph().add_run(name+' Hamburgefontsiv WWW iii 0123456789');run.font.name=name
path=repo/'public/fixtures/font-family-fallback-base.docx';doc.save(path)
with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
ns='http://schemas.openxmlformats.org/wordprocessingml/2006/main';tag=lambda n:'{'+ns+'}'+n
root=E.fromstring(parts['word/fontTable.xml'])
for i,(name,props) in enumerate(cases):
 if name=='no-entry':continue
 font=E.SubElement(root,tag('font'),{tag('name'):f'Codex Missing Face {i+1:03}'})
 for key,value in props:
  child=E.SubElement(font,tag(key))
  if value is not None:child.set(tag('val'),value)
parts['word/fontTable.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(path,'w',ZIP_DEFLATED) as z:
 for name,data in parts.items():z.writestr(name,data)
(folder/'base-cases.json').write_text(json.dumps(cases,indent=2))

r=E.fromstring(parts['word/theme/theme1.xml']);ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
r.find('.//a:majorFont/a:latin',ns).set('typeface','Courier New');r.find('.//a:minorFont/a:latin',ns).set('typeface','Arial')
parts['word/theme/theme1.xml']=E.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(repo/'public/fixtures/font-family-fallback-theme.docx','w',ZIP_DEFLATED) as z:
 for name,data in parts.items():z.writestr(name,data)
