"""Independent capitalization fixture and export checks using python-docx (MIT)."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/capitalization.docx'

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = section.page_height = Inches(6)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = 1
    caps = doc.styles.add_style('CapitalLetters', WD_STYLE_TYPE.PARAGRAPH); caps.base_style = normal; caps.font.all_caps = True
    small = doc.styles.add_style('SmallLetters', WD_STYLE_TYPE.PARAGRAPH); small.base_style = normal; small.font.small_caps = True
    child = doc.styles.add_style('CapitalOff', WD_STYLE_TYPE.PARAGRAPH); child.base_style = caps; child.font.all_caps = False
    char = doc.styles.add_style('CapitalCharacter', WD_STYLE_TYPE.CHARACTER); char.font.all_caps = True
    schar = doc.styles.add_style('SmallCharacter', WD_STYLE_TYPE.CHARACTER); schar.font.small_caps = True
    keep = doc.styles.add_style('SmallKeep', WD_STYLE_TYPE.CHARACTER); keep.font.small_caps = False
    table_style = doc.styles.add_style('CapitalTable', WD_STYLE_TYPE.TABLE); table_style.base_style = doc.styles['Table Grid']; table_style.font.small_caps = True
    doc.add_paragraph('Capitalization a\u0301ß😀')
    doc.add_paragraph('Direct caps').runs[0].font.all_caps = True
    doc.add_paragraph('Direct small').runs[0].font.small_caps = True
    doc.add_paragraph('Paragraph caps', style=caps)
    doc.add_paragraph('Child off', style=child)
    doc.add_paragraph('Character reversed', style=caps).runs[0].style = char
    doc.add_paragraph('Small reversed', style=small).runs[0].style = schar
    doc.add_paragraph('Small kept', style=small).runs[0].style = keep
    doc.add_paragraph('Switch small on caps', style=caps).runs[0].font.small_caps = True
    doc.add_paragraph('Caps off keeps small', style=small).runs[0].font.all_caps = False
    p = doc.add_paragraph(style=caps)
    for text, color in [('a', 'A00000'), ('vatar', '0000A0')]:
        r = p.add_run(text); r.font.color.rgb = RGBColor.from_string(color)
    r = doc.add_paragraph('Script a\u0301ß😀').runs[0]; r.font.small_caps = True; r.font.superscript = True; r.italic = True; r.bold = True; r.underline = True; r.font.size = Pt(24)
    p = doc.add_paragraph('Hello, world! 123 a\u0301ß😀', style=small)
    kern = OxmlElement('w:kern'); kern.set(qn('w:val'), '22'); p.runs[0]._r.get_or_add_rPr().append(kern)
    doc.add_paragraph('Tab\tTail' + ' ' * 100, style=caps)
    p = doc.add_paragraph(); mark = OxmlElement('w:rPr'); mark.append(OxmlElement('w:smallCaps')); p._p.get_or_add_pPr().append(mark)
    p = doc.add_paragraph('Numbered capitals'); p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 7
    root = doc.part.numbering_part.element; instance = next(n for n in root.findall(qn('w:num')) if n.get(qn('w:numId')) == '7'); abstract = next(n for n in root.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == instance.find(qn('w:abstractNumId')).get(qn('w:val')))
    level = abstract.find(qn('w:lvl')); level.find(qn('w:lvlText')).set(qn('w:val'), 'item %1.'); rpr = OxmlElement('w:rPr'); rpr.append(OxmlElement('w:smallCaps')); level.append(rpr)
    table = doc.add_table(rows=3, cols=2, style=table_style); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).text = 'Small header'; table.cell(0, 1).text = 'Caps header'; table.cell(0, 1).paragraphs[0].style = caps
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Table inherited'; table.cell(1, 1).text = 'Small capitals in continued table rows. ' * 40
    table.cell(2, 0).text = 'End cell'; table.cell(2, 1).text = 'Other end'
    doc.add_paragraph('End capitalization')
    doc.save(SOURCE); print(f'Generated independent capitalization fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file)
    roots = [doc.element, doc.styles.element]
    with ZipFile(file) as package:
        if 'word/numbering.xml' in package.namelist(): roots.append(ElementTree.fromstring(package.read('word/numbering.xml')))
    for root in roots:
        for rpr in root.iter(qn('w:rPr')):
            assert not (rpr.find(qn('w:caps')) is not None and rpr.find(qn('w:smallCaps')) is not None), 'Mutually exclusive declarations emitted together'
    if fresh:
        assert [p.text for p in doc.paragraphs] == ['Caps a\u0301ß😀', 'Small letters', 'Normal text', 'Future small']
        for p, mode in zip(doc.paragraphs, ['caps', 'small', 'normal', 'small']):
            for r in p.runs:
                assert bool(r.font.all_caps) == (mode == 'caps') and bool(r.font.small_caps) == (mode == 'small')
    else:
        old = Document(SOURCE)
        assert [p.text for p in doc.paragraphs] == [p.text for p in old.paragraphs]
        assert doc.paragraphs[1].runs[0].font.small_caps is True and doc.paragraphs[1].runs[0].font.all_caps is None
        assert not doc.paragraphs[2].runs[0].font.small_caps and not doc.paragraphs[2].runs[0].font.all_caps
        assert doc.styles['CapitalLetters'].font.all_caps is False
        assert doc.styles['CapitalCharacter'].font.all_caps is None
        def signature(root, omitted):
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in root.iter() if n not in omitted]
        def declarations(parent): return [n for n in parent.iter() if n.tag in [qn('w:caps'), qn('w:smallCaps')]]
        def edits(d): return [n for p in d.paragraphs[1:3] for n in declarations(p._p)]
        assert signature(old.element, edits(old)) == signature(doc.element, edits(doc)), 'Unrelated document XML changed'
        def edits(d): return [n for name in ['CapitalLetters', 'CapitalCharacter'] for n in declarations(d.styles[name].element)]
        assert signature(old.styles.element, edits(old)) == signature(doc.styles.element, edits(doc)), 'Unrelated style XML changed'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: independent capitalization modes, source text, styles and package preservation')
