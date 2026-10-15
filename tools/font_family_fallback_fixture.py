"""Original missing-font family classification controls."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
import json,sys
from docx import Document
from docx.shared import Pt
from lxml import etree as E
root=Path(__file__).resolve().parents[1]
folder=root/'tests/fixtures/font-family-fallback';folder.mkdir(exist_ok=True)
doc=Document();doc.settings.element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}zoom').set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}percent','100');doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(12)
cases=[(name,[('family',name)]) for name in ['auto','roman','swiss','modern','script','decorative']]+[('swiss-serif-panose',[('panose1','02020603050405020304'),('family','swiss')]),('roman-sans-panose',[('panose1','020B0604020202020204'),('family','roman')]),('modern-fixed',[('family','modern'),('pitch','fixed')]),('swiss-alt-unavailable',[('altName','Codex Absent Alternative'),('family','swiss')]),('empty-entry',[]),('self-alternative',[('altName','Codex Missing Face 012'),('family','swiss')])]
for i,(name,props) in enumerate(cases):
 run=doc.add_paragraph().add_run(f'{i+1} Hamburgefontsiv WWW iii 0123456789');run.font.name=f'Codex Missing Face {i+1:03}'
for name in ['Arial','Calibri','Times New Roman','Courier New']:
 run=doc.add_paragraph().add_run(name+' Hamburgefontsiv WWW iii 0123456789');run.font.name=name
path=root/'public/fixtures/font-family-fallback.docx';doc.save(path)
with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
ns='http://schemas.openxmlformats.org/wordprocessingml/2006/main';tag=lambda n:'{'+ns+'}'+n
root=E.fromstring(parts['word/fontTable.xml'])
for i,(name,props) in enumerate(cases):
 if name=='no-entry':continue
 font=E.SubElement(root,tag('font'),{tag('name'):f'Codex Missing Face {i+1:03}'})
 for key,value in props:
  child=E.SubElement(font,tag(key))
  if value is not None:child.set(tag('val'),value)
parts['word/fontTable.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(path,'w',ZIP_DEFLATED) as z:
 for name,data in parts.items():z.writestr(name,data)
(folder/'cases.json').write_text(json.dumps(cases,indent=2))
