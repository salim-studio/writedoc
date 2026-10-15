"""Original shared-style spacing fixture and independent python-docx checks."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/style-line-spacing.docx'

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(6)
    def style(name, base):
        s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH); s.base_style = base; return s
    base = style('Shared Base', normal); base.paragraph_format.line_spacing = Pt(24)
    base.paragraph_format.space_before = Pt(4); base.paragraph_format.space_after = Pt(8)
    child = style('Shared Child', base)
    reset = style('Shared Reset', base); reset.paragraph_format.line_spacing = 1.5
    omitted = style('Shared Omitted', base); omitted.paragraph_format.line_spacing = 1.25
    omitted.element.pPr.spacing.attrib.pop(qn('w:lineRule'))
    rule = style('Shared Rule', base)
    spacing = OxmlElement('w:spacing'); spacing.set(qn('w:lineRule'), 'atLeast'); rule.element.get_or_add_pPr().append(spacing)
    doc.add_paragraph('Shared style line spacing')
    doc.add_paragraph('Base A\u0301😀\nSecond soft line', base)
    doc.add_paragraph('Child follows parent\nSecond soft line', child)
    p = doc.add_paragraph('Direct minimum stays\nSecond soft line', child)
    p.paragraph_format.line_spacing = Pt(18); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    p.runs[0].italic = True
    doc.add_paragraph('Reset becomes inherited\nSecond soft line', reset)
    doc.add_paragraph('Omitted rule stays automatic', omitted)
    doc.add_paragraph('Rule inherits amount\nSecond soft line', rule)
    doc.add_paragraph('', child)
    p = doc.add_paragraph('Numbered paragraph', child)
    num = OxmlElement('w:numPr'); n = OxmlElement('w:numId'); n.set(qn('w:val'), '5'); num.append(n); p._p.get_or_add_pPr().append(num)
    table = doc.add_table(rows=2, cols=2, style='Table Grid'); table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i, text in enumerate(['Shared header', 'Other header']):
        p = table.cell(0, i).paragraphs[0]; p.style = child; p.add_run(text)
    for i in range(2):
        p = table.cell(1, i).paragraphs[0]; p.style = child
        p.add_run(('Flowing shared text. ' if not i else 'Neighbor content. ') * 45)
    doc.add_paragraph('Tail', child)
    doc.save(SOURCE); print(f'Generated independent style line-spacing fixture: {SOURCE}')

def verify(file, mode):
    doc = Document(file)
    def declarations(style):
        spacing = style.element.find('./' + qn('w:pPr') + '/' + qn('w:spacing'))
        return {} if spacing is None else {a: spacing.get(qn('w:' + a)) for a in ['line', 'lineRule'] if spacing.get(qn('w:' + a)) is not None}
    if mode == 'new':
        assert [p.text for p in doc.paragraphs] == ['First A\u0301😀', 'Second']
        assert declarations(doc.styles['Normal']) == {'line': '330', 'lineRule': 'auto'}
        assert all(p.paragraph_format.line_spacing is None for p in doc.paragraphs)
    else:
        expected = {'Shared Base': {'line': '601', 'lineRule': 'exact'}, 'Shared Reset': {}, 'Shared Omitted': {'line': '320', 'lineRule': 'auto'}, 'Shared Rule': {'lineRule': 'atLeast'}}
        for name, attrs in expected.items(): assert declarations(doc.styles[name]) == attrs, name
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/styles.xml', 'word/document.xml']: assert original.read(name) == exported.read(name), name
        prior = Document(SOURCE)
        def data(node):
            # Harmless xml:space may be added when the text has no edge spaces.
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (k == qn('xml:space') and n.tag == qn('w:t') and n.text == (n.text or '').strip())), n.text) for n in node.iter()]
        assert data(doc.element) == data(prior.element), 'Document text, direct formatting or table structure changed'
        for styles in [doc.styles, prior.styles]:
            for name in expected:
                node = styles[name].element.find('./' + qn('w:pPr') + '/' + qn('w:spacing'))
                if node is not None:
                    for attr in ['line', 'lineRule']: node.attrib.pop(qn('w:' + attr), None)
                    if not node.attrib: node.getparent().remove(node)
        assert data(doc.styles.element) == data(prior.styles.element), 'Other style declarations changed'
    print(f'PASS: {file}: shared spacing declarations, independent inheritance, source text and unrelated parts')
