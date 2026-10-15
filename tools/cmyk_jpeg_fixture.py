"""Original device-CMYK JPEG and RGB controls; no embedded ICC profiles."""
from pathlib import Path
from io import BytesIO
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
import json
from zipfile import ZipFile, ZIP_DEFLATED
root=Path(__file__).resolve().parents[1];out=root/'tests/fixtures/cmyk-jpeg';out.mkdir(exist_ok=True)
colors=[(0,0,0,0),(0,255,255,0),(255,0,255,0),(255,255,0,0),(0,0,0,255),(255,0,0,0),(0,255,0,0),(0,0,255,0),(32,64,96,0),(32,64,96,64),(128,128,128,0),(0,0,0,128),(64,128,192,32),(192,64,128,32),(128,192,64,32),(255,255,255,255)]
im=Image.new('CMYK',(256,128));draw=ImageDraw.Draw(im)
for i,c in enumerate(colors):draw.rectangle((i%4*64,i//4*32,i%4*64+63,i//4*32+31),fill=c)
d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
payloads=[]
for name,pic in [('device-cmyk',im),('naive-rgb',im.convert('RGB'))]:
 b=BytesIO();pic.save(b,'JPEG',quality=100,subsampling=0);(out/(name+'.jpg')).write_bytes(b.getvalue())
 payloads.append(b.getvalue())
 placeholder=Image.new('RGB',im.size,(len(payloads)*100,0,0));png=BytesIO();placeholder.save(png,'JPEG')
 d.add_paragraph(name);d.add_paragraph().add_run().add_picture(BytesIO(png.getvalue()),width=Pt(288),height=Pt(144))
buffer=BytesIO();d.save(buffer)
with ZipFile(BytesIO(buffer.getvalue())) as z:parts={n:z.read(n) for n in z.namelist()}
assert len([n for n in parts if n.startswith('word/media/')]) == 2, 'Placeholders must not deduplicate'
for i,b in enumerate(payloads):
 assert 'icc_profile' not in Image.open(BytesIO(b)).info
 parts[f'word/media/image{i+1}.jpg']=b
with ZipFile(out/'control.docx','w',ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
(out/'cases.json').write_text(json.dumps({'colors':colors,'names':['device-cmyk','naive-rgb']},indent=2)+'\n')

