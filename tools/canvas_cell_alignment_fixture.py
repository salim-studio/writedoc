"""Original cell-alignment fixture and independent python-docx export checks."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/cell-alignment.docx'

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1
    base = doc.styles.add_style('Canvas Centered Table', WD_STYLE_TYPE.TABLE); base.base_style = doc.styles['Table Grid']
    pr = OxmlElement('w:tcPr'); base.element.append(pr)
    node = OxmlElement('w:vAlign'); node.set(qn('w:val'), 'center'); pr.append(node)
    child = doc.styles.add_style('Canvas Bottom Table', WD_STYLE_TYPE.TABLE); child.base_style = base
    pr = OxmlElement('w:tcPr'); child.element.append(pr)
    node = OxmlElement('w:vAlign'); node.set(qn('w:val'), 'bottom'); pr.append(node)
    def table(rows, cols, style=base, height=72):
        result = doc.add_table(rows=rows, cols=cols, style=style); result.autofit = False
        for col in result.columns: col.width = Inches(5 / cols)
        for row in result.rows:
            row.height = Pt(height); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                cell.width = Inches(5 / cols)
                margins = OxmlElement('w:tcMar')
                for side, value in [('top', 90), ('left', 120), ('bottom', 150), ('right', 120)]:
                    n = OxmlElement('w:' + side); n.set(qn('w:w'), str(value)); n.set(qn('w:type'), 'dxa'); margins.append(n)
                cell._tc.get_or_add_tcPr().append(margins)
        return result
    doc.add_paragraph('Cell alignment A\u0301😀')
    first = table(3, 3)
    first.cell(0,0).text = 'Inherited center'; first.cell(0,1).text = 'Direct top'; first.cell(0,2).text = 'Direct bottom'
    first.cell(0,1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    first.cell(0,2).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    first.cell(1,0).text = 'Bordered center A\u0301😀'
    from canvas_paragraph_border_fixture import borders, shade
    pr = first.cell(1,0).paragraphs[0]._p.get_or_add_pPr()
    borders(pr, ['top', 'left', 'bottom', 'right'], color='224466', space='1'); shade(pr, 'DDEEFF')
    first.cell(1,1).text = 'First paragraph'; first.cell(1,1).add_paragraph('Second paragraph')
    first.cell(1,1).paragraphs[0].paragraph_format.space_before = Pt(4)
    first.cell(1,2).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    first.cell(2,0).merge(first.cell(2,1)); first.cell(2,0).text = 'Horizontal bottom'
    first.cell(2,0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    first.cell(2,2).text = 'Height reference\nSecond line'
    node = OxmlElement('w:vAlign'); node.set(qn('w:val'), 'both'); first.cell(2,2)._tc.get_or_add_tcPr().append(node)
    doc.add_paragraph('Between tables')
    merged = table(3, 2, child, 54)
    merged.cell(0,0).merge(merged.cell(2,0)); merged.cell(0,0).text = 'Vertical center'
    merged.cell(0,0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row, text in enumerate(['Inherited bottom', 'Middle neighbor', 'Last neighbor']): merged.cell(row,1).text = text
    doc.add_paragraph('Before repeated table')
    repeated = table(3, 2, base, 54)
    repeated.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    repeated.cell(0,0).text = 'Header center'; repeated.cell(0,1).text = 'Header bottom'
    repeated.cell(0,1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    repeated.cell(1,0).text = 'Flowing center. ' * 130
    repeated.cell(1,1).text = 'Short bottom'; repeated.cell(1,1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    repeated.rows[2].height = Inches(8)
    repeated.cell(2,0).text = 'Tall row center'; repeated.cell(2,1).text = 'Tall row bottom'
    repeated.cell(2,1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    doc.add_paragraph('Tail')
    doc.save(SOURCE); print(f'Generated independent cell-alignment fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file)
    cells = doc.element.findall('.//' + qn('w:tc'))
    text = lambda node: ''.join(n.text or '' for n in node.findall('.//' + qn('w:t')))
    by_text = {text(cell): cell for cell in cells}
    def alignment(label, value):
        node = by_text[label].find(qn('w:tcPr')).find(qn('w:vAlign'))
        assert (node.get(qn('w:val')) if node is not None else None) == value, label
    if fresh:
        assert [text(cell) for cell in cells] == ['New top', 'New center', 'New bottom']
        for label, value in [('New top', 'top'), ('New center', 'center'), ('New bottom', 'bottom')]: alignment(label, value)
    else:
        original_doc = Document(SOURCE); before = original_doc.element.findall('.//' + qn('w:tc'))
        assert [text(cell) for cell in cells] == [text(cell) for cell in before]
        changes = {'Inherited center': 'bottom', 'Direct bottom': None, 'Vertical center': 'bottom', 'Header center': 'bottom'}
        for label, value in changes.items(): alignment(label, value)
        for cell, original in zip(cells, before):
            def data(element):
                return [(n.tag, sorted(n.attrib.items()), n.text) for n in element.iter() if not (text(cell) in changes and n.tag == qn('w:vAlign') and n.getparent().tag == qn('w:tcPr'))]
            assert data(cell) == data(original), f'Unrelated cell properties or content changed: {text(cell)}'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: cell alignment, inheritance resets, text, merges, retained formatting and unrelated package parts')
