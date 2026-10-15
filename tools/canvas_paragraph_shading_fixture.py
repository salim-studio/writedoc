"""Original paragraph-fill fixture and independent python-docx export checks."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/paragraph-shading.docx'

def shade(parent, val='clear', **attrs):
    node = OxmlElement('w:shd'); node.set(qn('w:val'), val)
    for key, value in attrs.items(): node.set(qn('w:' + key), value)
    if parent.tag == qn('w:pPr'):
        parent.insert_element_before(node, 'w:tabs', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:jc', 'w:rPr', 'w:sectPr')
    else: parent.append(node)

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1
    defaults = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:pPrDefault')).find(qn('w:pPr'))
    shade(defaults, fill='F4F4F4')
    base = doc.styles.add_style('Canvas Paragraph Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    shade(base.element.get_or_add_pPr(), fill='DDEEFF')
    child = doc.styles.add_style('Canvas Paragraph Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    doc.add_paragraph('Paragraph shading A\u0301😀')
    doc.add_paragraph('Inherited first', child); doc.add_paragraph('Inherited second', child)
    p = doc.add_paragraph('Direct override with a first-line indent', child); shade(p._p.get_or_add_pPr(), fill='224466')
    p.paragraph_format.left_indent = Inches(.5); p.paragraph_format.right_indent = Inches(.25); p.paragraph_format.first_line_indent = Inches(.25)
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(9)
    p = doc.add_paragraph('No fill reveals the page', child); shade(p._p.get_or_add_pPr(), fill='auto')
    p = doc.add_paragraph('Atomic clear without fill', child); shade(p._p.get_or_add_pPr(), color='FF0000')
    p = doc.add_paragraph('Theme background '); shade(p._p.get_or_add_pPr(), fill='010203', themeFill='accent3', themeFillTint='99', themeFillShade='00')
    r = p.add_run('Highlighted'); node = OxmlElement('w:highlight'); node.set(qn('w:val'), 'yellow'); r._r.get_or_add_rPr().append(node)
    p.add_run(' '); r = p.add_run('Run shaded'); shade(r._r.get_or_add_rPr(), fill='CCAAEE')
    p = doc.add_paragraph('Solid foreground'); shade(p._p.get_or_add_pPr(), 'solid', color='0050C8', fill='FF0000')
    p = doc.add_paragraph('Automatic foreground'); shade(p._p.get_or_add_pPr(), 'solid', color='auto')
    p = doc.add_paragraph('Exact clipped line'); p.runs[0].font.size = Pt(24)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; p.paragraph_format.line_spacing = Pt(8)
    shade(p._p.get_or_add_pPr(), fill='66CCDD')
    p = doc.add_paragraph(); shade(p._p.get_or_add_pPr(), fill='55AA77')
    p = doc.add_paragraph('Hanging paragraph across wrapped lines. ' * 4, child)
    p.paragraph_format.left_indent = Inches(.5); p.paragraph_format.right_indent = Inches(.5); p.paragraph_format.first_line_indent = Inches(-.25)
    p = doc.add_paragraph('Flow across pages A\u0301😀. ' * 90); shade(p._p.get_or_add_pPr(), fill='FFE4CC')
    table_style = doc.styles.add_style('Canvas Shaded Table', WD_STYLE_TYPE.TABLE); table_style.base_style = doc.styles['Table Grid']
    shade(table_style.element.get_or_add_pPr(), fill='CCDDEE')
    table = doc.add_table(rows=3, cols=2, style=table_style); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(0,0).text = 'Shaded header'; table.cell(0,1).text = 'Other header'
    p = table.cell(0,0).paragraphs[0]; p.style = child
    shade(table.cell(1,0)._tc.get_or_add_tcPr(), fill='224466')
    p = table.cell(1,0).paragraphs[0]; p.add_run('Cell paragraph'); shade(p._p.get_or_add_pPr(), fill='FFDDAA')
    p = table.cell(1,0).add_paragraph('Cell no fill'); shade(p._p.get_or_add_pPr(), fill='auto')
    p = table.cell(1,0).add_paragraph('Cell automatic foreground'); shade(p._p.get_or_add_pPr(), 'solid', color='auto')
    table.cell(1,1).text = 'Flowing cell. ' * 110
    table.cell(2,0).text = 'Last row'; table.cell(2,1).text = 'Last cell'
    doc.add_paragraph('Tail')
    doc.save(SOURCE); print(f'Generated independent paragraph-shading fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file)
    paragraphs = doc.element.findall('.//' + qn('w:p'))
    text = lambda p: ''.join(n.text or '' for n in p.findall('.//' + qn('w:t')))
    by_text = {text(p): p for p in paragraphs}
    def fill(label, color):
        node = by_text[label].find(qn('w:pPr')).find(qn('w:shd'))
        assert node is not None and node.get(qn('w:val')) == 'clear'
        assert node.get(qn('w:fill')) == color
        assert node.get(qn('w:color')) == 'auto'
        assert not any('theme' in name for name in node.attrib)
    if fresh:
        assert [text(p) for p in paragraphs] == ['Blue paragraph', 'Plain paragraph', 'Last paragraph']
        fill('Blue paragraph', '224466'); fill('Plain paragraph', 'auto'); fill('Last paragraph', '55AA77')
    else:
        old = Document(SOURCE)
        assert [text(p) for p in paragraphs] == [text(p) for p in old.element.findall('.//' + qn('w:p'))]
        fill('Direct override with a first-line indent', 'AABBCC')
        fill('Theme background Highlighted Run shaded', 'auto')
        fill('Shaded header', '55AA77')
        changed = {'Direct override with a first-line indent', 'Theme background Highlighted Run shaded', 'Shaded header'}
        for p, before in zip(paragraphs, old.element.findall('.//' + qn('w:p'))):
            def data(element):
                return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in element.iter() if not (text(p) in changed and n.tag == qn('w:shd') and n.getparent().tag == qn('w:pPr'))]
            assert data(p) == data(before), f'Unrelated paragraph data changed: {text(p)}'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), f'Unrelated part changed: {name}'
    print(f'PASS: {file}: paragraph fills, explicit reset, text, theme/run data and unrelated parts')
