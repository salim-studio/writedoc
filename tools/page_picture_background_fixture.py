"""Original VML page-background controls; no external artwork."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import json
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree as E

root = Path(__file__).resolve().parents[1] / 'tests/fixtures/page-picture-background'
root.mkdir(parents=True, exist_ok=True)
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
V='urn:schemas-microsoft-com:vml'
O='urn:schemas-microsoft-com:office:office'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(12)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
doc.add_paragraph('Page picture background: first page')
doc.add_page_break();doc.add_paragraph('Page picture background: second page')
base=BytesIO();doc.save(base)
with ZipFile(base) as z:parts={n:z.read(n) for n in z.namelist()}
image=Image.new('RGB',(120,80),'white');draw=ImageDraw.Draw(image)
colors=['000000','404040','808080','C0C0C0','FFFFFF','FF0000','00FF00','0000FF','FFFF00','00FFFF','FF00FF','804020']
for i,color in enumerate(colors):
 x=(i%6)*20;y=(i//6)*40;draw.rectangle((x,y,x+19,y+39),fill='#'+color)
cases=[]
for name,kind,dpi,recolor in [('frame96','frame',96,False),('frame192','frame',192,False),('tile96','tile',96,False),('tile192','tile',192,False),('recolor96','frame',96,True)]:
 current=dict(parts);media=BytesIO();image.save(media,format='PNG',dpi=(dpi,dpi));current['word/media/background.png']=media.getvalue()
 document=E.fromstring(current['word/document.xml'])
 background=E.Element(f'{{{W}}}background',nsmap={'v':V,'o':O,'r':R});background.set(f'{{{W}}}color','FFFFFF')
 shape=E.SubElement(background,f'{{{V}}}background',id='_x0000_s1025')
 shape.set(f'{{{O}}}bwmode','white');shape.set(f'{{{O}}}targetscreensize','1024,768')
 fill=E.SubElement(shape,f'{{{V}}}fill',type=kind,color2='#80C0E0',recolor='t' if recolor else 'f')
 fill.set(f'{{{R}}}id','rIdBackground');document.insert(0,background)
 current['word/document.xml']=E.tostring(document,xml_declaration=True,encoding='UTF-8',standalone=True)
 rels=E.fromstring(current['word/_rels/document.xml.rels']);E.SubElement(rels,f'{{{REL}}}Relationship',Id='rIdBackground',Type=R+'/image',Target='media/background.png')
 current['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
 types=E.fromstring(current['[Content_Types].xml']);E.SubElement(types,f'{{{CT}}}Default',Extension='png',ContentType='image/png')
 current['[Content_Types].xml']=E.tostring(types,xml_declaration=True,encoding='UTF-8',standalone=True)
 with ZipFile(root/(name+'.docx'),'w',ZIP_DEFLATED) as z:
  for path,data in current.items():z.writestr(path,data)
 cases.append(dict(name=name,type=kind,dpi=dpi,recolor=recolor,color2='80C0E0',width=120,height=80,colors=colors))
(root/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
