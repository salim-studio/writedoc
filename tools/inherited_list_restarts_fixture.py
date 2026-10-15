"""Original controls for inherited list numbering after a direct instance change."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

def add(parent,name,**attrs):
 n=OxmlElement('w:'+name)
 for k,v in attrs.items():n.set(qn('w:'+k),str(v))
 parent.append(n);return n

def numpr(parent,num,level=0):
 pr=add(parent,'numPr');add(pr,'ilvl',val=level);add(pr,'numId',val=num)

doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12);normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.line_spacing=Pt(18)
s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.left_margin=s.right_margin=Inches(1)
numbering=doc.part.numbering_part.element
for case in range(18):
 style=f'ControlList{case}';abstract=100+case*2;base=100+case*3;override=base+1
 st=doc.styles.add_style(style,WD_STYLE_TYPE.PARAGRAPH);st.base_style=normal;numpr(st.element.get_or_add_pPr(),base)
 child=doc.styles.add_style(style+'Child',WD_STYLE_TYPE.PARAGRAPH);child.base_style=st
 for aid,linked in [(abstract,case!=1),(abstract+1,False)]:
  a=OxmlElement('w:abstractNum');a.set(qn('w:abstractNumId'),str(aid))
  numbering.insert(list(numbering).index(numbering.find(qn('w:num'))),a)
  add(a,'nsid',val=f'{aid:08X}');add(a,'multiLevelType',val='multilevel' if case>=14 else 'singleLevel');lvl=add(a,'lvl',ilvl=0)
  add(lvl,'start',val=1);add(lvl,'numFmt',val='decimal')
  if linked and case<10:add(lvl,'pStyle',val=style)
  add(lvl,'suff',val='tab');add(lvl,'lvlText',val='%1.');add(lvl,'lvlJc',val='left')
  pp=add(lvl,'pPr');add(add(pp,'tabs'),'tab',val='num',pos=360);add(pp,'ind',left=360,hanging=360)
  if case>=14:
   lvl=add(a,'lvl',ilvl=1);add(lvl,'start',val=1);add(lvl,'numFmt',val='decimal');add(lvl,'suff',val='tab');add(lvl,'lvlText',val='%1.%2.');add(lvl,'lvlJc',val='left');pp=add(lvl,'pPr');add(add(pp,'tabs'),'tab',val='num',pos=720);add(pp,'ind',left=720,hanging=360)
 for nid,aid,start in [(base,abstract,None),(override,abstract+1 if case==5 else abstract,None if case in [3,11,13] else 7 if case in [2,12] else 1)]:
  n=add(numbering,'num',numId=nid);add(n,'abstractNumId',val=aid)
  if case>=14 and nid==override:
   if case in [15,17]:add(add(n,'lvlOverride',ilvl=0),'startOverride',val=7)
   add(add(n,'lvlOverride',ilvl=1),'startOverride',val=5)
  elif start is not None:add(add(n,'lvlOverride',ilvl=0),'startOverride',val=start)
 heading=doc.add_paragraph(f'Numbering control {case}');heading.paragraph_format.page_break_before=bool(case)
 if case>=14:
  for text,nid,level in [('Alpha',base,0),('Beta',base,1),('Gamma',base,1),('Restart',override,1),('Following',base,1),('Parent',base,0),('Child',base,1),('Return',override,1)]:
   if case==16 and text=='Child':continue
   if case==17 and text=='Parent':nid=override
   p=doc.add_paragraph(text);numpr(p._p.get_or_add_pPr(),nid,level)
  continue
 if case>=10:
  for text,nid in [('Alpha',base),('Beta',base),('Restart',override),('Following',base),('Return',override),('Last',base)]:
   p=doc.add_paragraph(text,'Normal' if case in [10,11] else style)
   numpr(p._p.get_or_add_pPr(),nid)
  continue
 doc.add_paragraph('Alpha',style);doc.add_paragraph('Beta',style)
 p=doc.add_paragraph('Restart','Normal' if case==4 else style);numpr(p._p.get_or_add_pPr(),override)
 if case==6:
  p=doc.add_paragraph('Explicit base',style);numpr(p._p.get_or_add_pPr(),base)
 if case==8:doc.add_paragraph('Intervening ordinary paragraph')
 doc.add_paragraph('Following',style+'Child' if case==7 else style)
 if case==9:
  table=doc.add_table(rows=2,cols=1)
  for i,row in enumerate(table.rows):p=row.cells[0].paragraphs[0];p.style=style;p.add_run('Table '+str(i))
 doc.add_paragraph('Last',style)
root=Path(__file__).resolve().parents[1];doc.save(root/'public/fixtures/inherited-list-restarts.docx')
