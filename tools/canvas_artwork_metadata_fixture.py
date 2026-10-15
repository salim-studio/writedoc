"""Original fixture for retained drawing metadata and overridden shape styles.

Uses python-docx (MIT), lxml (BSD) and original XML/PNG source below and in
canvas_group_fixture.py. No user document or artwork is redistributed.
"""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
from copy import deepcopy
import struct
import zlib
from lxml import etree as E
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from canvas_group_fixture import NS, add, tag, shape

SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/artwork-metadata.docx'
NS['adec'] = 'http://schemas.microsoft.com/office/drawing/2017/decorative'

def decorative(parent, value='1'):
    ext = add(add(parent, 'a:extLst'), 'a:ext', uri='{C183D7F6-B498-43B3-948B-1728B52AA6E4}')
    node = E.SubElement(ext, tag('adec:decorative'), nsmap={'adec': NS['adec']})
    if value is not None: node.set('val', value)

def generate():
    doc = Document(); doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(6)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal'].paragraph_format.space_after = Pt(8)
    doc.styles['Normal'].paragraph_format.line_spacing = 1
    doc.add_paragraph('Artwork metadata Á😀')
    def chunk(t, data): return struct.pack('>I', len(data)) + t + data + struct.pack('>I', zlib.crc32(t + data) & 0xffffffff)
    png = b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 2, 2, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(b'\0\xd0\x20\x80\x20\x50\xc0' * 2)) + chunk(b'IEND', b'')
    inline = doc.add_paragraph().add_run().add_picture(BytesIO(png), width=Inches(.5))._inline
    picture = deepcopy(inline.find('.//' + tag('pic:pic')))
    fallback = deepcopy(inline.getparent())
    temporary = inline.getparent().getparent().getparent(); temporary.getparent().remove(temporary)
    host = doc.add_paragraph('Locked artwork anchor')
    host.paragraph_format.space_after = Pt(95)
    alt = E.SubElement(host.add_run()._r, tag('mc:AlternateContent'), nsmap={k: NS[k] for k in ['mc', 'wpg', 'wps', 'a', 'pic', 'r']})
    drawing = add(add(alt, 'mc:Choice', Requires='wpg'), 'w:drawing')
    anchor = add(drawing, 'wp:anchor', distT=0, distB=0, distL=114300, distR=114300, simplePos=0, relativeHeight=1, behindDoc=1, locked=1, layoutInCell=1, allowOverlap=1)
    add(anchor, 'wp:simplePos', x=0, y=0)
    add(add(anchor, 'wp:positionH', relativeFrom='column'), 'wp:posOffset', text='0')
    add(add(anchor, 'wp:positionV', relativeFrom='paragraph'), 'wp:posOffset', text='228600')
    width, height = 2743200, 914400
    add(anchor, 'wp:extent', cx=width, cy=height); add(anchor, 'wp:effectExtent', l=0, t=0, r=0, b=0)
    add(anchor, 'wp:wrapNone'); decorative(add(anchor, 'wp:docPr', id=20, name='Locked decorative group'))
    add(anchor, 'wp:cNvGraphicFramePr')
    group = add(add(add(anchor, 'a:graphic'), 'a:graphicData', uri=NS['wpg']), 'wpg:wgp')
    add(group, 'wpg:cNvGrpSpPr')
    def group_transform(parent, w, h):
        t = add(add(parent, 'wpg:grpSpPr'), 'a:xfrm')
        add(t, 'a:off', x=0, y=0); add(t, 'a:ext', cx=w, cy=h)
        add(t, 'a:chOff', x=0, y=0); add(t, 'a:chExt', cx=w, cy=h)
    group_transform(group, width, height)
    nested = add(group, 'wpg:grpSp'); decorative(add(nested, 'wpg:cNvPr', id=21, name='Decorative nested group'), 'true')
    add(nested, 'wpg:cNvGrpSpPr'); group_transform(nested, width, height)
    nested.append(picture)
    decorative(picture.find(tag('pic:nvPicPr')).find(tag('pic:cNvPr')))
    for ident, x, visible in [(22, 685800, True), (23, 0, False)]:
        node = shape(nested, ident, 914400 if visible else width, height, x, 0, '24A890')
        decorative(node.find(tag('wps:cNvPr')), '0' if visible else None)
        props = node.find(tag('wps:spPr'))
        if not visible:
            props.remove(props.find(tag('a:solidFill'))); props.insert(2, E.Element(tag('a:noFill')))
        style = E.Element(tag('wps:style')); node.insert(list(node).index(node.find(tag('wps:bodyPr'))), style)
        for name, index, color in [('lnRef', 2, 'accent1'), ('fillRef', 1, 'accent1'), ('effectRef', 0, 'accent1'), ('fontRef', 'minor', 'lt1')]:
            clr = add(add(style, 'a:' + name, idx=index), 'a:schemeClr', val=color)
            if name == 'lnRef': add(clr, 'a:shade', val=50000)
        body = node.find(tag('wps:bodyPr'))
        for k, v in {'rot': 0, 'numCol': 1, 'spcCol': 0, 'rtlCol': 0, 'fromWordArt': 0, 'forceAA': 0, 'compatLnSpc': 1}.items(): body.set(k, str(v))
        warp = E.Element(tag('a:prstTxWarp'), prst='textNoShape'); add(warp, 'a:avLst'); body.insert(0, warp)
    fallback.find('.//' + tag('wp:docPr')).set('id', '20'); add(alt, 'mc:Fallback').append(fallback)
    doc.add_paragraph('Inline decorative picture')
    inline = doc.add_paragraph().add_run().add_picture(BytesIO(png), width=Inches(.5))._inline
    inline.find(tag('wp:docPr')).set('id', '30'); decorative(inline.find(tag('wp:docPr')), 'false')
    decorative(inline.find('.//' + tag('pic:cNvPr')))
    doc.add_paragraph('Tail remains editable.'); doc.save(SOURCE); print(f'Generated {SOURCE}')

def verify(path):
    def signature(n): return [(x.tag, sorted(x.attrib.items()), x.text or '') for x in n.iter()]
    with ZipFile(SOURCE) as before, ZipFile(path) as after:
        old, new = [E.fromstring(z.read('word/document.xml')) for z in [before, after]]
        texts = lambda r: ''.join(r.xpath('//w:body/w:p/w:r/w:t/text()', namespaces=NS))
        assert texts(new) == 'Edited ' + texts(old)
        for name in ['mc:AlternateContent', 'w:drawing', 'w:sectPr']:
            assert [signature(n) for n in old.findall('.//' + tag(name))] == [signature(n) for n in new.findall('.//' + tag(name))], name
        assert set(before.namelist()) == set(after.namelist())
        for name in before.namelist():
            if name != 'word/document.xml': assert before.read(name) == after.read(name), name
    print(f'PASS {path}: locked anchor, metadata, styles, both drawing branches and unrelated package parts')

if __name__ == '__main__':
    import sys
    if len(sys.argv) == 1: generate()
    else:
        for path in sys.argv[1:]: verify(path)
