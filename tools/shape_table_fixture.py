"""Independent WordprocessingShape table controls; no corpus content or fonts."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
import canvas_group_fixture as g

ROOT = Path(__file__).resolve().parents[1] / 'tests/fixtures/shape-tables'

def generate():
    ROOT.mkdir(exist_ok=True)
    doc = Document()
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    for setting in doc.settings.element.iter(qn('w:compatSetting')):
        if setting.get(qn('w:name')) == 'compatibilityMode': setting.set(qn('w:val'), '15')
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'; normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    for i, (grouped, anchor, nested, height) in enumerate([
        (False, 't', False, 180), (True, 'ctr', False, 240),
        (False, 'b', True, 240), (False, 't', False, 60)
    ]):
        label = doc.add_paragraph(f'Case {i+1}: {anchor}, {"group" if grouped else "shape"}')
        label.paragraph_format.page_break_before = i > 0
        doc.add_paragraph('Before drawing')
        frame = g.add(g.add(doc.add_paragraph().add_run()._r, 'w:drawing'), 'wp:inline', distT=0, distB=0, distL=0, distR=0)
        width, high = 360*12700, height*12700
        g.add(frame, 'wp:extent', cx=width, cy=high)
        g.add(frame, 'wp:docPr', id=i+100, name=f'Table control {i+1}')
        g.add(frame, 'wp:cNvGraphicFramePr')
        data = g.add(g.add(frame, 'a:graphic'), 'a:graphicData', uri=g.NS['wpg' if grouped else 'wps'])
        parent = data
        if grouped:
            parent = g.add(data, 'wpg:wgp'); g.add(parent, 'wpg:cNvGrpSpPr')
            props = g.add(parent, 'wpg:grpSpPr'); g.transform(props, width, high, group=True)
            props[0].find(g.tag('a:chOff')).attrib.update({'x': '0', 'y': '0'})
        shape = g.shape(parent, i+200, width//(2 if grouped else 1), high//(2 if grouped else 1), 0, 0, 'F4F7FF', 'Heading')
        body = shape.find(g.tag('wps:bodyPr')); body.set('anchor', anchor)
        if i == 0: body.set('upright', '1')
        for side in ['lIns','rIns','tIns','bIns']: body.set(side, str(6*12700))
        content = shape.find('.//' + g.tag('w:txbxContent'))
        rp = content.find('.//' + g.tag('w:rPr')); rp.find(g.tag('w:sz')).set(g.tag('w:val'), '24')
        table = doc.add_table(rows=3, cols=2)
        table.autofit = False
        for column in table.columns: column.width = Inches(2)
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                cell.width = Inches(2)
                cell.text = f'R{row_index+1}C{col_index+1}'
                pr = cell._tc.get_or_add_tcPr()
                borders = g.add(pr, 'w:tcBorders')
                for side in ['top','left','bottom','right']:
                    g.add(borders, 'w:'+side, **{'w:val':'single','w:sz':2 if row_index == 2 and side == 'bottom' else 8,'w:color':'246080'})
                g.add(pr, 'w:shd', **{'w:val':'clear','w:fill':'D6EADF' if row_index % 2 == 0 else 'FFFFFF'})
                margins = g.add(pr, 'w:tcMar')
                for side in ['top','left','bottom','right']:
                    g.add(margins, 'w:'+side, **{'w:w':80,'w:type':'dxa'})
        table.cell(0,0).merge(table.cell(0,1)).text = 'Merged heading'
        if nested:
            cell = table.cell(1,1)
            inner = cell.add_table(rows=1, cols=2)
            inner.autofit = False
            for column in inner.columns: column.width = Inches(.7)
            inner.cell(0,0).text = 'Nested A'; inner.cell(0,1).text = 'Nested B'
        content.append(deepcopy(table._tbl)); table._tbl.getparent().remove(table._tbl)
        tail = g.add(content, 'w:p'); g.add(g.add(tail, 'w:r'), 'w:t', text='Tail paragraph')
        doc.add_paragraph('After drawing')
    doc.save(ROOT/'source.docx')

    symbols = Document()
    symbols.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    symbols.styles['Normal'].font.name = 'Arial'
    symbols.styles['Normal'].font.size = Pt(16)
    symbols.add_paragraph('Retained symbol controls')
    frame = g.add(g.add(symbols.add_paragraph().add_run()._r, 'w:drawing'), 'wp:inline', distT=0, distB=0, distL=0, distR=0)
    width, high = 360*12700, 220*12700
    g.add(frame, 'wp:extent', cx=width, cy=high)
    g.add(frame, 'wp:docPr', id=900, name='Symbols')
    g.add(frame, 'wp:cNvGraphicFramePr')
    data = g.add(g.add(frame, 'a:graphic'), 'a:graphicData', uri=g.NS['wps'])
    shape = g.shape(data, 901, width, high, 0, 0, 'FFFFFF', 'temporary')
    content = shape.find('.//' + g.tag('w:txbxContent')); content.clear()
    for i,(face,code) in enumerate([('Wingdings 2','F0A3'),('Wingdings 2','00A3'),('Arial','0041'),('Arial','F041'),('Arial','03A9')]):
        p=g.add(content,'w:p');r=g.add(p,'w:r');pr=g.add(r,'w:rPr')
        g.add(pr,'w:rFonts',**{'w:ascii':'Courier New','w:hAnsi':'Courier New'})
        g.add(r,'w:t',text=f'{i+1}Before ')
        g.add(r,'w:sym',**{'w:font':face,'w:char':code})
        g.add(r,'w:t',text=' After')
    symbols.add_paragraph('After symbols')
    symbols.save(ROOT/'symbols.docx')

if __name__ == '__main__': generate()
