"""Original reference controls for direct numbering cancellation and indent inheritance."""
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
from io import BytesIO
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
def setup(linked=False):
 p=ROOT/'public/fixtures/numbering-cancel-indent.docx'
 d=Document(); s=d.sections[0]; s.left_margin=s.right_margin=Inches(1)
 d.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
 d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
 for name,numbered in [('PlainIndent',False),('NumberedIndent',True)]:
  st=d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.base_style=d.styles['List Bullet' if numbered and linked else 'Normal'];st.paragraph_format.left_indent=Pt(18)
  if numbered:
   pr=st.element.get_or_add_pPr();pr.get_or_add_numPr().get_or_add_numId().val=1
 b=BytesIO();Image.new('RGB',(237,30),'#1255aa').save(b,format='PNG')
 return d,b,p

d,b,p=setup(linked=True)
for i,(style,cancel,where) in enumerate((style,cancel,where) for where in ['body','cell'] for style in ['PlainIndent','NumberedIndent'] for cancel in [False,True]):
 if i:d.add_page_break()
 d.add_paragraph(f'{style} cancel={cancel} {where}')
 if where=='cell':
  t=d.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Pt(237);t.cell(0,0).width=Pt(237);p0=t.cell(0,0).paragraphs[0]
 else:p0=d.add_paragraph()
 p0.style=style;p0.alignment=1
 if cancel:
  p0._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val=0
 p0.add_run().add_picture(BytesIO(b.getvalue()),width=Pt(237))
d.save(p);print(p.resolve())

d,b,p=setup()
st=d.styles['PlainIndent'];st.paragraph_format.right_indent=Pt(12);st.paragraph_format.first_line_indent=Pt(6)
for i,(cancel,direct) in enumerate((c,v) for c in [False,True] for v in ['none','left','right','first','all']):
 if i:d.add_page_break()
 d.add_paragraph(f'cancel={cancel} direct={direct}')
 p0=d.add_paragraph(style='PlainIndent');p0.alignment=0
 if cancel:
  p0._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val=0
 for k,v in [('left',30),('right',24),('first',9)]:
  if direct in [k,'all']:setattr(p0.paragraph_format,{'left':'left_indent','right':'right_indent','first':'first_line_indent'}[k],Pt(v))
 p0.add_run().add_picture(BytesIO(b.getvalue()),width=Pt(237))
d.save(ROOT/'public/fixtures/numbering-cancel-direct.docx')
