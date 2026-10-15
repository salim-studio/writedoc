"""Independent PNG/JPEG declaration mismatch controls; no editor code involved."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as ET
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

out = Path(__file__).resolve().parents[1] / 'tests/fixtures/image-content-types'
out.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(12)
doc.add_paragraph('Image content type controls')
for fmt in ['PNG', 'JPEG']:
    doc.add_paragraph(fmt + ' encoded picture')
    im = Image.new('RGB', (120, 60), '#2060c0')
    draw = ImageDraw.Draw(im)
    draw.rectangle((60, 0, 119, 59), fill='#e08020')
    stream = BytesIO()
    im.save(stream, format=fmt)
    stream.seek(0)
    doc.add_picture(stream, width=Inches(2))
buf = BytesIO()
doc.save(buf)
with ZipFile(buf) as z:
    parts = {n: z.read(n) for n in z.namelist()}
for name in ['correct', 'mismatched', 'jpeg-as-png', 'png-as-jpeg', 'jpeg-payload-png-name', 'png-payload-jpeg-name']:
    variant = dict(parts)
    if 'payload' in name:
        old = next(n for n, b in variant.items() if n.startswith('word/media/') and (b.startswith(b'\xff\xd8') if name.startswith('jpeg') else b.startswith(b'\x89PNG')))
        new = 'word/media/renamed.' + ('png' if name.startswith('jpeg') else 'jpeg')
        variant[new] = variant.pop(old)
        rels = ET.fromstring(variant['word/_rels/document.xml.rels'])
        for n in rels:
            if n.get('Target') == old.removeprefix('word/'):
                n.set('Target', new.removeprefix('word/'))
        variant['word/_rels/document.xml.rels'] = ET.tostring(rels, encoding='utf-8', xml_declaration=True)
    elif name != 'correct':
        types = ET.fromstring(variant['[Content_Types].xml'])
        for n in types:
            mime = n.get('ContentType')
            selected = {'mismatched': ('image/png', 'image/jpeg'), 'jpeg-as-png': ('image/jpeg',), 'png-as-jpeg': ('image/png',)}[name]
            if mime in selected:
                n.set('ContentType', 'image/jpeg' if mime == 'image/png' else 'image/png')
        variant['[Content_Types].xml'] = ET.tostring(types, encoding='utf-8', xml_declaration=True)
    with ZipFile(out / (name + '.docx'), 'w', ZIP_DEFLATED) as z:
        for n, b in variant.items():
            z.writestr(n, b)
