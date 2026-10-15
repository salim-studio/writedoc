"""Independent WordprocessingML pageBreakBefore controls inside retained shapes."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from lxml import etree as E
from canvas_group_fixture import NS,tag,add,shape
root=Path(__file__).resolve().parents[1]/'tests/fixtures/textbox-page-break-before';root.mkdir(parents=True,exist_ok=True)
doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100');s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.left_margin=s.right_margin=s.top_margin=s.bottom_margin=Inches(.5)
n=doc.styles['Normal'];n.font.name='Arial';n.font.size=Pt(12);n.paragraph_format.space_after=Pt(0)
st=doc.styles.add_style('BreakBefore',WD_STYLE_TYPE.PARAGRAPH);st.base_style=n;st.paragraph_format.page_break_before=True
cases=['none','first','middle','last','all','inherited','override-false']
for i,case in enumerate(cases):
 if i:doc.add_page_break()
 doc.add_paragraph('Text box break: '+case)
 p=doc.add_paragraph();drawing=E.SubElement(p.add_run()._r,tag('w:drawing'),nsmap={k:NS[k] for k in ['wps','a']});inline=add(drawing,'wp:inline');w,h=360*12700,180*12700
 add(inline,'wp:extent',cx=w,cy=h);add(inline,'wp:docPr',id=i+1,name=case);data=add(add(inline,'a:graphic'),'a:graphicData',uri=NS['wps']);node=shape(data,100+i,w,h,0,0,'EEEEEE','placeholder')
 content=node.find('.//'+tag('w:txbxContent'));content.clear()
 for j,text in enumerate(['Alpha','Beta','Gamma']):
  p=add(content,'w:p');pr=add(p,'w:pPr')
  if case in ['inherited','override-false']:add(pr,'w:pStyle',**{'w:val':'BreakBefore'})
  if case=='all' or case==['first','middle','last'][j]:add(pr,'w:pageBreakBefore')
  if case=='override-false':add(pr,'w:pageBreakBefore',**{'w:val':'0'})
  add(pr,'w:spacing',**{'w:after':'120'})
  r=add(p,'w:r');rp=add(r,'w:rPr');add(rp,'w:color',**{'w:val':'D02030'});add(r,'w:t',text=text)
 doc.add_paragraph('After the frame')
doc.save(root/'control.docx');print(root/'control.docx')
