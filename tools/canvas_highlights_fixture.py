"""Original highlight fixture and independent DOCX checks; no renderer code reused."""
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
import struct
import zlib
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'public/fixtures/text-highlights.docx'
PALETTE = ['yellow', 'green', 'cyan', 'magenta', 'blue', 'red', 'darkBlue', 'darkCyan', 'darkGreen', 'darkMagenta', 'darkRed', 'darkYellow', 'darkGray', 'lightGray', 'black', 'white']

def highlight(rpr, value):
    n = OxmlElement('w:highlight'); n.set(qn('w:val'), value); rpr.append(n)

def shading(rpr, value='clear', **attrs):
    n = OxmlElement('w:shd'); n.set(qn('w:val'), value)
    for k, v in attrs.items(): n.set(qn('w:' + k), v)
    rpr.append(n)

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(6)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    defaults = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr')); shading(defaults, fill='auto')
    base = doc.styles.add_style('Canvas Shade Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal; shading(base.element.get_or_add_rPr(), fill='DDEEFF')
    child = doc.styles.add_style('Canvas Shade Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    accent = doc.styles.add_style('Canvas Shade Accent', WD_STYLE_TYPE.CHARACTER); shading(accent.element.get_or_add_rPr(), fill='010203', themeFill='accent3', themeFillTint='99', themeFillShade='00')
    doc.add_paragraph('Highlights A\u0301😀')
    p = doc.add_paragraph()
    for text, color in [('Yellow ', 'yellow'), ('plain ', 'none'), ('dark ', 'black'), ('colored text', 'black')]:
        r = p.add_run(text); highlight(r._r.get_or_add_rPr(), color)
        if text == 'colored text':
            n = OxmlElement('w:color'); n.set(qn('w:val'), 'FF0000'); r._r.rPr.insert(0, n)
    p = doc.add_paragraph('Shaded then highlighted', style=base); highlight(p.runs[0]._r.get_or_add_rPr(), 'yellow')
    doc.add_paragraph('Inherited shading', style=child)
    p = doc.add_paragraph(); r = p.add_run('Character shading'); r.style = accent; highlight(r._r.get_or_add_rPr(), 'none')
    p = doc.add_paragraph('Solid foreground'); shading(p.runs[0]._r.get_or_add_rPr(), 'solid', color='445566', fill='FFFF00')
    p = doc.add_paragraph(); p.paragraph_format.tab_stops.add_tab_stop(Inches(3), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run('A\tB'); r.font.size = Pt(16); r.underline = True; highlight(r._r.get_or_add_rPr(), 'yellow'); p.add_run(' Small').font.size = Pt(10)
    p = doc.add_paragraph('Exact line clipped'); p.runs[0].font.size = Pt(20); highlight(p.runs[0]._r.get_or_add_rPr(), 'cyan'); p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; p.paragraph_format.line_spacing = Pt(6)
    p = doc.add_paragraph(); mark = OxmlElement('w:rPr'); highlight(mark, 'magenta'); p._p.get_or_add_pPr().append(mark)
    p = doc.add_paragraph('Marked list'); numpr = OxmlElement('w:numPr'); level = OxmlElement('w:ilvl'); level.set(qn('w:val'), '0'); numpr.append(level)
    num = OxmlElement('w:numId'); num.set(qn('w:val'), '7'); numpr.append(num); p._p.get_or_add_pPr().append(numpr)
    mark = OxmlElement('w:rPr'); highlight(mark, 'yellow'); p._p.pPr.append(mark)
    for row in range(4):
        p = doc.add_paragraph()
        for name in PALETTE[row * 4:row * 4 + 4]:
            r = p.add_run(name + ' '); r.font.size = Pt(10); highlight(r._r.get_or_add_rPr(), name)
    # Original 8x8 RGBA pixels: blue border, transparent center. The highlight
    # must show through the center without replacing or editing the image asset.
    raw = b''.join(b'\x00' + b''.join(bytes((0, 80, 200, 255 if x in [0, 7] or y in [0, 7] else 0)) for x in range(8)) for y in range(8))
    def chunk(kind, data): return struct.pack('!I', len(data)) + kind + data + struct.pack('!I', zlib.crc32(kind + data) & 0xffffffff)
    png = b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('!IIBBBBB', 8, 8, 8, 6, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')
    p = doc.add_paragraph('Image '); r = p.add_run(); r.add_picture(BytesIO(png), width=Inches(.4)); highlight(r._r.get_or_add_rPr(), 'magenta'); p.add_run(' after image')
    table = doc.add_table(rows=3, cols=2, style='Table Grid'); table.autofit = False
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).text = 'Highlighted header'; highlight(table.cell(0, 0).paragraphs[0].runs[0]._r.get_or_add_rPr(), 'yellow')
    table.cell(0, 1).text = 'Shaded header'; table.cell(0, 1).paragraphs[0].style = child
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Wrap highlight A\u0301😀 ' * 55; highlight(table.cell(1, 0).paragraphs[0].runs[0]._r.get_or_add_rPr(), 'green')
    table.cell(1, 1).text = 'Beside body'; table.cell(2, 0).text = 'Last row'; table.cell(2, 1).text = 'Details'
    doc.add_paragraph('Tail')
    doc.save(SOURCE); print(f'Generated independent highlight fixture: {SOURCE}')

def verify(file, fresh=False):
    doc = Document(file)
    def value(rpr, tag):
        n = None if rpr is None else rpr.find(qn('w:' + tag))
        return None if n is None else n.get(qn('w:val'))
    if fresh:
        assert [p.text for p in doc.paragraphs] == ['Yellow text', 'Plain text', 'Green typing']
        assert [value(p.runs[0]._r.rPr, 'highlight') for p in doc.paragraphs] == ['yellow', 'none', 'green']
    else:
        old = Document(SOURCE)
        assert [p.text for p in doc.paragraphs] == [p.text for p in old.paragraphs]
        assert value(doc.paragraphs[1].runs[0]._r.rPr, 'highlight') == 'cyan'
        assert value(doc.paragraphs[2].runs[0]._r.rPr, 'highlight') == 'none'
        assert value(doc.tables[0].cell(0, 0).paragraphs[0].runs[0]._r.rPr, 'highlight') == 'magenta'
        assert value(doc.paragraphs[8]._p.pPr.find(qn('w:rPr')), 'highlight') == 'magenta'
        assert value(doc.paragraphs[9]._p.pPr.find(qn('w:rPr')), 'highlight') == 'yellow'
        assert value(doc.paragraphs[14].runs[1]._r.rPr, 'highlight') == 'magenta'
        assert doc.paragraphs[5].runs[0]._r.rPr.find(qn('w:shd')).get(qn('w:color')) == '445566'
        for p, original in zip(doc.paragraphs, old.paragraphs):
            assert p.style.style_id == original.style.style_id
            assert [r.text for r in p.runs] == [r.text for r in original.runs]
            for r, old_r in zip(p.runs, original.runs):
                def unedited(n):
                    # The codec adds xml:space=preserve on rebuilt text nodes.
                    # That is equivalent when the unchanged text has no edge spaces.
                    return [(e.tag, sorted((k, v) for k, v in e.attrib.items() if not (e.tag == qn('w:t') and k == qn('xml:space') and e.text == (e.text or '').strip())), e.text) for e in n.iter() if e.tag != qn('w:highlight')]
                assert unedited(r._r) == unedited(old_r._r), 'Unrelated run data changed'
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), f'Unrelated part changed: {name}'
    print(f'PASS: {file}: highlight values, explicit removal, marks, inherited shading, original source/media and unrelated parts')
