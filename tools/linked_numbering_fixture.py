"""Independent numStyleLink controls, authored as OOXML rather than application edits."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
root=Path(__file__).resolve().parents[1]/'tests/fixtures/linked-numbering'
root.mkdir(parents=True,exist_ok=True)
doc=Document();doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(11)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
numbering=doc.part.numbering_part.element
for n in list(numbering):numbering.remove(n)
def node(name,**attrs):
 n=OxmlElement('w:'+name)
 for k,v in attrs.items():n.set(qn('w:'+k),str(v))
 return n
def level(index,text,start=1):
 n=node('lvl',ilvl=index)
 for a in [node('start',val=start),node('numFmt',val='decimal'),node('suff',val='tab'),node('lvlText',val=text),node('lvlJc',val='left')]:n.append(a)
 p=node('pPr');tabs=node('tabs');tabs.append(node('tab',val='num',pos=720+index*360));p.append(tabs);p.append(node('ind',left=720+index*360,hanging=360));n.append(p)
 r=node('rPr');r.append(node('rFonts',ascii='Arial',hAnsi='Arial'));n.append(r)
 return n
def abstract(aid,link=None,own=False,style_link=None):
 a=node('abstractNum',abstractNumId=aid);a.append(node('multiLevelType',val='multilevel'))
 if link:a.append(node('numStyleLink',val=link))
 if style_link:a.append(node('styleLink',val=style_link))
 if not link or own:
  a.append(level(0,'LOCAL-%1' if link else '%1.'));a.append(level(1,'%1.%2'))
 numbering.append(a)
def instance(nid,aid,start=None,replace=False):
 n=node('num',numId=nid);n.append(node('abstractNumId',val=aid))
 if start is not None or replace:
  o=node('lvlOverride',ilvl=0)
  if start is not None:o.append(node('startOverride',val=start))
  if replace:o.append(level(0,'OVERRIDE-%1',4))
  n.append(o)
 numbering.append(n)
def style(name,nid=None,base=None):
 s=node('style',type='numbering',styleId=name);s.append(node('name',val=name))
 if base:s.append(node('basedOn',val=base))
 if nid is not None:
  p=node('pPr');n=node('numPr');n.append(node('numId',val=nid));p.append(n);s.append(p)
 doc.styles.element.append(s)
abstract(100,style_link='Link');abstract(101,'Link');abstract(102,'Link');abstract(103,'OverrideLink');abstract(104,'Inherited');abstract(105,'Link',True);abstract(108,style_link='OverrideLink')
abstract(109,style_link='BackOnly');abstract(110,'BackOnly');abstract(111,style_link='WrongTarget');abstract(112,'WrongTarget')
for a in numbering:
 if a.get(qn('w:abstractNumId')) in ['109','111']:
  a.find(qn('w:lvl')).find(qn('w:lvlText')).set(qn('w:val'),('BACK-' if a.get(qn('w:abstractNumId'))=='109' else 'RIGHT-')+'%1')
instance(100,100);instance(101,101);instance(102,102);instance(103,108,7,True);instance(104,103);instance(105,101,20);instance(106,104);instance(107,105)
instance(110,110);instance(112,112)
style('Link',100);style('OverrideLink',103);style('Inherited',base='Link')
style('BackOnly');style('WrongTarget',100)
cases=[]
doc.add_paragraph('Linked numbering controls')
for nid,lev,label in [(100,0,'Direct first'),(101,0,'Linked first'),(102,0,'Second alias'),(101,0,'Linked second'),(100,0,'Direct second'),(101,1,'Linked child'),(102,1,'Other alias child'),(104,0,'Target override first'),(104,0,'Target override second'),(105,0,'Caller override first'),(105,0,'Caller override second'),(101,0,'Linked after restart'),(106,0,'Inherited style'),(107,0,'Local level alongside link'),(110,0,'Backlink without numId'),(112,0,'Conflicting numId')]:
 p=doc.add_paragraph(label);p.paragraph_format.space_after=Pt(6)
 n=node('numPr');n.append(node('ilvl',val=lev));n.append(node('numId',val=nid));p._p.get_or_add_pPr().insert(0,n)
 cases.append({'numId':str(nid),'level':lev,'label':label})
doc.save(root/'control.docx');(root/'cases.json').write_text(json.dumps(cases,indent=2)+'\n',encoding='utf-8')
for n in numbering:
 for link in list(n.findall(qn('w:styleLink'))):n.remove(link)
doc.save(root/'no-backlink.docx')
print(root/'control.docx')
