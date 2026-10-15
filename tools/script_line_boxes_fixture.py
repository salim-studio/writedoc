"""Independent controls separating script ink from line allocation."""
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures/script-line-boxes'
ROOT.mkdir(exist_ok=True)
doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12)
normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
for setting in doc.settings.element.iter(qn('w:compatSetting')):
 if setting.get(qn('w:name'))=='compatibilityMode':setting.set(qn('w:val'),'15')
cases=[]
for family in ['Arial','Calibri','Times New Roman']:
 for mixed in [False,True]:
  for script,shift in [('baseline',0),('superscript',0),('subscript',0),('superscript',14),('subscript',-14)]:
   cases.append((family,mixed,script,shift))
for i,(family,mixed,script,shift) in enumerate(cases):
 title=doc.add_paragraph(f'Case {i+1}: {family}, mixed={mixed}, {script}, position={shift}')
 title.paragraph_format.page_break_before=i>0
 doc.add_paragraph('Before')
 p=doc.add_paragraph()
 for text,active in ([('Start ',False)] if mixed else [])+[('Test',True)]+([(' end',False)] if mixed else []):
  r=p.add_run(text);r.font.name=family;r.font.size=Pt(20)
  if active:
   pr=r._r.get_or_add_rPr()
   if script!='baseline':E.SubElement(pr,qn('w:vertAlign'),{qn('w:val'):script})
   if shift:E.SubElement(pr,qn('w:position'),{qn('w:val'):str(shift)})
 doc.add_paragraph('After')
doc.save(ROOT/'script-line-boxes.docx')
print(f'Generated {len(cases)} line-allocation controls')
