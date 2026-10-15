"""Original inline picture controls with independently authored layout extents."""
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
import json
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
out=root/'tests/fixtures/inline-picture-extents';out.mkdir(parents=True,exist_ok=True)
with ZipFile(root/'public/fixtures/group-picture-flips.docx') as z: pixels=z.read('word/media/quadrants.png')
cases=[dict(angle=0,edges=[0,0,0,0]),dict(angle=90,edges=[0,0,0,0]),dict(angle=270,edges=[.3,0,0,0]),dict(angle=45,edges=[0,0,0,0]),dict(angle=90,edges=[30,30,30,30]),dict(angle=45,edges=[30,30,30,30]),dict(angle=90,edges=None),dict(angle=0,edges=[3,4,5,6])]
for angle in [44.999,45.001,134.999,135,135.001,224.999,225,225.001,314.999,315,315.001,-90,630]:
 cases.append(dict(angle=angle,edges=[0,0,0,0]))
cases.extend([
 dict(angle=90,edges=[3,4,5,6],flipH=True),
 dict(angle=45,edges=[0,0,0,0],flipV=True),
 dict(angle=270,edges=None,crop=[10000,20000,15000,5000]),
])
doc=Document();style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(12)
style.paragraph_format.space_before=Pt(0);style.paragraph_format.space_after=Pt(0);style.paragraph_format.line_spacing=1
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for i,c in enumerate(cases):
 p=doc.add_paragraph('Case '+str(i+1));p.paragraph_format.page_break_before=i in [4,6] or i>=8 and (i-8)%4==0
 p=doc.add_paragraph('Left ');p.paragraph_format.left_indent=Pt(45)
 shape=p.add_run().add_picture(BytesIO(pixels),width=Pt(40),height=Pt(70))
 xf=shape._inline.find('.//'+qn('a:xfrm'));xf.set('rot',str(round(c['angle']*60000)))
 for key in ['flipH','flipV']:
  if c.get(key):xf.set(key,'1')
 if c.get('crop'):
  fill=shape._inline.find('.//'+qn('pic:blipFill'));crop=OxmlElement('a:srcRect')
  for key,value in zip(['l','t','r','b'],c['crop']):crop.set(key,str(value))
  fill.insert(1,crop)
 if c['edges'] is not None:
  ext=OxmlElement('wp:effectExtent')
  for key,value in zip(['l','t','r','b'],c['edges']):ext.set(key,str(round(value*12700)))
  shape._inline.insert(1,ext)
 p.add_run(' Right');doc.add_paragraph('After '+str(i+1))
doc.save(out/'control.docx')
(out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
