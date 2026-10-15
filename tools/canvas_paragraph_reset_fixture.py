"""Original paragraph reset fixture and independent python-docx export checks."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'public/fixtures/clear-paragraph-formatting.docx'

def element(parent, tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items(): node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = Inches(6); section.page_height = Inches(7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name = 'Noto Sans'; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(6)
    base = doc.styles.add_style('Paragraph Reset Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    base.font.bold = True; base.font.size = Pt(13); base.font.color.rgb = RGBColor.from_string('224488')
    f = base.paragraph_format; f.alignment = WD_ALIGN_PARAGRAPH.LEFT; f.space_before = Pt(4); f.space_after = Pt(6); f.line_spacing = 1.25
    f.left_indent = Inches(.15); f.first_line_indent = Inches(.1); f.keep_with_next = False; f.keep_together = False; f.widow_control = True; f.page_break_before = False
    f.tab_stops.add_tab_stop(Inches(1.5), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.DOTS)
    props = base.element.get_or_add_pPr(); element(props, 'contextualSpacing', val=0); element(props, 'outlineLvl', val=2)
    element(element(props, 'pBdr'), 'bottom', val='single', sz=8, space=2, color='224488'); element(props, 'shd', val='clear', fill='EEF2FF')
    # Paragraph properties have ordered children in the official schema.
    order = ['keepNext', 'keepLines', 'pageBreakBefore', 'widowControl', 'pBdr', 'shd', 'tabs', 'spacing', 'ind', 'contextualSpacing', 'jc', 'outlineLvl']
    for child in sorted(list(props), key=lambda n: order.index(n.tag.split('}')[-1])): props.append(child)
    child = doc.styles.add_style('Paragraph Reset Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    numbered = doc.styles.add_style('Paragraph Reset Numbered', WD_STYLE_TYPE.PARAGRAPH); numbered.base_style = child
    element(element(numbered.element.get_or_add_pPr(), 'numPr'), 'numId', val=5)
    # The default template binds numId 5 to ListNumber. This fixture exercises
    # inheritance by our custom style, so give it reciprocal level ownership.
    # Otherwise the native reference suppresses its marker (numbering-style-bindings controls).
    numbering = doc.part.numbering_part.element
    abstract_id = numbering.xpath("./w:num[@w:numId='5']/w:abstractNumId")[0].get(qn('w:val'))
    binding = numbering.xpath(f"./w:abstractNum[@w:abstractNumId='{abstract_id}']/w:lvl[@w:ilvl='0']/w:pStyle")[0]
    binding.set(qn('w:val'), numbered.style_id)
    huge = doc.styles.add_style('Paragraph Reset Huge', WD_STYLE_TYPE.PARAGRAPH); huge.base_style = child; huge.paragraph_format.line_spacing = Pt(600)
    doc.add_paragraph('Clear direct paragraph formatting')
    p = doc.add_paragraph(style=child); p.add_run('Direct paragraph A\u0301😀').italic = True; p.add_run('\tamount').font.color.rgb = RGBColor.from_string('BB2200')
    f = p.paragraph_format; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT; f.left_indent = Inches(.5); f.first_line_indent = Inches(-.15)
    f.space_before = Pt(20); f.space_after = Pt(18); f.line_spacing = Pt(30); f.keep_with_next = True; f.widow_control = False; f.page_break_before = True
    f.tab_stops.add_tab_stop(Inches(1.5), WD_TAB_ALIGNMENT.CLEAR); f.tab_stops.add_tab_stop(Inches(2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DASHES)
    props = p._p.get_or_add_pPr(); element(props, 'contextualSpacing', val=1); element(props, 'outlineLvl', val=9)
    element(element(props, 'pBdr'), 'bottom', val='double', sz=12, space=3, color='BB2200'); element(props, 'shd', val='clear', fill='FFF0DD')
    mark = element(props, 'rPr'); element(mark, 'rStyle', val='Emphasis'); element(mark, 'sz', val=36); element(mark, 'lang', val='fr-CA')
    order = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'widowControl', 'pBdr', 'shd', 'tabs', 'spacing', 'ind', 'contextualSpacing', 'jc', 'outlineLvl', 'rPr']
    for prop in sorted(list(props), key=lambda n: order.index(n.tag.split('}')[-1])): props.append(prop)
    p = doc.add_paragraph('Following paragraph', child); p.runs[0].bold = False; p.runs[0].font.size = Pt(20)
    f = p.paragraph_format; f.alignment = WD_ALIGN_PARAGRAPH.CENTER; f.space_before = Pt(0); f.space_after = Pt(0); f.line_spacing = 2
    p = doc.add_paragraph(style=child); p.paragraph_format.left_indent = Pt(0); mark = element(p._p.get_or_add_pPr(), 'rPr'); element(mark, 'sz', val=44)
    p = doc.add_paragraph('Opaque paragraph property'); element(p._p.get_or_add_pPr(), 'suppressLineNumbers', val=1)
    p = doc.add_paragraph('Style-only paragraph', child); element(p.runs[0]._r.get_or_add_rPr(), 'lang', val='fr-CA')
    p = doc.add_paragraph('Style numbering returns', numbered); element(element(p._p.get_or_add_pPr(), 'numPr'), 'numId', val=0)
    p = doc.add_paragraph('Direct numbering removed', child); num = element(p._p.get_or_add_pPr(), 'numPr'); element(num, 'ilvl', val=0); element(num, 'numId', val=5)
    p = doc.add_paragraph('Style numbering switches', numbered); num = element(p._p.get_or_add_pPr(), 'numPr'); element(num, 'ilvl', val=0); element(num, 'numId', val=1)
    doc.add_paragraph('Inherited numbering', numbered)
    p = doc.add_paragraph(style=child); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; element(p._p, 'bookmarkStart', id=0, name='ParagraphResetLink')
    link = element(p._p, 'hyperlink'); link.set(qn('r:id'), doc.part.relate_to('https://example.org/paragraph-reset', RT.HYPERLINK, is_external=True))
    run = element(link, 'r'); element(element(run, 'rPr'), 'i'); element(run, 't').text = 'Linked paragraph'; element(p._p, 'bookmarkEnd', id=0)
    p = doc.add_paragraph('Breaks\tsoft\nline', child); p.paragraph_format.space_after = Pt(20); p.add_run().add_break(WD_BREAK.PAGE); p.add_run('After manual break')
    # Reuse the original test pixels from canvas_clear_format_fixture.py.
    with ZipFile(ROOT / 'public/fixtures/clear-text-formatting.docx') as images: png = images.read('word/media/image1.png')
    p = doc.add_paragraph('Picture: '); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.add_run().add_picture(BytesIO(png), width=Inches(.35))
    table = doc.add_table(rows=2, cols=2, style='Table Grid'); table.autofit = True; element(table.rows[0]._tr.get_or_add_trPr(), 'tblHeader')
    for i, text in enumerate(['Reset paragraph header', 'Other header']):
        p = table.cell(0, i).paragraphs[0]; p.style = child; p.add_run(text); f = p.paragraph_format; f.space_before = Pt(20); f.space_after = Pt(30); f.line_spacing = 2
    for i in range(2):
        p = table.cell(1, i).paragraphs[0]; p.style = child; p.add_run(('Flowing paragraph reset. ' if not i else 'Neighbor content. ') * 65)
        f = p.paragraph_format; f.line_spacing = 2; f.space_before = Pt(0); f.space_after = Pt(0); f.left_indent = Pt(0)
    doc.add_paragraph('Tail')
    doc.save(SOURCE); print(f'Generated independent paragraph reset fixture: {SOURCE}')

def verify(file, mode):
    doc = Document(file)
    def data(node):
        return None if node is None else [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (k == qn('xml:space') and n.tag == qn('w:t') and n.text == (n.text or '').strip())), n.text) for n in node.iter()]
    if mode == 'new':
        assert [p.text for p in doc.paragraphs] == ['First A\u0301😀', 'Second']
        assert all(p.style.name == 'Heading 1' for p in doc.paragraphs)
        assert all(r.italic is True for p in doc.paragraphs for r in p.runs)
        assert all(p._p.find('./' + qn('w:pPr') + '/' + qn('w:numPr')) is None for p in doc.paragraphs)
    else:
        before = Document(SOURCE); old = before.element.findall('.//' + qn('w:p')); new = doc.element.findall('.//' + qn('w:p')); assert len(old) == len(new)
        for a, b in zip(old, new):
            for p in [a, b]:
                props = p.find(qn('w:pPr'))
                if props is not None:
                    for n in list(props):
                        if n.tag not in [qn('w:pStyle'), qn('w:rPr'), qn('w:sectPr')]: props.remove(n)
                    if not len(props) and not props.attrib: p.remove(props)
            assert data(a) == data(b), 'Source text, text/mark formatting, links or markers changed'
        for tag in ['tblPr', 'tblGrid', 'trPr', 'tcPr', 'drawing', 'sectPr']:
            assert [data(n) for n in before.element.iter(qn('w:' + tag))] == [data(n) for n in doc.element.iter(qn('w:' + tag))]
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name != 'word/document.xml': assert original.read(name) == exported.read(name), name
    doc = Document(file)
    for props in doc.element.iter(qn('w:pPr')):
        assert all(n.tag in [qn('w:pStyle'), qn('w:rPr'), qn('w:sectPr')] for n in props), 'Direct paragraph properties survived'
    print(f'PASS: {file}: paragraph overrides removed, style/mark/text formatting and package structures retained')
