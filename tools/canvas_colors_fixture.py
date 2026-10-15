"""Original color fixture/verification using python-docx and lxml (independent of JS)."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'public/fixtures/text-colors.docx'

def color(rpr, value='auto', **attrs):
    n = OxmlElement('w:color'); n.set(qn('w:val'), value)
    for k, v in attrs.items(): n.set(qn('w:' + k), v)
    rpr.append(n)

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.line_spacing = 1; normal.paragraph_format.space_after = Pt(0)
    color(doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr')), '445566')
    base = doc.styles.add_style('Canvas Color Base', WD_STYLE_TYPE.PARAGRAPH); base.base_style = normal
    color(base.element.get_or_add_rPr(), '010203', themeColor='accent1')
    child = doc.styles.add_style('Canvas Color Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
    tint = doc.styles.add_style('Canvas Color Tint', WD_STYLE_TYPE.CHARACTER)
    color(tint.element.get_or_add_rPr(), '010203', themeColor='accent3', themeTint='99', themeShade='00')
    shade = doc.styles.add_style('Canvas Color Shade', WD_STYLE_TYPE.CHARACTER)
    color(shade.element.get_or_add_rPr(), '010203', themeColor='accent2', themeShade='BF')
    p = doc.add_paragraph('Colors ')
    color(p.add_run('red ')._r.get_or_add_rPr(), 'FF0000'); color(p.add_run('blue ')._r.get_or_add_rPr(), '0000FF'); p.add_run('A\u0301😀')
    doc.add_paragraph('Paragraph theme', style=base); doc.add_paragraph('Inherited theme', style=child)
    p = doc.add_paragraph(); p.add_run('Character tint').style = tint; p.add_run(' shaded').style = shade
    p = doc.add_paragraph('Automatic override', style=base); color(p.runs[0]._r.get_or_add_rPr())
    p = doc.add_paragraph(); r = p.add_run('Underlined\t text'); r.underline = True; color(r._r.get_or_add_rPr(), '008800')
    r = p.add_run(' black underline'); r.underline = True; color(r._r.get_or_add_rPr(), 'FF0000'); r._r.rPr.find(qn('w:u')).set(qn('w:color'), '000000')
    p = doc.add_paragraph('Numbered color', style=base)
    numpr = OxmlElement('w:numPr'); level = OxmlElement('w:ilvl'); level.set(qn('w:val'), '0'); numpr.append(level)
    num = OxmlElement('w:numId'); num.set(qn('w:val'), '7'); numpr.append(num); p._p.get_or_add_pPr().append(numpr)
    mark = OxmlElement('w:rPr'); color(mark, 'AA00AA'); p._p.pPr.append(mark)
    p = doc.add_paragraph(); mark = OxmlElement('w:rPr'); color(mark, '00AA00'); p._p.get_or_add_pPr().append(mark)
    table = doc.add_table(rows=3, cols=2, style='Table Grid'); table.autofit = False
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).text = 'Colored header'; table.cell(0, 0).paragraphs[0].style = child
    table.cell(0, 1).text = 'Dark automatic'; color(table.cell(0, 1).paragraphs[0].runs[0]._r.get_or_add_rPr())
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), '202020'); table.cell(0, 1)._tc.get_or_add_tcPr().append(shd)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Table color across pages. ' * 55; table.cell(1, 0).paragraphs[0].style = child
    table.cell(1, 1).text = 'Beside body'; table.cell(2, 0).text = 'Last row'; table.cell(2, 1).text = 'Details'
    doc.add_paragraph('Tail theme', style=child)
    doc.save(SOURCE)
    with ZipFile(SOURCE) as archive: parts = {name: archive.read(name) for name in archive.namelist()}
    settings = etree.fromstring(parts['word/settings.xml']); mapping = settings.find(qn('w:clrSchemeMapping')); mapping.set(qn('w:accent1'), 'accent2')
    parts['word/settings.xml'] = etree.tostring(settings)
    theme = etree.fromstring(parts['word/theme/theme1.xml']); theme.find('.//' + qn('a:accent3')).find(qn('a:srgbClr')).set('val', '4F81BD'); parts['word/theme/theme1.xml'] = etree.tostring(theme)
    numbering = etree.fromstring(parts['word/numbering.xml']); abstract_id = numbering.find(qn('w:num') + '[@' + qn('w:numId') + '="7"]').find(qn('w:abstractNumId')).get(qn('w:val')); abstract = numbering.find(qn('w:abstractNum') + '[@' + qn('w:abstractNumId') + '="' + abstract_id + '"]'); lvl = abstract.find(qn('w:lvl')); rpr = OxmlElement('w:rPr'); color(rpr, '010203', themeColor='accent3'); lvl.append(rpr); parts['word/numbering.xml'] = etree.tostring(numbering)
    with ZipFile(SOURCE, 'w', ZIP_DEFLATED) as archive:
        for name, data in parts.items(): archive.writestr(name, data)
    print(f'Generated independent text color fixture: {SOURCE}')

def verify(file, fresh=False):
    doc = Document(file)
    def declaration(rpr):
        n = None if rpr is None else rpr.find(qn('w:color'))
        return None if n is None else {etree.QName(k).localname: v for k, v in n.attrib.items()}
    if fresh:
        assert [p.text for p in doc.paragraphs] == ['Red text', 'Automatic text', 'Green typing']
        assert declaration(doc.paragraphs[0].runs[0]._r.rPr) == {'val': 'CC2200'}
        assert declaration(doc.paragraphs[1].runs[0]._r.rPr) == {'val': 'auto'}
        assert declaration(doc.paragraphs[2].runs[0]._r.rPr) == {'val': '008844'}
    else:
        assert doc.paragraphs[0].text == 'Colors red blue A\u0301😀'
        assert declaration(doc.paragraphs[0].runs[1]._r.rPr) == {'val': '663399'}
        styles = {s.style_id: s for s in doc.styles}
        assert declaration(styles['CanvasColorBase'].element.rPr) == {'val': '008844'}
        assert declaration(styles['CanvasColorTint'].element.rPr) == {'val': '010203', 'themeColor': 'accent3', 'themeTint': '99', 'themeShade': '00'}
        assert declaration(doc.paragraphs[4].runs[0]._r.rPr) == {'val': 'auto'}
        assert declaration(doc.paragraphs[5].runs[0]._r.rPr) == {'val': '008800'}
        assert doc.paragraphs[5].runs[1]._r.rPr.find(qn('w:u')).get(qn('w:color')) == '000000'
        assert declaration(doc.paragraphs[7]._p.pPr.find(qn('w:rPr'))) == {'val': '00AA00'}
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), f'Unrelated part changed: {name}'
        assert len(doc.tables) == 1 and doc.tables[0].cell(0, 0).text == 'Colored header'
    print(f'PASS: {file}: RGB, automatic and theme declarations, inheritance, underlines, source text and unrelated parts')
