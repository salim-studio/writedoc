"""Independent python-docx fixture for the editor's font picker and navigation.
Run with tools/requirements-conformance.txt installed; no application code used.
"""
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
doc = Document()
# python-docx's default template omits the schema-required zoom percentage.
zoom = doc.settings.element.find(qn('w:zoom'))
if zoom is not None:
    zoom.set(qn('w:percent'), '100')
doc.add_heading('Editor workspace', 1)
p = doc.add_paragraph()
r = p.add_run('Choose this font éΩ')
fonts = OxmlElement('w:rFonts')
for key, value in {'asciiTheme': 'minorAscii', 'hAnsiTheme': 'minorHAnsi', 'eastAsia': 'Noto Sans', 'cs': 'Noto Sans'}.items():
    fonts.set(qn('w:' + key), value)
r._r.get_or_add_rPr().append(fonts)
doc.add_paragraph('Body text for links, paragraph fill and spacing.')
doc.add_heading('Next steps', 2)
doc.add_paragraph('A second section stays visible in the outline.')
r = doc.add_paragraph().add_run('Missing font example')
r.font.name = 'Unavailable Editor Fixture Font'
doc.save(root / 'public/fixtures/editor-ux.docx')
