"""Original paired-picture controls for allowOverlap with wrapNone."""
from pathlib import Path
from io import BytesIO
import json
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def node(name, **attrs):
    element = OxmlElement('wp:' + name)
    for key, value in attrs.items(): element.set(key, str(value))
    return element

def picture(paragraph, color, x, y, behind, allow, ident):
    stream = BytesIO(); Image.new('RGB', (16, 16), color).save(stream, format='PNG'); stream.seek(0)
    inline = paragraph.add_run().add_picture(stream, width=Inches(1), height=Inches(1))._inline
    anchor = node('anchor', distT=0, distB=0, distL=0, distR=0, simplePos=0,
                  relativeHeight=ident, behindDoc=behind, locked=0, layoutInCell=1, allowOverlap=allow)
    anchor.append(node('simplePos', x=0, y=0))
    for axis, value in [('H', x), ('V', y)]:
        position = node('position' + axis, relativeFrom='page'); offset = node('posOffset')
        offset.text = str(value * 12700); position.append(offset); anchor.append(position)
    anchor.append(inline.find(qn('wp:extent')))
    anchor.append(node('effectExtent', l=0, t=0, r=0, b=0)); anchor.append(node('wrapNone'))
    properties = inline.find(qn('wp:docPr')); properties.set('id', str(ident)); properties.set('name', f'Control {ident}')
    anchor.append(properties)
    anchor.append(inline.find(qn('wp:cNvGraphicFramePr'))); anchor.append(inline.find(qn('a:graphic')))
    inline.getparent().replace(inline, anchor)

doc = Document(); doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
s = doc.sections[0]; s.page_width, s.page_height = Inches(8.5), Inches(11)
s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
style = doc.styles['Normal']; style.font.name = 'Arial'; style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(0); style.paragraph_format.line_spacing = Pt(15)
cases = [dict(behind=behind, red=red, blue=blue) for behind in [0, 1]
         for red, blue in [(1,1), (0,0), (0,1), (1,0)]]
cases += [dict(behind=0, red=0, blue=0, redXY=red, blueXY=blue)
          for red, blue in [((144,144),(144,144)), ((180,180),(144,144)),
                            ((144,180),(180,144)), ((500,144),(530,180)),
                            ((144,144),(216,144)), ((144,144),(144,216)),
                            ((144,144),(215,144)), ((144,144),(144,215))]]
cases += [dict(behind=0, red=0, blue=0, redXY=(x,144), blueXY=(x+30,180))
          for x in [400, 450, 468, 469]]
cases += [dict(behind=behind, blueBehind=1-behind, red=0, blue=0) for behind in [0,1]]
for i, case in enumerate(cases):
    heading = doc.add_paragraph('No-wrap overlap case ' + str(i)); heading.paragraph_format.page_break_before = bool(i)
    for _ in range(4): doc.add_paragraph()
    p = doc.add_paragraph('Text beneath the picture overlap region.')
    picture(p, '#CC3344', *case.get('redXY', (144,144)), case['behind'], case['red'], i * 2 + 1)
    picture(p, '#2255AA', *case.get('blueXY', (180,180)), case.get('blueBehind', case['behind']), case['blue'], i * 2 + 2)
    doc.add_paragraph('Following paragraph remains in normal flow.')
root = Path(__file__).resolve().parents[1]; target = root / 'public/fixtures/no-wrap-overlap.docx'
doc.save(target)
out = root / 'tests/fixtures/no-wrap-overlap'; out.mkdir(exist_ok=True)
(out / 'cases.json').write_text(json.dumps(cases, indent=2) + '\n')
print(target)
