"""Original vertical-merge border ownership and row reserve controls."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
root=Path(__file__).resolve().parents[1]
folder=root/'tests/fixtures/merged-border-ownership';folder.mkdir(parents=True,exist_ok=True)
doc=Document();doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
normal.paragraph_format.space_before=normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.line_spacing=Pt(12)
cases=[('baseline',{}),('owner-bottom',{(0,'bottom'):('single',36,'D02020')}),('middle-bottom',{(1,'bottom'):('single',36,'D02020')}),('last-bottom',{(2,'bottom'):('single',36,'2060D0')}),('owner-and-last',{(0,'bottom'):('single',36,'D02020'),(2,'bottom'):('single',16,'2060D0')}),('last-nil',{(0,'bottom'):('single',36,'D02020'),(2,'bottom'):('nil',0,'2060D0')}),('interior-opposed',{(0,'bottom'):('single',16,'D02020'),(1,'top'):('single',48,'208040')}),('table-inside',{}),('table-bottom',{}),('table-inside-owner-nil',{(0,'bottom'):('nil',0,'D02020')})]
def border(parent,side,value):
 n=OxmlElement('w:'+side)
 for k,v in zip(['val','sz','color'],value):n.set(qn('w:'+k),str(v))
 parent.append(n)
for index,(name,edges) in enumerate(cases):
 title=doc.add_paragraph(name);title.paragraph_format.page_break_before=bool(index)
 table=doc.add_table(rows=3,cols=2);table.autofit=False
 for c in table.columns:c.width=Inches(2)
 props=table._tbl.tblPr
 margins=OxmlElement('w:tblCellMar')
 for side in ['top','left','bottom','right']:
  n=OxmlElement('w:'+side);n.set(qn('w:w'),'0');n.set(qn('w:type'),'dxa');margins.append(n)
 props.insert(list(props).index(props.find(qn('w:tblLook'))),margins)
 if name.startswith('table-'):
  b=OxmlElement('w:tblBorders');border(b,'bottom' if name=='table-bottom' else 'insideH',('single',24,'8040A0'));props.insert(list(props).index(props.find(qn('w:tblLayout'))),b)
 for ri,row in enumerate(table.rows):
  for ci,cell in enumerate(row.cells):
   pr=cell._tc.get_or_add_tcPr();cell.width=Inches(2)
   if ci==0:
    m=OxmlElement('w:vMerge')
    if ri==0:m.set(qn('w:val'),'restart')
    pr.append(m)
    b=OxmlElement('w:tcBorders')
    for (r,side),value in edges.items():
     if ri==r:border(b,side,value)
    if len(b):pr.append(b)
   cell.paragraphs[0].text=f'R{ri}' if ci else ('Merged' if ri==0 else '')
 doc.add_paragraph('After')
doc.save(folder/'controls.docx')
(folder/'cases.json').write_text(json.dumps([name for name,_ in cases],indent=2)+'\n')
