"""Independent strikethrough fixture and export checks using python-docx (MIT)."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
SOURCE = Path(__file__).resolve().parents[1] / 'public/fixtures/strikethrough.docx'

def generate(factory):
    doc = factory(); section = doc.sections[0]
    section.page_width = section.page_height = Inches(6)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.5)
    normal = doc.styles['Normal']; normal.font.name, normal.font.size = 'Noto Sans', Pt(12)
    normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = 1
    single = doc.styles.add_style('StrikeSingle', WD_STYLE_TYPE.PARAGRAPH); single.base_style = normal; single.font.strike = True
    child = doc.styles.add_style('StrikeOff', WD_STYLE_TYPE.PARAGRAPH); child.base_style = single; child.font.strike = False
    double = doc.styles.add_style('StrikeDouble', WD_STYLE_TYPE.PARAGRAPH); double.base_style = normal; double.font.double_strike = True
    dchild = doc.styles.add_style('StrikeDoubleChild', WD_STYLE_TYPE.PARAGRAPH); dchild.base_style = double; dchild.font.double_strike = True
    char = doc.styles.add_style('StrikeCharacter', WD_STYLE_TYPE.CHARACTER); char.font.strike = True
    off = doc.styles.add_style('StrikeCharacterOff', WD_STYLE_TYPE.CHARACTER); off.font.strike = False
    table_style = doc.styles.add_style('StrikeTable', WD_STYLE_TYPE.TABLE); table_style.base_style = doc.styles['Table Grid']; table_style.font.double_strike = True
    doc.add_paragraph('Strikethrough A\u0301😀')
    doc.add_paragraph('Direct single').runs[0].font.strike = True
    doc.add_paragraph('Direct double').runs[0].font.double_strike = True
    doc.add_paragraph('Paragraph single', style=single)
    doc.add_paragraph('Child off', style=child)
    doc.add_paragraph('Character reversed', style=single).runs[0].style = char
    doc.add_paragraph('Character kept', style=single).runs[0].style = off
    doc.add_paragraph('Double child stays on', style=dchild)
    doc.add_paragraph('Double direct off', style=double).runs[0].font.double_strike = False
    doc.add_paragraph('Switch double on single', style=single).runs[0].font.double_strike = True
    doc.add_paragraph('Switch single on double', style=double).runs[0].font.strike = True
    doc.add_paragraph('Single off keeps double', style=double).runs[0].font.strike = False
    r = doc.add_paragraph('Script A\u0301😀').runs[0]; r.font.double_strike = True; r.font.superscript = True; r.italic = True; r.bold = True; r.underline = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor.from_string('A00000'); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    r = doc.add_paragraph('Tab\tTail' + ' ' * 100, style=single).runs[0]; r.font.highlight_color = WD_COLOR_INDEX.TURQUOISE
    p = doc.add_paragraph(); mark = OxmlElement('w:rPr'); mark.append(OxmlElement('w:dstrike')); p._p.get_or_add_pPr().append(mark)
    p = doc.add_paragraph('Numbered strike'); p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 7
    root = doc.part.numbering_part.element; instance = next(n for n in root.findall(qn('w:num')) if n.get(qn('w:numId')) == '7'); abstract = next(n for n in root.findall(qn('w:abstractNum')) if n.get(qn('w:abstractNumId')) == instance.find(qn('w:abstractNumId')).get(qn('w:val')))
    rpr = OxmlElement('w:rPr'); rpr.append(OxmlElement('w:dstrike')); abstract.find(qn('w:lvl')).append(rpr)
    table = doc.add_table(rows=3, cols=2, style=table_style); table.autofit = True
    for col in table.columns: col.width = Inches(2.5)
    for row in table.rows:
        for cell in row.cells: cell.width = Inches(2.5)
    table.cell(0, 0).text = 'Header double'; table.cell(0, 1).text = 'Header single'; table.cell(0, 1).paragraphs[0].style = single
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).text = 'Table inherited'; table.cell(1, 1).text = 'Strikethrough across continued table rows. ' * 80
    table.cell(2, 0).text = 'End cell'; table.cell(2, 1).text = 'Other end'
    doc.add_paragraph('End strikethrough')
    doc.save(SOURCE); print(f'Generated independent strikethrough fixture: {SOURCE}')

def verify(file, fresh):
    doc = Document(file)
    roots = [doc.element, doc.styles.element]
    with ZipFile(file) as package:
        if 'word/numbering.xml' in package.namelist(): roots.append(ElementTree.fromstring(package.read('word/numbering.xml')))
    for root in roots:
        for rpr in root.iter(qn('w:rPr')):
            assert not (rpr.find(qn('w:strike')) is not None and rpr.find(qn('w:dstrike')) is not None), 'Mutually exclusive declarations emitted together'
    if fresh:
        assert [p.text for p in doc.paragraphs] == ['Single A\u0301😀', 'Double strike', 'Off again', 'Future double']
        for p, mode in zip(doc.paragraphs, ['single', 'double', 'none', 'double']):
            for r in p.runs:
                assert bool(r.font.strike) == (mode == 'single') and bool(r.font.double_strike) == (mode == 'double')
    else:
        old = Document(SOURCE)
        assert [p.text for p in doc.paragraphs] == [p.text for p in old.paragraphs]
        assert doc.paragraphs[1].runs[0].font.double_strike is True and doc.paragraphs[1].runs[0].font.strike is None
        assert not doc.paragraphs[2].runs[0].font.double_strike and not doc.paragraphs[2].runs[0].font.strike
        assert doc.styles['StrikeSingle'].font.strike is False
        assert doc.styles['StrikeCharacter'].font.strike is None
        def signature(root, omitted):
            return [(n.tag, sorted((k, v) for k, v in n.attrib.items() if not (n.tag == qn('w:t') and k == qn('xml:space') and n.text == (n.text or '').strip())), n.text) for n in root.iter() if n not in omitted]
        def declarations(parent): return [n for n in parent.iter() if n.tag in [qn('w:strike'), qn('w:dstrike')]]
        def edits(d): return [n for p in d.paragraphs[1:3] for n in declarations(p._p)]
        assert signature(old.element, edits(old)) == signature(doc.element, edits(doc)), 'Unrelated document XML changed'
        def edits(d): return [n for name in ['StrikeSingle', 'StrikeCharacter'] for n in declarations(d.styles[name].element)]
        assert signature(old.styles.element, edits(old)) == signature(doc.styles.element, edits(doc)), 'Unrelated style XML changed'
        assert [[c.text for row in t.rows for c in row.cells] for t in doc.tables] == [[c.text for row in t.rows for c in row.cells] for t in old.tables]
        with ZipFile(SOURCE) as original, ZipFile(file) as exported:
            assert set(original.namelist()) == set(exported.namelist())
            for name in original.namelist():
                if name not in ['word/document.xml', 'word/styles.xml']: assert original.read(name) == exported.read(name), name
    print(f'PASS: {file}: independent strike modes, mutually exclusive XML, source text, styles and package parts')
