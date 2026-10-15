from pathlib import Path
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
folder=Path(__file__).resolve().parents[1]/'tests/fixtures/shifted-inline-leading';folder.mkdir(parents=True,exist_ok=True)
doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'),'100')
for n in doc.settings.element.iter(qn('w:compatSetting')):
 if n.get(qn('w:name'))=='compatibilityMode':n.set(qn('w:val'),'15')
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12);normal.paragraph_format.space_before=normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
png=BytesIO();Image.new('RGB',(8,8),(208,32,32)).save(png,format='PNG');data=png.getvalue();cases=[]
for shift,image_shift in [(0,0),(30,30),(-30,-30),(30,0),(0,30),(-30,0),(0,-30)]:
 for multiple in [.8,1,1.2,2]:
  name=f'Text {shift} Image {image_shift} Multiple {multiple}';title=doc.add_paragraph(name);title.paragraph_format.page_break_before=bool(cases)
  table=doc.add_table(rows=1,cols=1);table.autofit=False;table.columns[0].width=Inches(4)
  row=table.rows[0];row.height=Pt(216)
  cell=row.cells[0];cell.width=Inches(4)
  valign=OxmlElement('w:vAlign');valign.set(qn('w:val'),'center');cell._tc.get_or_add_tcPr().append(valign)
  p=cell.paragraphs[0];p.paragraph_format.line_spacing=multiple
  p.add_run('Left ');p.add_run().add_picture(BytesIO(data),width=Pt(72),height=Pt(72));p.add_run(' Right')
  for ri,run in enumerate(p.runs):
   pos=OxmlElement('w:position');pos.set(qn('w:val'),str((image_shift if ri==1 else shift)*2));run._r.get_or_add_rPr().append(pos)
  mark=OxmlElement('w:rPr');pos=OxmlElement('w:position');pos.set(qn('w:val'),str(shift*2));mark.append(pos);p._p.get_or_add_pPr().append(mark)
  doc.add_paragraph('After');cases.append(dict(name=name,shift=shift,imageShift=image_shift,multiple=multiple))
doc.save(folder/'controls.docx');(folder/'cases.json').write_text(json.dumps(cases,indent=2))
