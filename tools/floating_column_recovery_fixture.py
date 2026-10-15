"""Original paired rich-document controls for floating-table column recovery."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'tests/fixtures/floating-column-preview'
OUT.mkdir(parents=True, exist_ok=True)


def element(parent, name, **attrs):
    node = OxmlElement('w:' + name)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node


def picture(color):
    image = Image.new('RGB', (96, 36), '#FFFFFF')
    draw = ImageDraw.Draw(image)
    draw.rectangle((1, 1, 94, 34), fill=color)
    draw.ellipse((6, 6, 30, 30), fill='#FFFFFF')
    stream = BytesIO()
    image.save(stream, format='PNG')
    stream.seek(0)
    return stream


def table(doc, name, floating):
    t = doc.add_table(rows=2, cols=2)
    t.autofit = False
    pr = t._tbl.tblPr
    width = pr.find(qn('w:tblW'))
    width.set(qn('w:type'), 'dxa')
    width.set(qn('w:w'), '2880')
    if floating:
        pos = OxmlElement('w:tblpPr')
        for key, value in {'horzAnchor':'margin', 'vertAnchor':'text',
                           'tblpX':'240', 'tblpY':'120', 'rightFromText':'120'}.items():
            pos.set(qn('w:' + key), value)
        pr.insert(0, pos)
    for column in t.columns:
        column.width = Pt(72)
    for row, texts in zip(t.rows, [[name, 'Value'], ['Rich cell', 'Retained']]):
        for cell, text in zip(row.cells, texts):
            cell.width = Pt(72)
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
            element(cell._tc.get_or_add_tcPr(), 'shd', val='clear', fill='E4F1FA')
    doc.add_paragraph(name + ' anchor and following text.')


def write(mode, floating):
    doc = Document()
    doc.settings.element.find(qn('w:zoom')).set(qn('w:percent'), '100')
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.75)
    section.header_distance = section.footer_distance = Inches(.3)
    cols = section._sectPr.find(qn('w:cols'))
    cols.set(qn('w:num'), '2')
    cols.set(qn('w:space'), '360')
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Arial', Pt(11)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = Pt(14)
    for setting in doc.settings.element.iter(qn('w:compatSetting')):
        if setting.get(qn('w:name')) == 'compatibilityMode':
            setting.set(qn('w:val'), '15')
    section.header.paragraphs[0].text = 'Healthy page header'
    footer = section.footer.paragraphs[0]
    footer.add_run('Healthy page footer ')
    field = element(footer._p, 'fldSimple', instr=' PAGE ')
    element(element(field, 'r'), 't').text = '1'
    doc.add_paragraph('Rich column recovery').runs[0].bold = True
    if mode in ['before', 'multiple']:
        table(doc, 'First table', floating)
    doc.add_picture(picture('#267EAE'), width=Inches(1))
    doc.add_paragraph('Healthy formatted paragraph between objects.').runs[0].italic = True
    table(doc, 'Healthy table', False)
    doc.add_picture(picture('#327C42'), width=Inches(1))
    if mode in ['after', 'multiple']:
        table(doc, 'Last table', floating)
    doc.add_paragraph('Healthy column text. ' * 24)
    doc.add_paragraph('Final retained paragraph.').runs[0].bold = True
    data = BytesIO()
    doc.save(data)
    name = f'{mode}-' + ('unsupported' if floating else 'matched') + '.docx'
    # Stable ZIP timestamps make repeated source generation reproducible.
    archive = BytesIO()
    with ZipFile(data) as source, ZipFile(archive, 'w', ZIP_DEFLATED) as target:
        for path in sorted(source.namelist()):
            info = ZipInfo(path, (2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(path))
    path = OUT / name
    if not path.exists() or path.read_bytes() != archive.getvalue():
        path.write_bytes(archive.getvalue())


for case in ['before', 'after', 'multiple']:
    for active in [False, True]:
        write(case, active)
